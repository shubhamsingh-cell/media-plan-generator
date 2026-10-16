"""2026-10-16 sweep (follow-up) -- deck plurals + the last double-possessive sites.

Evidence: the 26-bundle sweep (``ledger/2026-10-16/B_sweep``).

1. DECK PLURALS. The k_tiny_nonprofit deck printed "Client goal: 1 hires --
   this plan projects 0" (slide 2) and "1 hires target" (slide 8). Every
   ``{n} hires`` / ``{n} channels`` in ppt_generator now goes through the same
   ``count_noun`` rule the workbook uses ("1 hire", "0 hires", "1,250 hires").

2. DOUBLE POSSESSIVES that ``display_format.possessive`` had not reached:
   ``insight_composer`` (the "presence in this pool" / "reputation with"
   skeletons), ``research`` (competitor intelligence) and
   ``google_search_pagespeed`` (the 403 career-page note). A name that already
   ends in "'s" or in "s" must never become "McDonald's's" / "Stores's".

Runs under pytest, or standalone:
``python3 tests/test_deck_plurals_and_possessive_sweep_2026_10_16.py``.
"""

from __future__ import annotations

import importlib.util
import io
import os
import re
import sys
from typing import Any

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import google_search_pagespeed  # noqa: E402
import insight_composer  # noqa: E402
import ppt_generator as ppt  # noqa: E402
import research  # noqa: E402
from pptx import Presentation  # noqa: E402


def _load_deck_qa(name: str):
    """Import scripts/deck_qa/<name>.py (a script dir, not a package)."""
    qa_dir = os.path.join(PROJECT_ROOT, "scripts", "deck_qa")
    if qa_dir not in sys.path:
        sys.path.insert(0, qa_dir)
    spec = importlib.util.spec_from_file_location(
        f"deck_qa_{name}", os.path.join(qa_dir, f"{name}.py")
    )
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


matrix = _load_deck_qa("matrix")


@pytest.fixture(autouse=True)
def _reset_active_currency():
    """_set_active_currency keeps the plan currency in a thread-local that
    outlives a single test in the same thread."""
    yield
    try:
        ppt._currency_tls.code = "USD"
    except AttributeError:
        pass


def _all_text(prs: Any) -> str:
    parts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    parts.extend(c.text for c in row.cells)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def _tiny_nonprofit_deck(n_channels: int = 4) -> str:
    """The k_tiny_nonprofit shape: one stated hire target ("1-10 hires" parses
    to a goal of 1), one location, one role, $3,000 for a month."""
    data = matrix._base_plan(
        channel_categories=matrix._channel_categories(n_channels),
        client_name="Riverbend Community Food Bank",
        industry="general_entry_level",
        industry_label="General / Entry-Level",
        budget="3,000",
        campaign_duration="1 month",
        hire_volume="1-10 hires",
        locations=[{"city": "Sacramento", "state": "CA", "country": "United States"}],
        roles=[{"title": "Warehouse Associate", "count": 1, "tier": "entry"}],
        target_roles=["Warehouse Associate"],
        experience_level="entry",
    )
    data = matrix._with_budget_allocation(data, n_channels=n_channels)
    return _all_text(Presentation(io.BytesIO(ppt.generate_pptx(data))))


# "1 hires", "1 channels", ... and the "(s)" hedge. "1" must be a whole token
# ("11 hires", "21 channels" are correct).
_BAD_SINGULAR_RE = re.compile(
    r"(?<![\d,.])1 (hires|locations|roles|channels|markets|platforms|competitors)\b"
    r"|\bhire\(s\)|\brole\(s\)|\blocation\(s\)|\bchannel\(s\)",
    re.IGNORECASE,
)
_DOUBLE_POSSESSIVE_RE = re.compile(r"[sS]['’]s\b")


# ===========================================================================
# 1. deck plurals
# ===========================================================================
class TestCountNounHelper:
    @pytest.mark.parametrize(
        "n, singular, expected",
        [
            (1, "hire", "1 hire"),
            (1.0, "hire", "1 hire"),
            ("1", "hire", "1 hire"),
            (0, "hire", "0 hires"),
            (2, "hire", "2 hires"),
            (1250, "hire", "1,250 hires"),
            (None, "hire", "0 hires"),
            (1, "channel", "1 channel"),
        ],
    )
    def test_rule(self, n, singular, expected):
        assert ppt._count_noun(n, singular) == expected

    def test_matches_the_workbook_helper_when_it_exists(self):
        """Deck and workbook must pluralise identically."""
        import gold_standard

        shared = getattr(gold_standard, "count_noun", None)
        if shared is None:
            pytest.skip("gold_standard.count_noun is not on this base yet")
        for n in (0, 1, 2, 1.0, "1", 1250, None, 0.5):
            assert ppt._count_noun(n, "hire") == shared(n, "hire")
        assert ppt._count_noun is shared


class TestDeckPlurals:
    def test_tiny_nonprofit_deck_has_no_singular_plural_or_hedge(self):
        text = _tiny_nonprofit_deck()
        assert _BAD_SINGULAR_RE.findall(text) == [], [
            m.group(0) for m in _BAD_SINGULAR_RE.finditer(text)
        ]

    def test_goal_lines_read_one_hire(self):
        text = _tiny_nonprofit_deck()
        # slide 2 executive-summary line and slide 8 goal band
        assert re.search(r"Client goal: 1 hire\b(?!s)", text)
        assert re.search(r"\b1 hire target\b", text)

    def test_one_channel_plan_reads_one_channel(self):
        text = _tiny_nonprofit_deck(n_channels=1)
        assert not re.search(r"\b1 channels\b", text)

    def test_plural_goal_is_unchanged(self):
        """A real multi-hire goal keeps its plural and thousands separator."""
        data = matrix._base_plan(hire_volume="1,500 hires")
        data = matrix._with_budget_allocation(data, n_channels=4)
        text = _all_text(Presentation(io.BytesIO(ppt.generate_pptx(data))))
        assert not _BAD_SINGULAR_RE.search(text)
        assert re.search(r"1,500 hires (target|goal)|Client goal: 1,500 hires", text)


# ===========================================================================
# 2. the last double-possessive sites
# ===========================================================================
NAMES = [
    ("Brightway Home Stores", "Brightway Home Stores'"),
    ("McDonald's", "McDonald's"),
    ("Barclays", "Barclays'"),
    ("Joveo", "Joveo's"),
]


class TestInsightComposerPossessive:
    """Ordinals 2 and 6 of the default skeleton bank carry the possessive
    ("<name> presence in this pool", "<name> reputation with")."""

    @pytest.mark.parametrize("name, poss", NAMES)
    def test_possessive_skeletons(self, name, poss):
        for ordinal, tail in ((2, "presence in this pool"), (6, "reputation with")):
            sentence = insight_composer.compose_counter_strategy(
                name, {"role": "Nurse", "city": "Dallas", "ordinal": ordinal}
            )
            assert f"{poss} {tail}" in sentence, sentence
            assert not _DOUBLE_POSSESSIVE_RE.search(sentence), sentence

    @pytest.mark.parametrize("name, _poss", NAMES)
    def test_no_skeleton_in_any_bank_doubles_the_possessive(self, name, _poss):
        for ctype in ("", "staffing_agency", "direct_employer", "gig_platform"):
            for ordinal in range(12):
                sentence = insight_composer.compose_counter_strategy(
                    name,
                    {
                        "role": "Nurse",
                        "city": "Dallas",
                        "ordinal": ordinal,
                        "competitor_type": ctype,
                    },
                )
                assert not _DOUBLE_POSSESSIVE_RE.search(sentence), sentence


class TestResearchPossessive:
    def test_unknown_competitor_intelligence(self):
        rows = research.get_client_competitor_intelligence(
            ["Brightway Home Stores", "Joe's"], "retail_consumer"
        )
        by_name = {r["competitor"]: r for r in rows}
        stores = by_name["Brightway Home Stores"]
        assert "Monitor Brightway Home Stores' Glassdoor profile" in (
            stores["employer_brand_strength"]
        )
        assert "Audit Brightway Home Stores' career site" in (
            stores["strategic_recommendation"]
        )
        joes = by_name["Joe's"]
        assert "Monitor Joe's Glassdoor profile" in joes["employer_brand_strength"]
        assert "Audit Joe's career site" in joes["strategic_recommendation"]
        for row in rows:
            for value in row.values():
                assert not _DOUBLE_POSSESSIVE_RE.search(str(value)), value

    @pytest.mark.parametrize(
        "name, poss",
        [("Stores", "Stores'"), ("McDonald's", "McDonald's"), ("Acme", "Acme's")],
    )
    def test_known_employer_glassdoor_recommendation(self, name, poss):
        out = research._generate_competitive_recommendation(
            name, {"glassdoor_rating": "3.2/5"}, "retail_consumer"
        )
        text = " ".join(out) if isinstance(out, (list, tuple)) else str(out)
        assert f"{poss} Glassdoor rating (3.2/5)" in text, text
        assert not _DOUBLE_POSSESSIVE_RE.search(text), text


class TestPagespeedPossessive:
    @pytest.mark.parametrize(
        "url, lead",
        [
            ("https://www.stores.com/careers", "Stores' career page"),
            ("https://careers.mcdonalds.com", "Mcdonalds' career page"),
            ("https://www.joveo.com/jobs", "Joveo's career page"),
        ],
    )
    def test_blocked_audit_note(self, url, lead):
        note = google_search_pagespeed._benchmark_based_audit(url, "mobile")["note"]
        assert note.startswith(lead), note
        assert not _DOUBLE_POSSESSIVE_RE.search(note), note


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
