"""The workbook's Executive Summary and the deck show applications, cost per
application and cost per hire as low-high (the hires range was already there),
and state that hires / cost per hire are estimated with no client hire history.
"""

from __future__ import annotations

import io
import logging
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
logging.disable(logging.CRITICAL)

import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

DASH = "–"


def _brief(role="Registered Nurse", industry="Healthcare & Medical", loc="Columbus, OH"):
    return {
        "client_name": "Probe Health",
        "industry": industry,
        "budget": "$60,000",
        "campaign_duration": "3 months",
        "hire_volume": "Not specified",
        "work_environment": "onsite",
        "locations": [loc],
        "roles": [role],
        "target_roles": [role],
    }


@pytest.fixture(scope="module")
def rn():
    return T.build_plan_data(_brief())


def _summary_strings(data):
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    if isinstance(raw, tuple):
        raw = raw[0]
    ws = load_workbook(io.BytesIO(raw))["Executive Summary"]
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def _deck_strings(data):
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append((i, shape.text_frame.text.replace("\n", " | ")))
    return out


def test_workbook_summary_prints_applications_cpa_and_cph_ranges(rn):
    cells = _summary_strings(rn)
    line = next((c for c in cells if "cost per application:" in c), "")
    assert line, "no projection range line on the Executive Summary"
    assert f"Applications: " in line and DASH in line.split("cost per application:")[0]
    cpa = line.split("cost per application:")[1].split(". ")[0]
    assert DASH in cpa
    assert "Cost per hire: " in line and "not media-only" in line
    assert "no client hire history" in line
    assert "Industry benchmarks run higher" in line  # RN: engine CPA < KB $35


def test_workbook_still_prints_the_hires_range_first(rn):
    cells = _summary_strings(rn)
    assert any(c.startswith("Projected hires: ") and DASH in c for c in cells)


def test_deck_slide_two_headline_shows_the_hires_and_cpa_ranges(rn):
    lines = [t for s, t in _deck_strings(rn) if s == 2]
    joined = " || ".join(lines)
    # the KPI strip values are ranges, not the optimistic point
    import re

    assert re.search(r"\d[\d,]*" + DASH + r"\d[\d,]*", joined), joined
    assert re.search(r"\$[\d.,]+" + DASH + r"\$[\d.,]+", joined), joined


def test_deck_states_the_hires_basis(rn):
    text = " || ".join(t for _, t in _deck_strings(rn))
    assert "no client hire history" in text
