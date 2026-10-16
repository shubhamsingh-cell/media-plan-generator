"""Residuals from the 2026-10-16 26-bundle sweep (origin/main 4f29917)
inside excel_v2 / ppt_generator:

B. Sources & Confidence D21:F25 printed the raw client name ("THE HERSHEY
   COMPANY") from data_synthesizer's location warnings while every other
   sheet printed "The Hershey Company" (10 bundle_qa criticals).
C. One deck priced the same role twice: slide 2's "Salary Range" read the
   plan-level salary_intelligence base ($150K), slide 7's Role Breakdown the
   city-adjusted gold-standard rows ($186K).
D. Every non-US plan carried 9-14 US-labelled stat cells in Market
   Intelligence ("US Remote Rate: 22.8% of US employees", ...).
E. The executive-narrative LLM calls capped output at 600 tokens; adaptive
   thinking spent the whole cap and the text came back empty.
"""

from __future__ import annotations

import ast
import copy
import io
import json
import os
import re
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import data_synthesizer  # noqa: E402
import excel_v2  # noqa: E402
import gold_standard  # noqa: E402
import ppt_generator as ppt  # noqa: E402
from bundle_qa import run_bundle_qa  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402


@pytest.fixture(autouse=True)
def _reset_active_currency():
    yield
    for mod in (ppt, excel_v2):
        try:
            mod._currency_tls.code = "USD"
        except AttributeError:
            pass


def _cells(ws):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                yield cell


# ---------------------------------------------------------------------------
# B. client name casing in the location-plausibility table
# ---------------------------------------------------------------------------
def test_location_warnings_print_the_display_client_name():
    raw_name = "THE HERSHEY COMPANY"
    data = {
        "client_name": raw_name,
        "industry": "food_beverage",
        "budget": "$90,000",
        "campaign_duration": "6 months",
        "locations": ["Memphis, TN"],
        "roles": ["Machine Operator"],
        "_synthesized": {
            "_validation": {
                "location_warnings": [
                    {
                        "location": "Memphis, TN",
                        "severity": "low",
                        "reason": (
                            "Memphis, TN is in a neighboring region (US Southeast) "
                            f"but not in {raw_name}'s known operating area."
                        ),
                        "known_states_display": "Pennsylvania",
                        "suggestion": (
                            f"Verify that {raw_name} operates in Memphis, TN. "
                            "Known locations: Pennsylvania."
                        ),
                    }
                ]
            }
        },
    }
    raw = excel_v2.generate_excel_v2(copy.deepcopy(data))
    ws = load_workbook(io.BytesIO(raw))["Sources & Confidence"]
    texts = [c.value for c in _cells(ws)]
    assert not [t for t in texts if raw_name in t], texts
    assert any("not in The Hershey Company's known operating area" in t for t in texts)
    assert any("Verify that The Hershey Company operates in Memphis" in t for t in texts)
    casing = [
        f for f in run_bundle_qa(None, raw, data)
        if f.get("code") == "client_name_wrong_casing"
    ]
    assert not casing, casing


# ---------------------------------------------------------------------------
# C. one salary per role per deck
# ---------------------------------------------------------------------------
_ROLES = ["Software Engineer", "Data Engineer", "Site Reliability Engineer", "Cook", "Registered Nurse"]


def _salary_band(median):
    return {
        "min": round(median * 0.65), "p25": round(median * 0.85), "median": median,
        "p75": round(median * 1.2), "max": round(median * 1.45),
        "sources": ["Industry Benchmark"], "confidence": 0.6, "currency": "USD",
    }


def test_situation_card_and_role_breakdown_show_one_median_per_role():
    intel = {"Software Engineer": _salary_band(150_000)}
    per_role = {
        "Software Engineer": dict(
            _salary_band(150_000), source="Industry Benchmark",
            confidence="benchmark", city_adjust=True,
        )
    }
    data = {
        "client_name": "ADP",
        "industry": "hospitality",
        "budget": "$1,200,000",
        "budget_period": "campaign",
        "campaign_duration": "12 months",
        "campaign_start_month": 1,
        "hire_volume": "240 hires",
        "locations": ["San Francisco, CA", "New York, NY"],
        "roles": list(_ROLES),
        "target_roles": list(_ROLES),
        "campaign_goals": ["volume_hiring"],
        "_synthesized": {"salary_intelligence": intel, "per_role_salaries": per_role},
    }
    gold_standard.apply_all_quality_gates(data)
    gold_median, _est = ppt._role_breakdown_median_salary(
        data["_gold_standard"], "Software Engineer"
    )
    assert gold_median and abs(gold_median - 150_000) > 5_000  # city-adjusted

    prs = Presentation(io.BytesIO(ppt.generate_pptx(data)))
    texts = [
        sh.text_frame.text
        for slide in prs.slides
        for sh in slide.shapes
        if sh.has_text_frame
    ]
    situation = [t for t in texts if "median" in t and "Software Engineer" in t]
    assert situation, "slide 2 Salary Range line missing"
    m = re.search(r"\$(\d+)K median", situation[0])
    assert m, situation[0]
    breakdown = None
    for slide in prs.slides:
        st = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        if "Channel Emphasis" in st and "Software Engineer" in st:
            breakdown = st[st.index("Software Engineer") + 2]
    assert breakdown, "slide 7 Role Breakdown missing"
    m7 = re.search(r"\$(\d+)K", breakdown)
    assert m7, breakdown
    assert m.group(1) == m7.group(1), (situation[0], breakdown)
    assert m.group(1) == str(round(gold_median / 1000)), (situation[0], gold_median)


def test_role_salary_rows_never_average_two_currencies():
    gold = {
        "city_level_data": {
            "Dallas": {"per_role_salary": {"Engineer": {"median": 120_000, "currency": "USD"}}},
            "London": {"per_role_salary": {"Engineer": {"median": 60_000, "currency": "GBP"}}},
        }
    }
    median, _est = ppt._role_breakdown_median_salary(gold, "Engineer")
    assert median == 120_000
    assert ppt._role_breakdown_salary_basis(gold, "Engineer") == (False, True)


# ---------------------------------------------------------------------------
# D. no US-labelled statistics on a plan with no US market
# ---------------------------------------------------------------------------
_US_LABEL = re.compile(r"(?<![\w$])US(?![\w$-])|U\.S\.|\bAmericans?\b|\bUnited States\b")
# The workbook's own notes explaining that US data was withheld/converted.
_US_EXPLAIN = re.compile(
    r"US-domiciled|US salary figures|US figures are not used|US industry average|"
    r"US\$ inputs|US-calibrated|US-only|not used for non-US|non-US|US data|US\$1 =",
    re.IGNORECASE,
)


def _workforce_plan(locations, currency, budget):
    with open(PROJECT_ROOT / "data" / "workforce_trends_intelligence.json", encoding="utf-8") as f:
        kb = {"workforce_trends": json.load(f)}
    return {
        "client_name": "Infosys Test",
        "industry": "tech_engineering",
        "budget": budget,
        "campaign_duration": "3 months",
        "campaign_start_month": 9,
        "roles": ["Software Engineer"],
        "target_roles": ["Software Engineer"],
        "locations": locations,
        "plan_currency": currency,
        "_synthesized": {
            "workforce_insights": data_synthesizer.fuse_workforce_insights(
                {}, kb, "tech_engineering"
            )
        },
    }


def _market_intelligence_text(data):
    raw = excel_v2.generate_excel_v2(copy.deepcopy(data))
    ws = load_workbook(io.BytesIO(raw))["Market Intelligence"]
    return [(c.coordinate, c.value) for c in _cells(ws)]


def test_non_us_plan_has_no_us_labelled_stat_rows():
    cells = _market_intelligence_text(
        _workforce_plan(["Bengaluru, Karnataka, India", "Pune, India"], "INR", "₹25,000,000")
    )
    leaked = [
        (coord, v[:160]) for coord, v in cells
        if _US_LABEL.search(v) and not _US_EXPLAIN.search(v)
    ]
    assert not leaked, leaked
    # Global lines of the same section still print.
    assert any("Hiring Growth: 88% year-on-year growth" in v for _c, v in cells)


def test_mixed_plan_keeps_us_labelled_stats_for_its_us_market():
    cells = _market_intelligence_text(
        _workforce_plan(["Austin, TX", "Bengaluru, Karnataka, India"], "USD", "$900,000")
    )
    assert any("US Remote Rate" in v for _c, v in cells)


# ---------------------------------------------------------------------------
# E. narrative output cap leaves room for adaptive thinking
# ---------------------------------------------------------------------------
def test_every_plan_narrative_call_allows_room_for_thinking():
    tree = ast.parse((PROJECT_ROOT / "excel_v2.py").read_text(encoding="utf-8"))
    caps = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        kw = {k.arg: k.value for k in node.keywords if k.arg}
        task = kw.get("task_type")
        if not (isinstance(task, ast.Name) and task.id == "TASK_PLAN_NARRATIVE"):
            continue
        cap = kw.get("max_tokens")
        assert cap is not None, f"line {node.lineno}: no max_tokens"
        if isinstance(cap, ast.Constant):
            value = cap.value
        elif isinstance(cap, ast.Name):
            value = getattr(excel_v2, cap.id)
        else:
            pytest.fail(f"line {node.lineno}: unresolvable max_tokens {ast.dump(cap)}")
        caps.append((node.lineno, value))
    assert len(caps) >= 3, caps
    assert all(v >= 2048 for _l, v in caps), caps
