"""Three truthfulness defects in the generated workbook/deck (2026-10-16,
reproduced on origin/main 4f29917 through the real generators):

1. Sources & Confidence said uploaded client history "takes highest
   precedence", but ``historical_performance`` is read by no live generator.
   Now: history is stored and shown for reference; projections use market
   benchmarks; a "Client history (as provided)" note only when it exists.
2. A plan whose user unticked Niche Boards still shipped the "Niche Board
   Matching" sheet, the Channels sheet's niche section, and a role table
   (workbook + deck) steering roles to "Niche / Industry Boards".
3. The activation calendar printed US holidays (Memorial Day, Thanksgiving,
   "July 4") on plans with no US market.
"""

from __future__ import annotations

import copy
import io
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import excel_v2  # noqa: E402
import gold_standard  # noqa: E402
import ppt_generator as ppt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

_ROLES = [
    "Registered Nurse",
    "Software Engineer",
    "Diesel Mechanic",
    "Data Scientist",
    "Cook",
]
_ALL_ON = {
    "programmatic_dsp": True,
    "global_boards": True,
    "niche_boards": True,
    "social_media": True,
    "regional_boards": True,
    "employer_branding": False,
}
_HISTORY = {
    "platforms": {
        "Indeed": {"spend": 5000.0, "applications": 400, "hires": 8},
    },
    "total_spend": 5000.0,
    "total_hires": 8,
    "total_applications": 400,
    "total_cpa": 12.5,
    "total_cph": 625.0,
    "platform_count": 1,
    "source_files": ["history.csv"],
}
_US_HOLIDAY = re.compile(
    r"memorial day|thanksgiving|july 4|fourth of july|independence day|"
    r"labor day|spring break",
    re.IGNORECASE,
)


@pytest.fixture(autouse=True)
def _reset_active_currency():
    """_set_active_currency stores the plan currency in a thread-local that
    outlives a single test in the same thread."""
    yield
    for mod in (ppt, excel_v2):
        try:
            mod._currency_tls.code = "USD"
        except AttributeError:
            pass


def _plan(**over):
    data = {
        "client_name": "Acme Hotels",
        "industry": "hospitality",
        "budget": "$120,000",
        "budget_period": "campaign",
        "campaign_duration": "12 months",
        "campaign_start_month": 1,
        "hire_volume": "60 hires",
        "work_environment": "onsite",
        "locations": ["Dallas, TX"],
        "roles": list(_ROLES),
        "target_roles": list(_ROLES),
        "campaign_goals": ["volume_hiring"],
        "channel_categories": dict(_ALL_ON),
    }
    data.update(over)
    return data


def _no_niche():
    return dict(_ALL_ON, niche_boards=False)


def _gated(data):
    data = copy.deepcopy(data)
    gold_standard.apply_all_quality_gates(data)
    return data


def _workbook(data):
    data = _gated(data)
    return data, load_workbook(io.BytesIO(excel_v2.generate_excel_v2(data)))


def _cells(ws):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                yield cell


def _all_text(wb):
    return [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for c in _cells(ws)]


# ---------------------------------------------------------------------------
# Defect 1: no "client history takes precedence" claim
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def history_wb():
    return _workbook(_plan(historical_performance=copy.deepcopy(_HISTORY)))[1]


def test_methodology_never_claims_client_history_drives_projections(history_wb):
    text = _all_text(history_wb)
    claims = [t for t in text if re.search(r"precedence", t[2], re.IGNORECASE)]
    assert not claims, f"workbook still ranks client history first: {claims}"

    ws = history_wb["Sources & Confidence"]
    rows = {
        c.value: ws.cell(row=c.row, column=c.column + 2).value
        for c in _cells(ws)
        if c.column == excel_v2.COL_START
    }
    method = rows.get("Client History") or ""
    assert "reference only" in method and "not adjusted" in method, method

    note = rows.get("Client history (as provided)") or ""
    assert "400 applications" in note and "8 hires" in note, note
    assert "history.csv" in note and "Indeed" in note, note
    assert "not used to adjust" in note, note
    # The upload declares no currency: no symbol may be assumed for spend.
    assert not re.search(r"[$£€₹¥]", note), note


def test_no_client_history_note_without_uploaded_history():
    _data, wb = _workbook(_plan(locations=["London, UK"]))
    labels = [t for t in _all_text(wb) if t[2] == "Client history (as provided)"]
    assert not labels, labels


# ---------------------------------------------------------------------------
# Defect 2: an unticked channel category is never recommended
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def no_niche_wb():
    return _workbook(_plan(channel_categories=_no_niche()))


def test_unticked_niche_boards_ship_no_niche_board_surfaces(no_niche_wb):
    data, wb = no_niche_wb
    assert "Niche Board Matching" not in wb.sheetnames, wb.sheetnames

    channels = [c.value for c in _cells(wb["Channels & Strategy"])]
    assert not [v for v in channels if v.startswith("Niche Channels")], channels
    assert "Industry Niche Board" not in channels

    emphases = [
        r["channel_emphasis"] for r in data["_gold_standard"]["difficulty_framework"]
    ]
    assert "niche_boards" not in emphases, emphases
    qi = [c.value for c in _cells(wb["Quality Intelligence"])]
    assert "Niche Boards" not in qi


def test_selected_or_untouched_niche_boards_keep_the_sheet():
    _d, wb_on = _workbook(_plan())
    assert "Niche Board Matching" in wb_on.sheetnames
    # Untouched wizard toggles omit channel_categories: the recommended mix.
    _d, wb_untouched = _workbook(_plan(channel_categories=None))
    assert "Niche Board Matching" in wb_untouched.sheetnames


def test_deck_role_table_names_only_selected_channels():
    data = _gated(_plan(channel_categories=_no_niche()))
    prs = Presentation(io.BytesIO(ppt.generate_pptx(data)))
    role_slides = []
    for slide in prs.slides:
        texts = [
            sh.text_frame.text
            for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        if "Channel Emphasis" in texts:
            role_slides.append(texts)
    assert role_slides, "role-breakdown slide not rendered for a 5-role plan"
    for texts in role_slides:
        assert not [t for t in texts if "Niche" in t], texts
        assert "Job Boards" in texts  # the hard roles' next-best selected channel


def test_senior_profile_description_drops_niche_channels_when_unticked():
    rows = gold_standard.classify_difficulty(
        {"roles": ["Senior Widget Associate"], "channel_categories": _no_niche()}
    )
    assert rows and rows[0]["classification_source"] == "seniority_keyword_match"
    assert rows[0]["channel_emphasis"] != "niche_boards"
    assert "niche" not in rows[0]["description"].lower(), rows[0]["description"]


# ---------------------------------------------------------------------------
# Defect 3: US holidays only when the plan has a US market
# ---------------------------------------------------------------------------
def _calendar_text(data):
    cal = data["_gold_standard"]["activation_calendar"]
    events = [e for m in cal["timeline"] for e in m["key_events"]]
    return events + list(cal.get("industry_events") or [])


def test_non_us_plans_print_no_us_holidays():
    data, wb = _workbook(_plan(locations=["London, UK"]))
    leaked = [e for e in _calendar_text(data) if _US_HOLIDAY.search(e)]
    assert not leaked, f"US holidays on a London plan: {leaked}"
    cells = [t for t in _all_text(wb) if _US_HOLIDAY.search(t[2])]
    assert not cells, cells

    plant = _gated(_plan(industry="manufacturing", locations=["Munich, Germany"]))
    july = next(
        m for m in plant["_gold_standard"]["activation_calendar"]["timeline"]
        if m["month"] == 7
    )
    # The shutdown claim survives; only the US-holiday aside goes.
    assert "Summer plant shutdown / maintenance week" in july["key_events"]
    assert not [e for e in july["key_events"] if _US_HOLIDAY.search(e)]

    generic = _gated(_plan(industry="general_entry_level", locations=["Toronto, ON, Canada"]))
    assert not [e for e in _calendar_text(generic) if _US_HOLIDAY.search(e)]

    # Canadian Thanksgiving is in October: US Thanksgiving in November must
    # not appear on a Canada-only plan (2026-10-16 sweep, x_canada_cad F73).
    canada = _gated(_plan(locations=["Toronto, ON, Canada", "Vancouver, BC, Canada"]))
    assert not [e for e in _calendar_text(canada) if _US_HOLIDAY.search(e)]


def test_us_plan_keeps_us_holidays_unlabelled():
    events = _calendar_text(_gated(_plan(locations=["Dallas, TX", "Remote"])))
    assert "Memorial Day staffing" in events
    assert "Thanksgiving staffing surge" in events


def test_mixed_plan_attributes_us_holidays_to_us_markets():
    data = _gated(_plan(locations=["Dallas, TX", "Toronto, ON, Canada"]))
    events = _calendar_text(data)
    assert "Memorial Day staffing (US markets)" in events
    assert "Thanksgiving staffing surge (US markets)" in events
    assert "Spring break peak hiring (Mar, US markets)" in events
    unlabelled = [e for e in events if _US_HOLIDAY.search(e) and "US markets" not in e]
    assert not unlabelled, unlabelled

    plant = _gated(
        _plan(industry="manufacturing", locations=["Dallas, TX", "Munich, Germany"])
    )
    july = next(
        m for m in plant["_gold_standard"]["activation_calendar"]["timeline"]
        if m["month"] == 7
    )
    assert (
        "Summer plant shutdown / maintenance week (often around July 4 in US markets)"
        in july["key_events"]
    ), july["key_events"]
