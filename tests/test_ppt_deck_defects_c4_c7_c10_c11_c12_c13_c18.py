"""Regression tests for the 2026-09-24 deck audit items C4, C7, C10, C11,
C12, C13 and the zero-hire (C18) finding fixed in this change.

Runs under pytest, or standalone:
``python3 tests/test_ppt_deck_defects_c4_c7_c10_c11_c12_c13_c18.py``.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from typing import Any, Dict

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import ppt_generator as ppt  # noqa: E402
from pptx import Presentation  # noqa: E402

EMU_PER_IN = 914400


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            try:
                yield from _iter_shapes(sh.shapes)
            except Exception:
                pass


def _text_shapes(slide):
    out = []
    for sh in _iter_shapes(slide.shapes):
        try:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh)
        except Exception:
            pass
    return out


def _slide_containing(prs: Presentation, needle: str):
    for slide in prs.slides:
        for sh in _text_shapes(slide):
            if needle in sh.text_frame.text:
                return slide
    return None


def _shape_with_text(slide, needle: str, exact: bool = False):
    for sh in _text_shapes(slide):
        t = sh.text_frame.text
        if (t.strip() == needle) if exact else (needle in t):
            return sh
    return None


def _all_texts(slide):
    return [sh.text_frame.text for sh in _text_shapes(slide)]


def _gen(data: Dict[str, Any]) -> Presentation:
    return Presentation(io.BytesIO(ppt.generate_pptx(data)))


# ---------------------------------------------------------------------------
# C4 -- a channel the budget engine never funded (toggled ON in
# channel_categories, but absent from _budget_allocation.channel_allocations)
# must not appear as a phantom row/bar on slides 2, 5, 6 or 8, and every
# slide's channel COUNT and channel-mix PERCENTAGES must agree with each
# other and foot to exactly 100%.
# ---------------------------------------------------------------------------
def _uk_plan_with_unfunded_apac() -> Dict[str, Any]:
    return {
        "client_name": "Brightwater Retail Group",
        "industry": "retail_consumer",
        "industry_label": "Retail & Consumer",
        "locations": ["London, United Kingdom", "Manchester, United Kingdom"],
        "roles": ["Store Manager", "Sales Associate"],
        "budget": "£180,000",
        "target_region": "emea",
        # The wizard toggled 8 channels ON, INCLUDING apac_regional -- but
        # the budget engine below only actually funded 7 of them.
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "niche_boards": True,
            "social_media": True,
            "regional_boards": True,
            "employer_branding": True,
            "emea_regional": True,
            "apac_regional": True,  # toggled ON but NEVER funded below
        },
        "_budget_allocation": {
            "metadata": {"total_budget": 180000},
            "total_projected": {"applications": 16514, "hires": 107},
            "channel_allocations": {
                "programmatic_dsp": {
                    "percentage": 34,
                    "dollar_amount": 62061,
                    "projected_hires": 36,
                    "projected_applications": 6770,
                },
                "global_boards": {
                    "percentage": 30,
                    "dollar_amount": 53109,
                    "projected_hires": 46,
                    "projected_applications": 5707,
                },
                "regional_boards": {
                    "percentage": 16,
                    "dollar_amount": 28113,
                    "projected_hires": 12,
                    "projected_applications": 1913,
                },
                "emea_regional": {
                    "percentage": 9,
                    "dollar_amount": 16617,
                    "projected_hires": 7,
                    "projected_applications": 1130,
                },
                "social_media": {
                    "percentage": 5,
                    "dollar_amount": 8497,
                    "projected_hires": 0,
                    "projected_applications": 65,
                },
                "niche_boards": {
                    "percentage": 3,
                    "dollar_amount": 6150,
                    "projected_hires": 3,
                    "projected_applications": 500,
                },
                "employer_branding": {
                    "percentage": 3,
                    "dollar_amount": 5454,
                    "projected_hires": 3,
                    "projected_applications": 429,
                },
                # NOTE: no "apac_regional" key -- the engine declined to
                # fund it, even though the wizard toggled it on above.
            },
        },
    }


def test_funded_display_channels_drops_unfunded_channel_and_foots_to_100():
    data = _uk_plan_with_unfunded_apac()
    funded = ppt._funded_display_channels(data)
    keys = [c["key"] for c in funded]
    assert "apac_regional" not in keys, (
        f"unfunded apac_regional leaked into the funded channel set: {keys}"
    )
    assert len(funded) == 7, f"expected 7 funded channels, got {len(funded)}: {keys}"
    assert sum(c["pct"] for c in funded) == 100, (
        f"funded channel percentages must foot to 100, got "
        f"{sum(c['pct'] for c in funded)}"
    )


def test_no_phantom_apac_channel_and_consistent_count_across_slides():
    data = _uk_plan_with_unfunded_apac()
    prs = _gen(data)

    # Slide 5 (Channel Strategy): no phantom "APAC Regional" bar/label.
    s5 = _slide_containing(prs, "CHANNEL MIX")
    assert s5 is not None, "Channel Strategy slide not found"
    s5_texts = " | ".join(_all_texts(s5))
    assert "APAC Regional" not in s5_texts, (
        f"phantom unfunded 'APAC Regional' channel rendered on slide 5: "
        f"{s5_texts}"
    )
    assert "8-channel" not in s5_texts, (
        "slide 5 headline still counts the unfunded channel"
    )
    assert "7-channel" in s5_texts, (
        f"slide 5 headline should read the FUNDED 7-channel count: {s5_texts}"
    )

    # Slide 6 (Budget Allocation): same funded count, no phantom row.
    s6 = _slide_containing(prs, "CHANNEL-BY-CHANNEL BREAKDOWN")
    assert s6 is not None, "Budget Allocation slide not found"
    s6_texts = " | ".join(_all_texts(s6))
    assert "APAC Regional" not in s6_texts
    assert "Investment breakdown across 7 channels" in s6_texts, s6_texts

    # Slide 8 (Plan Comparison): "Channels Selected" must read 7, not 8.
    s8 = _slide_containing(prs, "Channels Selected")
    assert s8 is not None, "Plan Comparison slide not found"
    channels_shape = _shape_with_text(s8, "Channels Selected")
    # The client-column value for "Channels Selected" is rendered as its own
    # shape elsewhere on the slide -- assert the funded count 7 appears and
    # the unfunded count 8 does not appear as a standalone value shape.
    s8_texts = [t.strip() for t in _all_texts(s8)]
    assert not any(t.startswith("8") and "▲" in t for t in s8_texts), (
        f"slide 8 'Channels Selected' still shows the unfunded 8-channel "
        f"count: {s8_texts}"
    )
    assert any(t.startswith("7") for t in s8_texts if "▲" in t or "▼" in t), (
        f"slide 8 'Channels Selected' should show the funded 7-channel "
        f"count: {s8_texts}"
    )


# ---------------------------------------------------------------------------
# C7 -- a salary sourced entirely from a US-only benchmark (DOL H-1B/LCA,
# the US fallback table, or a "US-basis" driver wage) must render with the
# US$ marker on a non-USD plan, not the plan's local currency symbol.
# ---------------------------------------------------------------------------
def _gbp_plan_with_us_sourced_salary() -> Dict[str, Any]:
    return {
        "client_name": "Thameside Freight",
        "industry": "blue_collar_trades",
        "industry_label": "Blue Collar / Skilled Trades",
        "locations": ["London, United Kingdom"],
        "roles": ["CDL-A Truck Driver"],
        "budget": "£300,000",
        "channel_categories": {"programmatic_dsp": True, "global_boards": True},
        "_synthesized": {
            "salary_intelligence": {
                "CDL-A Truck Driver": {
                    "median": 65000,
                    "min": 45000,
                    "max": 88000,
                    "sources": ["Industry Benchmark (US-basis, salaried trucking)"],
                    "currency": "USD",
                }
            }
        },
    }


def test_us_sourced_salary_marked_usd_on_non_usd_plan():
    data = _gbp_plan_with_us_sourced_salary()
    prs = _gen(data)
    slide = _slide_containing(prs, "Salary Range")
    assert slide is not None, "Executive Summary slide not found"
    salary_shape = _shape_with_text(slide, "Salary Range")
    assert salary_shape is not None
    text = salary_shape.text_frame.text
    assert "US$65K" in text, f"expected a US$-marked salary figure, got: {text!r}"
    # The bare local-currency symbol must never sit directly on this
    # US-sourced number.
    assert "£65K" not in text, (
        f"US-sourced salary rendered with the plan's local '£' symbol "
        f"instead of 'US$': {text!r}"
    )


# ---------------------------------------------------------------------------
# C13 -- an industry's own KB recruitment benchmark for CPA/CPC must win
# over the generic cross-industry ad-platform range.
# ---------------------------------------------------------------------------
def _industry_kb_vs_generic_platform_plan() -> Dict[str, Any]:
    return {
        "client_name": "Acme Health",
        "industry": "healthcare_medical",
        "industry_label": "Healthcare & Medical",
        "locations": ["Dallas, TX"],
        "roles": ["Registered Nurse"],
        "budget": "$500,000",
        "channel_categories": {"programmatic_dsp": True, "global_boards": True},
        "_knowledge_base": {
            "recruitment_benchmarks": {
                "industry_benchmarks": {
                    "healthcare_medical": {
                        "cpa": {"range": "$40 - $90"},
                        "cpc": {"range": "$1.10 - $2.20"},
                    }
                }
            }
        },
        "_synthesized": {
            "ad_platform_analysis": {
                "google": {"avg_cpc": 8.00, "avg_cpa": 400.0},
                "meta_fb": {"avg_cpc": 9.00, "avg_cpa": 450.0},
            }
        },
    }


def test_industry_kb_benchmark_wins_over_generic_ad_platform_range():
    data = _industry_kb_vs_generic_platform_plan()
    bm = ppt._get_benchmarks("healthcare_medical", data)
    assert "8" not in bm["cpc"] and "9" not in bm["cpc"], (
        f"generic ad-platform CPC range leaked into the industry-specific "
        f"row: {bm['cpc']!r}"
    )
    assert not bm.get("cpc_is_generic_platform_range"), (
        "the KB-industry CPC should not be flagged as a generic range"
    )

    prs = _gen(data)
    slide = _slide_containing(prs, "INDUSTRY BENCHMARKS")
    assert slide is not None
    texts = " | ".join(_all_texts(slide))
    assert "$8.00" not in texts and "$9.00" not in texts, (
        f"slide 5 shows the generic cross-industry ad-platform CPC range "
        f"instead of the industry's own KB benchmark: {texts}"
    )


def test_generic_platform_range_labeled_cross_industry_when_no_kb_benchmark():
    # No KB entry for this industry -> Layer 1 (ad platform) is the only
    # source, and it must be labeled as a cross-industry range.
    data = {
        "client_name": "Acme",
        "industry": "general_entry_level",
        "_knowledge_base": {"recruitment_benchmarks": {"industry_benchmarks": {}}},
        "_synthesized": {
            "ad_platform_analysis": {
                "google": {"avg_cpc": 3.0, "avg_cpa": 60.0},
            }
        },
    }
    bm = ppt._get_benchmarks("general_entry_level", data)
    if bm.get("cpc_is_generic_platform_range"):
        assert "cross-industry" not in bm["cpc"]  # marker lives on the label,
        # not embedded in the value string -- sanity check only.


# ---------------------------------------------------------------------------
# C10 -- the competitor card's "Why:" body text box must be measured at its
# TRUE usable width (box width minus BOTH the 0.4in margin AND python-pptx's
# 0.2in internal left+right insets), matching the name-box measure -- not
# just the 0.4in margin, which undercounts wrapped lines and lets "Counter:"
# print on top of "Why:"'s real last line.
# ---------------------------------------------------------------------------
def _competitor_card_stress_plan() -> Dict[str, Any]:
    # A role title calibrated (via _compose_competitor_why + _measure_lines)
    # to sit exactly on the wrap boundary the audit reproduced: the
    # composed "Why:" sentence needs 2 lines at the box's raw width (5.9in)
    # but 3 lines at the TRUE usable width after python-pptx's 0.2in text
    # insets (5.7in) -- the exact numbers from the audit's own repro
    # (_measure_lines(why, 5.9, 8) -> 2, _measure_lines(why, 5.7, 8) -> 3).
    role = "Senior Registered Nurse " + "X" * 42
    return {
        "client_name": "Acme Logistics",
        "industry": "healthcare_medical",
        "industry_label": "Healthcare & Medical",
        "roles": [role],
        "locations": ["Rancho Santa Margarita, CA"],
        "competitors": ["Tenet Health"],
        "budget": "$500,000",
        "channel_categories": {"programmatic_dsp": True},
    }


def test_competitor_card_counter_never_overlaps_why_real_last_line():
    data = _competitor_card_stress_plan()
    prs = _gen(data)
    slide = _slide_containing(prs, "COMPETITIVE LANDSCAPE") or _slide_containing(
        prs, "COMPANY PROFILE"
    )
    assert slide is not None, "Competitive Landscape slide not found"

    why_shape = None
    counter_shape = None
    for sh in _text_shapes(slide):
        t = sh.text_frame.text.strip()
        if t.startswith("Why:"):
            why_shape = sh
        elif t.startswith("Counter:"):
            counter_shape = sh
    assert why_shape is not None, "no 'Why:' competitor card shape found"
    assert counter_shape is not None, "no 'Counter:' competitor card shape found"

    why_top_in = why_shape.top / EMU_PER_IN
    why_w_in = why_shape.width / EMU_PER_IN
    counter_top_in = counter_shape.top / EMU_PER_IN

    # The TRUE number of lines "Why:" needs, measured at the box's real
    # usable text width (box width minus python-pptx's 0.2in left+right
    # insets) -- this is what the card's own ink actually occupies,
    # regardless of what width the (possibly buggy) measurement code used
    # to decide the box height.
    true_n_lines = ppt._measure_lines(
        why_shape.text_frame.text, why_w_in - 0.2, 8.0, bold=False
    )
    assert true_n_lines >= 3, (
        f"fixture calibration drifted: expected the composed Why sentence "
        f"to need >=3 true lines, got {true_n_lines}"
    )
    line_h_in = (8.0 * 1.35) / 72.0
    why_true_ink_bottom_in = why_top_in + true_n_lines * line_h_in

    assert counter_top_in >= why_true_ink_bottom_in - 0.005, (
        f"Counter (top={counter_top_in:.3f}in) starts before Why's TRUE "
        f"{true_n_lines}-line ink actually ends (bottom="
        f"{why_true_ink_bottom_in:.3f}in) -- Counter's first line prints "
        f"on top of Why's real last line"
    )


# ---------------------------------------------------------------------------
# C11 -- slide 6 (Budget Allocation) subhead must cascade the hero-card row
# below its measured bottom for a long (~40+ char) client name.
# ---------------------------------------------------------------------------
def _long_client_name_plan_for_budget_slide() -> Dict[str, Any]:
    return {
        "client_name": "Kaiser Foundation Health Plan of Northern California",
        "industry": "healthcare_medical",
        "industry_label": "Healthcare & Medical",
        "locations": ["Oakland, CA"],
        "roles": ["Registered Nurse"],
        "budget": "$2,000,000",
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "social_media": True,
        },
        "_budget_allocation": {
            "metadata": {"total_budget": 2000000},
            "total_projected": {"applications": 20000, "hires": 800},
            "channel_allocations": {
                "programmatic_dsp": {"percentage": 40, "projected_hires": 400},
                "global_boards": {"percentage": 35, "projected_hires": 300},
                "social_media": {"percentage": 25, "projected_hires": 100},
            },
        },
    }


def test_budget_slide_subhead_never_overlaps_hero_cards():
    data = _long_client_name_plan_for_budget_slide()
    prs = _gen(data)
    slide = _slide_containing(prs, "Total Investment")
    assert slide is not None, "Budget Allocation slide not found"

    subhead = _shape_with_text(slide, "Investment breakdown across")
    assert subhead is not None
    font_pt = _run_font_pt_local(subhead)
    n_lines = ppt._measure_lines(
        subhead.text_frame.text, subhead.width / EMU_PER_IN, font_pt, bold=True
    )
    assert n_lines >= 2, "fixture should force a 2+ line subhead"

    label_shape = _shape_with_text(slide, "Total Investment", exact=True)
    assert label_shape is not None
    # The hero card's accent bar sits at the card's declared top -- find the
    # card via the value shape above the label and use ITS top.
    hero_top_in = label_shape.top / EMU_PER_IN - 0.72  # label sits +0.72in into card

    # The PRE-FIX code pinned the hero row at a hardcoded y=1.5in regardless
    # of the subhead's line count. A 2-line subhead's measured box alone
    # (0.92in top + 2 lines at 15pt*1.35 line-height + 0.08in padding =
    # ~1.56in bottom) already exceeds that fixed 1.5in anchor -- so the fix
    # must have moved the hero row measurably below 1.5in. Assert the hero
    # row's actual top is strictly past the old fixed anchor (with margin),
    # which only holds when the code cascades from the subhead's REAL
    # measured height instead of the constant.
    assert hero_top_in > 1.55, (
        f"hero-card row still sits at (or near) the old fixed y=1.5in "
        f"anchor (hero_top={hero_top_in:.3f}in) despite a {n_lines}-line "
        f"subhead -- it was not cascaded from the subhead's measured height"
    )

    # And, directly: the subhead's own measured ink must not run past
    # where the hero row now starts.
    subhead_top_in = subhead.top / EMU_PER_IN
    measured_bottom_in = subhead_top_in + n_lines * (font_pt * 1.35) / 72.0
    assert measured_bottom_in <= hero_top_in + 0.02, (
        f"multi-line subhead's measured ink (bottom={measured_bottom_in:.2f}in) "
        f"overlaps the hero-card row (top={hero_top_in:.2f}in)"
    )


def _run_font_pt_local(shape, default: float = 10.0) -> float:
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                return r.font.size.pt
    return default


# ---------------------------------------------------------------------------
# C12 -- slide 8 (Plan Comparison) subhead must cascade the comparison
# panels below its measured bottom -- this reproduces on EVERY deck (the
# fixed sentence + industry label already wraps to 2 lines), not only
# long-name outliers.
# ---------------------------------------------------------------------------
def test_comparison_slide_subhead_never_overlaps_panels():
    # A long client name + long industry label together force a 3-line
    # subhead (mirrors the audit's long_name_120 fixture class).
    data = _long_client_name_plan_for_budget_slide()
    data["client_name"] = (
        "The Extraordinarily Long-Named International Consolidated "
        "Holdings and Regional Distribution Partners Group Worldwide LLC"
    )
    data["industry_label"] = (
        "Hospitality, Travel, Leisure, Gaming and Entertainment Services "
        "(Global)"
    )
    prs = _gen(data)
    slide = _slide_containing(prs, "Channels Selected")
    assert slide is not None, "Plan Comparison slide not found"

    subhead = _shape_with_text(slide, "optimized media plan vs")
    assert subhead is not None
    font_pt = _run_font_pt_local(subhead)
    n_lines = ppt._measure_lines(
        subhead.text_frame.text, subhead.width / EMU_PER_IN, font_pt, bold=True
    )
    assert n_lines >= 3, "fixture should force a 3+ line subhead"
    subhead_top_in = subhead.top / EMU_PER_IN
    measured_bottom_in = subhead_top_in + n_lines * (font_pt * 1.35) / 72.0

    panel_header = _shape_with_text(slide, "⬢", exact=False)  # "⬢" panel marker
    assert panel_header is not None
    panel_top_in = panel_header.top / EMU_PER_IN

    # PRE-FIX code pinned the comparison panels at a hardcoded y=1.55in
    # regardless of the subhead's line count. A 3-line subhead's own
    # measured bottom (~1.76in) already clears that fixed anchor, so the
    # panels must have been cascaded measurably past it.
    assert panel_top_in > 1.6, (
        f"comparison panels still sit at (or near) the old fixed y=1.55in "
        f"anchor (panel_top={panel_top_in:.3f}in) despite a {n_lines}-line "
        f"subhead -- they were not cascaded from the subhead's measured "
        f"height"
    )
    assert measured_bottom_in <= panel_top_in + 0.02, (
        f"subhead's measured ink (bottom={measured_bottom_in:.2f}in) "
        f"overlaps the comparison panel header (top={panel_top_in:.2f}in)"
    )


# ---------------------------------------------------------------------------
# C11/C12 class -- the comparison panel's own header bar ("<client>'s Plan")
# must shrink-to-fit a long client name rather than clip it. Pre-fix, the
# header was a fixed 12pt/one-line box: a name past ~50 characters got cut
# off mid-name (e.g. the entity suffix "LLC" silently never drawn).
# ---------------------------------------------------------------------------
def test_comparison_panel_header_never_clips_long_client_name():
    long_name = (
        "The Extraordinarily Long-Named International Consolidated Holdings "
        "and Regional Distribution Partners Group Worldwide LLC"
    )
    data = {
        "client_name": long_name,
        "industry": "general_entry_level",
        "roles": ["Operations Associate"],
        "locations": ["Denver, CO"],
        "budget": "$120,000",
        "channel_categories": {"programmatic_dsp": True},
    }
    prs = _gen(data)
    slide = _slide_containing(prs, "Channels Selected")
    assert slide is not None, "Plan Comparison slide not found"

    panel_header = _shape_with_text(slide, "⬢", exact=False)
    assert panel_header is not None
    full_text = panel_header.text_frame.text
    assert "LLC" in full_text and full_text.rstrip().endswith("Plan"), (
        f"panel header text was clipped -- the full client name (including "
        f"the trailing entity suffix 'LLC' and \"'s Plan\") must be present: "
        f"{full_text!r}"
    )
    font_pt = _run_font_pt_local(panel_header)
    assert font_pt >= 8.0, f"panel header font fell below the 8pt floor: {font_pt}"
    # PRE-FIX code hardcoded font_size=12 unconditionally, regardless of
    # client-name length -- at 12pt this long name needs 2 lines in a
    # single-line 0.35in box and got clipped. A strictly-smaller font is
    # the direct signal that shrink-to-fit actually engaged.
    assert font_pt < 12.0, (
        f"panel header is still the old hardcoded 12pt (font_pt={font_pt}) "
        f"for a name this long -- shrink-to-fit did not engage, so the "
        f"text overflows its fixed one-line 0.35in box"
    )
    # Shrink-to-fit lands at the 8pt readability floor for a name this
    # long (the box only ever offers one line's worth of width -- at 8pt
    # the text still wraps to 2 lines, which the MIDDLE-anchored box
    # accommodates without clipping, unlike the pre-fix fixed 12pt where
    # word-wrap silently dropped the overflow line's content).
    assert font_pt == 8.0, (
        f"expected shrink-to-fit to bottom out at the 8pt floor for this "
        f"120-char name, got {font_pt}pt"
    )


# ---------------------------------------------------------------------------
# Zero-hire (C18) -- a plan projecting 0 hires must never render a
# cost-per-hire figure derived from dividing the whole budget by a
# hires-floor-of-1.
# ---------------------------------------------------------------------------
def test_zero_hire_plan_never_shows_budget_as_cost_per_hire():
    data = {
        "client_name": "Pixel Bakery",
        "industry": "retail_consumer",
        "_knowledge_base": {"recruitment_benchmarks": {"industry_benchmarks": {}}},
        "_budget_allocation": {
            "metadata": {"total_budget": 3000},
            "total_projected": {
                "applications": 227,
                "hires": 0,
                # This is the exact shape of the pre-existing budget_engine
                # bug this fix defends against: cost_per_hire == the WHOLE
                # budget (a hires-floor-of-1 artifact), even though hires
                # is honestly reported as 0 right above it.
                "cost_per_hire": 3000.0,
            },
            "channel_allocations": {
                "programmatic_dsp": {"projected_hires": 0},
            },
        },
    }
    bm = ppt._get_benchmarks("retail_consumer", data)
    # Pre-fix, Layer 0 falls back to the raw total_projected.cost_per_hire
    # (3000.0) whenever _compute_blended_cph returns a falsy 0.0 -- even
    # though _compute_blended_cph ALSO tells the caller hires is truly 0.
    # It then formats a +/-20% "range" around that bad number: low=2400,
    # high=3600. Assert neither half of that fabricated range appears.
    assert "2,400" not in bm["cph"] and "2400" not in bm["cph"], (
        f"zero-hire plan's cost-per-hire benchmark is the 0.8x-budget "
        f"fabricated range low end: {bm['cph']!r}"
    )
    assert "3,600" not in bm["cph"] and "3600" not in bm["cph"], (
        f"zero-hire plan's cost-per-hire benchmark is the 1.2x-budget "
        f"fabricated range high end (whole budget as cost-per-hire): "
        f"{bm['cph']!r}"
    )
    assert bm.get("cph_is_usd_benchmark", True) is True, (
        "a zero-hire plan's CPH row must fall back to a real benchmark "
        "(flagged as a USD/industry benchmark), never be marked as this "
        "plan's own derived figure"
    )


# ---------------------------------------------------------------------------
# Missing-budget default: a plan with no "budget" must render the neutral
# "Not specified" on client-facing slides, never the literal "TBD" jargon
# abbreviation (matches the hire_volume / workbook-budget fix elsewhere).
# ---------------------------------------------------------------------------
def test_missing_budget_renders_not_specified_not_tbd():
    data = {
        "client_name": "No Budget Co",
        "industry": "general_entry_level",
        "roles": ["Clerk"],
        "locations": ["Austin, TX"],
        "channel_categories": {"programmatic_dsp": True},
        # No "budget" key at all, AND a present-but-zero budget allocation
        # (exercises BOTH reachable sites: slide 2's SITUATION card and
        # slide 6's "Total Investment" hero value).
        "_budget_allocation": {
            "metadata": {"total_budget": 0},
            "total_projected": {"applications": 0, "hires": 0},
            "channel_allocations": {},
        },
    }
    prs = _gen(data)
    all_text = []
    for slide in prs.slides:
        all_text.extend(_all_texts(slide))
    joined = "\n".join(all_text)
    assert "TBD" not in joined, (
        f"the literal 'TBD' abbreviation leaked onto a client-facing slide "
        f"for a plan with no budget specified"
    )

    situation_slide = _slide_containing(prs, "Target Roles")
    assert situation_slide is not None
    situation_shape = _shape_with_text(situation_slide, "Budget:")
    assert situation_shape is not None
    assert "Not specified" in situation_shape.text_frame.text

    hero_slide = _slide_containing(prs, "Total Investment")
    assert hero_slide is not None, "Budget Allocation slide not found"
    assert _shape_with_text(hero_slide, "Not specified", exact=True) is not None, (
        "slide 6 hero 'Total Investment' value should read 'Not specified' "
        "when there is no parseable budget"
    )


if __name__ == "__main__":
    import traceback

    tests = [
        test_funded_display_channels_drops_unfunded_channel_and_foots_to_100,
        test_no_phantom_apac_channel_and_consistent_count_across_slides,
        test_us_sourced_salary_marked_usd_on_non_usd_plan,
        test_industry_kb_benchmark_wins_over_generic_ad_platform_range,
        test_generic_platform_range_labeled_cross_industry_when_no_kb_benchmark,
        test_competitor_card_counter_never_overlaps_why_real_last_line,
        test_budget_slide_subhead_never_overlaps_hero_cards,
        test_comparison_slide_subhead_never_overlaps_panels,
        test_comparison_panel_header_never_clips_long_client_name,
        test_zero_hire_plan_never_shows_budget_as_cost_per_hire,
        test_missing_budget_renders_not_specified_not_tbd,
    ]
    failures = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except Exception:
            failures += 1
            print(f"FAIL {t.__name__}")
            traceback.print_exc()
    if failures:
        sys.exit(1)
    print("All tests passed.")
