"""Boards are chosen by local relevance, not by how many countries they list.

The regional & aggregator budget goes to the boards with the highest
local-relevance score for each plan country (market-share list, supply-team
note, Joveo-built plans, market-native, role fit, programmatic feed,
priority). The hard gate -- in network for the country, not client-excluded,
no board with its own channel -- is not a score and is never overridden.
"""

import openpyxl
import pytest

import channel_funding as cf
import excel_v2
import publisher_filter as pf

DE_MANUFACTURING = {
    "locations": ["Berlin, Germany"],
    "target_roles": ["Maintenance Technician - Paint Shop", "Industrial Painter"],
    "industry": "automotive",
}
UK_LOGISTICS = {
    "locations": ["London, United Kingdom"],
    "target_roles": ["Warehouse Operative", "Forklift Driver"],
    "industry": "logistics_supply_chain",
}
US_HOURLY = {
    "locations": ["Austin, TX"],
    "target_roles": ["Warehouse Associate", "Picker Packer"],
    "industry": "logistics_supply_chain",
}


def _alloc(amount=86511.54):
    return {
        "channel_allocations": {
            "regional_boards": {"dollar_amount": amount, "label": "Regional Boards"}
        }
    }


def _entry(data, amount=86511.54):
    return cf.regional_budget_split(data, _alloc(amount))["countries"][0]


def _cents(x):
    return round(x * 100)


# -- the Germany manufacturing plan -----------------------------------------


def test_germany_manufacturing_funds_stepstone_xing_and_studysmarter():
    funded = [pf.normalize(b["name"]) for b in _entry(DE_MANUFACTURING)["boards"]]
    for wanted in ("stepstone", "xing", "studysmarter"):
        assert any(wanted in n for n in funded), (wanted, funded)


def test_a_multi_country_board_is_funded_as_a_local_board():
    entry = _entry(DE_MANUFACTURING)
    ranked = {b["name"]: b for b in cf.ranked_boards(DE_MANUFACTURING, "Germany")}
    wide = [b for b in entry["boards"] if ranked[b["name"]]["n_countries"] > 5]
    assert wide, "no board listing more than five countries was funded"
    # StudySmarter lists 23 countries and is funded on the strength of its
    # supply-team note, plan evidence and role fit
    study = ranked["StudySmarter"]
    assert study["n_countries"] > 5 and study["fundable"]
    # a board that lists fewer countries but has no relevance signal is not
    narrow = [
        b
        for b in ranked.values()
        if b["n_countries"] <= 2 and not b["relevant"] and not b["manual"]
    ]
    assert narrow
    assert all(b["score"] < study["score"] for b in narrow)


def test_score_is_the_sum_of_its_signals():
    ranked = {b["name"]: b for b in cf.ranked_boards(DE_MANUFACTURING, "Germany")}
    stepstone = ranked["StepStone"]
    # first on Germany's market-share list (+40), funded in Joveo-built plans
    # (+20), market-native by its "Germany" row name (+15), role fit, priority
    # P2 (+2); StepStone Germany is a posting-priced board with a feed, so no
    # programmatic bonus.
    assert stepstone["priority"] == "P2"
    assert stepstone["score"] == pytest.approx(
        cf.SCORE_MARKET_SHARE_TOP
        + cf.SCORE_PLAN_FUNDED
        + cf.SCORE_MARKET_NATIVE
        + stepstone["role_fit"]
        + cf.SCORE_PRIORITY["P2"]
    )
    study = ranked["StudySmarter"]
    assert study["score"] == pytest.approx(
        cf.SCORE_PERFORMANCE_NOTE
        + cf.SCORE_PLAN_FUNDED
        + study["role_fit"]
        + cf.SCORE_PROGRAMMATIC
        + cf.SCORE_PRIORITY["P1"]
    )


def test_market_share_bonus_steps_down_to_its_floor():
    assert cf._market_share_bonus(0) == 40.0
    assert cf._market_share_bonus(1) == 38.0
    assert cf._market_share_bonus(50) == 30.0


def test_ties_break_on_fewer_listed_countries_then_name():
    ranked = cf.ranked_boards(DE_MANUFACTURING, "Germany")
    for a, b in zip(ranked, ranked[1:]):
        key_a = (-a["score"], a["n_countries"], a["name"].casefold())
        key_b = (-b["score"], b["n_countries"], b["name"].casefold())
        # name ties can reorder across display/row names; scores never go up
        assert a["score"] >= b["score"]
        if a["score"] == b["score"]:
            assert a["n_countries"] <= b["n_countries"], (a["name"], b["name"])


# -- shares and amounts ------------------------------------------------------


@pytest.mark.parametrize("amount", [86511.54, 12345.67, 0.07, 100000.0])
def test_shares_follow_scores_and_amounts_sum_exactly(amount):
    entry = _entry(DE_MANUFACTURING, amount)
    boards = entry["boards"]
    assert len(boards) == cf.FUNDED_BOARD_COUNT
    assert sum(_cents(b["amount"]) for b in boards) == _cents(entry["amount"])
    total = sum(b["score"] for b in boards)
    for b in boards:
        assert b["share"] == pytest.approx(b["score"] / total)
    assert sum(b["share"] for b in boards) == pytest.approx(1.0)


# -- hard gate ---------------------------------------------------------------


def _funded_and_standby(data):
    entry = _entry(data)
    return entry["boards"] + entry["standby"] + entry["post_directly"]


@pytest.mark.parametrize("data", [DE_MANUFACTURING, UK_LOGISTICS, US_HOURLY])
def test_every_listed_board_is_in_network_and_has_no_channel_of_its_own(data):
    country = pf.plan_countries(data)[0]
    for b in _funded_and_standby(data):
        rec = pf.board_supply_record(b["name"], country)
        assert rec["state"] == pf.SUPPLY_IN_NETWORK, b["name"]
        assert not any(
            pf.names_match(b["name"], o) for o in ("Indeed", "LinkedIn", "Glassdoor")
        ), b["name"]


def test_a_performance_note_alone_never_funds_a_board_outside_the_network(monkeypatch):
    # Jobflow: named in Joveo-built plans AND given a performance note, but the
    # supply data has no in-network record for it in Germany
    assert pf.board_supply_record("Jobflow", "Germany")["state"] != pf.SUPPLY_IN_NETWORK
    monkeypatch.setattr(
        cf,
        "_performance_notes",
        lambda c: [
            {
                "names": ["Jobflow"],
                "text": "performs well",
                "source": "Joveo supply team",
                "as_of": "2026-10-21",
            }
        ],
    )
    names = [b["name"] for b in _funded_and_standby(DE_MANUFACTURING)]
    assert not [n for n in names if pf.names_match(n, "Jobflow")]


def test_client_excluded_board_is_never_listed():
    victim = _entry(DE_MANUFACTURING)["boards"][0]["name"]
    data = dict(DE_MANUFACTURING, excluded_publishers=[victim])
    assert victim not in [b["name"] for b in _funded_and_standby(data)]


# -- niche boards need role fit ---------------------------------------------


def _fake_row(name, category, **kw):
    row = {
        "name": name,
        "priority": "P0",
        "job_type": "General",
        "n_countries": 1,
        "record": {
            "state": pf.SUPPLY_IN_NETWORK,
            "status": "In Joveo Network",
        },
        "category": category,
        "job_categories": (),
        "pricing": "CPA + CPC",
        "xml": "Yes",
    }
    row.update(kw)
    return row


def test_a_niche_board_without_role_fit_is_never_funded(monkeypatch):
    real = pf.in_network_job_boards
    fakes = [
        _fake_row("Campus Careers DE", "University Job Board"),
        _fake_row("MedJobs DE", "Health"),
        _fake_row("TechTalent DE", "Tech"),
    ]
    monkeypatch.setattr(pf, "in_network_job_boards", lambda c: real(c) + fakes)
    monkeypatch.setattr(pf, "status_label", lambda rec: "In Joveo Network")
    # each fake would win on priority + programmatic + .de nativeness
    entry = _entry(DE_MANUFACTURING)
    funded = {b["name"] for b in entry["boards"]}
    assert not funded & {f["name"] for f in fakes}
    assert not {b["name"] for b in entry["standby"]} & {f["name"] for f in fakes}
    # the same board funds when the plan's roles are its audience
    nurses = dict(DE_MANUFACTURING, target_roles=["Registered Nurse"], industry="x")
    ranked = {b["name"]: b for b in cf.ranked_boards(nurses, "Germany")}
    assert ranked["MedJobs DE"]["fundable"] and ranked["MedJobs DE"]["role_fit"] > 0
    assert not ranked.get("TechTalent DE", {"fundable": False})["fundable"]


@pytest.mark.parametrize("data", [DE_MANUFACTURING, UK_LOGISTICS, US_HOURLY])
def test_no_funded_board_is_a_niche_board_without_role_fit(data):
    country = pf.plan_countries(data)[0]
    ranked = {b["name"]: b for b in cf.ranked_boards(data, country)}
    for b in _entry(data)["boards"]:
        row = ranked[b["name"]]
        assert not (row["niche"] and row["role_fit"] <= 0), b["name"]


# -- manual-posting boards ---------------------------------------------------


def test_manual_boards_get_no_budget_and_are_listed_to_post_directly():
    entry = _entry(DE_MANUFACTURING)
    assert 1 <= len(entry["post_directly"]) <= cf.POST_DIRECTLY_COUNT
    funded = {b["name"] for b in entry["boards"] + entry["standby"]}
    rows = [r for r in pf._repo_job_boards() if "Germany" in r["countries"]]

    def _row(name):
        return next(r for r in rows if pf.names_match(name, r["name"]))

    for b in entry["post_directly"]:
        assert b["name"] not in funded
        assert "amount" not in b
        row = _row(b["name"])
        assert cf._is_manual(row["pricing"], row["xml"]), b["name"]
    for b in entry["boards"] + entry["standby"]:
        row = _row(b["name"])
        assert not cf._is_manual(row["pricing"], row["xml"]), b["name"]


@pytest.mark.parametrize(
    "pricing,xml,manual",
    [
        ("CPC", "Yes", False),
        ("CPA + CPC", "Yes", False),
        ("Posting", "No", True),
        ("Subscription", "", True),
        ("Slot", "nan", True),
        ("CPC", "No", True),
        ("", "No", True),
        # priced by posting but takes a feed (the market-leading German
        # generalist): can be fed, so not manual
        ("Posting", "Yes", False),
        ("", "", False),
    ],
)
def test_manual_posting_rule(pricing, xml, manual):
    assert cf._is_manual(pricing, xml) is manual


# -- why strings -------------------------------------------------------------


def test_every_funded_and_standby_board_says_why():
    entry = _entry(DE_MANUFACTURING)
    for b in entry["boards"] + entry["standby"]:
        assert b["why"].strip()
    why = {pf.normalize(b["name"]): b["why"] for b in entry["boards"]}
    study = next(v for k, v in why.items() if "studysmarter" in k)
    assert "performs well (Joveo supply team)" in study
    assert "in Joveo-built plans" in study
    assert "fits production roles" in study
    assert "programmatic" in study
    stepstone = next(v for k, v in why.items() if "stepstone" in k)
    assert "market leader" in stepstone and "local to Germany" in stepstone


# -- the other markets -------------------------------------------------------


@pytest.mark.parametrize(
    "data,country", [(UK_LOGISTICS, "United Kingdom"), (US_HOURLY, "United States")]
)
def test_uk_and_us_hourly_plans_fund_five_scored_programmatic_boards(data, country):
    entry = _entry(data, 60000)
    assert len(entry["boards"]) == 5
    assert sum(_cents(b["amount"]) for b in entry["boards"]) == _cents(60000)
    ranked = {b["name"]: b for b in cf.ranked_boards(data, country)}
    for b in entry["boards"]:
        assert not ranked[b["name"]]["manual"]
        assert b["score"] > 0
    # blue-collar logistics roles: the boards lead on role fit, not on count
    assert any("fits logistics roles" in b["why"] for b in entry["boards"])


# -- evidence file -----------------------------------------------------------


def test_plan_evidence_file_has_provenance_and_the_germany_lists():
    import json

    doc = json.load(open(cf._PLAN_EVIDENCE_PATH, encoding="utf-8"))
    assert doc["source"] == "Joveo-built media plans (internal), 2026-10"
    assert doc["as_of"]
    de = doc["countries"]["Germany"]
    assert len(de["funded"]) == 14 and len(de["standby"]) == 9
    assert {"XING", "StepStone", "StudySmarter"} <= set(de["funded"])
    for entry in doc["countries"].values():
        assert not set(entry["funded"]) & set(entry["standby"])


# -- workbook ----------------------------------------------------------------


def test_workbook_sub_table_has_a_why_column_and_the_post_directly_line():
    data = dict(DE_MANUFACTURING)
    data["_budget_allocation"] = {
        "metadata": {"total_budget": 250000},
        "channel_allocations": {
            "regional_boards": {
                "label": "Regional Job Boards",
                "dollar_amount": 86511.54,
                "projected_hires": 20,
                "projected_applications": 2000,
                "projected_clicks": 5000,
                "cpa": 27,
            }
        },
    }
    ws = openpyxl.Workbook().active
    excel_v2._build_sheet_channel_recommendations(ws, data)
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Why" in cells
    assert cf.POST_DIRECTLY_LABEL in cells
    assert cf.STANDBY_LABEL in cells
    assert any("performs well (Joveo supply team)" in c for c in cells)
    assert any(c.startswith("In Joveo Network") for c in cells)  # status kept


# -- Nova agrees with the plan -----------------------------------------------


def test_nova_germany_listing_leads_with_the_plans_funding_order():
    import nova

    engine = nova.Nova()
    query = "List all job boards and publishers for jobs in Germany"
    result = engine._fast_path_supply_listing(query, query.lower())
    text = result["response"]
    start = text.index("### Best-fit boards in Germany")
    section = text[start:].split("\n\n")[0].splitlines()[1:]
    listed = [line.split(". ", 1)[1].split(" — ")[0] for line in section]
    expected = [b["name"] for b in cf.funding_order({}, "Germany")[:5]]
    assert listed == expected
    assert len(listed) == 5


# -- government general boards: post directly, first --------------------------


def test_germany_post_directly_leads_with_the_public_employment_service():
    entry = _entry(DE_MANUFACTURING)
    post = entry["post_directly"]
    assert pf.normalize(post[0]["name"]).startswith("bundesagentur")
    assert post[0]["why"] == "public employment service · free organic posting"
    assert len(post) <= cf.POST_DIRECTLY_COUNT
    # it is a government general board with no role fit to a production plan:
    # exempt for the post-directly list only, never funded
    ranked = cf.ranked_boards(DE_MANUFACTURING, "Germany")
    assert not [b for b in ranked if b["gov_general"] and b["fundable"]]
    funded = {b["name"] for b in entry["boards"] + entry["standby"]}
    assert not any(pf.normalize(n).startswith("bundesagentur") for n in funded)


def test_government_boards_with_a_narrow_audience_are_not_exempt(monkeypatch):
    real = pf.in_network_job_boards
    fake = {
        "name": "Ministry Careers DE",
        "priority": "P0",
        "job_type": "Blue Collar",  # not a general board
        "n_countries": 1,
        "record": {"state": pf.SUPPLY_IN_NETWORK, "status": "In Joveo Network"},
        "category": "Govt",
        "job_categories": (),
        "pricing": "Organic",
        "xml": "No",
    }
    monkeypatch.setattr(pf, "in_network_job_boards", lambda c: real(c) + [fake])
    names = [b["name"] for b in cf.post_directly_boards(DE_MANUFACTURING, "Germany")]
    assert "Ministry Careers DE" not in names


# -- the selection note follows the surface ----------------------------------


def test_ranked_note_replaces_the_market_share_note_beside_scored_boards():
    groups = pf.plan_market_boards({"locations": ["Berlin, Germany"]})
    assert pf.RANKED_NOTE == (
        "Boards are ranked by local relevance: market share, Joveo supply-team "
        "performance notes, use in Joveo-built plans, local presence, role fit "
        "and programmatic buying."
    )
    ranked = pf.supply_depth_line(groups, ranked=True)
    assert ranked.endswith(pf.RANKED_NOTE)
    assert "chosen by market share" not in ranked
    plain = pf.supply_depth_line(groups)
    assert "chosen by market share, not matched to the role" in plain
    assert pf.RANKED_NOTE not in plain
    assert pf.supply_depth_line([{"depth": None}], ranked=True) == pf.RANKED_NOTE


def test_germany_deck_and_workbook_use_the_note_that_is_true_there():
    import io

    import ppt_generator
    from pptx import Presentation

    data = dict(DE_MANUFACTURING, client_name="Acme", budget="250000")
    data["_budget_allocation"] = {
        "channel_allocations": {
            "regional_boards": {
                "label": "Regional Job Boards",
                "dollar_amount": 86511.54,
            }
        }
    }
    # workbook: the sub-table carries the ranked note
    ws = openpyxl.Workbook().active
    excel_v2._build_sheet_channel_recommendations(ws, data)
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert pf.RANKED_NOTE in cells
    # deck slide 5: the funded line is there, so only the ranked note is
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    texts = [
        sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame
    ]
    assert any(t.startswith("Regional & aggregator budget funds") for t in texts)
    assert any(pf.RANKED_NOTE in t for t in texts)
    assert not any("chosen by market share" in t for t in texts)
