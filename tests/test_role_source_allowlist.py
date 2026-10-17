"""ENRICH_ROLE_SOURCES is a per-label allowlist (2026-10-17).

Only BLS and Adzuna run by default; the other 15 role-based sources stay
off until each is audited (see the block comment above ROLE_SOURCES_ENV in
api_enrichment.py). Also covers the shared Adzuna request quota
(adzuna_quota) and proves Adzuna's curated fallback never reads as a live
pull.

No test here touches the network: every role fetcher is replaced by a
recorder, HTTP helpers are patched, and the suite's external-network guard
(tests/conftest.py) is the backstop.
"""

from __future__ import annotations

import io
import logging
import sys
import uuid
from pathlib import Path
from typing import Any, Callable, Dict, List

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import adzuna_quota  # noqa: E402
import api_enrichment  # noqa: E402
import api_integrations  # noqa: E402
import data_synthesizer  # noqa: E402
import excel_v2  # noqa: E402
import market_signals  # noqa: E402
from tests.test_enrich_geopolitical_budget import (  # noqa: E402,F401
    isolated_breakers,
    stub_sources,
)

# Spelled out (not read from the module) so no assertion can pass vacuously
# against an empty declaration. label -> the api_enrichment function it calls.
_FETCHER_OF: Dict[str, str] = {
    "BLS": "fetch_salary_data",
    "GoogleTrends": "fetch_search_trends",
    "O*NET": "fetch_onet_occupation_data",
    "DataUSA-Occ": "fetch_datausa_occupation_stats",
    "GoogleAds": "fetch_google_ads_data",
    "MetaAds": "fetch_meta_ads_data",
    "BingAds": "fetch_bing_ads_data",
    "TikTokAds": "fetch_tiktok_ads_data",
    "LinkedInAds": "fetch_linkedin_ads_data",
    "CareerOneStop": "fetch_careeronestop_data",
    "COS-Projections": "fetch_cos_occupation_projections",
    "H1B-Wages": "fetch_h1b_wage_benchmarks",
    "H1B-Salary-Intel": "_fetch_h1b_salary_intelligence",
    "Adzuna": "fetch_job_market",
    "Jooble": "fetch_jooble_data",
    "RegionalLabour": "fetch_regional_labour_data",
    "CareerJet": "fetch_careerjet_data",
}
_ALL = set(_FETCHER_OF)
_DEFAULT = {"BLS", "Adzuna"}

_WIZARD = {
    "client_name": "Acme Health",
    "industry": "healthcare_medical",
    "target_roles": ["Registered Nurse", "Medical Assistant"],
    "locations": ["Denver, CO"],
}


@pytest.fixture(autouse=True)
def _clean_state(monkeypatch: pytest.MonkeyPatch):
    """Fresh quota + warn-once state per test; the env starts unset."""
    adzuna_quota.reset()
    monkeypatch.delenv("ENRICH_ROLE_SOURCES", raising=False)
    monkeypatch.setattr(
        api_enrichment, "_warned_unknown_role_labels", set(), raising=False
    )
    monkeypatch.setenv("NOVA_GEOPOLITICAL_TIMEOUT_S", "0")
    yield
    adzuna_quota.reset()


class _Recorder:
    """Stands in for every role fetcher; remembers which ones were invoked."""

    def __init__(self) -> None:
        self.invoked: List[str] = []

    def install(self, monkeypatch: pytest.MonkeyPatch) -> None:
        fetchers = [
            n
            for n in dir(api_enrichment)
            if (n.startswith("fetch_") or n == "_fetch_h1b_salary_intelligence")
            and callable(getattr(api_enrichment, n))
        ]
        stub_sources(monkeypatch, fetchers)
        for label, fn_name in _FETCHER_OF.items():
            monkeypatch.setattr(api_enrichment, fn_name, self._make(label))

    def _make(self, label: str) -> Callable[..., Dict[str, Any]]:
        def _fn(*_args: Any, **_kwargs: Any) -> Dict[str, Any]:
            self.invoked.append(label)
            return {}

        return _fn


def _run(monkeypatch: pytest.MonkeyPatch, request: Dict[str, Any]) -> "tuple[set, set]":
    """Returns (role labels dispatched, role fetchers actually invoked)."""
    rec = _Recorder()
    rec.install(monkeypatch)
    result = api_enrichment.enrich_data(dict(request))
    called = set(result["enrichment_summary"]["apis_called"])
    return called & _ALL, set(rec.invoked)


# ─── the effective set ───────────────────────────────────────────────────────


@pytest.mark.parametrize("value", [None, "", "   "])
def test_unset_or_empty_gives_the_default_set(monkeypatch, value):
    if value is not None:
        monkeypatch.setenv("ENRICH_ROLE_SOURCES", value)
    assert set(api_enrichment.role_sources_allowed()) == _DEFAULT
    assert tuple(api_enrichment.DEFAULT_ROLE_SOURCE_LABELS) == ("BLS", "Adzuna")


def test_default_dispatches_exactly_bls_and_adzuna(monkeypatch, isolated_breakers):
    dispatched, invoked = _run(monkeypatch, _WIZARD)
    assert dispatched == _DEFAULT, dispatched
    assert invoked == _DEFAULT, invoked


def test_default_without_locations_dispatches_bls_only(monkeypatch, isolated_breakers):
    request = {k: v for k, v in _WIZARD.items() if k != "locations"}
    dispatched, invoked = _run(monkeypatch, request)
    assert dispatched == {"BLS"}
    assert invoked == {"BLS"}


@pytest.mark.parametrize("value", ["0", "false", "no", "off", "none", "OFF", " None "])
def test_kill_switch_dispatches_no_role_task(monkeypatch, isolated_breakers, value):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", value)
    assert api_enrichment.role_sources_allowed() == ()
    assert api_enrichment.role_sources_enabled() is False
    dispatched, invoked = _run(monkeypatch, _WIZARD)
    assert dispatched == set()
    assert invoked == set()


def test_kill_switch_also_covers_an_explicit_roles_key(monkeypatch, isolated_breakers):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", "off")
    request = dict(_WIZARD, roles=_WIZARD["target_roles"])
    request.pop("target_roles")
    dispatched, invoked = _run(monkeypatch, request)
    assert dispatched == set() and invoked == set()


def test_kill_switch_logs_the_sources_that_would_have_run(
    monkeypatch, isolated_breakers, caplog
):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", "0")
    with caplog.at_level(logging.INFO, logger="api_enrichment"):
        _run(monkeypatch, _WIZARD)
    lines = [
        r.getMessage()
        for r in caplog.records
        if "ENRICH_ROLE_SOURCES" in r.getMessage()
    ]
    assert len(lines) == 1, lines
    assert all(label in lines[0] for label in _ALL), lines[0]


@pytest.mark.parametrize("value", ["1", "true", "YES", "on", "all", " All "])
def test_all_values_enable_all_seventeen(monkeypatch, isolated_breakers, value):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", value)
    assert set(api_enrichment.role_sources_allowed()) == _ALL
    assert len(api_enrichment.role_sources_allowed()) == 17
    dispatched, invoked = _run(monkeypatch, _WIZARD)
    assert dispatched == _ALL
    assert invoked == _ALL


@pytest.mark.parametrize(
    "value, expected",
    [
        ("BLS", {"BLS"}),
        ("bls,adzuna", {"BLS", "Adzuna"}),
        ("  O*NET ,  jooble ", {"O*NET", "Jooble"}),
        ("BLS,,Adzuna,", {"BLS", "Adzuna"}),
        ("careerjet,GoogleAds,h1b-wages", {"CareerJet", "GoogleAds", "H1B-Wages"}),
        ("BLS,BLS", {"BLS"}),
    ],
)
def test_comma_list_runs_exactly_those_labels(
    monkeypatch, isolated_breakers, value, expected
):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", value)
    assert set(api_enrichment.role_sources_allowed()) == expected
    dispatched, invoked = _run(monkeypatch, _WIZARD)
    assert dispatched == expected
    assert invoked == expected


def test_unknown_label_warns_once_and_is_ignored(monkeypatch, caplog):
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", "BLS, Monster ,adzuna")
    with caplog.at_level(logging.WARNING, logger="api_enrichment"):
        first = api_enrichment.role_sources_allowed()
        second = api_enrichment.role_sources_allowed()
    assert set(first) == set(second) == {"BLS", "Adzuna"}
    warnings = [
        r.getMessage()
        for r in caplog.records
        if r.levelno == logging.WARNING and "Monster" in r.getMessage()
    ]
    assert len(warnings) == 1, warnings


def test_only_unknown_labels_leaves_no_role_source(monkeypatch, caplog):
    """A typo fails closed (nothing runs) and is loud about it."""
    monkeypatch.setenv("ENRICH_ROLE_SOURCES", "bsl")
    with caplog.at_level(logging.WARNING, logger="api_enrichment"):
        assert api_enrichment.role_sources_allowed() == ()
    assert any("bsl" in r.getMessage() for r in caplog.records)


def test_allowlist_is_read_on_every_call(monkeypatch, isolated_breakers):
    seen = []
    for value in (None, "off", "all", "O*NET", "", "bls"):
        if value is None:
            monkeypatch.delenv("ENRICH_ROLE_SOURCES", raising=False)
        else:
            monkeypatch.setenv("ENRICH_ROLE_SOURCES", value)
        seen.append(_run(monkeypatch, _WIZARD)[0])
    assert seen == [
        _DEFAULT,
        set(),
        _ALL,
        {"O*NET"},
        _DEFAULT,
        {"BLS"},
    ]


def test_skipped_labels_are_still_logged_as_would_have_run(
    monkeypatch, isolated_breakers, caplog
):
    with caplog.at_level(logging.INFO, logger="api_enrichment"):
        _run(monkeypatch, _WIZARD)
    lines = [
        r.getMessage()
        for r in caplog.records
        if "ENRICH_ROLE_SOURCES" in r.getMessage()
    ]
    assert len(lines) == 1, lines
    for label in _ALL - _DEFAULT:
        assert label in lines[0], (label, lines[0])


def test_every_declared_role_source_has_a_fetcher_mapping():
    assert set(api_enrichment.role_sources_for(True)) == _ALL


# ─── Adzuna request quota ────────────────────────────────────────────────────


def test_adzuna_caps_are_per_worker_shares_of_the_service_budget():
    assert adzuna_quota.DAILY_CAP == 20
    assert adzuna_quota.MONTHLY_CAP == 600
    # four gunicorn workers must stay inside the 80/day and 2,400/month budget
    assert 4 * adzuna_quota.DAILY_CAP <= 80
    assert 4 * adzuna_quota.MONTHLY_CAP <= 2400
    # ... which itself sits under the provider ToS (250/day, 2,500/month)
    assert 80 <= 250 and 2400 <= 2500


def test_adzuna_daily_cap_refuses_the_next_request():
    results = [adzuna_quota.try_acquire() for _ in range(adzuna_quota.DAILY_CAP + 5)]
    assert results[: adzuna_quota.DAILY_CAP] == [True] * adzuna_quota.DAILY_CAP
    assert results[adzuna_quota.DAILY_CAP :] == [False] * 5
    assert adzuna_quota.exhausted() is True


class _Clock:
    def __init__(self) -> None:
        self.now = 1_000_000.0

    def time(self) -> float:
        return self.now


def test_adzuna_daily_window_reopens_after_a_day(monkeypatch):
    clock = _Clock()
    monkeypatch.setattr(adzuna_quota, "time", clock)
    for _ in range(adzuna_quota.DAILY_CAP):
        assert adzuna_quota.try_acquire()
    assert not adzuna_quota.try_acquire()
    clock.now += adzuna_quota.DAILY_WINDOW_S + 1
    assert adzuna_quota.try_acquire()


def test_adzuna_monthly_cap_binds_even_when_the_daily_window_has_room(monkeypatch):
    clock = _Clock()
    monkeypatch.setattr(adzuna_quota, "time", clock)
    monkeypatch.setattr(
        adzuna_quota, "_WINDOWS", (("daily", 5, 100), ("monthly", 8, 1000))
    )
    assert [adzuna_quota.try_acquire() for _ in range(6)] == [True] * 5 + [False]
    clock.now += 101  # daily window reopens, monthly (8) still has 3 left
    assert [adzuna_quota.try_acquire() for _ in range(4)] == [True] * 3 + [False]
    # the refusal above was the monthly cap: the daily window is not full
    assert adzuna_quota.status()["daily"]["calls_used"] == 3
    clock.now += 101
    assert adzuna_quota.try_acquire() is False  # still inside the month
    clock.now += 1000
    assert adzuna_quota.try_acquire() is True  # month rolled over


def test_a_refused_request_burns_no_quota_in_the_other_window(monkeypatch):
    monkeypatch.setattr(
        adzuna_quota, "_WINDOWS", (("daily", 2, 100), ("monthly", 1, 1000))
    )
    assert adzuna_quota.try_acquire()
    assert not adzuna_quota.try_acquire()
    assert adzuna_quota.status()["daily"]["calls_used"] == 1


def test_real_windows_never_let_more_than_the_monthly_cap_through(monkeypatch):
    clock = _Clock()
    monkeypatch.setattr(adzuna_quota, "time", clock)
    admitted = 0
    for _day in range(29):  # 29 daily steps stay inside one 30-day window
        for _ in range(adzuna_quota.DAILY_CAP + 10):
            admitted += adzuna_quota.try_acquire()
        clock.now += adzuna_quota.DAILY_WINDOW_S + 1
    assert admitted == 29 * adzuna_quota.DAILY_CAP <= adzuna_quota.MONTHLY_CAP


@pytest.fixture
def adzuna_http(monkeypatch):
    """Keys present, every Adzuna transport patched; counts what is 'sent'."""
    monkeypatch.setenv("ADZUNA_APP_ID", "test-id")
    monkeypatch.setenv("ADZUNA_APP_KEY", "test-key")
    sent: List[str] = []

    def _json(url: str, *args: Any, **kwargs: Any) -> Dict[str, Any]:
        sent.append(url)
        return {"count": 4321, "mean": 51000}

    monkeypatch.setattr(api_enrichment, "_http_get_json", _json)
    monkeypatch.setattr(api_enrichment, "_get_cached", lambda key: None)
    monkeypatch.setattr(api_enrichment, "_set_cached", lambda key, value: None)
    monkeypatch.setattr(api_integrations, "_http_get", _json)
    monkeypatch.setattr(api_integrations, "_get_cached", lambda key: None)
    monkeypatch.setattr(api_integrations, "_set_cached", lambda key, value: None)

    def _urlopen(*args: Any, **kwargs: Any) -> Any:
        sent.append("market_signals")
        raise OSError("blocked in test")

    monkeypatch.setattr(market_signals.urllib.request, "urlopen", _urlopen)
    return sent


def test_every_adzuna_call_site_draws_from_one_counter(monkeypatch, adzuna_http):
    monkeypatch.setattr(
        adzuna_quota, "_WINDOWS", (("daily", 4, 86400), ("monthly", 40, 2592000))
    )
    roles = [f"role-{uuid.uuid4().hex[:8]}" for _ in range(2)]

    # 2 requests: api_enrichment.fetch_job_market
    out = api_enrichment.fetch_job_market(roles, ["Denver, CO"])
    assert set(out) == set(roles)
    # 1 request: the AdzunaClient the sync /api/generate path uses
    client = api_integrations.AdzunaClient()
    assert client.get_job_count(f"role-{uuid.uuid4().hex[:8]}", "us")
    # 1 request: market_signals' own Adzuna call (fails after the send)
    market_signals._adzuna_live_cpc(f"role-{uuid.uuid4().hex[:8]}")
    assert len(adzuna_http) == 4

    # quota is now full: none of the four sites may send another request
    before = len(adzuna_http)
    assert (
        api_enrichment.fetch_job_market([f"r-{uuid.uuid4().hex[:8]}"], ["Denver, CO"])
        == {}
    )
    assert client.get_job_count(f"role-{uuid.uuid4().hex[:8]}", "us") is None
    assert client.search_jobs(f"role-{uuid.uuid4().hex[:8]}", "us") is None
    assert client.get_salary_histogram(f"role-{uuid.uuid4().hex[:8]}", "us") is None
    assert client.get_top_companies(f"role-{uuid.uuid4().hex[:8]}", "us") is None
    assert market_signals._adzuna_live_cpc(f"role-{uuid.uuid4().hex[:8]}") is None
    assert len(adzuna_http) == before


def test_over_quota_adzuna_is_skipped_not_ok(monkeypatch, adzuna_http):
    monkeypatch.setattr(api_enrichment, "_rate_limiter_state", {})
    monkeypatch.setattr(api_enrichment, "_circuit_breaker_check", lambda label: False)
    for _ in range(adzuna_quota.DAILY_CAP):
        adzuna_quota.try_acquire()

    calls: List[str] = []

    def _would_succeed() -> Dict[str, Any]:
        calls.append("ran")
        return {"r": {"posting_observation": "observed", "posting_count": 1}}

    result, status, meta = api_enrichment._safe_call(_would_succeed, "Adzuna")
    assert (result, status) == (None, "rate_limited")
    assert meta["success"] is False and meta["source"] == "rate_limited"
    assert calls == []  # the source never ran


def test_over_quota_adzuna_lands_in_apis_skipped_in_the_summary(
    monkeypatch, adzuna_http
):
    monkeypatch.setattr(api_enrichment, "_rate_limiter_state", {})
    monkeypatch.setattr(api_enrichment, "_circuit_breaker_check", lambda label: False)
    for _ in range(adzuna_quota.DAILY_CAP):
        adzuna_quota.try_acquire()
    fetchers = [
        n
        for n in dir(api_enrichment)
        if n.startswith("fetch_") and n != "fetch_job_market"
    ]
    stub_sources(monkeypatch, fetchers)
    summary = api_enrichment.enrich_data(dict(_WIZARD))["enrichment_summary"]
    assert "Adzuna" in summary["apis_called"]
    assert "Adzuna" not in summary["apis_succeeded"]
    assert "Adzuna" in summary["apis_skipped"]
    assert not adzuna_http  # no request left the process


def test_rate_limit_status_reports_both_adzuna_windows():
    status = api_enrichment._rate_limit_status()
    assert status["Adzuna:daily"]["max_calls"] == 20
    assert status["Adzuna:monthly"]["max_calls"] == 600


# ─── Adzuna's curated fallback never reads as a live pull ───────────────────


@pytest.fixture
def adzuna_without_keys(monkeypatch):
    monkeypatch.delenv("ADZUNA_APP_ID", raising=False)
    monkeypatch.delenv("ADZUNA_APP_KEY", raising=False)
    monkeypatch.setattr(api_enrichment, "_rate_limiter_state", {})
    monkeypatch.setattr(api_enrichment, "_circuit_breaker_check", lambda label: False)


def _enrich_with_real_adzuna(monkeypatch: pytest.MonkeyPatch) -> Dict[str, Any]:
    fetchers = [
        n
        for n in dir(api_enrichment)
        if n.startswith("fetch_") and n != "fetch_job_market"
    ]
    stub_sources(monkeypatch, fetchers)
    return api_enrichment.enrich_data(
        {
            "roles": ["Warehouse Associate"],
            "locations": ["Manchester, UK"],
            "industry": "logistics",
        }
    )


def test_adzuna_fallback_is_estimated_not_ok(adzuna_without_keys):
    result, status, meta = api_enrichment._safe_call(
        lambda: api_enrichment.fetch_job_market(["Warehouse Associate"], ["London"]),
        "Adzuna",
    )
    assert result["Warehouse Associate"]["posting_observation"] == "estimated"
    assert status == "estimated"
    assert meta["success"] is False and meta["source"] == "benchmark"


def test_adzuna_fallback_is_not_counted_as_a_success_or_as_confidence(
    monkeypatch, adzuna_without_keys
):
    enriched = _enrich_with_real_adzuna(monkeypatch)
    summary = enriched["enrichment_summary"]

    assert "Adzuna" in summary["apis_estimated"]
    assert "Adzuna" not in summary["apis_succeeded"]
    assert api_enrichment.enrichment_confidence(summary) == 0.0
    assert summary["confidence_score"] == 0.0
    # the payload is kept (it is a usable estimate) ...
    assert enriched["job_market"]["Warehouse Associate"]["posting_count"] > 0
    # ... but is stamped so downstream cannot mistake it for an observation
    assert (
        enriched["job_market"]["Warehouse Associate"]["posting_observation"]
        == "estimated"
    )

    import app as app_module

    warning, n_ok, _conf = app_module._enrichment_quality_warning(enriched)
    assert n_ok == 0 and warning  # the quality warning still fires


def test_adzuna_fallback_is_not_a_collector_payload_in_the_workbook(
    monkeypatch, adzuna_without_keys
):
    summary = _enrich_with_real_adzuna(monkeypatch)["enrichment_summary"]
    n_returned, n_applicable = excel_v2._collector_payload_coverage(summary)
    assert n_returned == 0
    assert n_applicable >= 1


def test_adzuna_fallback_never_renders_live_postings_or_market_temp(
    monkeypatch, adzuna_without_keys
):
    from pptx import Presentation

    import ppt_generator

    enriched = _enrich_with_real_adzuna(monkeypatch)
    request = {
        "roles": ["Warehouse Associate"],
        "locations": ["Manchester, UK"],
    }
    demand = data_synthesizer.fuse_job_market_demand(enriched, {}, request)
    role = demand["Warehouse Associate"]
    assert role["total_postings"] == 0
    assert role["posting_sources"] == []

    deck_data = {
        "client_name": "Fallback Test",
        "industry": "logistics",
        "industry_label": "Logistics",
        "budget": "$150,000",
        "budget_period": "campaign",
        "campaign_duration": "3 months",
        "campaign_start_month": 9,
        "hire_volume": "35 hires",
        "work_environment": "onsite",
        "locations": [{"city": "Columbus", "state": "OH", "country": "United States"}],
        "roles": [{"title": "Warehouse Associate", "count": 20, "tier": "mid"}],
        "_enriched": enriched,
        "_synthesized": {"job_market_demand": demand},
    }
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(deck_data)))
    text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Live Postings" not in text
    assert "Market Temp:" not in text
