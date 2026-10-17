"""Client story: a plain-language summary built ONLY from the plan's own fields.

Honesty guards (the point of the feature):
* every number in the story traces to a field of the plan (the allowlist below
  is built from the plan, never from the story);
* no superlative / promise phrases;
* a market with no local cost-per-hire benchmark carries no hire figure;
* a client's excluded boards are never named, and the story reaches the
  workbook, the deck and the results summary unchanged.
"""

from __future__ import annotations

import io
import os
import re
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import budget_engine  # noqa: E402
import client_story  # noqa: E402
import excel_v2  # noqa: E402
import plan_projection  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402


def _brief(name, industry, budget, locations, roles, **extra):
    brief = {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": "100 hires",
        "work_environment": "onsite",
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
    }
    brief.update(extra)
    return brief


@pytest.fixture(scope="module")
def plans():
    briefs = {
        "us_rn": _brief(
            "Probe Health",
            "Healthcare & Medical",
            "$250,000",
            ["Dallas, TX", "Austin, TX"],
            ["Registered Nurse", "Pharmacist"],
        ),
        "uk": _brief(
            "Probe UK",
            "Healthcare & Medical",
            "£200,000",
            ["London, UK", "Manchester, UK"],
            ["Registered Nurse"],
        ),
        "india": _brief(
            "Probe Infotech",
            "tech_engineering",
            "₹25,000,000",
            ["Bengaluru, Karnataka, India", "Pune, Maharashtra, India"],
            ["Java Developer", "QA Engineer"],
        ),
        "mixed": _brief(
            "Probe Mixed",
            "Healthcare & Medical",
            "$200,000",
            ["Dallas, TX", "London, UK"],
            ["Registered Nurse"],
        ),
        "exclusions": _brief(
            "Probe Excl",
            "Healthcare & Medical",
            "$250,000",
            ["Dallas, TX"],
            ["Registered Nurse"],
            excluded_publishers=["Alpha Jobs Exclusion", "Beta Careers Exclusion"],
        ),
        "enterprise": _brief(
            "Probe Hourly",
            "Healthcare & Medical",
            "$250,000",
            ["Dallas, TX", "Austin, TX", "Denver, CO", "Boston, MA", "Miami, FL"],
            ["Registered Nurse", "Pharmacist", "Technician", "Aide"],
            enterprise_pack=True,
            enterprise_system_of_record="Example ATS",
            enterprise_request_lock_day="friday",
            enterprise_setup_sla_days=5,
        ),
    }
    return {k: T.build_plan_data(b) for k, b in briefs.items()}


# --------------------------------------------------------------------------
# allowlist built from the PLAN, never from the story
# --------------------------------------------------------------------------
_NUM = re.compile(r"\d[\d,]*(?:\.\d+)?")


def story_numbers(text: str) -> list:
    text = re.sub(r"/10\b", "", text)  # the "9/10" scale denominator
    return [float(m.rstrip(",").replace(",", "")) for m in _NUM.findall(text)]


def _add(allowed: set, value) -> None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        allowed.update({float(value), float(round(value)), round(float(value), 2)})


def plan_numbers(plan: dict) -> set:
    alloc = plan["_budget_allocation"]
    allowed: set = set()
    _add(allowed, alloc["metadata"]["total_budget"])
    hp = alloc["headline_projection"]
    for key in ("applications", "hires", "cpa", "cph"):
        for bound in ("low", "mid", "high"):
            _add(allowed, (hp.get(key) or {}).get(bound))
    chans = alloc["channel_allocations"]
    funded = {k: v for k, v in chans.items() if (v.get("dollar_amount") or 0) > 0}
    _add(allowed, len(funded))
    avg = budget_engine.plan_average_cpa(chans)
    for entry in funded.values():
        for field in ("percentage", "roi_score", "trend_pct_yoy"):
            _add(allowed, entry.get(field))
        allowed.update(story_numbers(budget_engine.channel_why(entry, avg)))
    n_locs = len(plan.get("locations") or [])
    n_roles = len(plan.get("target_roles") or plan.get("roles") or [])
    _add(allowed, max(n_locs - client_story._MAX_MARKETS, 0))
    _add(allowed, max(n_roles - client_story._MAX_ROLES, 0))
    _add(allowed, len(plan.get("excluded_publishers") or []))
    _add(allowed, plan.get("enterprise_setup_sla_days"))
    _add(allowed, plan.get("campaign_weeks"))
    allowed.update(story_numbers(str(plan.get("campaign_duration") or "")))
    return allowed


def untraceable(plan: dict, text: str) -> list:
    allowed = plan_numbers(plan)
    return [n for n in story_numbers(text) if n not in allowed]


PLAN_KEYS = ["us_rn", "uk", "india", "mixed", "exclusions", "enterprise"]


@pytest.mark.parametrize("key", PLAN_KEYS)
def test_story_present_and_sized(plans, key):
    story = client_story.build_client_story(plans[key])
    assert story["ok"], story
    assert [s["key"] for s in story["sections"]] == [
        "what_we_will_do",
        "what_to_expect",
        "pacing",
        "needs",
        "next_steps",
    ]
    assert 120 <= story["word_count"] <= 280, story["word_count"]
    assert story["text"].startswith("What we will do")


@pytest.mark.parametrize("key", PLAN_KEYS)
def test_every_number_traces_to_the_plan(plans, key):
    story = client_story.build_client_story(plans[key])
    assert untraceable(plans[key], story["text"]) == []


def test_traceability_check_can_fail(plans):
    """Planted defect: an invented figure must be caught by the same check."""
    story = client_story.build_client_story(plans["us_rn"])
    planted = story["text"] + " We expect 98,765 applications."
    assert untraceable(plans["us_rn"], planted) == [98765.0]


_FORBIDDEN = [
    re.compile(r"(?<!\w)" + re.escape(p) + r"(?!\w)", re.I)
    for p in client_story.FORBIDDEN_PHRASES
]


def forbidden_hits(text: str) -> list:
    return [p.pattern for p in _FORBIDDEN if p.search(text)]


@pytest.mark.parametrize("key", PLAN_KEYS)
def test_no_superlatives_or_promises(plans, key):
    assert forbidden_hits(client_story.build_client_story(plans[key])["text"]) == []


def test_forbidden_phrase_check_can_fail():
    for phrase in (
        "guaranteed results",
        "the best channel",
        "ranked #1",
        "we will hire 40",
    ):
        assert forbidden_hits(phrase), phrase
    assert forbidden_hits("modelled estimates, not commitments") == []


def test_no_local_benchmark_market_carries_no_hire_figure(plans):
    text = client_story.build_client_story(plans["uk"])["text"]
    assert plan_projection.NO_LOCAL_HIRES_NOTE in text
    low = text.lower()
    assert "per hire" not in low and "estimated hires" not in low
    assert not re.search(r"\d[\d,\u2013-]*\s+hires?\b", low)
    # applications and cost per application still show, in the plan's own currency
    assert "applications at" in text and "\u00a3" in text and "$" not in text


def test_market_with_a_local_benchmark_keeps_its_own_currency(plans):
    plan = plans["india"]
    hp = plan["_budget_allocation"]["headline_projection"]
    assert not hp["assumptions"]["hires_suppressed"]
    text = client_story.build_client_story(plan)["text"]
    assert (
        "Estimated hires: " + plan_projection.range_text(hp["hires"], "count", "\u20b9")
        in text
    )
    assert "\u20b9" in text and "$" not in text


def test_us_plan_states_hires_with_their_basis(plans):
    story = client_story.build_client_story(plans["us_rn"])
    text = story["text"]
    hp = plans["us_rn"]["_budget_allocation"]["headline_projection"]
    expected = plan_projection.range_text(hp["hires"], "count", "$")
    assert f"Estimated hires: {expected}" in text
    assert "no client hire history provided" in text
    assert "not commitments" in text


def test_mixed_market_states_its_basis(plans):
    text = client_story.build_client_story(plans["mixed"])["text"]
    assert "Estimated hires" in text
    assert plan_projection.MIXED_MARKETS_LABEL in text


def test_exclusions_are_counted_never_named(plans):
    story = client_story.build_client_story(plans["exclusions"])
    assert "exclusion list (2 boards)" in story["text"]
    for name in plans["exclusions"]["excluded_publishers"]:
        assert name.lower() not in story["text"].lower()


def test_enterprise_plan_names_its_own_settings(plans):
    text = client_story.build_client_story(plans["enterprise"])["text"]
    assert "Example ATS as the system of record" in text
    assert "(Friday)" in text and "5-day setup window" in text
    assert "(and 2 more)" in text  # 5 markets, 3 named
    # a standard plan asks for the system of record instead of inventing one
    std = client_story.build_client_story(plans["us_rn"])["text"]
    assert "Please name the system of record" in std and "Friday" not in std


def test_pacing_months_come_from_the_activation_calendar(plans):
    plan = plans["us_rn"]
    timeline = plan["_gold_standard"]["activation_calendar"]["timeline"]
    high = [m["month_name"] for m in timeline if m["hiring_intensity"] == "high"]
    text = client_story.build_client_story(plan)["text"]
    for month in high[:3]:
        assert month in text
    assert "starting in " + timeline[0]["month_name"] in text


def test_deterministic_and_never_raises(plans):
    a = client_story.build_client_story(plans["us_rn"])
    b = client_story.build_client_story(plans["us_rn"])
    assert a == b
    for bad in (
        None,
        {},
        {"_budget_allocation": {}},
        {"_budget_allocation": {"x": 1}},
        "str",
    ):
        assert client_story.build_client_story(bad)["ok"] is False


def test_remote_is_not_phrased_as_a_place():
    plan = {
        "client_name": "Probe Remote",
        "locations": ["Remote", "Austin, TX"],
        "roles": ["Software Engineer"],
    }
    rows = client_story._markets(plan)
    assert rows == ["Austin, TX"] and client_story._has_remote(plan)


def test_stray_enterprise_fields_do_not_change_a_standard_story(plans):
    plan = dict(
        plans["us_rn"],
        enterprise_system_of_record="Example ATS",
        enterprise_setup_sla_days=5,
    )
    assert client_story.build_client_story(plan) == client_story.build_client_story(
        plans["us_rn"]
    )


# --------------------------------------------------------------------------
# outputs: workbook sheet, deck slide, results summary
# --------------------------------------------------------------------------
def _workbook(plan):
    raw = excel_v2.generate_excel_v2(dict(plan), load_kb_fn=load_knowledge_base)
    if isinstance(raw, tuple):
        raw = raw[0]
    return load_workbook(io.BytesIO(raw))


def _deck(plan):
    return Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(plan))))


def _sheet_text(ws) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)


def _slide_text(slide) -> str:
    return "\n".join(
        sh.text_frame.text
        for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text
    )


@pytest.mark.parametrize("key", ["us_rn", "uk", "exclusions"])
def test_workbook_sheet_sits_after_the_executive_summary(plans, key):
    wb = _workbook(plans[key])
    assert wb.sheetnames[:2] == ["Executive Summary", "Client Story"]
    story = client_story.build_client_story(plans[key])
    text = _sheet_text(wb["Client Story"])
    for sec in story["sections"]:
        assert sec["heading"] in text and sec["text"] in text
    assert story["text"] in text  # the copy-friendly block


def test_workbook_story_for_suppressed_market_has_no_hire_figure(plans):
    text = _sheet_text(_workbook(plans["uk"])["Client Story"]).lower()
    assert "per hire" not in text and "estimated hires" not in text
    assert plan_projection.NO_LOCAL_HIRES_NOTE.lower() in text


@pytest.mark.parametrize("key", ["us_rn", "uk", "exclusions", "enterprise"])
def test_deck_slide_after_the_summary_with_zero_collisions(plans, key):
    sys.path.insert(0, os.path.join(PROJECT_ROOT, "scripts", "deck_qa"))
    import collide

    prs = _deck(plans[key])
    slide = prs.slides[2]
    text = _slide_text(slide)
    story = client_story.build_client_story(plans[key])
    assert "Client Story" in text or "CLIENT STORY" in text.upper()
    for sec in story["sections"]:
        assert sec["text"] in text
    assert collide.analyze_slide(slide, 3) == []
    if key == "exclusions":
        for name in plans[key]["excluded_publishers"]:
            assert name.lower() not in text.lower()
    if key == "uk":
        assert "per hire" not in text.lower()


def test_results_summary_carries_the_story(plans):
    import app

    summary = app._plan_results_summary(plans["us_rn"])
    story = client_story.build_client_story(plans["us_rn"])
    assert summary["client_story"]["text"] == story["text"]
    assert [s["key"] for s in summary["client_story"]["sections"]] == [
        s["key"] for s in story["sections"]
    ]
    assert "client_story" not in app._plan_results_summary({"_budget_allocation": {}})


def test_results_screen_has_an_accessible_copy_button():
    path = os.path.join(
        PROJECT_ROOT, "templates", "partials", "index", "body_app_js.html"
    )
    src = open(path, encoding="utf-8").read()
    assert (
        'type="button" class="plan-story-btn" id="copyClientStoryBtn">Copy client story'
        in src
    )
    handler = src[src.index('getElementById("copyClientStoryBtn")') :]
    handler = handler[: handler.index("Render Sankey")]
    assert "navigator.clipboard.writeText" in handler and "showToast(" in handler
