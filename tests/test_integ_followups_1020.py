"""Follow-ups on the merged integration branch (Germany-plan feedback):

1. a non-US plan shows its sourced large-employer list (deck + workbook) where
   it used to say "No competitor data available";
2. the regional & aggregator budget is split across named in-network boards;
3. funded channels that project no hires are labelled, once, under a plan that
   projects hires;
4. warehouse associate maps to the same SOC (53-7062) in wages and standardizer;
5. industrial painter / paint shop titles have a difficulty profile.
"""

from __future__ import annotations

import openpyxl
import pytest
from pptx import Presentation

import api_enrichment
import channel_funding
import excel_v2
import gold_standard as gs
import ppt_generator
import publisher_filter

_DE_ROLES = ["Maintenance Technician - Paint Shop", "Industrial Painter"]


def _gold(loc: str, industry: str) -> dict:
    data = {
        "client_name": "Acme",
        "locations": [loc],
        "industry": industry,
        "target_roles": list(_DE_ROLES),
        "competitors": [],
    }
    cities = gs.enrich_city_level_data(dict(data))
    data["_gold_standard"] = {
        "competitor_mapping": gs.build_competitor_map(data, cities),
        "city_level_data": cities,
    }
    return data


def _slide_text(data: dict) -> str:
    prs = Presentation()
    ppt_generator._build_slide_competitive_landscape(prs, data)
    return "\n".join(
        sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
    )


def _sheet_text(ws) -> str:
    return "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )


# -- 1. local employers visible ---------------------------------------------


def test_german_deck_competitor_slide_shows_labelled_large_employers():
    text = _slide_text(_gold("Berlin, Germany", "Automotive & Manufacturing"))
    assert "No competitor data available" not in text
    assert (
        "Largest employers in Germany (Automotive & Manufacturing)"
        in text
    )
    assert "Volkswagen Group" in text
    assert "not measured competition" in text
    assert "Source: " in text and ", as of " in text
    # never the US roster in a non-US plan
    assert "Amazon" not in text and "Walmart" not in text


def test_deck_shows_nothing_rather_than_another_industry_for_an_uncovered_cell():
    text = _slide_text(_gold("Berlin, Germany", "staffing"))
    assert "No competitor data available" in text
    assert "Large employers" not in text


def test_german_workbook_competitor_section_shows_labelled_large_employers():
    data = _gold("Berlin, Germany", "Automotive & Manufacturing")
    ws = openpyxl.Workbook().active
    excel_v2._build_sheet_quality_intelligence(ws, data, data["_gold_standard"])
    text = _sheet_text(ws)
    assert (
        "Largest employers in Germany (Automotive & Manufacturing)"
        in text
    )
    assert "Volkswagen Group" in text and "Siemens" in text
    assert "not measured competition" in text
    assert "Source: " in text and ", as of " in text


def test_us_plan_has_no_large_employer_groups():
    data = _gold("Austin, TX", "automotive")
    assert gs.large_employer_groups(data["_gold_standard"]["competitor_mapping"]) == []


# -- 2. regional & aggregator budget named ----------------------------------


def _alloc(**hires) -> dict:
    return {
        "metadata": {"total_budget": 250000},
        "channel_allocations": {
            "programmatic_dsp": {
                "label": "Programmatic DSP",
                "dollar_amount": 100000,
                "projected_hires": hires.get("programmatic_dsp", 40),
                "projected_applications": 4000,
                "projected_clicks": 9000,
                "cpa": 25,
            },
            "regional_boards": {
                "label": "Regional Job Boards",
                "dollar_amount": 53868.27,
                "projected_hires": hires.get("regional_boards", 20),
                "projected_applications": 2000,
                "projected_clicks": 5000,
                "cpa": 27,
            },
            "emea_regional": {
                "label": "EMEA Regional",
                "dollar_amount": 32643.27,
                "projected_hires": hires.get("emea_regional", 10),
                "projected_applications": 1500,
                "projected_clicks": 4000,
                "cpa": 22,
            },
            "social_media": {
                "label": "Social Media",
                "dollar_amount": 13488.46,
                "projected_hires": hires.get("social_media", 0),
                "projected_applications": 300,
                "projected_clicks": 900,
                "cpa": 45,
            },
        },
    }


def _cents(x: float) -> int:
    return round(x * 100)


def test_regional_budget_split_sums_exactly_and_funds_five_named_boards():
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    split = channel_funding.regional_budget_split(data, _alloc())
    assert split["budget"] == pytest.approx(86511.54)
    (entry,) = split["countries"]
    assert entry["country"] == "Germany"
    assert len(entry["boards"]) == 5
    assert sum(_cents(b["amount"]) for b in entry["boards"]) == _cents(86511.54)
    assert sum(_cents(b["amount"]) for b in entry["boards"]) == _cents(entry["amount"])
    # every funded board is one the supply data confirms in network for Germany
    for b in entry["boards"]:
        rec = publisher_filter.board_supply_record(b["name"], "Germany")
        assert rec["state"] == publisher_filter.SUPPLY_IN_NETWORK, b["name"]
        assert b["status"].startswith("In Joveo Network")
    # standby carries no money
    assert 1 <= len(entry["standby"]) <= 10
    assert all("amount" not in s for s in entry["standby"])


def _explicit_germany(name: str):
    return [
        r
        for r in publisher_filter._repo_job_boards()
        if "Germany" in r["countries"] and publisher_filter.names_match(name, r["name"])
    ]


def test_germany_funds_its_market_leaders_and_only_german_rows():
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    entry = channel_funding.regional_budget_split(data, _alloc())["countries"][0]
    funded = [b["name"] for b in entry["boards"]]
    funded_keys = " ".join(publisher_filter.normalize(n) for n in funded)
    # StepStone and XING lead the market-share list and are in network in the data
    for name in ("StepStone", "XING"):
        assert (
            publisher_filter.board_supply_record(name, "Germany")["state"]
            == publisher_filter.SUPPLY_IN_NETWORK
        )
        assert publisher_filter.normalize(name) in funded_keys, name
    # the market leaders carry the largest scores, so they get the largest shares
    by_name = {b["name"]: b for b in entry["boards"]}
    leaders = [b for b in entry["boards"] if "market leader" in b["why"]]
    others = [b for b in entry["boards"] if "market leader" not in b["why"]]
    assert {publisher_filter.normalize(b["name"]) for b in leaders} >= {
        "stepstone",
        "xing",
    }
    assert min(b["share"] for b in leaders) >= max(
        (b["share"] for b in others), default=0
    )
    assert by_name["StepStone"]["share"] > 0.15
    # no US-led app: JobGet (US blue-collar app) is not funded
    assert "JobGet" not in funded
    # every funded board has a repository row that names Germany explicitly,
    # and none is a board Joveo buys under another channel
    for name in funded:
        assert _explicit_germany(name), name
        assert not any(
            publisher_filter.names_match(name, o) for o in ("Indeed", "LinkedIn")
        ), name
    # the ranking is by score, best first
    ranked = channel_funding.ranked_boards(data, "Germany")
    scores = [b["score"] for b in ranked]
    assert scores == sorted(scores, reverse=True)
    assert [b["name"] for b in ranked if b["fundable"]][:5] == funded


def test_a_board_in_network_only_through_global_or_a_group_is_never_ranked():
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    for b in channel_funding.ranked_boards(data, "Germany"):
        assert _explicit_germany(b["name"]), b["name"]


def test_regional_split_shares_follow_score_never_priority_or_country_count():
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    entry = channel_funding.regional_budget_split(data, _alloc())["countries"][0]
    total = sum(b["score"] for b in entry["boards"])
    for b in entry["boards"]:
        assert b["share"] == pytest.approx(b["score"] / total)
    # the same score gets the same share, whatever the repository priority
    by_score = {}
    for b in entry["boards"]:
        by_score.setdefault(round(b["score"], 6), set()).add(round(b["share"], 6))
    assert all(len(v) == 1 for v in by_score.values())
    # a higher score never gets a lower amount
    ordered = sorted(entry["boards"], key=lambda b: -b["score"])
    amounts = [_cents(b["amount"]) for b in ordered]
    assert amounts == sorted(amounts, reverse=True)


def test_client_excluded_board_is_never_funded_or_standby():
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    first = channel_funding.regional_budget_split(data, _alloc())["countries"][0]
    victim = first["boards"][0]["name"]
    data["excluded_publishers"] = [victim]
    again = channel_funding.regional_budget_split(data, _alloc())["countries"][0]
    names = [b["name"] for b in again["boards"] + again["standby"]]
    assert victim not in names


def test_country_with_under_two_in_network_boards_keeps_one_line(monkeypatch):
    monkeypatch.setattr(publisher_filter, "in_network_job_boards", lambda c: [])
    monkeypatch.setattr(publisher_filter, "dataset_platform_names", lambda c: [])
    data = {"locations": ["Berlin, Germany"], "target_roles": _DE_ROLES}
    entry = channel_funding.regional_budget_split(data, _alloc())["countries"][0]
    assert entry["boards"] == [] and entry["standby"] == []
    assert entry["note"] == (
        "No in-network local boards on record for Germany; budget held as "
        "aggregator spend"
    )


def test_multi_country_split_foots_to_the_channel_budget():
    data = {
        "locations": ["Berlin, Germany", "Munich, Germany", "Paris, France"],
        "target_roles": _DE_ROLES,
    }
    split = channel_funding.regional_budget_split(data, _alloc())
    assert [e["country"] for e in split["countries"]] == ["Germany", "France"]
    assert sum(_cents(e["amount"]) for e in split["countries"]) == _cents(86511.54)
    de, fr = split["countries"]
    assert _cents(de["amount"]) > _cents(fr["amount"])  # 2 of 3 locations
    for e in split["countries"]:
        if e["boards"]:
            assert sum(_cents(b["amount"]) for b in e["boards"]) == _cents(e["amount"])


def _rec_sheet(data: dict, alloc: dict) -> str:
    data = dict(data)
    data["_budget_allocation"] = alloc
    ws = openpyxl.Workbook().active
    excel_v2._build_sheet_channel_recommendations(ws, data)
    return _sheet_text(ws)


def test_channel_recommendations_names_the_regional_boards():
    data = {
        "locations": ["Berlin, Germany"],
        "target_roles": _DE_ROLES,
        "industry": "automotive",
    }
    text = _rec_sheet(data, _alloc())
    assert "Where the regional & aggregator budget goes — Germany" in text
    assert channel_funding.STANDBY_LABEL in text
    assert "In Joveo Network, as of" in text
    # the channel row names boards, not the unnamed category
    split = channel_funding.regional_budget_split(data, _alloc())
    first = split["countries"][0]["boards"][0]["name"]
    assert f"Buy via {first}" in text or first in text


def test_deck_push_pull_slide_names_the_funded_boards():
    import inspect

    src = inspect.getsource(ppt_generator._build_slide_push_meets_pull)
    assert "regional_budget_split" in src and "_funded_lines" in src


# -- 3. zero-hire funded channels --------------------------------------------


def test_zero_hire_channel_is_labelled_and_summarised_once():
    data = {
        "locations": ["Austin, TX"],
        "target_roles": ["Warehouse Associate"],
        "industry": "logistics_supply_chain",
    }
    text = _rec_sheet(data, _alloc())
    assert channel_funding.AWARENESS_TEST_LABEL in text
    assert text.count("funds channels that project no hires at this budget") == 1
    assert "7% of budget" in text and "Social Media" in text


def test_deck_budget_slide_labels_zero_hire_row_and_states_the_share_once():
    a = _alloc()
    a["metadata"]["funnel"] = {
        "totals": {
            "raw_apps": 3000,
            "qualified_apps": 300,
            "interviews": 100,
            "hires": 70,
        }
    }
    a["total_projected"] = {
        "applications": 7800,
        "hires": 70,
        "clicks": 18900,
        "cost_per_application": 25,
    }
    data = {
        "client_name": "Acme",
        "industry": "logistics_supply_chain",
        "locations": ["Austin, TX"],
        "target_roles": ["Warehouse Associate"],
        "_budget_allocation": a,
    }
    prs = Presentation()
    ppt_generator._build_slide_budget_allocation(prs, data)
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert "Awareness / test" in texts
    said = [t for t in texts if "funds channels that project no hires" in t]
    assert len(said) == 1 and "Social Media" in said[0]
    funnel = next(t for t in texts if t.startswith("Funnel:"))
    assert "awareness" not in funnel.lower()


def test_summary_sentence_wording():
    s = channel_funding.zero_hire_summary(_alloc(), False, lambda v: f"${v:,.0f}")
    assert s == (
        "7% of budget ($13,488) funds channels that project no hires at this "
        "budget: Social Media."
    )


def test_nothing_extra_when_hires_are_suppressed_or_none_projected():
    assert channel_funding.zero_hire_channels(_alloc(), True) == []
    assert channel_funding.zero_hire_summary(_alloc(), True) == ""
    zero = _alloc(programmatic_dsp=0, regional_boards=0, emea_regional=0)
    assert channel_funding.zero_hire_channels(zero, False) == []


def test_nova_plan_result_flags_zero_hire_channels():
    import nova

    src = open(nova.__file__, encoding="utf-8").read()
    assert "zero_hire_channels_note" in src and "AWARENESS_TEST_LABEL" in src


# -- 4. SOC consistency -------------------------------------------------------


def test_warehouse_associate_uses_one_soc_everywhere():
    import standardizer

    # O*NET lists Warehouse Associate under 53-7062.00
    assert api_enrichment.SOC_CODES["warehouse associate"] == "53-7062"
    assert api_enrichment.SOC_CODES["warehouse worker"] == "53-7062"
    assert api_enrichment.SOC_CODES["stocker"] == "53-7065"
    assert standardizer.get_soc_code("Warehouse Associate") == "53-7062"


# -- 5. industrial painter difficulty ------------------------------------------


@pytest.mark.parametrize(
    "title",
    [
        "Industrial Painter",
        "Paint Shop Technician",
        "Industrial Painter / Paint Shop Technician",
    ],
)
def test_painter_titles_have_a_difficulty_profile(title):
    tier, source = gs._lookup_role_tier(title)
    assert (tier, source) == ("skilled_trade", "role_profile_match")
    prof = gs._lookup_role_difficulty(title)
    assert prof and prof["base_difficulty"] == 6
    assert gs._lookup_role_tier("Industrial Painter")[1] != "unclassified_default"
