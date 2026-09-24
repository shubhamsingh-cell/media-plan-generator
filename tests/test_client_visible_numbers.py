"""Numbers a client reads must be internally consistent and honest.

Three defects the 2026-09 quality sweep found on real generated bundles, each
now pinned:

  1. Deck slide 6 "Total" row must FOOT to the "Total Investment" tile above
     it. It printed 109% / $981,000 on a $900,000 plan (a display channel the
     engine never funded was given a dollar figure from its static profile
     percentage) and $8,019,032 as "100%" of a $12.5M plan (channels the
     engine DID fund but the display profile omits were dropped from rows and
     Total alike).
  2. A local-currency CPC basis must be ONE unit. India's basis covers only
     job_board + social; every other category fell through to the USD cascade
     and sat in the same ₹ column -- ₹0.62 beside ₹13.37 -- so a ~83x unit
     error steered ~80% of an ₹18M budget to the "cheap" channels.
  3. A plan projecting zero hires has an UNDEFINED cost per hire, not $0. The
     workbook hero printed "Cost / Hire: $0" and the deck multiplied the 0.0
     into "scaling path: ~$0 additional".

Offline: drives the real budget engine and generators. No network / LLM.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import budget_engine  # noqa: E402
import display_format  # noqa: E402
import intl_benchmark_lookup  # noqa: E402
import ppt_generator as ppt  # noqa: E402
from pptx import Presentation  # noqa: E402


def _money(text: str) -> float:
    return float(re.sub(r"[^\d.]", "", text or "") or 0)


def _slide6_total_row(pptx_bytes: bytes) -> tuple[str, list[str]]:
    """(Total Investment tile value, the 6 cells of the Total row)."""
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide = prs.slides[5]
    texts = [
        sh.text_frame.text.strip()
        for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
    tile = next(
        (texts[i - 1] for i, t in enumerate(texts) if t == "Total Investment" and i),
        "",
    )
    idx = texts.index("Total")
    return tile, texts[idx : idx + 6]


def _multi_country_plan() -> dict:
    """The sweep's q_multi_country shape: engine funds apac/emea regional
    channels the static display profile for technology omits."""
    locations = [
        {"city": "Austin", "state": "TX", "country": "United States"},
        {"city": "London", "state": "", "country": "United Kingdom"},
        {"city": "Bangalore", "state": "", "country": "India"},
        {"city": "Berlin", "state": "", "country": "Germany"},
    ]
    channel_pcts = {
        "programmatic_dsp": 22,
        "global_boards": 18,
        "niche_boards": 14,
        "social_media": 12,
        "regional_boards": 10,
        "employer_branding": 6,
        "apac_regional": 9,
        "emea_regional": 9,
    }
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=900_000,
        roles=[
            {"title": "Software Engineer", "count": 60, "tier": "senior"},
            {"title": "Data Engineer", "count": 30, "tier": "mid"},
            {"title": "Product Manager", "count": 12, "tier": "senior"},
        ],
        locations=locations,
        industry="technology_engineering",
        channel_percentages=channel_pcts,
        collar_type="white",
        campaign_start_month=5,
    )
    return {
        "client_name": "Helio Global Tech",
        "industry": "technology_engineering",
        "locations": [f"{l['city']}, {l['country']}" for l in locations],
        "roles": ["Software Engineer", "Data Engineer", "Product Manager"],
        "budget": "$900,000",
        "channel_categories": {k: True for k in channel_pcts},
        "_budget_allocation": alloc,
    }


class TestSlide6TotalFoots:
    def test_total_row_equals_total_investment_tile(self):
        data = _multi_country_plan()
        tile, row = _slide6_total_row(ppt.generate_pptx(data))
        budget = data["_budget_allocation"]["metadata"]["total_budget"]
        assert row[0] == "Total"
        assert row[1] == "100%", f"Total row percentage is {row[1]!r}"
        assert abs(_money(row[2]) - budget) < 1.0, (
            f"Total row investment {row[2]!r} does not foot to the plan "
            f"budget {budget:,.0f}"
        )
        assert abs(_money(tile) - _money(row[2])) < 1.0, (
            f"'Total Investment' tile {tile!r} contradicts Total row {row[2]!r} "
            "on the same slide"
        )

    def test_every_funded_channel_is_in_the_table(self):
        """Channels the engine funds but the display profile omits must not
        vanish from the rows -- that is where the missing 36% went."""
        data = _multi_country_plan()
        alloc = data["_budget_allocation"]["channel_allocations"]
        funded_dollars = sum(
            c.get("dollar_amount") or 0 for c in alloc.values() if isinstance(c, dict)
        )
        _tile, row = _slide6_total_row(ppt.generate_pptx(data))
        assert abs(_money(row[2]) - funded_dollars) < 1.0


class TestLocalCurrencyBasisIsOneUnit:
    def test_local_basis_carries_the_datasets_own_rate(self):
        basis = intl_benchmark_lookup.get_locale_cpc_basis(["India"], "INR")
        assert basis and basis["basis"] == "local"
        assert basis.get("currency") == "INR"
        rate = basis.get("usd_per_local")
        assert rate and 0.005 < rate < 0.05, f"INR usd_per_local={rate!r}"

    def test_usd_blend_basis_carries_no_rate(self):
        """Only a local basis needs -- or may claim -- a single conversion."""
        basis = intl_benchmark_lookup.get_locale_cpc_basis(
            ["United Kingdom", "Australia"], "GBP"
        )
        assert basis and basis["basis"] == "usd_blend"
        assert "usd_per_local" not in basis

    def test_inr_plan_cpcs_share_one_unit(self):
        """A category outside India's local coverage (programmatic) must be
        converted into ₹, not left as a ~$0.62 USD constant beside ₹13.37."""
        alloc = budget_engine.calculate_budget_allocation(
            total_budget=18_000_000,
            roles=[{"title": "Delivery Driver", "count": 600, "tier": "entry"}],
            locations=[{"city": "Mumbai", "state": "", "country": "India"}],
            industry="logistics_supply_chain",
            channel_percentages={
                "global_boards": 30,
                "programmatic_dsp": 30,
                "social_media": 20,
                "niche_boards": 20,
            },
            collar_type="blue",
            campaign_start_month=8,
            plan_currency="INR",
        )
        chans = alloc["channel_allocations"]
        cpcs = {k: float(c.get("cpc") or 0) for k, c in chans.items()}
        assert all(v >= 1.0 for v in cpcs.values()), (
            f"sub-₹1 CPC in an INR plan -- a USD figure leaked into the ₹ "
            f"column: {cpcs}"
        )
        # (No spread heuristic here: niche boards legitimately cost ~30x a mass
        # job board per click in USD too, so ₹394 beside ₹13 is a correct
        # conversion, not a leak. The leak signature is a sub-₹1 value.)
        converted = [
            k
            for k, c in chans.items()
            if str(c.get("cpc_source", "")).endswith("->INR")
        ]
        assert converted, (
            "no channel records a USD->INR conversion; the fallthrough path "
            f"was not exercised: {[c.get('cpc_source') for c in chans.values()]}"
        )
        # Loose sanity bound only -- NOT the leak guard (that's the cpc >= 1.0
        # assertion above; a currency-leak simulation on this exact scenario
        # produced ~53% concentration, well under this cap either side of the
        # fix, so a leak would slip past a tighter cap too). F5 FIX
        # (2026-09-24) made industry_avg_cph currency-correct for roi_score
        # too, so a *legitimate* INR plan now properly differentiates ROI
        # (global_boards' real ₹13.37 CPC vs. ₹52-394 for the rest) and
        # rebalance_low_roi_channels correctly concentrates spend there --
        # observed ~64%. This just guards against a single channel taking
        # the near-entire budget (e.g. a >=90% collapse would still trip it).
        total = sum(c.get("dollar_amount") or 0 for c in chans.values())
        shares = {k: (c.get("dollar_amount") or 0) / total for k, c in chans.items()}
        assert max(shares.values()) < 0.7, shares


class TestZeroHireCostPerHire:
    @staticmethod
    def _zero_hire_plan() -> dict:
        alloc = budget_engine.calculate_budget_allocation(
            total_budget=1_200,
            roles=[{"title": "Barista", "count": 2, "tier": "entry"}],
            locations=[{"city": "Boise", "state": "ID", "country": "United States"}],
            industry="retail_hospitality",
            channel_percentages={"global_boards": 50, "social_media": 50},
            collar_type="blue",
            campaign_start_month=1,
        )
        return {
            "client_name": "Cedar & Co Coffee",
            "industry": "retail_hospitality",
            "locations": ["Boise, ID"],
            "roles": ["Barista"],
            "budget": "$1,200",
            "hire_volume": "2 hires",
            "_budget_allocation": alloc,
        }

    def test_goal_gap_reports_unknown_budget_when_cph_is_unknown(self):
        gap = display_format.goal_gap(0, 2, 0.0)
        assert gap is not None
        assert gap["projected"] == 0 and gap["goal"] == 2
        assert gap["additional_budget"] is None

    def test_goal_gap_still_computes_when_cph_is_known(self):
        gap = display_format.goal_gap(1, 4, 500.0)
        assert gap["additional_budget"] == 1500.0

    def test_deck_states_the_gap_without_a_zero_dollar_scaling_path(self):
        data = self._zero_hire_plan()
        hires = sum(
            c.get("projected_hires") or 0
            for c in data["_budget_allocation"]["channel_allocations"].values()
        )
        if hires > 0:
            pytest.skip("engine projected hires for the $1,200 plan; scenario moot")
        prs = Presentation(io.BytesIO(ppt.generate_pptx(data)))
        deck = "\n".join(
            sh.text_frame.text
            for s in prs.slides
            for sh in s.shapes
            if sh.has_text_frame
        )
        assert not re.search(
            r"~?\$0(?:\.00)?\s+additional", deck
        ), "deck still prints a $0 scaling path for a zero-hire plan"
        goal_line = next((l for l in deck.splitlines() if "Client goal" in l), "")
        assert goal_line and "projects 0" in goal_line, goal_line

    def test_workbook_hero_cost_per_hire_is_dash_not_zero(self):
        import openpyxl
        from excel_v2 import generate_excel_v2

        data = self._zero_hire_plan()
        hires = sum(
            c.get("projected_hires") or 0
            for c in data["_budget_allocation"]["channel_allocations"].values()
        )
        if hires > 0:
            pytest.skip("engine projected hires for the $1,200 plan; scenario moot")
        out = generate_excel_v2(dict(data))
        xlsx = out[0] if isinstance(out, tuple) else out
        ws = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True)[
            "Executive Summary"
        ]
        for row in ws.iter_rows(min_row=1, max_row=40):
            for cell in row:
                if cell.value == "Cost / Hire":
                    shown = ws.cell(cell.row - 1, cell.column).value
                    assert shown == "--", f"Cost / Hire hero shows {shown!r}, not '--'"
                    return
        pytest.fail("'Cost / Hire' hero label not found on Executive Summary")
