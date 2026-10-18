"""Supply freshness footers (international audit 2026-10-18, item #5).

The footers said "supply network 2026-10-17": that date is the newest single
board edit on top of the supply repository, not a refresh of anything, and
the deck footer dropped the repository's own date (2026-03-31).  A footer now
states the repository date and, separately, the latest confirmed exception, in
plain words, on the workbook and the deck.

Every expected date is read from the data files in the test, never typed.
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from typing import Any, List, Set

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402
from pptx import Presentation  # noqa: E402

import excel_v2  # noqa: E402
import plan_trust  # noqa: E402
import ppt_generator  # noqa: E402

DATA_DIR = PROJECT_ROOT / "data"
REPO = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
GROUPS = json.loads((DATA_DIR / "joveo_network_groups.json").read_text())


def _plan() -> dict:
    return {
        "client_name": "Fixture Co",
        "company_name": "Fixture Co",
        "industry": "general_entry_level",
        "budget": "$100,000",
        "locations": ["Dallas, TX"],
        "roles": ["Agent"],
        "target_roles": ["Agent"],
        "campaign_duration": "3 months",
        "hire_volume": "50",
        "work_environment": "onsite",
        "_enriched": {},
        "_synthesized": {},
        "_budget_allocation": {},
    }


def _workbook(plan: dict):
    return openpyxl.load_workbook(io.BytesIO(excel_v2.generate_excel_v2(dict(plan))))


def _cells(ws) -> List[str]:
    return [
        str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None
    ]


def _deck_texts(plan: dict) -> List[str]:
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(plan))))
    return [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


def _latest_exception_date() -> str:
    """Newest dated single-board edit on top of the repository (read from the
    three files that carry them)."""
    found: List[str] = []

    def walk(node: Any, keys: Set[str]) -> None:
        if isinstance(node, dict):
            for k, v in node.items():
                if k in keys and isinstance(v, str):
                    found.append(v[:10])
                walk(v, keys)
        elif isinstance(node, list):
            for v in node:
                walk(v, keys)

    walk(GROUPS, {"as_of"})
    walk(json.loads((DATA_DIR / "publisher_availability.json").read_text()), {"as_of"})
    walk(REPO, {"network_as_of"})
    return max(found)


# ---------------------------------------------------------------------------
# #5: footers state the repository date and, separately, the latest exception
# ---------------------------------------------------------------------------


def test_footers_say_what_each_date_is():
    info = plan_trust.data_asof({})
    repo_date = REPO["last_updated"]
    exception = _latest_exception_date()
    assert exception > repo_date  # else the two dates could not be told apart
    assert f"supply repository last updated {repo_date}" in info["headline"]
    assert f"board exceptions confirmed up to {exception}" in info["headline"]
    assert info["footer_headline"] == (
        f"Supply repository {repo_date} | latest confirmed exception {exception}"
    )
    for text in (info["headline"], info["footer_headline"]):
        assert "supply network" not in text.lower()
        assert "refresh" not in text.lower()


def test_exception_family_is_not_called_an_update_of_the_network():
    info = plan_trust.data_asof({})
    fam = next(
        f for f in info["families"] if f["label"] == plan_trust.SUPPLY_NETWORK_LABEL
    )
    assert fam["as_of"] == f"latest confirmed change {_latest_exception_date()}"
    assert "updated" not in fam["as_of"]


def test_workbook_footer_states_both_dates_in_plain_words():
    wb = _workbook(_plan())
    text = "\n".join(_cells(wb["Executive Summary"]))
    assert f"supply repository last updated {REPO['last_updated']}" in text
    assert f"board exceptions confirmed up to {_latest_exception_date()}" in text
    assert "supply network" not in text.lower()
    sources = "\n".join(_cells(wb["Sources & Confidence"]))
    assert "not a refresh of the repository" in sources
    assert "Last Refreshed" not in sources


def test_deck_footer_names_both_dates_and_stays_clear_of_the_credit():
    deck = _deck_texts(_plan())
    footer = (
        f"Supply repository {REPO['last_updated']} | "
        f"latest confirmed exception {_latest_exception_date()}"
    )
    assert footer in deck
    # measured with real Poppins advances, not a character count: the line
    # starts at 2.0in and must end before the credit text begins
    sys.path.insert(0, str(PROJECT_ROOT / "scripts" / "deck_qa"))
    import _textlayout as tl  # noqa: E402

    credit = next(t for t in deck if t.startswith("Created by"))
    credit_left = 12.78 - 0.1 - tl.adv_in(credit, 9, False)
    footer_right = 1.9 + 0.1 + tl.adv_in(footer, 9, False)
    assert footer_right < credit_left - 0.2, (footer_right, credit_left)
    assert len(footer) <= ppt_generator._ASOF_FOOTER_MAX_CHARS
