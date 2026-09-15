"""Closure tests for the last two MPG plan-output quality-audit defects.

DEFECT 1 -- currency convert-vs-declare.
The wizard has no currency field, so plan currency was resolved purely by
guessing from the location string. A client who typed "$2,000,000" against a
London office had it rendered back as "£2M": their own declaration silently
overwritten by a guess, with NO conversion applied. Proven by generating the
same plan twice with only the market changed -- 17 of 23 money figures came
back numerically identical, i.e. the generator relabels and never converts.
The guess also flipped on the ORDER of the locations list ("Dallas, London"
-> $, "London, Dallas" -> £).

  THE RULE, now enforced here:
    1. An explicit currency/currency_code field wins.
    2. Otherwise the symbol the CLIENT typed into the budget wins. Location can
       never contradict it -- only disambiguate a symbol shared by several
       codes ("$" -> CAD for a Toronto plan, "¥" -> CNY for a Shanghai plan).
    3. Only with nothing declared may the market fill the gap, and only when
       every market agrees -- otherwise USD, because anything else would make
       the answer depend on list order.
    4. Nothing is ever FX-converted. US-calibrated benchmarks stay "US$", and a
       non-USD deck states its currency basis on the surface the client reads.

DEFECT 2 -- slide 5 / slide 7 collisions.
Slide 7 ("Competitive Landscape"): a prior wave made "Counter:" cascade from
"Why:"'s measured height but left "Why:" pinned to a constant 0.3in, which
assumes the competitor NAME above it is one line. The name is 10pt bold in a
3.0in box, so any name past ~30 chars wraps and its second line printed
straight through "Why:" -- reproduced on "Universal Health Services Behavioral
Division", "Encompass Health Rehabilitation Hospital Group" and "Select
Medical Critical Illness Recovery Holdings".

Slide 1 (cover, checked because a reflow fix can push a defect one slide over):
_estimate_lines models wrapping with ONE average glyph advance, but wrapping is
decided by the words that actually fill a line. An 80-char legal name averages
0.526 em (under _AVG_CHAR_EM's 0.53, so the estimator said 2 lines) while its
line-filling words run 0.55-0.58 em and it truly needs 3. The hero was sized
for 2 and its third line printed through the Industry subtitle.

Slide 5 ("Channel Strategy & Investment") is asserted CLEAN here rather than
fixed: its benchmark table and attribution cards already measure per row and
drop trailing rows, and reproduce clean across the whole envelope below.

MEASUREMENT INDEPENDENCE: the layout assertions deliberately do NOT reuse
ppt_generator's own _measure_lines/_estimate_lines -- a generator checked with
its own ruler cannot be falsified. They lay text out here from the shipped
Poppins .ttf via PIL glyph advances and collide the resulting ink rectangles.

Runs under pytest, or standalone:
``python3 tests/test_plan_output_audit_closure.py``.
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import excel_v2  # noqa: E402
import plan_currency  # noqa: E402
import ppt_generator as ppt  # noqa: E402
from pptx import Presentation  # noqa: E402

EMU = 914400.0
FONTS_DIR = os.path.join(PROJECT_ROOT, "fonts")
_INK_TOL_IN = 0.015  # ignore sub-hairline touching


@pytest.fixture(autouse=True)
def _reset_active_currency():
    """_set_active_currency stores the plan currency in a thread-local that
    outlives a single test in the same thread."""
    yield
    try:
        ppt._currency_tls.code = "USD"
        excel_v2._currency_tls.code = "USD"
    except AttributeError:
        pass


# ---------------------------------------------------------------------------
# Independent text measurement (PIL + the shipped Poppins faces)
# ---------------------------------------------------------------------------
def _pil_font(size_pt: float, bold: bool):
    from PIL import ImageFont

    name = "Poppins-Bold.ttf" if bold else "Poppins-Regular.ttf"
    return ImageFont.truetype(os.path.join(FONTS_DIR, name), max(int(size_pt * 4), 4))


def _adv_in(text: str, size_pt: float, bold: bool) -> float:
    if not text:
        return 0.0
    return _pil_font(size_pt, bold).getlength(text) / (4.0 * 72.0)


def _wrap(text: str, size_pt: float, bold: bool, avail_in: float):
    if avail_in <= 0:
        return [text] if text else []
    lines, cur = [], ""
    for word in text.split():
        trial = (cur + " " + word).strip()
        if not cur or _adv_in(trial, size_pt, bold) <= avail_in:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines or [""]


def _ink_rects(shape):
    """Per-line glyph ink rectangles (inches, absolute) for one text frame.

    Uses real glyph bounding boxes so line LEADING is never mistaken for ink --
    the reason a naive box-overlap check reports false collisions between any
    heading and the line beneath it.
    """
    tf = shape.text_frame
    left, top = (shape.left or 0) / EMU, (shape.top or 0) / EMU
    width, height = (shape.width or 0) / EMU, (shape.height or 0) / EMU
    ml = (tf.margin_left if tf.margin_left is not None else 91440) / EMU
    mr = (tf.margin_right if tf.margin_right is not None else 91440) / EMU
    mt = (tf.margin_top if tf.margin_top is not None else 45720) / EMU
    avail = max(width - ml - mr, 0.05)
    wrap_on = tf.word_wrap is not False

    plan, total = [], 0.0
    for para in tf.paragraphs:
        runs = para.runs
        text = "".join(r.text for r in runs)
        size = next((r.font.size.pt for r in runs if r.font.size), None)
        if size is None:
            size = para.font.size.pt if para.font.size else 18.0
        bold = any(r.font.bold for r in runs)
        spacing = para.line_spacing
        line_h = size * (spacing if isinstance(spacing, float) else 1.0) * 1.2 / 72.0
        before = (para.space_before.pt if para.space_before else 0) / 72.0
        after = (para.space_after.pt if para.space_after else 0) / 72.0
        lines = (
            _wrap(text, size, bold, avail) if wrap_on else ([text] if text else [""])
        )
        plan.append((para, lines, size, bold, line_h, before, after))
        total += before + after + line_h * len(lines)

    anchor = str(tf.vertical_anchor or "")
    y = top + mt
    if "MIDDLE" in anchor:
        y = top + max((height - total) / 2.0, 0)
    elif "BOTTOM" in anchor:
        y = top + max(height - total, 0)

    rects = []
    for para, lines, size, bold, line_h, before, after in plan:
        y += before
        align = str(para.alignment or "")
        for line in lines:
            if line.strip():
                bbox = _pil_font(size, bold).getbbox(line)
                bbox_in = tuple(v / (4.0 * 72.0) for v in bbox)
                w = _adv_in(line, size, bold)
                x = left + ml
                if "CENTER" in align:
                    x = left + ml + max((avail - w) / 2.0, 0)
                elif "RIGHT" in align:
                    x = left + ml + max(avail - w, 0)
                rects.append(
                    (
                        x + bbox_in[0],
                        y + bbox_in[1],
                        x + bbox_in[2],
                        y + bbox_in[3],
                        line,
                        size,
                    )
                )
            y += line_h
        y += after
    return rects


def _ink_collisions(slide):
    """Every pair of text frames on ``slide`` whose painted glyphs overlap."""
    frames = []
    for idx, shape in enumerate(slide.shapes):
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        rects = _ink_rects(shape)
        if rects:
            frames.append((idx, rects))
    hits = []
    for i in range(len(frames)):
        for j in range(i + 1, len(frames)):
            for ra in frames[i][1]:
                for rb in frames[j][1]:
                    ox = min(ra[2], rb[2]) - max(ra[0], rb[0])
                    oy = min(ra[3], rb[3]) - max(ra[1], rb[1])
                    if ox > _INK_TOL_IN and oy > _INK_TOL_IN:
                        hits.append((ra[4], rb[4], round(ox, 3), round(oy, 3)))
    return hits


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
LONG_COMPETITORS = [
    "Universal Health Services Behavioral Division",
    "Encompass Health Rehabilitation Hospital Group",
    "Select Medical Critical Illness Recovery Holdings",
    "Kindred Healthcare Transitional Care Systems",
]
LONG_CLIENT = (
    "Consolidated Transcontinental Healthcare & Rehabilitation Partners International"
)


def _plan(**over):
    data = {
        "client_name": "Mercy Health Partners",
        "industry": "healthcare",
        "industry_label": "Healthcare",
        "budget": "$150,000",
        "budget_period": "campaign",
        "campaign_duration": "3 months",
        "campaign_start_month": 9,
        "hire_volume": "90 hires",
        "work_environment": "onsite",
        "locations": [{"city": "Dallas", "state": "TX", "country": "United States"}],
        "roles": [{"title": "Registered Nurse", "count": 40, "tier": "mid"}],
        "target_roles": ["Registered Nurse"],
        "campaign_goals": ["volume_hiring"],
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "niche_boards": True,
            "social_media": True,
        },
        "experience_level": "mid",
    }
    data.update(over)
    return data


def _with_competitor_intel(names):
    return {
        "competitive_intelligence": {
            "company_profile": {
                "name": "Mercy Health Partners",
                "industry_sector": "Healthcare",
                "employee_count": "25,000-50,000",
            },
            "competitors": {
                n: {"employee_count": "10,000+", "hiring_velocity": "high"}
                for n in names
            },
        }
    }


def _deck(data):
    return Presentation(io.BytesIO(ppt.generate_pptx(data)))


# ===========================================================================
# DEFECT 1 -- currency: declare, never convert
# ===========================================================================
class TestCurrencyDeclareNotConvert:
    def test_typed_dollar_is_not_overridden_by_a_london_office(self):
        """The headline regression: the client's own symbol outranks the guess."""
        data = _plan(budget="$2,000,000", locations=["London, United Kingdom"])
        assert ppt._plan_currency_code(data) == "USD"
        assert data["_currency_basis"] == "declared"

    def test_typed_pound_is_honoured(self):
        data = _plan(budget="£2,000,000", locations=["London, United Kingdom"])
        assert ppt._plan_currency_code(data) == "GBP"
        assert data["_currency_basis"] == "declared"

    def test_typed_pound_wins_even_against_a_us_market(self):
        data = _plan(budget="£2,000,000", locations=["Dallas, TX, United States"])
        assert ppt._plan_currency_code(data) == "GBP"

    def test_result_does_not_depend_on_location_order(self):
        """ "Dallas, London" used to give $ and "London, Dallas" £ -- same plan."""
        a = _plan(
            budget="$2,000,000",
            locations=["Dallas, TX, United States", "London, United Kingdom"],
        )
        b = _plan(
            budget="$2,000,000",
            locations=["London, United Kingdom", "Dallas, TX, United States"],
        )
        assert ppt._plan_currency_code(a) == ppt._plan_currency_code(b) == "USD"

    def test_market_may_fill_a_gap_when_nothing_was_declared(self):
        data = _plan(budget="2,000,000", locations=["London, United Kingdom"])
        assert ppt._plan_currency_code(data) == "GBP"
        assert data["_currency_basis"] == "market"

    def test_disagreeing_markets_fall_back_to_usd(self):
        data = _plan(
            budget="2,000,000",
            locations=["Dallas, TX, United States", "London, United Kingdom"],
        )
        assert ppt._plan_currency_code(data) == "USD"
        assert data["_currency_basis"] == "default"

    def test_explicit_currency_code_outranks_a_typed_symbol(self):
        data = _plan(
            budget="£2,000,000",
            locations=["London, United Kingdom"],
            currency_code="usd",
        )
        assert ppt._plan_currency_code(data) == "USD"
        assert data["_currency_basis"] == "explicit"

    def test_ambiguous_dollar_is_disambiguated_by_a_single_market(self):
        """ "$" is shared by USD/CAD/AUD/...: the symbol constrains, one
        unambiguous market picks within it -- it never overrides it."""
        data = _plan(budget="$2,000,000", locations=["Toronto, Canada"])
        assert ppt._plan_currency_code(data) == "CAD"
        unambiguous = _plan(budget="C$2,000,000", locations=["Toronto, Canada"])
        assert ppt._plan_currency_code(unambiguous) == "CAD"

    def test_ambiguous_dollar_falls_back_to_usd_when_markets_disagree(self):
        """Caught by rendering the deck, not by the linter: a London+Singapore
        plan resolved a "$" budget to SGD purely because Singapore shares the
        glyph and appeared in the list. Unqualified "$" means USD."""
        data = _plan(
            budget="$2,000,000",
            locations=["London, United Kingdom", "Singapore, Singapore"],
        )
        assert ppt._plan_currency_code(data) == "USD"
        flipped = _plan(
            budget="$2,000,000",
            locations=["Singapore, Singapore", "London, United Kingdom"],
        )
        assert ppt._plan_currency_code(flipped) == "USD"

    @pytest.mark.parametrize(
        "budget,expected",
        [
            ("£2,000,000", ("GBP",)),
            ("€2.000.000", ("EUR",)),
            ("₹200,00,000", ("INR",)),
            ("NZ$2,000,000", ("NZD",)),
            ("C$2,000,000", ("CAD",)),
            ("US$2,000,000", ("USD",)),
            ("2,000,000", ()),
            ("", ()),
        ],
    )
    def test_symbol_parsing(self, budget, expected):
        assert plan_currency.currency_codes_from_symbol(budget) == expected

    def test_longer_symbols_win_over_bare_dollar(self):
        """ "NZ$" must not be read as "$"."""
        for prefix, code in (("NZ$", "NZD"), ("HK$", "HKD"), ("MX$", "MXN")):
            assert plan_currency.currency_codes_from_symbol(f"{prefix}1,000") == (code,)

    def test_deck_and_workbook_never_disagree(self):
        """They ship in one bundle; a "$" deck beside a "£" workbook is worse
        than either alone."""
        for budget, locs in (
            ("$2,000,000", ["London, United Kingdom"]),
            ("£2,000,000", ["London, United Kingdom"]),
            ("2,000,000", ["London, United Kingdom"]),
            ("2,000,000", ["Dallas, TX, United States", "London, United Kingdom"]),
        ):
            data = _plan(budget=budget, locations=locs)
            assert ppt._plan_currency_code(dict(data)) == excel_v2._plan_currency_code(
                dict(data)
            )

    def test_no_fx_conversion_is_ever_applied(self):
        """The budget the client entered is the budget the deck states."""
        data = _plan(budget="£2,000,000", locations=["London, United Kingdom"])
        ppt._set_active_currency(data)
        assert ppt._format_budget_display("£2,000,000") == "£2M"
        assert ppt._parse_budget_number("£2,000,000") == 2_000_000

    def test_non_usd_deck_states_its_currency_basis(self):
        data = _plan(budget="£2,000,000", locations=["London, United Kingdom"])
        ppt._set_active_currency(data)
        note = ppt._currency_basis_note(data)
        assert "GBP" in note
        assert "not" in note and "FX-converted" in note
        # must fit the single 8pt line it is rendered into
        assert len(note) <= 150

    def test_inferred_currency_says_so(self):
        data = _plan(budget="2,000,000", locations=["London, United Kingdom"])
        ppt._set_active_currency(data)
        note = ppt._currency_basis_note(data)
        assert "inferred" in note.lower()
        assert len(note) <= 150

    def test_usd_plan_discloses_nothing_new(self):
        """The common case must render exactly as before."""
        data = _plan(budget="$150,000")
        ppt._set_active_currency(data)
        assert ppt._currency_basis_note(data) == ""

    def test_rendered_non_usd_deck_carries_the_note_on_the_money_page(self):
        data = _plan(
            budget="£2,000,000",
            locations=["London, United Kingdom"],
            client_name="Uber",
        )
        prs = _deck(data)
        pages = set()
        for i, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame and "FX-converted" in shape.text_frame.text:
                    pages.add(i)
        assert pages, "non-USD deck rendered without any currency-basis note"

    def test_us_calibrated_benchmarks_stay_us_dollar_on_a_gbp_plan(self):
        """A never-converted USD benchmark must never wear a bare "£"."""
        data = _plan(
            budget="£2,000,000",
            locations=["London, United Kingdom"],
            client_name="Uber",
        )
        prs = _deck(data)
        blob = "\n".join(
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        assert "US$" in blob, "USD benchmarks lost their US$ marker"


# ===========================================================================
# DEFECT 2 -- slide 5 / slide 7 collisions
# ===========================================================================
class TestSlideCollisions:
    def test_slide7_long_competitor_names_do_not_overprint_why(self):
        data = _plan(_synthesized=_with_competitor_intel(LONG_COMPETITORS))
        slide = _deck(data).slides[6]
        hits = _ink_collisions(slide)
        assert not hits, f"slide 7 text collides: {hits}"

    def test_slide7_short_competitor_names_still_clean(self):
        data = _plan(_synthesized=_with_competitor_intel(["HCA", "CVS", "Tenet"]))
        assert not _ink_collisions(_deck(data).slides[6])

    def test_cover_long_client_name_does_not_overprint_industry(self):
        data = _plan(client_name=LONG_CLIENT)
        hits = _ink_collisions(_deck(data).slides[0])
        assert not hits, f"cover text collides: {hits}"

    def test_slide5_is_clean_across_the_channel_envelope(self):
        """Slide 5 is asserted clean, not fixed -- it already measures per row."""
        for cats in (
            ["programmatic_dsp"],
            ["programmatic_dsp", "social_media"],
            [
                "programmatic_dsp",
                "global_boards",
                "niche_boards",
                "regional_boards",
                "social_media",
                "employer_branding",
                "apac_regional",
                "emea_regional",
            ],
        ):
            data = _plan(
                channel_categories={k: True for k in cats},
                locations=[
                    {"city": "London", "state": "", "country": "United Kingdom"},
                    {"city": "Singapore", "state": "", "country": "Singapore"},
                ],
            )
            hits = _ink_collisions(_deck(data).slides[4])
            assert not hits, f"slide 5 collides with {len(cats)} channels: {hits}"

    def test_slides_adjacent_to_5_and_7_are_clean(self):
        """A collision fix that reflows content can push the defect one slide
        over, so check the neighbours on the worst fixture."""
        data = _plan(
            client_name=LONG_CLIENT,
            _synthesized=_with_competitor_intel(LONG_COMPETITORS),
        )
        prs = _deck(data)
        for idx in (3, 4, 5, 6, 7):  # slides 4-8, 0-based
            if idx < len(prs.slides):
                hits = _ink_collisions(prs.slides[idx])
                assert not hits, f"slide {idx + 1} collides: {hits}"

    def test_whole_deck_is_collision_free_on_the_worst_fixture(self):
        data = _plan(
            client_name=LONG_CLIENT,
            budget="£2,000,000",
            locations=["London, United Kingdom"],
            _synthesized=_with_competitor_intel(LONG_COMPETITORS),
        )
        prs = _deck(data)
        for i, slide in enumerate(prs.slides, 1):
            hits = _ink_collisions(slide)
            assert not hits, f"slide {i} collides: {hits}"

    def test_short_client_name_keeps_the_original_cover_geometry(self):
        """The fix must not reflow the common case."""
        prs = _deck(_plan(client_name="Mercy Health Partners"))
        hero = [
            s
            for s in prs.slides[0].shapes
            if s.has_text_frame and s.text_frame.text.strip() == "Mercy Health Partners"
        ]
        assert hero, "cover hero not found"
        assert abs(hero[0].top / EMU - 3.48) < 0.01
        assert abs(hero[0].height / EMU - 1.0) < 0.01  # the old fixed 1.0in box


class TestMeasurementHelper:
    """_measure_lines is what makes the cover fix correct; pin its behaviour."""

    def test_measures_the_case_the_estimator_got_wrong(self):
        assert ppt._estimate_lines(LONG_CLIENT, 11.8, 40) == 2  # the old answer
        assert ppt._measure_lines(LONG_CLIENT, 11.8, 40, bold=True) == 3  # the truth

    def test_falls_back_when_the_font_is_unavailable(self, monkeypatch):
        monkeypatch.setattr(ppt, "_measure_font", lambda bold: None)
        assert ppt._measure_lines(
            LONG_CLIENT, 11.8, 40, bold=True
        ) == ppt._estimate_lines(LONG_CLIENT, 11.8, 40)

    def test_short_strings_are_one_line(self):
        assert ppt._measure_lines("Mercy Health Partners", 11.8, 42, bold=True) == 1
        assert ppt._measure_lines("HCA", 2.8, 10, bold=True) == 1


class TestQaGateStillBites:
    """The delivery gate had to learn the new rule; prove it did not go blind.

    bundle_qa's currency_symbol_mixing check now accepts a figure that states
    its own denomination ("£60,000 (GBP)") because a local-market salary may
    legitimately differ from the plan currency. That must not become a hole:
    the symbol has to AGREE with the code it declares.
    """

    @staticmethod
    def _run(text, plan):
        import bundle_qa

        findings: list = []
        bundle_qa._check_currency_symbol_mixing(
            [bundle_qa._TextUnit(text, "Sheet1!A1", top=1, left=1)], plan, findings
        )
        return [f for f in findings if f.get("code") == "currency_symbol_mixing"]

    def test_correctly_denominated_local_figure_is_accepted(self):
        plan = _plan(budget="$150,000", locations=["London, United Kingdom"])
        assert self._run("£60,000 - £97,500 (GBP)", plan) == []

    def test_symbol_contradicting_its_own_code_is_still_critical(self):
        plan = _plan(budget="$150,000", locations=["London, United Kingdom"])
        assert self._run("$42,000 (GBP)", plan), "self-contradictory figure not flagged"

    def test_mismatched_symbol_and_code_is_still_critical(self):
        plan = _plan(budget="$150,000", locations=["London, United Kingdom"])
        assert self._run("€50,000 (GBP)", plan), "€ declared as GBP not flagged"

    def test_unmarked_stray_symbol_is_still_critical(self):
        plan = _plan(budget="$150,000", locations=["London, United Kingdom"])
        assert self._run("£99 with no denomination", plan), "stray glyph not flagged"

    def test_gate_agrees_with_the_generators_on_plan_currency(self):
        """The gate resolving currency differently from the generators is what
        turned a correctly-built USD/London bundle into 50 criticals."""
        import bundle_qa

        for budget, locs in (
            ("$150,000", ["London, United Kingdom"]),
            ("£150,000", ["London, United Kingdom"]),
            ("150,000", ["London, United Kingdom"]),
        ):
            plan = _plan(budget=budget, locations=locs)
            gate_code, _sym = bundle_qa._resolve_plan_currency(plan)
            assert gate_code == ppt._plan_currency_code(dict(plan))


class TestFontEmbedding:
    """Fonts must stay embedded -- an unembedded deck substitutes a serif face
    on the reader's machine and every measurement above becomes fiction."""

    def test_generated_pptx_embeds_poppins(self):
        import zipfile

        blob = ppt.generate_pptx(_plan())
        names = zipfile.ZipFile(io.BytesIO(blob)).namelist()
        fonts = [n for n in names if n.startswith("ppt/fonts/")]
        assert fonts, "no embedded fonts in generated pptx"


# ===========================================================================
# deck_qa closure verification -- slide 8 CPA-band status + currency notes
#
# These three tests round out the audit closure: B1 proves the "Projected
# CPA" row on slide 8 never fabricates a beating/trailing claim by comparing
# a plan-currency figure straight against a US$ benchmark with no FX
# conversion (it must render the same neutral "none" status as "Projected
# Hires"); B2 proves the currency-basis disclosure actually reaches every
# money slide (2, 5, 6, 8); B3 pins the disclosure's own text contrast to
# WCAG AA so the fix doesn't ship unreadable.
# ===========================================================================
import budget_engine  # noqa: E402

_DEFAULT_CHANNEL_PCTS = {
    "programmatic_dsp": 25,
    "global_boards": 25,
    "niche_boards": 25,
    "social_media": 25,
}


def _with_budget_allocation(data, channel_pcts=None, collar_type="white"):
    """Attach a REAL ``_budget_allocation`` via budget_engine, the same way
    app.py wires it before ppt_generator sees the plan. Without this, slide
    8's "Projected CPA"/"Projected Hires" rows never render at all (no data
    to project from) -- that is the row this whole closure verifies."""
    try:
        budget_val = float(
            str(data.get("budget", "0")).replace(",", "").lstrip("$£€¥").strip() or 0
        )
    except ValueError:
        budget_val = 150000.0
    result = budget_engine.calculate_budget_allocation(
        total_budget=budget_val,
        roles=data.get("roles") or [],
        locations=data.get("locations") or [],
        industry=data.get("industry", "healthcare"),
        channel_percentages=dict(channel_pcts or _DEFAULT_CHANNEL_PCTS),
        synthesized_data=data.get("_synthesized"),
        collar_type=collar_type,
        campaign_start_month=data.get("campaign_start_month", 9),
    )
    data["_budget_allocation"] = result
    return data


def _synth_5platform():
    """5-platform ad_platform_analysis + job_market_demand -- drives both the
    slide-5 benchmark table and the slide-6/8 funnel strip, so a currency
    note placed near that content is genuinely exercised, not just present
    on an otherwise-empty slide."""
    platforms = {}
    for i in range(1, 6):
        platforms[f"platform_{i}"] = {
            "platform_name": f"Platform {i}",
            "CPC": round(1.2 + i * 0.3, 2),
            "CPA": round(18 + i * 4, 2),
            "estimated_reach": 50000 * i,
            "fit_score": round(0.6 + i * 0.05, 2),
            "deep_intelligence": {
                "monthly_visitors": 1_000_000 * i,
                "best_for": [
                    "Registered Nurse",
                    "Clinical Coordinator",
                    "Care Manager",
                ],
            },
        }
    return {
        "ad_platform_analysis": platforms,
        "job_market_demand": {
            "Registered Nurse": {
                "total_postings": 12000,
                "avg_salary": 78000,
                "market_temperature": "hot",
                "posting_sources": ["Adzuna", "Jooble"],
            }
        },
    }


def _slide8_shape_texts(prs):
    slide = prs.slides[7]
    return [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


def _run_rgb(run):
    """RGBColor of a run's font color, or None if it carries no explicit RGB."""
    try:
        color = run.font.color
        if color is not None and color.type is not None:
            return color.rgb
    except (AttributeError, KeyError, ValueError):
        pass
    return None


def _projected_cpa_row_shapes(slide):
    """Every text shape sharing a top position with a 'Projected CPA' label
    -- i.e. the client-panel and industry-panel value boxes for that row,
    the only shapes the CPA-band fix actually touches. Scoped to the row
    rather than "any shape on the slide" because other rows (Channels
    Selected, Programmatic Allocation, Geographic Coverage) are independent,
    data-driven comparisons that can legitimately show a beating/trailing
    arrow on any given fixture -- that is not part of this defect."""
    labels = [
        s
        for s in slide.shapes
        if s.has_text_frame and s.text_frame.text.strip() == "Projected CPA"
    ]
    assert labels, "no 'Projected CPA' row found on slide 8 -- fixture didn't drive one"
    row_shapes = []
    for lbl in labels:
        for s in slide.shapes:
            if s.has_text_frame and abs((s.top or 0) - (lbl.top or 0)) <= 1000:
                row_shapes.append(s)
    return row_shapes


class TestSlide8ProjectedCpaNeverFabricatesABeat:
    def test_gbp_plan_never_claims_beating_a_usd_benchmark(self):
        """The headline slide-8 regression: a US$-calibrated CPA benchmark
        compared with no FX conversion against a plan-currency figure must
        never render a fabricated 'beating'/'trailing' claim -- only the
        same neutral 'none' status 'Projected Hires' already uses.

        Channel mix is skewed all-in on Global Job Boards (cheapest channel
        in this fixture) so the raw client-currency CPA figure -- never
        FX-converted, so numerically it is the same digits a USD plan would
        show -- lands BELOW the US$12-US$50 healthcare benchmark band. Pre-fix
        this genuinely rendered "£10.04  ▲" (a fabricated "beating a
        US benchmark" claim); an even channel split coincidentally lands
        inside the band on both sides of the fix and would prove nothing."""
        data = _plan(
            budget="£2,000,000",
            locations=[{"city": "London", "state": "", "country": "United Kingdom"}],
            client_name="Uber",
            _synthesized=_synth_5platform(),
        )
        _with_budget_allocation(data, channel_pcts={"global_boards": 100})
        prs = _deck(data)
        slide8 = prs.slides[7]

        row_shapes = _projected_cpa_row_shapes(slide8)
        for shape in row_shapes:
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                assert (
                    "▲" not in text
                ), f"Projected CPA row shows a beat arrow: {text!r}"
                for run in para.runs:
                    rgb = _run_rgb(run)
                    assert (
                        rgb != ppt.GREEN
                    ), f"Projected CPA row run is GREEN: {run.text!r}"

        # And: the ordinary USD Dallas plan's slide 8 must render exactly as
        # before -- the currency-basis fix must not touch the common case.
        # Self-contained (no external baseline file / env var): assert the
        # USD-specific invariants directly instead of comparing to a fixture
        # that silently skipped this half of the test on CI.
        usd_data = _with_budget_allocation(_plan())
        usd_prs = _deck(usd_data)
        usd_slide8 = usd_prs.slides[7]

        for shape in usd_slide8.shapes:
            if shape.has_text_frame:
                assert (
                    "FX-converted" not in shape.text_frame.text
                ), f"USD deck discloses a currency basis on slide 8: {shape.text_frame.text!r}"

        usd_legend = next(
            (
                sh
                for sh in usd_slide8.shapes
                if sh.has_text_frame
                and "Beating benchmark" in sh.text_frame.text
                and "On par / within range" in sh.text_frame.text
            ),
            None,
        )
        assert usd_legend is not None, (
            "USD deck's slide 8 legend must keep its full original wording "
            "('Beating benchmark' / 'On par / within range')"
        )

        usd_row_shapes = _projected_cpa_row_shapes(usd_slide8)
        value_texts = [
            "".join(r.text for r in para.runs)
            for shape in usd_row_shapes
            for para in shape.text_frame.paragraphs
            if "".join(r.text for r in para.runs).strip()
            and "".join(r.text for r in para.runs).strip() != "Projected CPA"
        ]
        assert value_texts, "no Projected CPA value text found on the USD slide 8"
        for text in value_texts:
            stripped = text.strip()
            assert stripped.startswith("$") and not stripped.startswith(
                "US$"
            ), f"USD Projected CPA value should read '$...', not {stripped!r}"


class TestCurrencyNoteOnEveryMoneySlide:
    def test_currency_note_present_on_every_money_slide(self):
        for synth in ({}, _synth_5platform()):
            gbp_data = _with_budget_allocation(
                _plan(
                    budget="£2,000,000",
                    locations=[
                        {"city": "London", "state": "", "country": "United Kingdom"}
                    ],
                    client_name="Uber",
                    _synthesized=synth,
                )
            )
            prs = _deck(gbp_data)
            pages_with_note = set()
            for i, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.has_text_frame and "FX-converted" in shape.text_frame.text:
                        pages_with_note.add(i)
            for page in (2, 5, 6, 8):
                assert page in pages_with_note, (
                    f"GBP plan (synth={'yes' if synth else 'no'}) missing the "
                    f"currency-basis note on slide {page}: found on {sorted(pages_with_note)}"
                )

        usd_data = _with_budget_allocation(_plan())
        prs = _deck(usd_data)
        for i, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    assert "FX-converted" not in shape.text_frame.text, (
                        f"USD plan discloses a currency basis on slide {i} -- "
                        "the common case must render exactly as before"
                    )


class TestCurrencyNoteContrast:
    """WCAG luminance contrast, inline (no extra dependency)."""

    @staticmethod
    def _srgb_to_linear(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    @classmethod
    def _relative_luminance(cls, rgb):
        r, g, b = rgb[0], rgb[1], rgb[2]
        rl, gl, bl = (cls._srgb_to_linear(v) for v in (r, g, b))
        return 0.2126 * rl + 0.7152 * gl + 0.0722 * bl

    @classmethod
    def _contrast_ratio(cls, rgb_a, rgb_b):
        la = cls._relative_luminance(rgb_a) + 0.05
        lb = cls._relative_luminance(rgb_b) + 0.05
        return max(la, lb) / min(la, lb)

    def test_currency_note_clears_aa_contrast(self):
        dark_text = (ppt.DARK_TEXT[0], ppt.DARK_TEXT[1], ppt.DARK_TEXT[2])
        lavender_50 = (ppt.LAVENDER_50[0], ppt.LAVENDER_50[1], ppt.LAVENDER_50[2])
        light_teal = (ppt.TEAL_LIGHT[0], ppt.TEAL_LIGHT[1], ppt.TEAL_LIGHT[2])
        navy = (ppt.NAVY[0], ppt.NAVY[1], ppt.NAVY[2])

        caption_ratio = self._contrast_ratio(dark_text, lavender_50)
        assert (
            caption_ratio >= 4.5
        ), f"DARK_TEXT caption on LAVENDER_50 fails AA contrast: {caption_ratio:.2f}:1"
        band_ratio = self._contrast_ratio(light_teal, navy)
        assert (
            band_ratio >= 4.5
        ), f"TEAL_LIGHT band paragraph on NAVY fails AA contrast: {band_ratio:.2f}:1"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
