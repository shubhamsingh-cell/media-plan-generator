"""Healthcare apply-rate re-level (CPA-realism C1): the contract around it.

The re-level changes APPLICATIONS and CPA only. These tests pin what a review
of the first two commits found around that contract:

* the optimizer's funnel-based "improve projected hires by N%" recommendation
  must be exactly what the un-re-levelled engine says (it used to read the
  lowered apply rate and quote a different number to Nova);
* a small plan never shows a priced channel with spend and clicks but "0
  applications" and a CPA equal to its whole spend: a fraction of an
  application is not rounded to nothing, and a channel with truly none has no
  CPA (n/a), which no blended figure averages in;
* the industry scope of the re-level is a deliberate list, not an accident.
"""

from __future__ import annotations

import contextlib

import pytest

import budget_engine as be
import plan_projection as pp

ROLES = ["Registered Nurse", "Pharmacist", "Certified Nursing Assistant", "Physician"]
CITIES = [("Houston", "TX"), ("Chicago", "IL")]


@contextlib.contextmanager
def _generic_rates():
    """The same engine with the healthcare apply-rate re-level switched off."""
    saved = be._JOB_BOARD_APPLY_RATE_KB_KEY.pop("healthcare_medical")
    try:
        yield
    finally:
        be._JOB_BOARD_APPLY_RATE_KB_KEY["healthcare_medical"] = saved


def _run(budget, role, city, state, month=7, industry="healthcare_medical"):
    return be.calculate_budget_allocation(
        total_budget=float(budget),
        roles=[{"title": role, "count": 1, "tier": "Professional"}],
        locations=[{"city": city, "state": state, "country": "United States"}],
        industry=industry,
        channel_percentages=pp.default_channel_percentages(industry),
        knowledge_base=None,
        campaign_start_month=month,
    )


def _both(budget, role, city, state, month=7):
    relevelled = _run(budget, role, city, state, month)
    with _generic_rates():
        generic = _run(budget, role, city, state, month)
    return relevelled, generic


# -- F1: the optimizer section is the un-re-levelled engine's ---------------


@pytest.mark.parametrize("month", [1, 7])
@pytest.mark.parametrize("city,state", CITIES)
@pytest.mark.parametrize("role", ROLES)
def test_optimizer_recommendation_is_unchanged_by_the_relevel(role, city, state, month):
    seen_recommendation = False
    for budget in (2_500, 9_000, 28_000, 60_000, 250_000):
        r, g = _both(budget, role, city, state, month)
        case = (role, city, month, budget)
        assert r["recommendations"] == g["recommendations"], case
        assert r["warnings"] == g["warnings"], case
        assert r["optimized"]["improvement"] == g["optimized"]["improvement"], case
        assert r["optimized"]["changes"] == g["optimized"]["changes"], case
        ro, go = (
            r["optimized"]["optimized_allocations"],
            g["optimized"]["optimized_allocations"],
        )
        assert set(ro) == set(go), case
        for name, entry in go.items():
            for key in ("projected_clicks", "projected_hires", "percentage"):
                assert ro[name][key] == entry[key], (case, name, key)
            # applications of the optimized split follow the same re-level as the rows
            assert (
                ro[name]["projected_applications"] <= entry["projected_applications"]
            ), (case, name)
        seen_recommendation |= any(
            "Optimised allocation" in t for t in g["recommendations"]
        )
    del seen_recommendation  # presence is asserted once, below


def test_the_optimizer_recommendation_is_present_in_the_sweep_so_the_check_is_not_vacuous():
    texts = []
    for role in ROLES:
        for budget in (9_000, 28_000, 60_000):
            _, g = _both(budget, role, "Houston", "TX", 7)
            texts += [t for t in g["recommendations"] if "Optimised allocation" in t]
    assert texts, "no budget in the sweep triggers the optimizer recommendation"


# -- F2: no priced channel with spend and clicks but "0 applications" -------


@pytest.mark.parametrize("budget", [500, 1_000, 2_000])
@pytest.mark.parametrize("role", ROLES)
def test_small_plans_have_no_zero_application_channel_with_a_cpa(budget, role):
    r = _run(budget, role, "Houston", "TX")
    priced = [
        (name, ch)
        for name, ch in r["channel_allocations"].items()
        if (ch.get("dollar_amount") or 0) > 0 and (ch.get("projected_clicks") or 0) > 0
    ]
    assert priced
    for name, ch in priced:
        apps, cpa, dollars = (
            ch["projected_applications"],
            ch.get("cpa"),
            ch["dollar_amount"],
        )
        if apps > 0:
            assert cpa == pytest.approx(dollars / apps, abs=0.01), (name, ch)
        else:
            # truly no application expected: no CPA at all, never the whole spend
            assert not cpa, (name, ch)


def test_a_fraction_of_an_application_is_not_rounded_to_none():
    """$500 RN plan: employer branding is ~$40 / 44 clicks at 1.44% = 0.63
    expected applications. The generic engine shows 1 application there; the
    re-level must not turn that into 0 with the whole spend as its CPA."""
    r = _run(500, "Registered Nurse", "Houston", "TX")
    with _generic_rates():
        g = _run(500, "Registered Nurse", "Houston", "TX")
    for name, gch in g["channel_allocations"].items():
        rch = r["channel_allocations"][name]
        exact = (rch["projected_clicks"] or 0) * rch["apply_rate"]
        if exact >= 0.5:
            assert rch["projected_applications"] >= 1, (name, exact)
    assert any(
        ch["projected_applications"] == 1 and ch["dollar_amount"] > 0
        for ch in r["channel_allocations"].values()
    )


def test_zero_application_channels_stay_out_of_the_blended_cpa():
    r = _run(500, "Registered Nurse", "Houston", "TX")
    allocs = r["channel_allocations"]
    zero = [
        n
        for n, ch in allocs.items()
        if ch["projected_applications"] == 0 and ch["dollar_amount"] > 0
    ]
    assert zero, "fixture must contain a channel with no expected application"
    spend = sum(
        ch["dollar_amount"]
        for n, ch in allocs.items()
        if n not in zero and ch.get("channel_role") != "brand"
    )
    apps = sum(
        ch["projected_applications"]
        for n, ch in allocs.items()
        if n not in zero and ch.get("channel_role") != "brand"
    )
    assert be.plan_average_cpa(allocs) == pytest.approx(spend / apps)


# -- F7: the scope of the re-level is a deliberate list ----------------------


def _kind(industry, roles):
    return pp.classify_brief_industry(industry, "Probe Co", roles)[0]


def test_relevel_scope_is_pinned():
    assert be._APPLY_RELEVEL_APPLICATIONS_ONLY == frozenset({"healthcare_medical"})
    assert set(be._JOB_BOARD_APPLY_RATE_KB_KEY) == {
        "healthcare_medical",
        "blue_collar_trades",
    }
    in_scope = [
        "hospital",
        "clinic",
        "dental",
        "senior living",
        "veterinary",
        "medical devices",
        "life sciences",
        "home health",
        "telehealth",
        "nursing home",
        "Healthcare & Medical",
        "healthcare_medical",
    ]
    for industry in in_scope:
        assert _kind(industry, ["Registered Nurse"]) == "healthcare_medical", industry
    # "general" is "not stated": the roles decide (an RN title is healthcare)
    assert _kind("general", ["Registered Nurse"]) == "healthcare_medical"
    # deliberately OUT of scope for an RN brief
    assert _kind("pharma", ["Registered Nurse"]) == "pharma_biotech"
    assert _kind("pharmaceutical", ["Registered Nurse"]) == "pharma_biotech"
    assert _kind("biotech", ["Registered Nurse"]) == "pharma_biotech"
    assert _kind("behavioral health", ["Registered Nurse"]) == "mental_health"
    assert _kind("health insurance", ["Registered Nurse"]) == "finance_banking"
    # the scope follows the role mix: a pharmacist brief classifies these as healthcare
    assert _kind("behavioral health", ["Pharmacist"]) == "healthcare_medical"
    assert _kind("health insurance", ["Pharmacist"]) == "healthcare_medical"
    assert _kind("pharma", ["Pharmacist"]) == "pharma_biotech"


# -- F2 on the client surfaces: a channel with no application reads n/a ------


@pytest.fixture(scope="module")
def tiny_plan():
    """A $500 healthcare plan through the real workbook and deck generators."""
    import io
    import logging

    import excel_v2
    import ppt_generator
    import tools_regen_bundles as T
    from kb_loader import load_knowledge_base
    from openpyxl import load_workbook
    from pptx import Presentation

    logging.disable(logging.CRITICAL)
    data = T.build_plan_data(
        {
            "client_name": "Probe Health",
            "industry": "Healthcare & Medical",
            "budget": "$500",
            "campaign_duration": "3 months",
            "hire_volume": "Not specified",
            "work_environment": "onsite",
            "locations": ["Columbus, OH"],
            "roles": ["Registered Nurse"],
            "target_roles": ["Registered Nurse"],
        }
    )
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    wb = load_workbook(io.BytesIO(raw))
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    allocs = data["_budget_allocation"]["channel_allocations"]
    none_ch = [
        n
        for n, c in allocs.items()
        if c["dollar_amount"] > 0 and not c["projected_applications"]
    ]
    assert none_ch == ["social_media"], none_ch  # the fixture's no-application channel
    return data, wb, prs, allocs


def _rows_named(wb, sheet, label):
    return [
        [c.value for c in row if c.value is not None]
        for row in wb[sheet].iter_rows()
        if any(c.value == label for c in row)
    ]


def test_workbook_tables_print_na_not_zero_for_a_channel_with_no_application(tiny_plan):
    _, wb, _, _ = tiny_plan
    for sheet in ("Executive Summary", "Channel Recommendations"):
        rows = _rows_named(wb, sheet, "Social Media")
        assert rows, sheet
        assert "n/a" in rows[0], (sheet, rows[0])
    # the channels that do have applications keep a real CPA beside them
    prog = _rows_named(wb, "Executive Summary", "Programmatic (DSP)")[0]
    assert any(isinstance(v, float) and 20 < v < 60 for v in prog), prog


def test_roi_sheet_does_not_invent_applications_for_that_channel(tiny_plan):
    _, wb, _, _ = tiny_plan
    exec_row = _rows_named(wb, "Executive Summary", "Social Media")[0]
    roi_row = _rows_named(wb, "ROI Projections", "Social Media")[0]
    assert exec_row[4] == 0  # [label, share, spend, clicks, APPLICATIONS, ...]
    # ROI Projections row: [label, spend, APPLICATIONS, hires, ...]: 0, not a
    # knowledge-base-CPA guess of 2
    assert roi_row[2] == 0, roi_row


def test_deck_channel_table_prints_dashes_for_that_channel(tiny_plan):
    _, _, prs, _ = tiny_plan
    slide = next(
        s
        for s in prs.slides
        if any(
            sh.has_text_frame and "CHANNEL-BY-CHANNEL BREAKDOWN" in sh.text_frame.text
            for sh in s.shapes
        )
    )
    texts = [
        sh.text_frame.text
        for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
    i = texts.index("Social Media")
    # Allocation %, Investment, then Proj. Apps / Proj. Hires / CPA all dashes
    assert texts[i + 3 : i + 6] == ["--", "--", "--"], texts[i : i + 6]
