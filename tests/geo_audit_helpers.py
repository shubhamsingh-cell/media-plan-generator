"""Shared fixtures for the geography output-audit regression tests (G1-G9).

The RTI / Hershey audit (2026-10-11) found one location string read five
different ways across research / gold_standard / data_synthesizer / excel_v2 /
ppt_generator. These helpers build a plan the way the request boundary does
(plan_location resolves every entry, ``data["locations"]`` is rewritten to the
canonical display names, the verdict is kept as ``_location_resolution``) and
read back the workbook tables, so the tests exercise the real code paths.

Only functions that exist on the pre-fix parent are used here, so the same
tests can be run against a throwaway pre-fix worktree to prove they fail there.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List

import plan_location

RTI_LOCATIONS: List[str] = [
    "Remote",
    "Nationwide",
    "MD",
    "NY",
    "NC",
    "OH",
    "PA",
    "DC",
    "Philadelphia",
    "Jacksonville",
    "Washington",
    "Pittsburgh",
    "Baltimore",
    "Cleveland",
    "Columbus",
    "Jackson",
]

HERSHEY_LOCATIONS: List[str] = [
    "Hershey, PA",
    "Hazleton, PA",
    "Lancaster, PA",
    "Memphis, TN",
    "Stuarts Draft, VA",
    "Robinson, IL",
    "Ogden, UT",
    "Oakdale, CA",
    "Reading, PA",
    "Harrisburg, PA",
]


def canonical_plan_locations(raw: List[str]) -> tuple[List[str], List[Dict[str, Any]]]:
    """(rewritten ``data["locations"]``, ``_location_resolution`` sidecar), the
    same shape app._resolve_and_rewrite_locations produces."""
    results = plan_location.resolve_locations(raw)

    def _display(orig: str, res: Any) -> str:
        # mirrors app._canonical_location_display (resolved / corrected only;
        # ambiguous and unresolved entries stay VERBATIM)
        if res.status not in ("resolved", "corrected"):
            return orig
        if res.kind == "city":
            name = res.display_name
            return f"{name}, {res.state_usps}" if res.state_usps else name
        return res.display_name or orig

    display = [_display(orig, r) for orig, r in zip(raw, results)]
    return display, [r.to_dict() for r in results]


def build_plan(
    raw_locations: List[str],
    *,
    industry: str = "healthcare_medical",
    role: str = "Registered Nurse",
    with_resolution: bool = True,
) -> Dict[str, Any]:
    import budget_engine
    import data_synthesizer

    locations, sidecar = canonical_plan_locations(raw_locations)
    roles = [{"title": role, "count": 4, "tier": "mid"}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=150_000,
        roles=roles,
        locations=locations,
        industry=industry,
        channel_percentages={"programmatic_dsp": 50, "global_boards": 30, "social_media": 20},
    )
    data: Dict[str, Any] = {
        "client_name": "Geo Audit Fixture",
        "industry": industry,
        "budget": "150000",
        "campaign_duration": "3 months",
        "campaign_weeks": 13,
        "roles": [role],
        "target_roles": roles,
        "locations": locations,
        "_budget_allocation": alloc,
        "_enriched": {},
    }
    if with_resolution:
        data["_location_resolution"] = sidecar
    data["_synthesized"] = {
        "location_profiles": data_synthesizer.fuse_location_profiles({}, {}, data)
    }
    return data


def workbook(data: Dict[str, Any]):
    import openpyxl

    import excel_v2
    import gold_standard
    import research

    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    return openpyxl.load_workbook(
        io.BytesIO(excel_v2.generate_excel_v2(data, research_mod=research))
    )


def _table_rows(ws, start_marker: str, stop_markers: tuple[str, ...]) -> Dict[str, Dict[str, str]]:
    rows: Dict[str, Dict[str, str]] = {}
    header: Dict[str, int] = {}
    in_section = False
    for row in ws.iter_rows():
        cells = {c.column: c.value for c in row if c.value not in (None, "")}
        label = " ".join(str(v) for v in cells.values())
        if not in_section:
            if start_marker.upper() in label.upper():
                in_section = True
            continue
        if not header:
            if "Location" in cells.values():
                header = {str(v): k for k, v in cells.items()}
            continue
        if not cells:
            if rows:
                break
            continue
        if any(m.upper() in label.upper() for m in stop_markers):
            break
        loc = str(cells.get(header["Location"]) or "")
        if loc:
            rows[loc] = {name: str(cells.get(col) or "") for name, col in header.items()}
    return rows


def context_table(wb) -> Dict[str, Dict[str, str]]:
    """Market Intelligence > Location Economic Context, keyed by location."""
    return _table_rows(
        wb["Market Intelligence"],
        "Location Economic Context",
        ("LOCATION INTELLIGENCE",),
    )


def intelligence_table(wb) -> Dict[str, Dict[str, str]]:
    """Market Intelligence > LOCATION INTELLIGENCE, keyed by location."""
    return _table_rows(
        wb["Market Intelligence"],
        "LOCATION INTELLIGENCE",
        ("Macro Economic", "COMPETITIVE", "Industry Hiring"),
    )


def sheet_text(ws) -> List[str]:
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def deck_texts(data: Dict[str, Any]) -> List[tuple]:
    """[(slide number, text)] for every text shape of the generated deck."""
    from pptx import Presentation

    import gold_standard
    import ppt_generator

    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(data)))
    return [
        (i, sh.text_frame.text.strip())
        for i, slide in enumerate(prs.slides, 1)
        for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
