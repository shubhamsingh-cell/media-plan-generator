"""2026-10-16 sweep -- double possessives and two deck layout defects.

Evidence: the 26-bundle sweep (``ledger/2026-10-16/B_sweep``).

1. DOUBLE POSSESSIVES. Every ``f"{client}'s ..."`` bolted a bare ``'s`` onto
   the client name: "McDonald's's" x10 on a McDonald's plan (slides 2, 6-11),
   "Stores's" / "Analytics's" / "Technologies's" / "Logistics's" / "Services's"
   / "Partners's" / "Systems's" ~11 each, "Barclays's" and "Werner
   Enterprises's" on the competitor risk row. ``display_format.possessive`` is
   now the one rule and every site (deck, workbook, footprint warning) uses it.

2. COVER INDUSTRY LINE. "Automotive & Manufacturing · NAICS 3363 — Motor
   Vehicle Parts Manufacturing" was sized with the AVERAGE glyph advance
   against the box's OUTER width, so it wrapped and its second line printed
   through the company blurb (3M x2, Helix Biopharma). It now measures the real
   Poppins glyphs against the inner width and budgets its lines.

3. SLIDE-5 BENCHMARK HEADER. The "<industry label> Range" header sat in a
   single-line 0.38in cell; a NAICS-suffixed label wrapped to three lines and
   the first body row clipped it (a_hershey).

MEASUREMENT INDEPENDENCE: layout assertions use ``scripts/deck_qa`` (its own
PIL/Poppins ruler -- it deliberately does not reuse ppt_generator's
``_measure_lines``), so the generator is not graded with its own ruler.

Runs under pytest, or standalone:
``python3 tests/test_deck_possessive_and_cover_fit_2026_10_16.py``.
"""

from __future__ import annotations

import importlib.util
import io
import os
import re
import sys
from typing import Any

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import display_format  # noqa: E402
import ppt_generator as ppt  # noqa: E402
from pptx import Presentation  # noqa: E402


def _load_deck_qa(name: str):
    """Import scripts/deck_qa/<name>.py (a script dir, not a package)."""
    qa_dir = os.path.join(PROJECT_ROOT, "scripts", "deck_qa")
    if qa_dir not in sys.path:
        sys.path.insert(0, qa_dir)
    spec = importlib.util.spec_from_file_location(
        f"deck_qa_{name}", os.path.join(qa_dir, f"{name}.py")
    )
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


collide = _load_deck_qa("collide")
matrix = _load_deck_qa("matrix")
textlayout = _load_deck_qa("_textlayout")


@pytest.fixture(autouse=True)
def _reset_active_currency():
    """_set_active_currency keeps the plan currency in a thread-local that
    outlives a single test in the same thread."""
    yield
    try:
        ppt._currency_tls.code = "USD"
    except AttributeError:
        pass


NAICS_LABEL_3M = "Automotive & Manufacturing · NAICS 3363 — Motor Vehicle Parts Manufacturing"
NAICS_LABEL_HERSHEY = (
    "Food & Beverage · NAICS 311351 — Chocolate and Confectionery "
    "Manufacturing from Cacao Beans"
)
# Two sentences: the cover prints the first (company_tagline); the second names
# the plan's sector so company_blurb's entity validation accepts the article.
BLURB_3M = (
    "The 3M Company is an American multinational conglomerate operating in "
    "the fields of industry, worker safety, and consumer goods. It "
    "manufactures adhesives, abrasives and automotive parts."
)
BLURB_3M_TAGLINE = BLURB_3M.split(". ")[0] + "."


def _synth_competitor(name: str) -> dict:
    return {
        "competitive_intelligence": {
            "competitors": {
                name: {
                    "employee_count": "10,000+",
                    "hiring_velocity": "high",
                    "hiring_intensity": "aggressive",
                    "competitor_type": "national",
                    "evidence": "fixture: synthesized intel record",
                }
            }
        }
    }


def _deck(**over) -> Any:
    data = matrix._base_plan(**over)
    data = matrix._with_budget_allocation(data, n_channels=4)
    return Presentation(io.BytesIO(ppt.generate_pptx(data)))


def _all_text(prs: Any) -> str:
    parts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    parts.extend(c.text for c in row.cells)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


# A doubled possessive: an s/S already carrying its apostrophe-s ("Stores's",
# "McDonald's's", "Barclays's"). No legitimate English phrase in a deck puts
# "s's" together.
_DOUBLE_POSSESSIVE_RE = re.compile(r"[sS]['’]s\b")


# ===========================================================================
# 1. possessive()
# ===========================================================================
class TestPossessiveRule:
    @pytest.mark.parametrize(
        "name, expected",
        [
            ("Joveo", "Joveo's"),
            ("3M", "3M's"),
            ("AT&T", "AT&T's"),
            ("Hershey", "Hershey's"),
            # already possessive-looking: unchanged (the McDonald's's bug)
            ("McDonald's", "McDonald's"),
            ("Macy's", "Macy's"),
            ("Kohl’s", "Kohl’s"),  # typographic apostrophe
            ("DENNY'S", "DENNY'S"),
            # s / S endings take the bare apostrophe
            ("Stores", "Stores'"),
            ("Brightway Home Stores", "Brightway Home Stores'"),
            ("Barclays", "Barclays'"),
            ("Werner Enterprises", "Werner Enterprises'"),
            ("Services", "Services'"),
            ("UPS", "UPS'"),
            # already a bare-apostrophe possessive
            ("Stores'", "Stores'"),
            # whitespace / emptiness
            ("  Joveo  ", "Joveo's"),
            ("", ""),
            ("   ", ""),
            (None, ""),
        ],
    )
    def test_rule(self, name, expected):
        assert display_format.possessive(name) == expected

    @pytest.mark.parametrize(
        "name",
        ["Joveo", "McDonald's", "Stores", "Barclays", "Kohl’s", "3M", "UPS"],
    )
    def test_idempotent(self, name):
        once = display_format.possessive(name)
        assert display_format.possessive(once) == once

    def test_never_produces_a_double_possessive(self):
        for name in ("McDonald's", "Stores", "Analytics", "Technologies", "Logistics"):
            assert not _DOUBLE_POSSESSIVE_RE.search(display_format.possessive(name))


# ===========================================================================
# 2. every possessive site, end to end
# ===========================================================================
class TestNoDoublePossessiveInDeck:
    @pytest.mark.parametrize(
        "client",
        ["McDonald's", "Brightway Home Stores", "Acme Analytics", "Delta Logistics"],
    )
    def test_whole_deck_has_no_double_possessive(self, client):
        prs = _deck(
            client_name=client,
            _synthesized=_synth_competitor("Barclays"),
        )
        text = _all_text(prs)
        hits = sorted(set(_DOUBLE_POSSESSIVE_RE.findall(text)))
        bad = [
            m.group(0)
            for m in re.finditer(r"\S*[sS]['’]s\b\S*", text)
        ]
        assert not bad, f"{client}: doubled possessives in deck text: {bad[:6]}"
        # vacuity guard: the possessive sites really rendered for this client
        assert display_format.possessive(client) in text, (
            f"{client}: the possessive form never rendered (vacuous test); "
            f"hits={hits}"
        )

    def test_competitor_risk_row_uses_the_competitor_possessive(self):
        prs = _deck(client_name="Joveo Demo", competitors=["Barclays"])
        text = _all_text(prs)
        assert "Monitor Barclays' job posting volumes weekly" in text, (
            "competitor possessive missing -- risk row vacuous or still "
            "printing \"Barclays's\""
        )

    def test_competitor_why_template_possessive(self):
        # "{name}'s scale as a ..." -- a competitor already named "Lowe's"
        # must not become "Lowe's's".
        for ordinal in range(4):
            out = ppt._compose_competitor_why(
                "Lowe's",
                {"competitor_type": "national"},
                {"vertical_type": "talent_market", "role": "Cashier", "city": "Dallas"},
                ordinal,
            )
            assert not _DOUBLE_POSSESSIVE_RE.search(out), out


class TestNoDoublePossessiveInWorkbookAndWarnings:
    def test_workbook_low_competition_row(self):
        """excel_v2 'Lower competition in <city> -- favorable market for
        <client>'s talent acquisition' (the sweep's 'Stores's' x8)."""
        import excel_v2
        import openpyxl

        data = {
            "client_name": "Brightway Home Stores",
            "industry": "retail_consumer",
            "industry_label": "Retail & Consumer",
            "locations": ["Columbus, OH"],
            "roles": ["Sales Associate"],
            "target_roles": ["Sales Associate"],
            "budget": "$200,000",
            "competitors": [],
            "_synthesized": {},
            "_enriched": {},
            "_budget_allocation": {},
            "_gold_standard": {
                "competitor_mapping": {
                    "Ohio": {"top_employers": ["Target"], "hiring_intensity": "low"},
                }
            },
        }
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        excel_v2._build_sheet_quality_intelligence(ws, data, data["_gold_standard"])
        text = "\n".join(
            str(c.value)
            for row in ws.iter_rows()
            for c in row
            if c.value not in (None, "")
        )
        assert "favorable market for" in text, "row did not render (vacuous test)"
        assert "Brightway Home Stores' talent acquisition" in text
        assert not _DOUBLE_POSSESSIVE_RE.search(text), text

    def test_footprint_warning_reason(self):
        import data_synthesizer

        data = {
            "client_name": "Brightway Home Stores",
            "locations": ["Seattle, WA"],
        }
        enriched = {
            "company_metadata": {"headquarters": "Raleigh, North Carolina"},
            "company_info": {
                "description": "A home goods retailer based in North Carolina."
            },
        }
        warns = data_synthesizer._validate_location_plausibility(data, enriched, {})
        reasons = " ".join(w["reason"] for w in warns)
        assert "known operating" in reasons, "no warning rendered (vacuous test)"
        assert "Brightway Home Stores' known operating" in reasons
        assert not _DOUBLE_POSSESSIVE_RE.search(reasons), reasons


# ===========================================================================
# 3. cover industry line
# ===========================================================================
def _cover_data(industry_label: str, client: str = "3M", **over) -> dict:
    over.setdefault("industry", "automotive")
    over.setdefault("target_roles", ["Manufacturing Engineer"])
    over.setdefault(
        "roles", [{"title": "Manufacturing Engineer", "count": 40, "tier": "mid"}]
    )
    over.setdefault(
        "_enriched",
        {"company_info": {"description": BLURB_3M, "wiki_title": "3M"}},
    )
    return dict(client_name=client, industry_label=industry_label, **over)


class TestCoverIndustryLine:
    def test_naics_label_does_not_overprint_the_blurb(self):
        prs = _deck(**_cover_data(NAICS_LABEL_3M))
        slide = prs.slides[0]
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert BLURB_3M_TAGLINE in texts, (
            "cover blurb did not render -- fixture is vacuous"
        )
        findings = collide.analyze_slide(slide, 1)
        assert not findings, findings

    def test_industry_label_stays_on_one_line_when_it_fits(self):
        prs = _deck(**_cover_data(NAICS_LABEL_3M))
        shape = next(
            s
            for s in prs.slides[0].shapes
            if s.has_text_frame and s.text_frame.text == NAICS_LABEL_3M
        )
        recs = textlayout.layout_textframe(shape)
        assert len(recs) == 1, [r["text"] for r in recs]

    def test_pathological_label_is_line_budgeted_and_ellipsized_on_a_word(self):
        long_label = (
            "Technology & Engineering · NAICS 541512 — "
            + "Computer Systems Design and Related Professional Services "
            * 6
        ).strip()
        prs = _deck(**_cover_data(long_label))
        slide = prs.slides[0]
        shape = next(
            s
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text.startswith("Technology & Eng")
        )
        recs = textlayout.layout_textframe(shape)
        assert len(recs) <= 2, f"{len(recs)} lines: {[r['text'] for r in recs]}"
        shown = shape.text_frame.text
        assert shown.endswith("…") and len(shown) < len(long_label)
        # cut on a word boundary: the kept text is a whole-word prefix
        assert long_label.startswith(shown[:-1].rstrip())
        assert not collide.analyze_slide(slide, 1)

    def test_worst_case_long_client_plus_two_line_label_is_clean(self):
        """A wrapped hero name AND a two-line industry label cascade the block
        below twice; the whole cover must still be ink-collision free and on
        the canvas."""
        long_label = (
            "Technology & Engineering · NAICS 541512 — "
            + "Computer Systems Design and Related Professional Services " * 3
        ).strip()
        prs = _deck(**_cover_data(long_label, client=matrix.LONG_CLIENT))
        assert not collide.analyze_slide(prs.slides[0], 1)

    @pytest.mark.parametrize("client", ["3M", "Helix Biopharma Global"])
    def test_cover_clean_for_short_and_normal_clients(self, client):
        prs = _deck(**_cover_data(NAICS_LABEL_3M, client=client))
        assert not collide.analyze_slide(prs.slides[0], 1)

    def test_short_industry_label_keeps_the_original_geometry(self):
        """The common case must not reflow: 18pt, anchored 0.08in under the
        hero name's 1.0in box (3.48 + 1.0 + 0.08 = 4.56in)."""
        prs = _deck(**_cover_data("Healthcare", client="Mercy Health Partners"))
        shape = next(
            s
            for s in prs.slides[0].shapes
            if s.has_text_frame and s.text_frame.text == "Healthcare"
        )
        assert shape.top / 914400.0 == pytest.approx(4.56, abs=0.01)
        assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(18)


# ===========================================================================
# 4. slide-5 benchmark header
# ===========================================================================
class TestBenchmarkHeader:
    def _header_shape(self, slide):
        return next(
            s
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text.rstrip().endswith("Range")
        )

    def test_hershey_naics_label_header_is_one_line_and_unclipped(self):
        prs = _deck(
            client_name="The Hershey Company",
            industry="food_beverage",
            industry_label=NAICS_LABEL_HERSHEY,
        )
        slide = prs.slides[4]
        header = self._header_shape(slide)
        assert len(textlayout.layout_textframe(header)) == 1, header.text_frame.text
        findings = collide.analyze_slide(slide, 5)
        assert not findings, findings

    @pytest.mark.parametrize(
        "label",
        [
            NAICS_LABEL_HERSHEY,
            "Healthcare & Medical · NAICS 622110 — General Medical and Surgical Hospitals",
            "Retail & Consumer",
            NAICS_LABEL_3M,
            "Technology & Engineering · NAICS 541512 — Computer Systems Design Services",
            "",
        ],
    )
    def test_header_never_wraps_whatever_the_label(self, label):
        prs = _deck(industry="food_beverage", industry_label=label)
        header = self._header_shape(prs.slides[4])
        assert len(textlayout.layout_textframe(header)) == 1, header.text_frame.text
        assert not collide.analyze_slide(prs.slides[4], 5)

    def test_head_name_steps_down_to_fit_instead_of_the_generic_word(self):
        """"Automotive & Manufacturing Range" is 2.32in at 9pt in a 2.3in cell:
        it must keep the industry name at 8.5pt, not fall back to
        "Industry Range" (the 3M / VW sweep decks) -- and never go under the
        8pt floor."""
        prs = _deck(industry="automotive", industry_label=NAICS_LABEL_3M)
        header = self._header_shape(prs.slides[4])
        assert header.text_frame.text == "Automotive & Manufacturing Range"
        size = header.text_frame.paragraphs[0].runs[0].font.size.pt
        assert 8.0 <= size < 9.0, size
        assert len(textlayout.layout_textframe(header)) == 1
        assert not collide.analyze_slide(prs.slides[4], 5)

    def test_short_label_header_text_is_unchanged(self):
        prs = _deck(industry="retail_consumer", industry_label="Retail & Consumer")
        assert self._header_shape(prs.slides[4]).text_frame.text == "Retail & Consumer Range"

    def test_industry_head_name(self):
        assert ppt._industry_head_name(NAICS_LABEL_HERSHEY) == "Food & Beverage"
        assert ppt._industry_head_name("Retail & Consumer") == "Retail & Consumer"
        assert ppt._industry_head_name("Wi-Fi Services") == "Wi-Fi Services"
        assert ppt._industry_head_name("") == "Industry"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
