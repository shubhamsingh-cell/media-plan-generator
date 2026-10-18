"""Phase B of the client-facing benchmark provenance fix (2026-10-18).

Phase A (data only, live) stamped every ``industry_benchmarks`` row with a
``provenance`` string and removed the false attributions. These tests pin the
code half -- what the deliverables PRINT of those rows:

1. the workbook's Executive Summary "Industry Benchmarks" block prints a key
   WHITELIST of the KB row (not every key, with its inline "Appcast 2026" /
   "Recruitics" asides, trend notes and sub-statistics) and ONE footnote taken
   from ``plan_trust.BENCHMARK_FOOTNOTE`` -- the same string the deck prints;
   merged-cell value rows are given a height that shows every line;
2. the deck's benchmark slide no longer cites "Joveo Recruitment Benchmarks KB
   (2026)" (the table citing itself) -- it prints the shared footnote -- and
   the competitive-landscape slide's knowledge-base fallback rows get the same
   note instead of an invented "curated knowledge base" source;
3. no "Industry Cost-per-Hire" is ever the media-only
   ``recruitment_marketing_only`` figure (all-in total or nothing);
4. "Model vs. benchmark" compares to the CPA median (see
   tests/test_model_vs_benchmark_currency.py);
5. the deck's plan-vs-industry column is labelled a Nova working estimate, not
   an "<industry> Average";
6. no deliverable calls the assumed all-in Nova cost-per-hire table a
   "benchmark".
"""

from __future__ import annotations

import copy
import io
import logging
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation

import budget_engine
import bundle_qa
import excel_v2
import kb_loader
import plan_projection
import plan_trust
import ppt_generator
import tools_regen_bundles as regen
import us_only_markers

_SECTION_FOOTNOTE = getattr(plan_trust, "BENCHMARK_FOOTNOTE", None)

# KB keys whose rows may reach the client in this block (the whitelist).
_ALLOWED_LABELS = {
    "CPA",
    "CPC",
    "Industry Cost-per-Hire",
    "Apply Rate",
    "Time To Fill",
    "Seasonal Patterns",
}
# Rows the block used to print that are free text / sub-statistics.
_RETIRED_LABELS = (
    "Provenance",
    "Offer Acceptance Rate",
    "Quality Of Hire",
    "Source Of Hire",
)


def _bundle(brief: Dict[str, Any]) -> Dict[str, Any]:
    logging.disable(logging.CRITICAL)
    try:
        return regen.generate_bundle(dict(brief), Path(tempfile.mkdtemp()), "t")
    finally:
        logging.disable(logging.NOTSET)


@pytest.fixture(scope="module")
def healthcare() -> Dict[str, Any]:
    return _bundle(dict(regen.ATRIA_BRIEF, industry="Healthcare & Medical"))


@pytest.fixture(scope="module")
def telecom() -> Dict[str, Any]:
    return _bundle(dict(regen.ATRIA_BRIEF, industry="Telecommunications"))


@pytest.fixture()
def telecom_media_only() -> Dict[str, Any]:
    """A telecom plan whose KB row carries ONLY the media-only cost per hire
    ($800-$1,500): neither surface may print it as the industry cost per hire.
    Function-scoped on purpose: the patched loader must be undone before any
    module-scoped bundle (``telecom``) is built."""
    real = kb_loader.load_knowledge_base

    def patched(*args, **kwargs):
        kb = copy.deepcopy(real(*args, **kwargs))
        row = kb["recruitment_benchmarks"]["industry_benchmarks"]["telecommunications"]
        row["cph"] = {"recruitment_marketing_only": "$800-$1,500", "notes": "media spend only"}
        return kb

    mp = pytest.MonkeyPatch()
    mp.setattr(kb_loader, "load_knowledge_base", patched)
    try:
        yield _bundle(dict(regen.ATRIA_BRIEF, industry="Telecommunications"))
    finally:
        mp.undo()


@pytest.fixture(scope="module")
def berlin_no_figures() -> Dict[str, Any]:
    """A non-US plan whose benchmark block can print only "No local
    benchmark for this market" (no industry row for it, no US figures)."""
    brief = dict(regen.ATRIA_BRIEF)
    brief.update(
        {
            "client_name": "Schwarzwald Automotive Systems GmbH",
            "industry": "Automotive & Manufacturing",
            "budget": "\u20ac300,000",
            "locations": ["Berlin, Germany", "Munich, Germany"],
            "roles": ["Assembly Technician"],
            "target_roles": [{"title": "Assembly Technician", "count": 20, "tier": "Hourly"}],
            "hire_volume": "50 hires",
            "competitors": [],
        }
    )
    return _bundle(brief)


@pytest.fixture(scope="module")
def uk_plan() -> Dict[str, Any]:
    brief = dict(regen.ATRIA_BRIEF)
    brief.update(
        {
            "client_name": "London NHS Trust",
            "industry": "Healthcare & Medical",
            "budget": "£2,000,000",
            "locations": ["London, UK"],
            "roles": ["Registered Nurse"],
            "target_roles": [{"title": "Registered Nurse", "count": 40, "tier": "Hourly"}],
            "hire_volume": "40 hires",
            "competitors": [],
        }
    )
    return _bundle(brief)


def _block_rows(xlsx_bytes: bytes) -> List[Tuple[str, str]]:
    """(label, value) of every row of the Executive Summary's "Industry
    Benchmarks" block, footnote rows included (they carry text in column B)."""
    ws = load_workbook(io.BytesIO(xlsx_bytes))["Executive Summary"]
    rows: List[Tuple[str, str]] = []
    inside = False
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        value = ws.cell(r, 4).value
        if isinstance(label, str) and label.startswith("Industry Benchmarks"):
            inside = True
            continue
        if inside:
            if not label:
                break
            rows.append((str(label), str(value or "")))
    return rows


def _deck_slide(pptx_bytes: bytes, title: str):
    prs = Presentation(io.BytesIO(pptx_bytes))
    for slide in prs.slides:
        if title in [s.text_frame.text for s in slide.shapes if s.has_text_frame]:
            return slide
    raise AssertionError(f"no slide titled {title!r}")


def _texts(slide) -> List[str]:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


# ---------------------------------------------------------------- 1. workbook
def test_workbook_block_prints_only_whitelisted_rows(healthcare):
    rows = _block_rows(healthcare["xlsx_bytes"])
    labels = [lbl for lbl, _ in rows if not lbl.startswith(("Industry ranges", "Model vs."))]
    assert not [lbl for lbl in labels if lbl in _RETIRED_LABELS], labels
    assert all(
        any(lbl.startswith(ok) for ok in _ALLOWED_LABELS) for lbl in labels
    ), labels


def test_workbook_block_values_carry_no_free_text_asides(healthcare):
    by_label = {lbl.split(" (")[0]: v for lbl, v in _block_rows(healthcare["xlsx_bytes"])}
    cpa, cpc, apply_rate = by_label["CPA"], by_label["CPC"], by_label["Apply Rate"]
    for text in (cpa, cpc, apply_rate):
        assert "Notes" not in text and "Hard To Fill" not in text, text
    assert "Median" in cpa and "Range" in cpa, cpa
    assert "Search CPC" not in cpc and "Social CPC" not in cpc and "Easy To Fill" not in cpc, cpc
    assert "Yoy Change" not in apply_rate and "2024" not in apply_rate, apply_rate
    # the deck's CPC row prints the same YoY trend, so the workbook still quotes it
    assert "Trend Yoy" in cpc, cpc


def test_workbook_prints_one_shared_footnote_under_the_block(healthcare):
    rows = _block_rows(healthcare["xlsx_bytes"])
    assert not [lbl for lbl, _ in rows if lbl == "Provenance"], rows
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    notes = [lbl for lbl, _ in rows if lbl == _SECTION_FOOTNOTE]
    assert len(notes) == 1, rows
    # directly under the benchmark rows, before the model-vs-benchmark note
    assert rows.index((_SECTION_FOOTNOTE, "")) >= 1
    first_note = next(i for i, (lbl, _) in enumerate(rows) if lbl.startswith(("Industry ranges", "Model vs.")))
    assert rows[first_note][0] == _SECTION_FOOTNOTE, rows


def test_telecom_without_an_all_in_row_never_shows_the_media_only_figure(telecom_media_only):
    block = _block_rows(telecom_media_only["xlsx_bytes"])
    cph_rows = [v for lbl, v in block if "Cost-per-Hire" in lbl]
    assert cph_rows, block
    assert not [v for v in cph_rows if "800" in v or "1,500" in v], cph_rows
    deck = " | ".join(
        _texts(_deck_slide(telecom_media_only["pptx_bytes"], "Channel Strategy & Investment"))
    )
    assert "$800" not in deck and "1,500" not in deck, deck


def test_telecom_cph_row_is_the_all_in_total_on_both_surfaces(telecom):
    block = _block_rows(telecom["xlsx_bytes"])
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    assert (_SECTION_FOOTNOTE, "") in block, block  # the row sits under the shared note
    rows = dict(block)
    assert "5,000" in rows["Industry Cost-per-Hire"], rows
    deck = " | ".join(_texts(_deck_slide(telecom["pptx_bytes"], "Channel Strategy & Investment")))
    assert "5,000" in deck and "$800-$1,500" not in deck, deck


def test_write_kv_row_sizes_a_wrapped_merged_value():
    ws = Workbook().active
    for col, width in zip("ABCDEFGHI", (3, 14, 14, 18, 18, 18, 18, 18, 3)):
        ws.column_dimensions[col].width = width
    excel_v2._write_kv_row(ws, 1, "Short", "one line")
    excel_v2._write_kv_row(ws, 2, "Long", "x " * 150)
    assert ws.row_dimensions[1].height is None
    # 300 chars over a ~90-unit merge is 4 lines: the row must show them all
    assert ws.row_dimensions[2].height is not None and ws.row_dimensions[2].height >= 40, (
        ws.row_dimensions[2].height
    )


def test_write_kv_row_sizes_a_long_label_over_the_key_cells():
    ws = Workbook().active
    for col, width in zip("ABCDEFGHI", (3, 18, 18, 18, 18, 18, 18, 18, 3)):
        ws.column_dimensions[col].width = width
    # the key sits in the merged B:C (36 units): 41 bold characters wrap to two
    # lines, the one-line value does not -- the row must still show both lines
    excel_v2._write_kv_row(ws, 1, "Apply Rate (US benchmark \u2014 no local data)", "9.5-10%")
    excel_v2._write_kv_row(ws, 2, "Industry Cost-per-Hire", "$9,000")
    assert ws.row_dimensions[1].height is not None and ws.row_dimensions[1].height >= 29, (
        ws.row_dimensions[1].height
    )
    assert ws.row_dimensions[2].height is None


def test_benchmark_block_figure_detector():
    ws = Workbook().active

    def put(r, label, value):
        ws.cell(r, 2, label)
        ws.cell(r, 4, value)

    put(1, "Industry Cost-per-Hire", "No local benchmark for this market")
    put(2, "Seasonal Patterns", "Peak Months: October; Basis: 2026-10 calendar (October to December)")
    put(3, "Industry CPA / CPC", "No local benchmark for this market (US figures are not shown).")
    assert excel_v2._benchmark_block_has_figure(ws, 1, 4) is False
    put(4, "CPA", "Range: $35-$58+; Median: $35.00")
    assert excel_v2._benchmark_block_has_figure(ws, 1, 5) is True


def test_footnote_is_not_orphaned_under_a_block_with_no_figure(berlin_no_figures):
    rows = _block_rows(berlin_no_figures["xlsx_bytes"])
    assert rows and all(not lbl.startswith("Industry ranges") for lbl, _ in rows), rows
    assert [v for lbl, v in rows if "No local benchmark" in v], rows


# ------------------------------------------------- footnote / marker contract
def test_footnotes_clear_the_us_only_marker_gate():
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    assert not us_only_markers.has_us_only_marker(_SECTION_FOOTNOTE)
    assert not us_only_markers.has_us_only_marker(plan_trust.INDUSTRY_REFERENCE_LABEL)
    # wording the verifier fixed (2026-10-18): says the ranges are estimates and
    # that only a figure naming its source traces to a report -- it must not
    # claim the ranges are informed by a report (or by Appcast apply-rate data:
    # the apply rates beside it, telecom 5.5-7.5% / finance 9.5-10% /
    # government 7-11%, exceed anything Appcast reports; max 7.14%).
    assert "Nova working estimates" in _SECTION_FOOTNOTE
    assert "not independently audited" in _SECTION_FOOTNOTE
    assert "Only figures that name a source" in _SECTION_FOOTNOTE
    assert "informed by" not in _SECTION_FOOTNOTE
    assert "apply-rate data" not in _SECTION_FOOTNOTE


def test_non_us_plan_prints_the_footnote_and_passes_the_us_marker_gate(uk_plan):
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    assert any(lbl == _SECTION_FOOTNOTE for lbl, _ in _block_rows(uk_plan["xlsx_bytes"]))
    deck = _deck_slide(uk_plan["pptx_bytes"], "Channel Strategy & Investment")
    assert _SECTION_FOOTNOTE in _texts(deck)
    findings = bundle_qa.run_bundle_qa(
        uk_plan["pptx_bytes"], uk_plan["xlsx_bytes"], uk_plan["data"]
    )
    assert [f for f in findings if f.get("code") == "us_data_on_non_us_plan"] == []
    # the omitted-US-figure row speaks of Nova's estimate, not "the knowledge base"
    block = " | ".join(f"{lbl} {val}" for lbl, val in _block_rows(uk_plan["xlsx_bytes"]))
    assert "Nova's industry estimate for CPA/CPC is a US figure" in block, block
    assert "knowledge base" not in block, block


# ------------------------------------------------------------------- 2. deck
@pytest.mark.parametrize("fixture", ["healthcare", "telecom"])
def test_benchmark_slide_source_line_is_the_shared_footnote(fixture, request):
    bundle = request.getfixturevalue(fixture)
    slide = _deck_slide(bundle["pptx_bytes"], "Channel Strategy & Investment")
    texts = _texts(slide)
    assert not [t for t in texts if "Joveo Recruitment Benchmarks KB" in t], texts
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    note = next((s for s in slide.shapes if s.has_text_frame and s.text_frame.text == _SECTION_FOOTNOTE), None)
    assert note is not None, texts
    # the table kept its Apply Rate row and the note clears the attribution heading
    assert "Apply Rate" in texts, texts
    heading = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "CHANNEL CATEGORY ATTRIBUTION")
    assert (note.top + note.height) / 914400 <= heading.top / 914400 + 0.02, (
        (note.top + note.height) / 914400,
        heading.top / 914400,
    )


def test_slide_10_calls_nova_s_own_constant_an_estimate_not_a_benchmark(healthcare):
    texts = _texts(_deck_slide(healthcare["pptx_bytes"], "Plan Comparison & Implementation"))
    assert "Industry CPA (Nova estimate)" in texts, texts
    assert not [t for t in texts if "benchmark" in t.lower() and "not an industry benchmark" not in t], texts
    legend = next(t for t in texts if "On par" in t)
    assert "Ahead" in legend and "Behind" in legend, legend
    assert "Beating" not in legend and "Trailing" not in legend, legend


@pytest.mark.parametrize("confidence", ["curated", "platform_benchmark", "market_intelligence_kb"])
def test_unnamed_source_confidences_all_print_the_shared_footnote(confidence):
    for ad_rows in (False, True):
        text, lines = ppt_generator._bench_source_text(
            {"confidence": confidence}, ad_rows, [], 4.9, 8.0
        )
        assert text == _SECTION_FOOTNOTE, text
        assert "Industry benchmarks" not in text and "Ad Platform" not in text
        assert lines == 2


def test_named_sources_keep_their_line_and_name_ad_platform_only_when_a_row_shows():
    without, n0 = ppt_generator._bench_source_text({"confidence": "trend_engine"}, False, [], 4.9, 8.0)
    with_ad, n1 = ppt_generator._bench_source_text({"confidence": "trend_engine"}, True, [], 4.9, 8.0)
    assert without == "Sources: Nova Trend Engine (platform CPC/CPA data)", without
    assert with_ad == without + ", Nova Ad Platform Intelligence", with_ad
    assert n0 == n1 == 1


def test_static_table_deck_prints_the_footnote_and_no_invented_source():
    real = ppt_generator._get_benchmarks

    def static(industry, data=None):
        out = dict(real(industry, data))
        out["confidence"] = "curated"
        return out

    mp = pytest.MonkeyPatch()
    mp.setattr(ppt_generator, "_get_benchmarks", static)
    try:
        bundle = _bundle(dict(regen.ATRIA_BRIEF, industry="Healthcare & Medical"))
    finally:
        mp.undo()
    slide_text = " | ".join(_texts(_deck_slide(bundle["pptx_bytes"], "Channel Strategy & Investment")))
    assert "Industry benchmarks 20" not in slide_text, slide_text
    assert "Ad Platform Intelligence" not in slide_text, slide_text
    assert _SECTION_FOOTNOTE in slide_text, slide_text
    # the confidence caption under the table says the same thing in four words
    assert "Nova estimate (unaudited)" in slide_text, slide_text
    assert "Curated" not in slide_text, slide_text


def test_static_confidence_captions_say_nova_estimate():
    for confidence in ("curated", ppt_generator.PLATFORM_BENCHMARK_CONFIDENCE, "knowledge_base"):
        assert ppt_generator._confidence_color(confidence)[1] == "Nova estimate (unaudited)", confidence
    assert ppt_generator._confidence_color("live_api")[1] == "Live Data"


def test_competitive_landscape_no_longer_names_a_curated_knowledge_base():
    ctx = {
        "company_profile": {"wikipedia_url": "https://en.wikipedia.org/wiki/X"},
        "hiring_trends": {"kb_context": {"outlook": "x"}, "market_context": {"market_state": "y"}},
        "_meta": {"competitor_evidence_count": 0},
    }
    line = ppt_generator._competitive_landscape_sources_line(ctx)
    assert "curated knowledge base" not in line, line
    assert "Nova industry notes (unaudited estimates)" in line, line
    assert "Wikipedia" in line and "no competitor-level sources" in line, line


def test_competitive_landscape_fallback_rows_get_the_shared_footnote():
    fallback = {"hiring_trends": {"avg_time_to_fill": {"average_days": 49}, "offer_acceptance_rate": "78-85%"}}
    line = ppt_generator._competitive_landscape_sources_line(fallback)
    assert "curated knowledge base" not in line, line
    assert _SECTION_FOOTNOTE, "plan_trust.BENCHMARK_FOOTNOTE is missing"
    assert line == _SECTION_FOOTNOTE, line
    # a synthesized profile with real sources keeps naming them
    real = {
        "company_profile": {"wikipedia_url": "https://en.wikipedia.org/wiki/X"},
        "hiring_trends": {"kb_context": "x"},
    }
    assert "Wikipedia" in ppt_generator._competitive_landscape_sources_line(real)


# ------------------------------------------------------- 3. media-only never
_MEDIA_ONLY_ROW = {"cph": {"recruitment_marketing_only": "$800-$1,500", "notes": "media spend"}}


def _kb_with(row: Dict[str, Any]) -> Dict[str, Any]:
    return {"recruitment_benchmarks": {"industry_benchmarks": {"x_industry": row}}}


def test_kb_cph_selector_never_returns_the_media_only_figure():
    assert budget_engine.kb_cph_node_text(_MEDIA_ONLY_ROW["cph"]) == ""
    both = {"recruitment_marketing_only": "$800-$1,500", "total_cost_per_hire": "$5,000-$8,000"}
    assert budget_engine.kb_cph_node_text(both) == "$5,000-$8,000"


def test_engine_kb_cph_row_is_none_for_a_media_only_node():
    assert budget_engine._kb_industry_cph_row(_kb_with(_MEDIA_ONLY_ROW), "x_industry") is None


def test_deck_kb_extract_range_never_returns_the_media_only_figure():
    assert ppt_generator._kb_extract_range({"recruitment_marketing_only": "$800-$1,500"}) == ""
    bm = ppt_generator._kb_recruitment_industry_benchmark(
        "x_industry",
        {"_knowledge_base": _kb_with({**_MEDIA_ONLY_ROW, "cpa": {"range": "$20-$55"}})},
    )
    assert bm is not None and bm["cph"] == "", bm
    assert bm["cpa"] == "$20-$55"


# ------------------------------------------------- 5. comparison panel label
def test_comparison_slide_labels_its_industry_column_a_nova_estimate(healthcare):
    slide = _deck_slide(healthcare["pptx_bytes"], "Plan Comparison & Implementation")
    texts = _texts(slide)
    assert not [t for t in texts if t.rstrip().endswith("Average")], texts
    assert not [t for t in texts if "industry averages" in t], texts
    assert any(plan_trust.INDUSTRY_REFERENCE_LABEL in t for t in texts if t.lstrip("○ ").startswith(plan_trust.INDUSTRY_REFERENCE_LABEL)), texts


# ---------------------------------------- 6. the Nova table is not a benchmark
def _alloc(info: Dict[str, Any]) -> Dict[str, Any]:
    return {"metadata": {"industry_cph": info}}


_KIND_INFOS = {
    "nova_table": {"basis": "us_benchmark", "industry_matched": True, "low": 9000, "high": 12000, "currency": "USD"},
    "kb_row": {
        "basis": "us_benchmark",
        "industry_matched": False,
        "low": 6000,
        "high": 9000,
        "currency": "USD",
        "display_row": {
            "text": "$5,000-$8,000 (general roles; technical roles $10,000-$20,000)",
            "low": 5000.0,
            "high": 8000.0,
            "mid": 6500.0,
        },
    },
    "cross_industry_default": {
        "basis": "us_benchmark",
        "industry_matched": False,
        "low": 6000,
        "high": 9000,
        "currency": "USD",
    },
    "local_kb": {
        "basis": "local_kb",
        "currency": "EUR",
        "low": 48000,
        "high": 64000,
        "source_names": ["Eurojob-Consulting"],
        "value_label": "midpoint of cited range",
    },
}
_KIND_FIGURE = {
    "nova_table": "$9,000\u2013$12,000",
    "kb_row": "$5,000\u2013$8,000",
    "cross_industry_default": "$6,000\u2013$9,000",
    "local_kb": "\u20ac48,000\u2013\u20ac64,000",
}


@pytest.fixture
def _plan_money():
    """Set the active plan currency the way the workbook does, restore after."""
    def _set(code: str, budget: str) -> None:
        excel_v2._set_active_currency({"currency": code, "budget": budget})

    yield _set
    if hasattr(excel_v2._currency_tls, "code"):
        del excel_v2._currency_tls.code


@pytest.mark.parametrize("kind", sorted(_KIND_INFOS))
def test_anchor_clause_is_one_plain_clause_per_basis_kind(kind, _plan_money):
    info = _KIND_INFOS[kind]
    _plan_money("EUR" if kind == "local_kb" else "USD", "\u20ac100,000" if kind == "local_kb" else "$100,000")
    assert plan_projection.cph_basis(info)["kind"] == kind
    clause = excel_v2._cph_anchor_clause(info)
    assert _KIND_FIGURE[kind] in clause, clause
    # one parenthetical, no internal wording, never a "benchmark"
    assert clause.count("(") == 1 and clause.count(")") == 1, clause
    for jargon in ("knowledge base", "floor", "cross-industry default", "benchmark", "row"):
        assert jargon not in clause, (jargon, clause)
    if kind == "local_kb":
        assert clause.endswith("(Eurojob-Consulting)"), clause
        assert "cited local cost-per-hire range" in clause, clause
    else:
        assert clause.startswith("based on an assumed "), clause
        assert clause.endswith("(Nova estimate)"), clause
    # the same words the methodology line and the niche-board purpose line print
    method = excel_v2._roi_methodology_note(_alloc(info))
    assert f"hire totals are {clause}, not on this funnel;" in method, method
    assert method.count("(") == 1, method


def test_methodology_wording_for_the_four_kinds_matches_the_wizard_basis(_plan_money):
    """The workbook clause never contradicts plan_projection.cph_basis, which the
    wizard hire_basis label and the plan_trust checklist print."""
    _plan_money("USD", "$100,000")
    for kind in ("nova_table", "kb_row", "cross_industry_default"):
        basis = plan_projection.cph_basis(_KIND_INFOS[kind])
        assert basis["kind"] == kind
        assert "Nova" in basis["label"] or "Nova" in basis["phrase"]
        assert "(Nova estimate)" in excel_v2._cph_anchor_clause(_KIND_INFOS[kind])
    local = plan_projection.cph_basis(_KIND_INFOS["local_kb"])
    assert "Eurojob-Consulting" in local["label"]
    assert "Nova" not in excel_v2._cph_anchor_clause(_KIND_INFOS["local_kb"])


def test_telecom_copy_in_the_bundle_is_one_plain_clause(telecom):
    wb = load_workbook(io.BytesIO(telecom["xlsx_bytes"]))
    cells = [
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str)
    ]
    method = next(t for t in cells if t.startswith("Methodology: hire totals are"))
    assert "based on an assumed industry cost per hire of $5,000\u2013$8,000 (Nova estimate), not on this funnel" in method, method
    purpose = next(t for t in cells if t.startswith("Specialty job boards matched"))
    assert "total hires are based on an assumed industry cost per hire of $5,000\u2013$8,000 (Nova estimate)" in purpose, purpose
    for t in (method, purpose):
        assert "knowledge base" not in t and "floor:" not in t and "anchored to" not in t, t
    assert not [t for t in cells if "cost-per-hire benchmarks" in t]


def test_narrative_facts_call_nova_figures_estimates_and_still_ground():
    kb = kb_loader.load_knowledge_base()["recruitment_benchmarks"]["industry_benchmarks"]
    info = {"basis": "us_benchmark", "industry_matched": True, "value": 10500.0, "currency": "USD"}
    facts = excel_v2._build_narrative_facts_block(
        {"ind_bench": kb["healthcare_medical"], "industry_cph_info": info,
         "client_name": "X", "budget_num": 250000}
    )
    industry_lines = [ln for ln in facts.splitlines() if ln.startswith(("Industry", "Assumed"))]
    assert industry_lines, facts
    assert not [ln for ln in industry_lines if "enchmark" in ln], industry_lines
    assert "Assumed industry cost per hire (Nova table, unsourced): $10,500" in industry_lines
    # the untraceable-figure gate reads the numbers, not the labels: a draft
    # citing the CPH estimate still grounds, an invented figure does not
    allowed = excel_v2._allowed_numbers_from_facts_text(facts)
    ok, bad = excel_v2._narrative_is_grounded(
        "Industry cost per hire is assumed at about $10,500 and CPA runs $35.", allowed
    )
    assert ok and not bad, bad
    ok, bad = excel_v2._narrative_is_grounded("The industry benchmark is $11,900 per hire.", allowed)
    assert not ok and bad, (ok, bad)


def test_funnel_assumption_note_does_not_call_the_table_a_benchmark():
    channels = {
        "programmatic_dsp": {"projected_applications": 8565, "projected_hires": 18, "category": "programmatic"},
        "niche_boards": {"projected_applications": 1461, "projected_hires": 6, "category": "niche_board"},
    }
    note = budget_engine.compute_funnel_stages(channels, 24)["assumption_note"]
    assert "cost-per-hire benchmarks" not in note, note
    assert "assumed all-in industry cost per hire" in note, note


def test_workbook_cph_readers_never_take_the_media_only_figure():
    kb = _kb_with(_MEDIA_ONLY_ROW)
    assert excel_v2._kb_industry_cph_benchmark("x_industry", kb=kb) == 0.0
    both = _kb_with({"cph": {"recruitment_marketing_only": "$800", "total_cost_per_hire": "$5,000-$8,000"}})
    assert excel_v2._kb_industry_cph_benchmark("x_industry", kb=both) == 6500.0
    facts = excel_v2._build_narrative_facts_block(
        {"ind_bench": _MEDIA_ONLY_ROW, "client_name": "X", "budget_num": 0}
    )
    assert "Industry Cost/Hire" not in facts, facts


def _plan_with(info: Dict[str, Any]) -> Dict[str, Any]:
    return {"_budget_allocation": {"metadata": {"industry_cph": info}}}


@pytest.mark.parametrize("kind", sorted(_KIND_INFOS))
def test_basis_wording_has_no_internal_vocabulary_in_any_reader(kind):
    """The checklist (which the deck's slide-1 notes and the workbook's
    "Review before sending" print) and the wizard label read the SAME plain
    words for every kind: one clause, no stacked brackets, no jargon."""
    info = _KIND_INFOS[kind]
    basis = plan_projection.cph_basis(info)
    assert basis["kind"] == kind
    first = plan_trust.review_items(_plan_with(info))[0]
    if kind == "nova_table":
        # unchanged on purpose: the help page quotes this sentence verbatim
        # (tests/test_help_media_plans.py equality test)
        assert "Nova's assumed industry cost-per-hire table (an all-in figure" in first, first
    else:
        assert basis["phrase"] in first, (basis, first)
    for text in (first, basis["label"], basis["phrase"]):
        for jargon in ("knowledge base", "floor: half", "Nova working estimate from", "cross-industry default"):
            assert jargon not in text, (jargon, text)
        assert "))" not in text and ")(" not in text and ") (" not in text, text
    if kind in ("kb_row", "cross_industry_default"):
        assert first.startswith(
            "Projected hires and costs are estimates: media budget divided by an assumed "
        ), first
        assert basis["label"].endswith("(" + basis["note"] + ")"), basis
        assert basis["phrase"] == "an " + basis["label"]


def test_kb_row_names_its_scope_only_when_the_row_does():
    named = plan_projection.cph_basis(
        {"basis": "us_benchmark", "industry_matched": False,
         "display_row": {"text": "$3,000-$5,000 (food manufacturing)", "mid": 4000.0}}
    )
    assert named["label"] == (
        "assumed industry cost per hire of $3,000\u2013$5,000 (Nova estimate, food-manufacturing row)"
    )
    # a row whose bracket is a note (digits, a semicolon) names nothing
    note_only = plan_projection.cph_basis(_KIND_INFOS["kb_row"])
    assert note_only["label"] == "assumed industry cost per hire of $5,000\u2013$8,000 (Nova estimate)"
    assert "general roles" not in note_only["label"]


# ----------------------------------------- observed rows beat a long caveat
def test_short_note_is_one_line_unaudited_and_marker_free():
    short = plan_trust.BENCHMARK_FOOTNOTE_SHORT
    assert short == "Industry ranges: Nova working estimates, unaudited."
    assert not us_only_markers.has_us_only_marker(short)
    text, lines = ppt_generator._bench_source_text(
        {"confidence": "curated"}, False, [], 4.9, 8.0, short=True
    )
    assert text == short and lines == 1
    # a cited local publisher still follows it
    local, _ = ppt_generator._bench_source_text(
        {"confidence": "curated"}, False, ["Eurojob-Consulting (Germany)"], 4.9, 8.0, short=True
    )
    assert local.startswith(short.rstrip(".")) and "local CPH: Eurojob-Consulting" in local, local


def _slide6_two_role(observed: bool = True):
    import tests.test_deck_roles_round2 as roles

    return roles, roles._slide5(roles._deck(roles._plan(roles.US_TWO_ROLE, observed=observed)))


def test_dense_two_role_plan_keeps_both_observed_rows_with_the_one_line_note():
    roles, slide = _slide6_two_role()
    texts = [s.text_frame.text for s in roles._shapes_with_text(slide)]
    assert any(t.startswith("Live Postings: Registered Nurse") for t in texts), texts
    assert any(t.startswith("Live Postings: Travel Nurse") for t in texts), texts
    assert plan_trust.BENCHMARK_FOOTNOTE_SHORT in texts, texts
    assert _SECTION_FOOTNOTE not in texts, texts
    # the one-line note sits clear of the attribution heading
    note = next(s for s in roles._shapes_with_text(slide) if s.text_frame.text == plan_trust.BENCHMARK_FOOTNOTE_SHORT)
    heading = next(s for s in roles._shapes_with_text(slide) if s.text_frame.text == "CHANNEL CATEGORY ATTRIBUTION")
    assert (note.top + note.height) / 914400 <= heading.top / 914400, (note.top, heading.top)


def test_full_two_line_note_is_used_whenever_the_table_has_room(healthcare):
    texts = _texts(_deck_slide(healthcare["pptx_bytes"], "Channel Strategy & Investment"))
    assert _SECTION_FOOTNOTE in texts, texts
    assert plan_trust.BENCHMARK_FOOTNOTE_SHORT not in texts, texts
    assert "Apply Rate" in texts, texts


# ------------------ narrative facts carry the plan's RESOLVED cost per hire
def _capture_narrative_prompt(build):
    """Run ``build()`` with the LLM router blocked and return
    ``(FACTS block the model would be given, build's result)``."""
    import llm_router

    captured: Dict[str, str] = {}

    def blocked(*args, **kwargs):
        captured["prompt"] = kwargs["messages"][0]["content"]
        return {"text": "", "error": "blocked", "attempts": []}

    mp = pytest.MonkeyPatch()
    mp.setattr(llm_router, "call_llm", blocked)
    try:
        result = build()
    finally:
        mp.undo()
    prompt = captured.get("prompt", "")
    head = "invent any other number):\n"
    facts = prompt.split(head, 1)[1].split("\n\nWrite as", 1)[0] if head in prompt else ""
    return facts, result


def _brief_case(industry: str, **over):
    brief = dict(regen.ATRIA_BRIEF, industry=industry)
    brief.update(over)
    return lambda: _bundle(brief)


def _berlin_local_case():
    from tests.test_plan_trust import _berlin_plan
    import research

    def build():
        data = _berlin_plan()
        return {"data": data, "xlsx_bytes": excel_v2.generate_excel_v2(
            data, research_mod=research, load_kb_fn=kb_loader.load_knowledge_base)}

    return build


_FACT_RE = r"^(Assumed industry cost per hire \(.+?\)|Assumed cross-industry cost per hire \(.+?\)|Cited local cost per hire \(.+?\)): (.+)$"
_LEGAL = dict(
    client_name="Harbor & Finch Legal LLP",
    roles=["Paralegal", "Legal Assistant"],
    target_roles=[{"title": "Paralegal", "count": 10, "tier": "Hourly"}],
)
_NARRATIVE_CASES = {
    "legal": (_brief_case("Legal Services", **_LEGAL), "nova_table"),
    "retail": (_brief_case("Retail & Consumer"), "nova_table"),
    "healthcare": (_brief_case("Healthcare & Medical"), "nova_table"),
    "food_beverage": (_brief_case("Food & Beverage"), "kb_row"),
    "telecom": (_brief_case("Telecommunications"), "kb_row"),
    "berlin_local": (_berlin_local_case(), "local_kb"),
}
_KIND_LABEL = {
    "nova_table": "(Nova table, unsourced)",
    "kb_row": "(Nova estimate",
    "local_kb": "Cited local cost per hire (",
}


@pytest.mark.parametrize("case", sorted(_NARRATIVE_CASES))
def test_narrative_cph_fact_is_the_resolved_figure_the_summary_prints(case, _plan_money):
    import re

    build, kind = _NARRATIVE_CASES[case]
    facts, bundle = _capture_narrative_prompt(build)
    data = bundle["data"]
    info = data["_budget_allocation"]["metadata"]["industry_cph"]
    assert plan_projection.cph_basis(info)["kind"] == kind, (case, info)
    lines = [ln for ln in facts.splitlines() if re.match(_FACT_RE, ln)]
    assert len(lines) == 1, (case, facts)
    label, figure = re.match(_FACT_RE, lines[0]).groups()
    assert _KIND_LABEL[kind] in label, (case, label)
    # the figure is the RESOLVED one -- the same text Executive Summary B19 prints
    _plan_money(str(info.get("currency") or "USD") if kind == "local_kb" else "USD",
                "\u20ac80,000" if kind == "local_kb" else "$100,000")
    resolved = excel_v2._resolved_industry_cph(info)
    assert figure == excel_v2._fmt_industry_cph(resolved, info), (case, figure, resolved)
    wb = load_workbook(io.BytesIO(bundle["xlsx_bytes"]))
    b19 = [c.value for row in wb["Executive Summary"].iter_rows() for c in row
           if isinstance(c.value, str) and "if every hire costs" in c.value]
    assert b19 and f"if every hire costs {figure}" in b19[0], (case, figure, b19)
    # and the grounding gate accepts it
    allowed = excel_v2._allowed_numbers_from_facts_text(facts)
    ok, bad = excel_v2._narrative_is_grounded(f"Hires are costed at {figure} each.", allowed)
    assert ok, (case, bad)
    if case == "legal":
        # the plan's Nova-table midpoint, not the knowledge-base row's $9,000
        assert figure == "$8,000", figure
        ok, bad = excel_v2._narrative_is_grounded("The industry cost per hire is $9,000.", allowed)
        assert not ok and bad, (ok, bad)
    if case == "retail":
        assert figure == "$3,350", figure  # not the KB row's $2,700


def test_narrative_has_no_cph_fact_for_a_market_with_no_local_benchmark():
    facts, _ = _capture_narrative_prompt(
        _brief_case(
            "Automotive & Manufacturing",
            client_name="Schwarzwald Automotive Systems GmbH",
            budget="\u20ac300,000",
            locations=["Berlin, Germany", "Munich, Germany"],
            roles=["Assembly Technician"],
            target_roles=[{"title": "Assembly Technician", "count": 20, "tier": "Hourly"}],
            competitors=[],
        )
    )
    assert facts, "the narrative prompt was never built"
    assert "cost per hire" not in facts.lower() or "no local" in facts.lower(), facts
    assert not [ln for ln in facts.splitlines() if ln.startswith(("Assumed", "Cited"))], facts

