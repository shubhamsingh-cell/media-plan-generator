"""B-39: a non-US plan must not carry the US "National Economic Snapshot".

Shipped defect: excel_v2's Market Intelligence sheet printed the US
snapshot (JOLTS-derived openings, US unemployment, US participation -- the
hardcoded fallback "~8.0M" / "62.5%") on every plan, while the sibling
Industry Metrics block was already gated on plan_geo.is_us_plan. A non-US
plan now gets the cited, country-matched indicators the knowledge base holds
for its market (industry_reports_2026.json), or an honest note.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

import excel_v2
import research as research_mod
import tools_regen_bundles as T


def _snapshot(locations, budget, roles, industry) -> str:
    brief = {
        "client_name": "Northwind Health",
        "budget": budget,
        "campaign_duration": "3-6 months",
        "industry": industry,
        "locations": list(locations),
        "roles": list(roles),
        "target_roles": [{"title": r, "count": 5} for r in roles],
    }
    data = T.build_plan_data(brief)
    xlsx = excel_v2.generate_excel_v2(dict(data), research_mod=research_mod)
    xlsx = xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx
    ws = load_workbook(io.BytesIO(xlsx))["Market Intelligence"]
    rows = []
    for row in ws.iter_rows(max_row=14):
        vals = [str(c.value) for c in row if c.value is not None]
        if vals:
            rows.append(" | ".join(vals))
    return "\n".join(rows)


@pytest.mark.parametrize(
    "location,budget,roles,industry,cited",
    [
        ("London, UK", "£500,000", ["Software Engineer"], "technology", "(ONS, 2026)"),
        ("Bangalore, India", "₹20,000,000", ["Registered Nurse"], "healthcare", "Naukri"),
        ("Auckland, New Zealand", "NZ$400,000", ["Registered Nurse"], "healthcare", "SEEK"),
    ],
)
def test_non_us_plan_gets_cited_local_indicators(location, budget, roles, industry, cited):
    snap = _snapshot([location], budget, roles, industry)
    assert "National Economic Snapshot" not in snap
    assert "8.0M" not in snap and "62.5%" not in snap
    assert "Local Market Indicators" in snap
    assert cited in snap


def test_non_us_plan_without_sourced_local_indicator_gets_a_note():
    snap = _snapshot(["Sao Paulo, Brazil"], "R$500,000", ["Cashier"], "retail")
    assert "National Economic Snapshot" not in snap
    assert "8.0M" not in snap
    assert "No sourced national labour-market indicators" in snap


def test_us_plan_keeps_national_economic_snapshot():
    snap = _snapshot(["Columbus, OH"], "$250,000", ["Cashier"], "retail")
    assert "National Economic Snapshot" in snap
    assert "Local Market Indicators" not in snap
