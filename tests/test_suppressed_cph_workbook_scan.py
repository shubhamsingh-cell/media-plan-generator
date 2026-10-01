"""Regression: a plan with NO verified local cost-per-hire benchmark prints
no per-hire figure anywhere in the workbook (numbers verifier round 2,
2026-10-01, item 1; design-judge round 2, item 2).

Pre-fix (branch at 2f92a64), Brazil R$1.2M: the Executive Summary card
said "--" but ROI Projections D4 printed 8759.12 ("Avg Cost/Hire") and
H8:H12 a per-channel cost per hire, the Confidence Intervals sheet printed
pessimistic/expected/optimistic cost per hire (D:F, 15 cells), and the
methodology line (ROI Projections) and Niche Board Matching D3 said hire
totals are "anchored to industry cost-per-hire benchmarks" -- for a market
that has none. Every such figure is bounded by an FX-translated US value.
"""

from __future__ import annotations

import io
import json
import os
import re
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import excel_v2  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

_PER_HIRE = re.compile(r"cost per hire|cost/hire|\bCPH\b|cost-per-hire", re.I)


def _brief(name, industry, budget, locations, roles, hire_volume):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": hire_volume,
        "work_environment": "onsite",
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
    }


_BRIEFS = {
    "brazil": _brief(
        "Probe Contact BR", "general_entry_level", "R$1,200,000",
        ["São Paulo, Brazil", "Rio de Janeiro, Brazil"],
        ["Customer Service Representative", "Telemarketing Agent"], "500+ hires"),
    "uk": _brief(
        "Probe Bank", "finance_banking", "£150,000", ["London, UK"],
        ["Financial Analyst", "Accountant"], "50 hires"),
    "japan": _brief(
        "Probe Denki", "tech_engineering", "¥45,000,000", ["Tokyo, Japan"],
        ["Software Engineer"], "60 hires"),
    "germany_auto": _brief(
        "Probe Auto", "automotive_manufacturing", "€400,000", ["Stuttgart, Germany"],
        ["Production Technician", "Quality Engineer"], "80 hires"),
}


def _enriched(data):
    """What the live /api/generate path attaches and T.build_plan_data does
    not: the international benchmarks dataset (Intl Benchmarks sheet) and
    the synthesizer's workforce insights carrying the US KB's Appcast
    occupation / full-funnel costs (Market Intelligence) -- the two places
    the numbers verifier (round 3, N1) found per-hire figures."""
    path = os.path.join(PROJECT_ROOT, "data", "international_benchmarks_2026.json")
    with open(path, encoding="utf-8") as fh:
        data["_intl_benchmarks"] = json.load(fh)
    synth = dict(data.get("_synthesized") or {})
    synth["workforce_insights"] = {
        "appcast_2026_benchmarks": {
            "data_year": 2025,
            "occupation_benchmarks": {
                "occupation_key": "finance", "cpa": "US$15.55", "cph": "US$1,244",
            },
            "full_funnel_costs": {
                "overall_median_cpa": "US$19.20", "overall_median_cph": "US$1,053",
            },
        }
    }
    data["_synthesized"] = synth
    return data


@pytest.fixture(scope="module", params=sorted(_BRIEFS))
def suppressed_wb(request):
    data = T.build_plan_data(_BRIEFS[request.param])
    assert data["_budget_allocation"]["metadata"]["industry_cph"]["claim_suppressed"]
    _enriched(data)
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    return load_workbook(io.BytesIO(raw))


# A money amount in a per-hire context, inside ONE text cell:
# "CPH: US$1,244", "Overall Median CPH: US$1,053", "R$8,759/hire",
# "Entry Level: $3,175" under a "CPH by Tier" header (checked separately).
_MONEY = r"(?:US\$|R\$|[$£€¥₹])\s?\d[\d,.]*\s*[KMk]?"
_PER_HIRE_TEXT = re.compile(
    rf"(?:\bCPH\b|cost[- ]per[- ]hire|cost per hire)[^;|\n]{{0,25}}?{_MONEY}"
    rf"|{_MONEY}\s*(?:/hire|per hire)",
    re.I,
)


def _per_hire_text_cells(wb):
    hits = []
    for ws in wb.worksheets:
        grid = {
            (c.row, c.column): c.value
            for row in ws.iter_rows()
            for c in row
            if c.value not in (None, "")
        }
        for (r, col), v in grid.items():
            if not isinstance(v, str):
                continue
            hdr = next(
                (grid[(rr, col)] for rr in range(r - 1, max(0, r - 60), -1)
                 if isinstance(grid.get((rr, col)), str) and len(grid[(rr, col)]) < 40),
                "",
            )
            if _PER_HIRE_TEXT.search(v) or (
                _PER_HIRE.search(hdr) and re.search(_MONEY, v)
            ):
                hits.append((ws.title, ws.cell(row=r, column=col).coordinate, v[:120]))
    return hits


def test_no_per_hire_figure_in_any_text_cell(suppressed_wb):
    assert _per_hire_text_cells(suppressed_wb) == []


def test_no_sheet_promises_cph_detail(suppressed_wb):
    texts = [t[2] for t in _all_text(suppressed_wb)]
    assert not any("CPC/CPA/CPH detail" in t for t in texts)


def _per_hire_numbers(wb):
    """Numeric cells whose column header (nearest string above, <=40 rows)
    or row label (nearest string to the left) names a per-hire metric."""
    hits = []
    for ws in wb.worksheets:
        grid = {
            (c.row, c.column): c.value
            for row in ws.iter_rows()
            for c in row
            if c.value not in (None, "")
        }
        for (r, col), v in grid.items():
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            hdr = next(
                (grid[(rr, col)] for rr in range(r - 1, max(0, r - 40), -1)
                 if isinstance(grid.get((rr, col)), str)),
                "",
            )
            lab = next(
                (grid[(r, cc)] for cc in range(col - 1, 0, -1)
                 if isinstance(grid.get((r, cc)), str)),
                "",
            )
            if _PER_HIRE.search(hdr) or _PER_HIRE.search(lab):
                hits.append((ws.title, ws.cell(row=r, column=col).coordinate, v))
    return hits


def _all_text(wb):
    return [
        (ws.title, c.coordinate, c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str)
    ]


def test_no_per_hire_number_on_any_sheet(suppressed_wb):
    assert _per_hire_numbers(suppressed_wb) == []


def test_no_money_per_hire_in_text(suppressed_wb):
    # an amount quoted per hire: "R$8,759/hire", "£3,261 per hire",
    # "blended cost of R$8,759.12 per hire", "cost per hire of ¥789,474"
    per_hire_amount = re.compile(
        r"(?:US\$|R\$|[$£€¥₹])\s?[\d][\d,.]*\s*[KMk]?\s*(?:/hire|per hire)"
        r"|cost[- ]per[- ]hire (?:of|is|at) (?:US\$|R\$|[$£€¥₹])",
        re.I,
    )
    bad = [t for t in _all_text(suppressed_wb) if per_hire_amount.search(t[2])]
    assert bad == []


def test_methodology_does_not_claim_a_benchmark(suppressed_wb):
    texts = [t[2] for t in _all_text(suppressed_wb)]
    assert not any("anchored to industry cost-per-hire" in t for t in texts)
    assert not any("anchors total hires to cost-per-hire" in t for t in texts)
    assert any(
        t.startswith("Methodology: this market has no verified local cost-per-hire")
        for t in texts
    )


def test_one_indicative_note_per_sheet(suppressed_wb):
    for name in ("ROI Projections", "Confidence Intervals"):
        notes = [
            c.value for row in suppressed_wb[name].iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("Cost per hire: —")
        ]
        assert len(notes) == 1, (name, notes)
        assert suppressed_wb["ROI Projections"]["D4"].value == "—"
