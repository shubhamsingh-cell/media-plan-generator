"""Regression: projected hires carry their range, a runaway single-metro
projection is flagged, and the deck/workbook state the cost-per-hire basis
actually used (audit 2026-10-01, F §3.1 / §3.2 / §4.1 / §4.2).

Pre-fix (f99beef):
* ``total_projected`` had a single hires point estimate; nothing said it is
  budget / (0.5 x industry average) in ~98% of plans.
* A $5M single-metro trades plan projected 2,197 hires with ``warnings: []``.
* Every non-USD deck printed "US-calibrated benchmarks (US$) not
  FX-converted -- hire and CPH projections assume parity", false whenever
  the engine had divided by the market's rate (India: US$0.012/INR).
* A plan whose market has no local cost-per-hire benchmark still printed
  "N hires at X/hire" with X an FX-translated US constant.
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import budget_engine as be  # noqa: E402
import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

_CHANNELS = {
    "programmatic_dsp": 35,
    "global_boards": 20,
    "niche_boards": 15,
    "social_media": 12,
    "regional_boards": 8,
    "employer_branding": 5,
}


def _alloc(industry, role, city, state, country, budget, currency="USD"):
    return be.calculate_budget_allocation(
        total_budget=budget,
        roles=[{"title": role, "count": 1, "tier": "Professional"}],
        locations=[{"city": city, "state": state, "country": country}],
        industry=industry,
        channel_percentages=dict(_CHANNELS),
        knowledge_base=None,
        plan_currency=currency,
    )


class TestHiresRange:
    def test_range_brackets_the_headline(self):
        tp = _alloc(
            "healthcare_medical", "Registered Nurse", "Dallas", "TX",
            "United States", 250_000.0,
        )["total_projected"]
        assert tp["hires_low"] == int(250_000 / 10_500)  # at the average
        assert tp["hires_high"] == int(250_000 / 5_250)  # at the floor
        assert tp["hires_low"] <= tp["hires"] <= tp["hires_high"]

    def test_range_clamped_to_contain_a_one_hire_minimum(self):
        tp = _alloc(
            "healthcare_medical", "Registered Nurse", "Dallas", "TX",
            "United States", 1_000.0,
        )["total_projected"]
        assert tp["hires_low"] <= tp["hires"] <= tp["hires_high"]


class TestDepthGuard:
    def test_runaway_single_metro_projection_is_flagged(self):
        res = _alloc(
            "blue_collar_trades", "Machine Operator", "Cleveland", "OH",
            "United States", 5_000_000.0,
        )
        assert res["total_projected"]["hires"] > 1_000
        assert any("depth guard" in w for w in res["warnings"]), res["warnings"]

    def test_ordinary_plan_not_flagged(self):
        res = _alloc(
            "blue_collar_trades", "Machine Operator", "Cleveland", "OH",
            "United States", 250_000.0,
        )
        assert not any("depth guard" in w for w in res["warnings"])


def _brief(name, industry, budget, loc, roles, hire_volume="Not specified"):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": hire_volume,
        "work_environment": "onsite",
        "locations": [loc],
        "roles": roles,
        "target_roles": roles,
    }


def _deck_lines(data):
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append((i, shape.text_frame.text.replace("\n", " | ")))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append((i, " ;; ".join(c.text for c in row.cells)))
    return out


def _xlsx_text(data):
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    if isinstance(raw, tuple):
        raw = raw[0]
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Executive Summary"]
    return [
        [c.value for c in row if c.value not in (None, "")]
        for row in ws.iter_rows()
    ]


@pytest.fixture(scope="module")
def dallas():
    return T.build_plan_data(
        _brief("Probe Health", "Healthcare & Medical", "$250,000", "Dallas, TX",
               ["Registered Nurse", "Pharmacist"], "60 hires")
    )


@pytest.fixture(scope="module")
def bangalore():
    return T.build_plan_data(
        _brief("Probe Care India", "Healthcare & Medical", "₹20,000,000",
               "Bangalore, India", ["Registered Nurse", "Staff Nurse"])
    )


@pytest.fixture(scope="module")
def tokyo():
    return T.build_plan_data(
        _brief("Probe Support JP", "Customer Service", "¥30,000,000",
               "Tokyo, Japan", ["Customer Service Representative"])
    )


class TestDeck:
    def test_slide2_prints_the_range(self, dallas):
        lines = [t for s, t in _deck_lines(dallas) if s == 2]
        assert any(
            "23 at industry-avg cost | 47 at plan efficiency" in t for t in lines
        ), lines

    def test_slide5_row_is_the_one_benchmark(self, dallas):
        lines = [t for s, t in _deck_lines(dallas) if s == 5]
        # one-line compact form of "$9,000–$12,000 (midpoint $10,500)"
        assert any("$9K–$12K (midpoint $10.5K)" in t for t in lines), lines

    def test_local_market_footnote_names_the_rate_not_parity(self, bangalore):
        notes = [t for _, t in _deck_lines(bangalore) if t.startswith("Figures in INR")]
        assert notes, "no currency-basis footnote on an INR deck"
        for t in notes:
            assert "assume parity" not in t, t
            assert "US$1 = ₹" in t and "local-market" in t, t

    def test_no_local_benchmark_drops_the_per_hire_claim(self, tokyo):
        lines = _deck_lines(tokyo)
        thesis = [t for s, t in lines if s == 2 and "This plan projects" in t]
        assert thesis and "/hire" not in thesis[0], thesis
        assert any("No local CPH benchmark" in t for _, t in lines)


class TestWorkbook:
    def test_summary_prints_the_range(self, dallas):
        rows = _xlsx_text(dallas)
        assert any(
            isinstance(v, str) and v.startswith("Projected hires: 23–47.")
            for r in rows
            for v in r
        ), rows[:20]

    def test_suppressed_cost_per_hire_card(self, tokyo):
        rows = _xlsx_text(tokyo)
        labels = next(i for i, r in enumerate(rows) if "Cost / Hire" in r)
        values = rows[labels - 1]
        assert values[-1] == "--", values
