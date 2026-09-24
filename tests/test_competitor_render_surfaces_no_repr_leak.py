"""Every client-facing rendering surface must render a competitor entry as
its NAME -- never as Python repr text -- whatever shape the entry arrived in.

Hershey Company defect chain (2026-09-24), round 6. Rounds 1-5 each fixed
one rendering surface (excel_v2 Market Intelligence, gold_standard padding,
excel_v2 Quality Intelligence, the app.py request boundary, ppt_generator's
metadata survival) and each follow-up review then found the SAME leak on a
surface nobody had looked at. Round 5 found it on the Google Slides deck
(joveo_slides_template._slide_benchmarking_1 did ``f"- {c}"`` on the raw
entry), which is Tier 1 of deck_generator.DeckGenerator -- the deck real
clients get whenever Google Slides succeeds in production.

This file closes the class instead of the instance, three ways:

1. ``test_surface_renders_competitors_without_repr_leak`` renders the WHOLE
   output of every live surface -- the workbook (excel_v2), its legacy
   fallback (archive/excel_legacy, used when excel_v2 raises), deck Tier 1
   (Google Slides, driven through the real DeckGenerator with a faked Google
   client), deck Tier 2 (python-pptx), and the PDF export -- for each input
   shape a caller can send, and scans ALL rendered text for repr markers.
   It does not look only where previous rounds looked.
2. ``test_deck_tier_ladder_is_fully_covered`` pins deck_generator's tier
   list, so adding a tier fails here until that tier is added to (1).
3. ``test_every_competitors_read_site_is_normalized_or_justified`` scans the
   source of every production module for a read of a ``competitors`` key.
   Each read must either go through the shared_utils helpers or be listed in
   an allowlist with the reason it is safe. A NEW unnormalized read site --
   the thing that produced rounds 2-6 -- fails this test when it is written,
   not when a client sees it.

Runs under pytest, or standalone:
``python3 tests/test_competitor_render_surfaces_no_repr_leak.py``.
"""

from __future__ import annotations

import base64
import copy
import io
import json
import re
import sys
import types
from pathlib import Path
from typing import Any, Callable, Dict, List

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402
from pptx import Presentation  # noqa: E402

import shared_utils  # noqa: E402
import tools_regen_bundles as trb  # noqa: E402

# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
_MARS_DESC = "Global confectionery and pet-care maker"

# What a direct API caller can send today: dict entries carrying metadata,
# plus plain strings.
_DICT_ENTRIES: List[Any] = [
    {
        "name": "Mars Wrigley",
        "description": _MARS_DESC,
        "domain": "mars.com",
        "competitor_type": "direct",
    },
    {"name": "Mondelez International", "domain": "mondelezinternational.com"},
    "Ferrero",
]

# Malformed-but-possible shapes that used to be stringified or dropped.
_HOSTILE_ENTRIES: List[Any] = [
    {"name": "Mars Wrigley", "description": _MARS_DESC, "domain": "mars.com"},
    ["Mondelez International"],  # nested list -> used to become "['Mondelez...']"
    {"name": {"en": "Nested Name Co"}},  # non-string name -> repr if str()'d
    {"name": "Ferrero", "description": {"long": "nested metadata"}},
    "  ",
    None,
]

# A mapping instead of a list (name -> metadata). Used to be dropped to [].
_MAPPING_INPUT: Dict[str, Any] = {
    "Mars Wrigley": {"description": _MARS_DESC, "domain": "mars.com"},
    "Mondelez International": {"domain": "mondelezinternational.com"},
    "Ferrero": {},
}

_EXPECTED_NAMES = ("Mars Wrigley", "Mondelez International")

# (label, raw competitors value, apply the app.py request boundary?)
_INPUT_CASES = [
    ("dict_entries_via_boundary", _DICT_ENTRIES, True),
    ("dict_entries_direct_call", _DICT_ENTRIES, False),
    ("hostile_entries_via_boundary", _HOSTILE_ENTRIES, True),
    ("hostile_entries_direct_call", _HOSTILE_ENTRIES, False),
    ("mapping_via_boundary", _MAPPING_INPUT, True),
    ("mapping_direct_call", _MAPPING_INPUT, False),
]

# Text that can only come from str()/repr() of a dict or list.
_REPR_LEAK_RE = re.compile(
    r"\{'|\['|'\}|'\]|'name'|'description'|'domain'|'competitor_type'"
    r"|'en'|'long'|\bNested Name Co\b|nested metadata"
)


def _hershey_brief(competitors: Any) -> Dict[str, Any]:
    return {
        "client_name": "The Hershey Company",
        "requester_name": "QA",
        "requester_email": "qa@joveo.com",
        "industry": "Consumer Packaged Goods",
        "budget": "$250,000",
        "campaign_duration": "6 months",
        "hire_volume": "100-500 hires",
        "work_environment": "onsite",
        "locations": ["Hershey, PA"],
        "roles": ["Production Operator"],
        "target_roles": [{"title": "Production Operator", "count": 120, "tier": "Hourly"}],
        "competitors": copy.deepcopy(competitors),
    }


_PLAN_CACHE: Dict[str, Dict[str, Any]] = {}


def _plan_data(label: str, raw: Any, via_boundary: bool) -> Dict[str, Any]:
    """The ``data`` dict app.py hands to the generators, built by the same
    pipeline replication the bundle-QA regression tests use. With
    ``via_boundary`` the competitors value first goes through the exact call
    app.py's /api/generate handler makes at its request boundary."""
    if label not in _PLAN_CACHE:
        brief = _hershey_brief(raw)
        if via_boundary:
            brief["competitors"] = shared_utils.clean_competitor_entries(
                brief["competitors"]
            )
        _PLAN_CACHE[label] = trb.build_plan_data(brief)
    return copy.deepcopy(_PLAN_CACHE[label])


# ---------------------------------------------------------------------------
# Text extraction per surface
# ---------------------------------------------------------------------------
def _xlsx_texts(blob: bytes) -> List[str]:
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    out = []
    for ws in wb.worksheets:
        out.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if val is not None:
                    out.append(str(val))
    return out


def _pptx_texts(blob: bytes) -> List[str]:
    prs = Presentation(io.BytesIO(blob))
    out = []

    def _walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6:  # group
                _walk(shape.shapes)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows:
                    for c in r.cells:
                        out.append(c.text)

    for slide in prs.slides:
        _walk(slide.shapes)
        if slide.has_notes_slide:
            out.append(slide.notes_slide.notes_text_frame.text)
    return out


def _pdf_texts(blob: bytes) -> List[str]:
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(blob))
    return [page.extract_text() or "" for page in reader.pages]


# ---------------------------------------------------------------------------
# Surfaces
# ---------------------------------------------------------------------------
def _render_excel_v2(data: Dict[str, Any], _mp) -> List[str]:
    import excel_v2
    from kb_loader import load_knowledge_base

    return _xlsx_texts(excel_v2.generate_excel_v2(data, load_kb_fn=load_knowledge_base))


def _render_excel_legacy(data: Dict[str, Any], _mp) -> List[str]:
    # app.py's sync /api/generate path falls back to this generator whenever
    # excel_v2 raises (and the async path uses it when excel_v2 is absent).
    from archive.excel_legacy import generate_excel

    return _xlsx_texts(generate_excel(data))


class _FakeCall:
    def __init__(self, result: Any) -> None:
        self._result = result

    def execute(self) -> Any:
        return self._result


class _FakeSlides:
    def __init__(self, sink: List[dict]) -> None:
        self._sink = sink

    def presentations(self) -> "_FakeSlides":
        return self

    def create(self, body: dict) -> _FakeCall:
        return _FakeCall({"presentationId": "fake", "slides": [{"objectId": "s0"}]})

    def batchUpdate(self, presentationId: str, body: dict) -> _FakeCall:  # noqa: N802
        self._sink.extend(body["requests"])
        return _FakeCall({})


class _FakeDrive:
    def files(self) -> "_FakeDrive":
        return self

    def export(self, fileId: str, mimeType: str) -> _FakeCall:  # noqa: N803
        return _FakeCall(b"PK" + b"\0" * 4096)

    def delete(self, fileId: str) -> _FakeCall:  # noqa: N803
        return _FakeCall(None)


def _install_fake_google_client(mp, sink: List[dict]) -> None:
    """Stand in for google-api-python-client / google-auth (installed on
    Render per requirements.txt, absent locally) so DeckGenerator's REAL
    Tier-1 code path runs end to end and hands its batchUpdate requests --
    i.e. every string the Google Slides deck will contain -- to ``sink``."""
    discovery = types.ModuleType("googleapiclient.discovery")
    discovery.build = lambda name, _v, credentials=None: (
        _FakeSlides(sink) if name == "slides" else _FakeDrive()
    )
    gapi = types.ModuleType("googleapiclient")
    gapi.discovery = discovery
    service_account = types.ModuleType("google.oauth2.service_account")

    class _Creds:
        @classmethod
        def from_service_account_info(cls, info, scopes=None):
            return cls()

    service_account.Credentials = _Creds
    oauth2 = types.ModuleType("google.oauth2")
    oauth2.service_account = service_account
    mp.setitem(sys.modules, "googleapiclient", gapi)
    mp.setitem(sys.modules, "googleapiclient.discovery", discovery)
    try:
        import google  # noqa: F401  (real namespace package, if present)
    except ImportError:
        google_pkg = types.ModuleType("google")
        google_pkg.__path__ = []  # type: ignore[attr-defined]
        mp.setitem(sys.modules, "google", google_pkg)
    mp.setitem(sys.modules, "google.oauth2", oauth2)
    mp.setitem(sys.modules, "google.oauth2.service_account", service_account)
    mp.setenv(
        "GOOGLE_SLIDES_CREDENTIALS_B64",
        base64.b64encode(json.dumps({"type": "service_account"}).encode()).decode(),
    )


def _render_deck_tier1_google_slides(data: Dict[str, Any], mp) -> List[str]:
    import deck_generator

    sink: List[dict] = []
    _install_fake_google_client(mp, sink)
    _blob, provider = deck_generator.DeckGenerator().generate(data)
    assert provider == "google_slides", (
        f"Tier 1 did not run (provider={provider!r}); the fake Google client "
        "no longer matches deck_generator's calls -- this surface is untested"
    )
    return [r["insertText"]["text"] for r in sink if "insertText" in r]


def _render_deck_tier2_pptx(data: Dict[str, Any], mp) -> List[str]:
    import deck_generator

    mp.delenv("GOOGLE_SLIDES_CREDENTIALS_B64", raising=False)
    mp.delenv("GOOGLE_SLIDES_CREDENTIALS", raising=False)
    blob, provider = deck_generator.DeckGenerator().generate(data)
    assert provider == "pptx", f"expected the python-pptx tier, got {provider!r}"
    return _pptx_texts(blob)


def _render_pdf_export(data: Dict[str, Any], _mp) -> List[str]:
    # POST /api/export/pdf renders the posted plan payload directly -- it
    # never passes through app.py's /api/generate request boundary.
    import pdf_report

    return _pdf_texts(
        pdf_report.generate_pdf_report(
            plan_data=data, client_name="The Hershey Company", industry="CPG"
        )
    )


# (surface id, renderer, tier key it covers or None)
_SURFACES: List[tuple] = [
    ("excel_v2_workbook", _render_excel_v2, None),
    ("excel_legacy_fallback_workbook", _render_excel_legacy, None),
    ("deck_tier1_google_slides", _render_deck_tier1_google_slides, "google_slides"),
    ("deck_tier2_python_pptx", _render_deck_tier2_pptx, "pptx"),
    ("pdf_export", _render_pdf_export, None),
]


@pytest.mark.parametrize(
    "surface_id,render", [(s[0], s[1]) for s in _SURFACES], ids=[s[0] for s in _SURFACES]
)
@pytest.mark.parametrize(
    "label,raw,via_boundary", _INPUT_CASES, ids=[c[0] for c in _INPUT_CASES]
)
def test_surface_renders_competitors_without_repr_leak(
    surface_id: str,
    render: Callable,
    label: str,
    raw: Any,
    via_boundary: bool,
    monkeypatch,
):
    data = _plan_data(label, raw, via_boundary)
    texts = render(data, monkeypatch)
    blob = "\n".join(texts)

    leaks = [t for t in texts if _REPR_LEAK_RE.search(t)]
    assert not leaks, (
        f"{surface_id} rendered Python repr text for a {label} competitor "
        f"input: {[t[:200] for t in leaks[:5]]}"
    )
    # Non-vacuous: the client's own competitors must actually be on the
    # surface (a surface that silently drops them would pass the leak check).
    for name in _EXPECTED_NAMES:
        assert name in blob, f"{surface_id} ({label}) is missing {name!r}"


@pytest.mark.parametrize(
    "label,raw,via_boundary", _INPUT_CASES, ids=[c[0] for c in _INPUT_CASES]
)
def test_deck_tier2_cards_keep_competitor_metadata(label, raw, via_boundary):
    """Metadata must survive the repr-safety cleaning where a surface can
    show it. Google Slides slide 4's "Competitive Landscape" box is a
    4.4in x 1.8in name list (5 x 11pt lines), so it renders names only; the
    python-pptx competitor cards DO show a description and must still get
    it, for every input shape that carries one."""
    import ppt_generator

    data = _plan_data(label, raw, via_boundary)
    blob = "\n".join(_pptx_texts(ppt_generator.generate_pptx(data)))
    assert _MARS_DESC in blob, f"{label}: Mars Wrigley's description was lost"


def test_pdf_export_escapes_competitor_markup():
    """The PDF's Competitive Landscape section feeds competitor text into
    ReportLab Paragraph markup. Unescaped, a description with an unbalanced
    tag ("<b>x") raised ValueError (a 500 from POST /api/export/pdf) and a
    name like "M&M Foods <Canada>" silently vanished from the page. Names
    and descriptions are client text, never markup."""
    import pdf_report

    plan = {
        "budget": 100000,
        "competitors": [
            {"name": "Mars Wrigley", "description": "<b>unbalanced tag description"},
            "M&M Foods <Canada>",
            {"name": "Procter & Gamble <Global>", "description": "a < b & c > d"},
        ],
    }
    blob = "\n".join(
        _pdf_texts(
            pdf_report.generate_pdf_report(
                plan_data=plan, client_name="The Hershey Company", industry="CPG"
            )
        )
    )
    assert "<b>unbalanced tag description" in blob
    assert "M&M Foods <Canada>" in blob
    assert "Procter & Gamble <Global>" in blob
    assert "a < b & c > d" in blob


def test_deck_tier_ladder_is_fully_covered():
    """If deck_generator grows a tier, it must be added to _SURFACES above --
    otherwise it is a client-facing surface this file never renders."""
    import deck_generator

    tiers = [key for key, _name in deck_generator._TIERS]
    covered = [s[2] for s in _SURFACES if s[2]]
    assert tiers == covered, (
        f"deck_generator tiers {tiers} != tiers rendered by this test {covered}"
    )


def test_joveo_slides_targeting_builder_is_safe_too():
    """_slide_targeting is not in build_joveo_slides' current 8-slide list,
    but it is a builder in the same module one list edit away from live --
    it must not reintroduce the leak when re-enabled."""
    import joveo_slides_template as jst

    _sid, reqs = jst._slide_targeting(
        {"competitors": copy.deepcopy(_HOSTILE_ENTRIES), "roles": ["Operator"]}
    )
    texts = [r["insertText"]["text"] for r in reqs if "insertText" in r]
    assert not [t for t in texts if _REPR_LEAK_RE.search(t)], texts
    assert "Mars Wrigley" in texts and "Mondelez International" in texts


# ---------------------------------------------------------------------------
# Secondary products that take a competitor list from their own API input
# ---------------------------------------------------------------------------
def test_competitive_intel_accepts_dict_and_string_competitor_inputs(monkeypatch):
    """POST /api/competitive/analyze and Nova's analyze_competitors tool both
    pass caller-supplied competitor lists straight in; a dict entry used to
    raise AttributeError ('dict' has no attribute 'strip') and a bare string
    was iterated character by character."""
    import competitive_intel as ci

    seen: Dict[str, Any] = {}
    monkeypatch.setattr(ci, "analyze_company", lambda n: {"name": n})

    out = ci.analyze_competitors("The Hershey Company", copy.deepcopy(_HOSTILE_ENTRIES))
    assert [p["name"] for p in out["competitors"]] == [
        "Mars Wrigley",
        "Mondelez International",
        "Ferrero",
    ]

    def _capture(name):
        def _fn(*args, **kwargs):
            seen[name] = args
            return {}

        return _fn

    for fn in ("analyze_competitors", "compare_hiring_activity", "get_market_trends"):
        monkeypatch.setattr(ci, fn, _capture(fn))
    monkeypatch.setattr(ci, "compare_ad_benchmarks", lambda *a, **k: {})
    monkeypatch.setattr(ci, "generate_competitive_brief", lambda r: {"ok": True})
    result = ci.run_full_analysis("The Hershey Company", "Mars Wrigley, Ferrero")
    assert result["competitors"] == ["Mars Wrigley", "Ferrero"]
    assert seen["analyze_competitors"][1] == ["Mars Wrigley", "Ferrero"]


def test_market_intel_report_normalizes_competitor_input(monkeypatch):
    """market_intel_reports.generate_report iterated data["competitors"]
    raw: a comma string became one "competitor" per CHARACTER and a dict
    entry crashed the section (``hash(dict)``)."""
    import market_intel_reports as mir

    seen: Dict[str, Any] = {}

    def _capture(competitors, industry):
        seen["competitors"] = competitors
        return {"source": "estimate", "competitors": []}

    monkeypatch.setattr(mir, "_collect_competitor_analysis", _capture)
    for name in (
        "_collect_market_overview",
        "_collect_salary_benchmarks",
        "_collect_channel_performance",
        "_collect_talent_supply",
        "_collect_seasonal_trends",
    ):
        monkeypatch.setattr(mir, name, lambda *a, **k: {"source": "estimate"})
    mir.generate_report({"industry": "technology", "competitors": "Mars Wrigley, Ferrero"})
    assert seen["competitors"] == ["Mars Wrigley", "Ferrero"]
    mir.generate_report(
        {"industry": "technology", "competitors": copy.deepcopy(_HOSTILE_ENTRIES)}
    )
    assert seen["competitors"] == ["Mars Wrigley", "Mondelez International", "Ferrero"]


# ---------------------------------------------------------------------------
# Structural guard: every read of a "competitors" key is normalized or
# explicitly justified.
# ---------------------------------------------------------------------------
_READ_RE = re.compile(
    r"""(\.get\(\s*["']competitors["']|\[\s*["']competitors["']\s*\](?!\s*=[^=]))"""
)
_HELPER_RE = re.compile(r"normalize_competitor_names\(|clean_competitor_entries\(")

# (path, stripped source line) -> why that read is safe without the helper.
_JUSTIFIED_READS: Dict[tuple, str] = {
    ("app.py", '_competitors = data.get("competitors") or []'): (
        "request validation (list length cap) only, never rendered; runs "
        "after the boundary clean_competitor_entries call"
    ),
    ("archive/excel_legacy.py", 'style_body_cell(ws_trends, row, 3, comp.get("competitors") or "")'): (
        "research.get_competitors KB row (a static string), not client input"
    ),
    ("archive/excel_legacy.py", '_competitors = _comp_intel.get("competitors", {})'): (
        "synthesized competitive intelligence (enrichment-keyed dict)"
    ),
    ("competitive_intel.py", 'if comp_data.get("competitors"):'): (
        "data_orchestrator enrichment output, not client input"
    ),
    ("competitive_intel.py", 'result["industry_competitors"] or comp_data["competitors"]'): (
        "data_orchestrator enrichment output, not client input"
    ),
    ("competitive_intel.py", '"competitors": (brief.get("competitors") or [])[:5],  # Cap for token budget'): (
        "server-built competitor profile dicts, serialized as JSON for an LLM"
    ),
    ("competitive_intel.py", 'comp_count = len(brief.get("competitors") or [])'): (
        "length only"
    ),
    ("excel_v2.py", 'if ctx.get("competitors"):'): (
        "ctx['competitors'] is normalize_competitor_names output"
    ),
    ("excel_v2.py", """f"Named Competitors: {', '.join(str(c) for c in ctx['competitors'][:5])}\""""): (
        "ctx['competitors'] is normalize_competitor_names output"
    ),
    ("excel_v2.py", '_ci_competitors_raw = comp_intel.get("competitors")'): (
        "synthesized competitive intelligence, only name-matched to the brief"
    ),
    ("market_intel_reports.py", 'data["competitors"].append(entry)'): (
        "the report's own output dict"
    ),
    ("market_intel_reports.py", """f"Analysis covers {len(data['competitors'])} competitor(s) in the \""""): (
        "length of the report's own output list"
    ),
    ("market_intel_reports.py", 'comps = report.get("competitor_analysis", {}).get("competitors") or []'): (
        "server-built report section"
    ),
    ("market_intel_reports.py", 'comps = report_data.get("competitor_analysis", {}).get("competitors") or []'): (
        "server-built report section"
    ),
    ("nova.py", 'if precomputed.get("competitors"):'): (
        "enrichment cache, returned to the LLM as JSON"
    ),
    ("nova.py", 'pc_result["top_competitors"] = precomputed["competitors"]'): (
        "enrichment cache, returned to the LLM as JSON"
    ),
    ("nova.py", 'if enriched.get("competitors"):'): (
        "enrichment output, returned to the LLM as JSON"
    ),
    ("nova.py", 'result["top_competitors"] = enriched["competitors"]'): (
        "enrichment output, returned to the LLM as JSON"
    ),
    ("ppt_generator.py", '_ci_competitors = comp_intel.get("competitors", {})'): (
        "synthesized competitive intelligence; card loop skips non-dict values"
    ),
    ("precompute.py", '"competitors": enriched.get("competitors"),'): (
        "caches enrichment output; not rendered here"
    ),
    ("research.py", 'data["competitors"], company_name'): (
        "INDUSTRY_COMPETITORS knowledge-base row, not client input"
    ),
    ("research.py", 'intl_categories[industry]["competitors"], company_name'): (
        "international KB row, not client input"
    ),
    ("research.py", 'comp_str = comp_entry.get("competitors") or ""'): (
        "KB-derived row"
    ),
    ("research.py", 'if not result or all(not (r.get("competitors") or "").strip() for r in result):'): (
        "KB-derived rows"
    ),
    ("routes/competitive.py", 'competitors=data.get("competitors") or [],'): (
        "competitive_intel.run_full_analysis normalizes its competitors arg"
    ),
}

# Dev tooling / docs generators never serve a client. Hidden dirs are pruned
# too -- in the main checkout .claude/worktrees/ holds other sessions' trees.
_EXCLUDED_TOP_DIRS = {"tests", "scripts", "docs", "node_modules", "venv", "out"}
_EXCLUDED_FILES = {"tools_regen_bundles.py"}


def _production_python_files() -> List[Path]:
    import os

    out = []
    for root, dirs, files in os.walk(PROJECT_ROOT):
        rel_root = Path(root).relative_to(PROJECT_ROOT)
        dirs[:] = sorted(
            d
            for d in dirs
            if not d.startswith((".", "__"))
            and not (rel_root == Path(".") and d in _EXCLUDED_TOP_DIRS)
            and "site-packages" not in d
        )
        for name in sorted(files):
            if not name.endswith(".py"):
                continue
            if rel_root == Path(".") and name in _EXCLUDED_FILES:
                continue
            out.append(Path(root) / name)
    return out


def test_every_competitors_read_site_is_normalized_or_justified():
    unjustified = []
    used = set()
    for path in _production_python_files():
        rel = path.relative_to(PROJECT_ROOT).as_posix()
        try:
            lines = path.read_text(encoding="utf-8").splitlines()
        except (OSError, UnicodeDecodeError):
            continue
        for i, line in enumerate(lines):
            if not _READ_RE.search(line) or line.lstrip().startswith("#"):
                continue
            # The helper may wrap the read on the same line or open on one
            # of the two lines above (black splits long calls).
            if any(_HELPER_RE.search(l) for l in lines[max(0, i - 2) : i + 1]):
                continue
            key = (rel, line.strip())
            if key in _JUSTIFIED_READS:
                used.add(key)
                continue
            unjustified.append(f"{rel}:{i + 1}: {line.strip()}")
    assert not unjustified, (
        "Read(s) of a 'competitors' key that neither go through "
        "shared_utils.normalize_competitor_names / clean_competitor_entries "
        "nor carry a justification in _JUSTIFIED_READS -- a dict-shaped "
        "competitor entry reaching str()/f-string/join renders Python repr "
        "text on a client surface:\n  " + "\n  ".join(unjustified)
    )
    stale = sorted(set(_JUSTIFIED_READS) - used)
    assert not stale, f"_JUSTIFIED_READS entries no longer in the code: {stale}"


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
