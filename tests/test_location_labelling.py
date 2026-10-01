"""Regression tests: area-level figures are labelled, never passed off as city data.

* Non-US workbooks repeated the NATIONAL population on every city row (UK / India /
  Germany / Japan / Brazil sweep bundles): the country figure now travels as
  ``country_population`` / ``area_population`` and the Location Intelligence table
  says "country-level ... not city" (a real city value, e.g. GeoNames, wins).
* ``plan_validator _check_location_sanity: Duplicate city data detected for
  ['Hershey','Hazleton','Lancaster']`` (5 of 5 runs, 0 auto-corrected). ROOT CAUSE:
  not a demographics bug -- ``gold_standard.enrich_city_level_data`` gives every
  city with no city-level entry its STATE's multiplier / difficulty, so same-state
  towns carry byte-identical rows that the workbook printed as separate "city"
  rows. Every row now says what it describes (``geo_basis``), the validator reports
  a shared state-level fallback as informational (and still flags a real
  duplicate), and the workbook marks those rows "[PA state-level estimate]".
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from typing import Any, Dict

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import api_enrichment  # noqa: E402
import data_synthesizer  # noqa: E402
import excel_v2  # noqa: E402
import gold_standard  # noqa: E402
import plan_validator  # noqa: E402
import public_data_sources as pds  # noqa: E402
from tests.public_data_fakes import FakeNet  # noqa: E402

PA_TOWNS = ["Hershey, PA", "Hazleton, PA", "Lancaster, PA"]


@pytest.fixture
def net(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(api_enrichment, "_get_cached", lambda key: None)
    monkeypatch.setattr(api_enrichment, "_set_cached", lambda key, data: None)
    monkeypatch.delenv("CENSUS_API_KEY", raising=False)
    pds.clear_caches()
    fake = FakeNet().install(monkeypatch)
    yield fake
    pds.clear_caches()



def _rows(locations: list, roles: list | None = None) -> Dict[str, Any]:
    return gold_standard.enrich_city_level_data(
        {"locations": locations, "roles": roles or ["Cook"]}
    )


def _fingerprint(row: Dict[str, Any]) -> tuple:
    return tuple(
        row[k]
        for k in (
            "salary_multiplier",
            "estimated_salary",
            "hiring_difficulty",
            "supply_tier",
            "cost_of_living_index",
        )
    )


def test_same_state_towns_share_a_state_level_row_and_say_so() -> None:
    rows = _rows(PA_TOWNS)

    # the (real) reason the validator saw duplicates ...
    assert len({_fingerprint(rows[t.split(",")[0]]) for t in PA_TOWNS}) == 1
    # ... and the label that now discloses it
    for town in PA_TOWNS:
        row = rows[town.split(",")[0]]
        assert row["geo_basis"] == "state"
        assert row["geo_basis_area"] == "PA"
        assert row["fallback_uniform"] is False  # it IS a state signal, just not city


def test_geo_basis_distinguishes_city_state_country_and_generic() -> None:
    rows = _rows(["Philadelphia, PA", "Hershey, PA", "India", "Nowheresville"])

    assert rows["Philadelphia"]["geo_basis"] == "city"
    assert rows["Philadelphia"]["geo_basis_area"] == ""
    assert rows["Hershey"]["geo_basis"] == "state"
    assert rows["India"]["geo_basis"] == "country"
    assert rows["India"]["geo_basis_area"]
    assert rows["Nowheresville"]["geo_basis"] == "generic"
    assert rows["Nowheresville"]["fallback_uniform"] is True


def test_validator_reports_a_shared_state_fallback_as_informational() -> None:
    data = {"_gold_standard": {"city_level_data": _rows(PA_TOWNS)}}

    findings = plan_validator._check_location_sanity(data)

    assert len(findings) == 1
    finding = findings[0]
    assert finding["severity"] == "low"  # not the medium "templated data" alarm
    assert "state-level" in finding["message"] and "PA" in finding["message"]
    assert "templated" not in finding["message"]
    assert sorted(finding["cities"]) == ["Hazleton", "Hershey", "Lancaster"]
    assert finding["geo_basis"] == ["state"]


def test_validator_still_flags_a_real_duplicate() -> None:
    row = {
        "salary_multiplier": 1.1,
        "estimated_salary": 80000,
        "hiring_difficulty": 6.0,
        "supply_tier": "balanced",
        "cost_of_living_index": 110.0,
        "geo_basis": "city",
    }
    for rows in (
        {"A": dict(row), "B": dict(row)},  # two CITY-level rows that match
        {"A": {k: v for k, v in row.items() if k != "geo_basis"}, "B": dict(row)},  # unlabeled
    ):
        findings = plan_validator._check_location_sanity(
            {"_gold_standard": {"city_level_data": rows}}
        )
        assert [f["severity"] for f in findings] == ["medium"]
        assert "templated/duplicated" in findings[0]["message"]


def test_workbook_suffix_marks_only_city_rows_that_fell_back_to_their_state() -> None:
    rows = _rows(["Philadelphia, PA", "Hershey, PA", "India"])

    assert excel_v2._geo_basis_suffix(rows["Philadelphia"]) == ""
    assert excel_v2._geo_basis_suffix(rows["Hershey"]) == " [PA state-level estimate]"
    # a market that IS a country has its own country figures: not a fallback
    assert excel_v2._geo_basis_suffix(rows["India"]) == ""
    assert excel_v2._geo_basis_suffix({}) == ""  # legacy rows: no claim either way


def _quality_sheet(locations: list):
    import budget_engine
    import openpyxl

    roles = [{"title": "Cook", "count": 4, "tier": "mid"}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=90_000,
        roles=roles,
        locations=locations,
        industry="food_beverage",
        channel_percentages={"programmatic_dsp": 50, "global_boards": 30, "social_media": 20},
    )
    data = {
        "client_name": "Label Fixture Foods",
        "industry": "food_beverage",
        "budget": "$90,000",
        "campaign_duration": "9 months",
        "campaign_weeks": 39,
        "roles": ["Cook"],
        "target_roles": roles,
        "locations": locations,
        "_budget_allocation": alloc,
    }
    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    wb = openpyxl.load_workbook(io.BytesIO(excel_v2.generate_excel_v2(data)))
    return wb["Quality Intelligence"]


def _quality_sheet_text(locations: list) -> list:
    ws = _quality_sheet(locations)
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def test_real_workbook_labels_state_level_rows_in_both_city_tables() -> None:
    text = _quality_sheet_text(PA_TOWNS)

    # the supply-demand table AND the per-role salary table
    labelled = [t for t in text if "state-level estimate]" in t and t.startswith(("Hershey", "Hazleton", "Lancaster"))]
    assert len(labelled) >= 6, labelled  # 3 towns x 2 tables
    assert any("[PA state-level estimate]" in t for t in labelled)
    # and the table footnote explains it
    assert any("share identical figures" in t for t in text)


def test_real_workbook_does_not_label_a_city_level_row() -> None:
    text = _quality_sheet_text(["Philadelphia, PA"])
    assert not any("state-level estimate]" in t for t in text)


# ---------------------------------------------------------------------------
# 6. Synthesis + workbook: area-level figures are labelled, never city numbers
# ---------------------------------------------------------------------------


def _synth(enriched: Dict[str, Any], locations: list) -> Dict[str, Any]:
    return data_synthesizer.fuse_location_profiles(enriched, {}, {"locations": locations})


def test_country_figure_is_not_the_city_population_but_geonames_city_value_is_used(
    net: FakeNet,
) -> None:
    demo = api_enrichment.fetch_location_demographics(["London, UK"])
    enriched = {
        "location_demographics": demo,
        "geonames_data": {"locations": {"London, UK": {"population": 8_908_081, "latitude": "51.5", "longitude": "-0.12"}}},
    }
    profile = _synth(enriched, ["London, UK"])["London, UK"]

    assert profile["population"] == 8_908_081  # the CITY, from GeoNames
    assert profile["area_population_scope"] == "country"
    assert profile["area_population"] == demo["London, UK"]["country_population"]


def test_country_only_figure_is_labelled_not_repeated_as_city(net: FakeNet) -> None:
    demo = api_enrichment.fetch_location_demographics(["London, UK", "Manchester, UK"])
    profiles = _synth({"location_demographics": demo}, ["London, UK", "Manchester, UK"])

    for loc in ("London, UK", "Manchester, UK"):
        assert "population" not in profiles[loc]
        assert profiles[loc]["area_population_scope"] == "country"


def _workbook_location_rows(data: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    import openpyxl

    import excel_v2

    wb = openpyxl.load_workbook(io.BytesIO(excel_v2.generate_excel_v2(data)))
    ws = wb["Market Intelligence"]
    rows: Dict[str, Dict[str, str]] = {}
    header: Dict[str, int] = {}
    in_section = False
    for row in ws.iter_rows():
        cells = {c.column: c.value for c in row if c.value not in (None, "")}
        label = " ".join(str(v) for v in cells.values())
        if "LOCATION INTELLIGENCE" in label.upper():
            in_section = True
            continue
        if not in_section:
            continue
        if "Location" in cells.values() and "Population" in cells.values():
            header = {str(v): k for k, v in cells.items()}
            continue
        if header and cells:
            if "COMPETITIVE" in label.upper():
                break
            rows[str(cells.get(header["Location"]) or "")] = {
                name: str(cells.get(col) or "") for name, col in header.items()
            }
    return rows


def _data_for(locations: list, demo: Dict[str, Any], extra: Dict[str, Any] | None = None) -> Dict[str, Any]:
    import budget_engine

    roles = [{"title": "Financial Analyst", "count": 4, "tier": "mid"}]
    channels = {"programmatic_dsp": 40, "global_boards": 40, "social_media": 20}
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=150_000,
        roles=roles,
        locations=locations,
        industry="finance_banking",
        channel_percentages=channels,
    )
    enriched: Dict[str, Any] = {"location_demographics": demo, **(extra or {})}
    return {
        "client_name": "Brief Fixture Bank",
        "industry": "finance_banking",
        "budget": "150000",
        "campaign_duration": "3 months",
        "campaign_weeks": 13,
        "roles": [r["title"] for r in roles],
        "target_roles": roles,
        "locations": locations,
        "_budget_allocation": alloc,
        "_enriched": enriched,
        "_synthesized": {
            "location_profiles": data_synthesizer.fuse_location_profiles(
                enriched, {}, {"locations": locations}
            )
        },
    }


def test_non_us_workbook_never_repeats_the_national_number_as_a_city_population(
    net: FakeNet,
) -> None:
    """Real workbook over a UK brief (the sweep's f_uk_finserv shape): the
    Location Intelligence table used to print one national figure per city."""
    locations = ["London, UK", "Manchester, UK", "Edinburgh, UK"]
    demo = api_enrichment.fetch_location_demographics(locations)
    national = f"{demo['London, UK']['country_population']:,}"

    rows = _workbook_location_rows(_data_for(locations, demo))

    assert set(rows) >= set(locations)
    for loc in locations:
        cell = rows[loc]["Population"]
        assert cell != national, f"{loc}: national population printed as the city's"
        assert "country-level" in cell and "not city" in cell


def test_us_workbook_shows_each_citys_own_population(net: FakeNet) -> None:
    demo = api_enrichment.fetch_location_demographics(["Hershey, PA", "Lancaster, PA"])
    rows = _workbook_location_rows(_data_for(["Hershey, PA", "Lancaster, PA"], demo))

    assert rows["Hershey, PA"]["Population"] == "14,242"
    assert rows["Lancaster, PA"]["Population"] == "57,719"  # was PA's 12,961,683


# ---------------------------------------------------------------------------
# A wrapped footnote in a MERGED row needs an explicit height (Excel never
# auto-fits merged cells, so the 265-char Quality Intelligence footnote clipped)
# ---------------------------------------------------------------------------


def test_long_quality_intelligence_footnote_row_has_a_height_that_fits_its_lines() -> None:
    import math

    from openpyxl.utils import get_column_letter

    ws = _quality_sheet(PA_TOWNS)
    cell = next(
        c
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and "share identical figures" in c.value
    )
    assert len(cell.value) > 250
    assert cell.alignment.wrap_text is True
    assert any(
        m.min_row == cell.row == m.max_row and m.min_col == 2 and m.max_col == 8
        for m in ws.merged_cells.ranges
    )

    merged_width = sum(
        ws.column_dimensions[get_column_letter(c)].width for c in range(2, 9)
    )
    # independent, deliberately generous bound: even at 1.3 chars per width unit
    # the text needs this many 11 pt lines
    min_lines = math.ceil(len(cell.value) / (merged_width * 1.3))
    height = ws.row_dimensions[cell.row].height
    assert height is not None, "merged wrapped footnote left at the default one-line height"
    assert height >= min_lines * 11.5 and min_lines >= 2


def test_footnote_height_helper_leaves_short_footnotes_alone() -> None:
    import openpyxl

    ws = openpyxl.Workbook().active
    for col, width in zip("BCDEFGH", (22, 18, 18, 18, 18, 18, 18)):
        ws.column_dimensions[col].width = width
    assert excel_v2._footnote_row_height(ws, "A short note.") is None
    assert excel_v2._footnote_row_height(ws, "x" * 265) >= 26.0
    assert excel_v2._footnote_row_height(ws, "x" * 600) > excel_v2._footnote_row_height(
        ws, "x" * 265
    )

