"""Regression coverage for observed versus benchmark job-posting evidence.

The collectors may return curated salary and market benchmarks when credentials
are unavailable.  Those values remain useful estimates, but must never be
promoted to live postings, live market heat, or a render-time freshness date.
"""

from __future__ import annotations

import io
import importlib

import api_enrichment
import data_synthesizer
import excel_v2
import openpyxl
import ppt_generator
import plan_validator
import pytest
from pptx import Presentation


def _fuse(adzuna: dict, jooble: dict) -> dict:
    return data_synthesizer.fuse_job_market_demand(
        {"job_market": adzuna, "jooble_data": jooble},
        {},
        {"roles": ["Warehouse Associate"], "locations": ["London, UK"]},
    )["Warehouse Associate"]


def test_collectors_mark_curated_posting_counts_as_benchmarks() -> None:
    adzuna = api_enrichment._adzuna_benchmark_fallback(
        ["Warehouse Associate"], ["London, UK"]
    )
    jooble = api_enrichment._get_fallback_data("Warehouse Associate", "London, UK")

    assert adzuna["Warehouse Associate"]["posting_observation"] == "estimated"
    assert jooble["posting_observation"] == "estimated"


def test_fusion_does_not_promote_collector_benchmarks_to_live_postings() -> None:
    """Provider-looking fallback names are still benchmark data, not observations."""
    demand = _fuse(
        {
            "Warehouse Associate": {
                "posting_count": 95_000,
                "source": "Adzuna (curated benchmark)",
                "posting_observation": "benchmark",
                "posting_provider_id": "adzuna_benchmark",
            }
        },
        {
            "source": "Jooble Market Benchmarks",
            "job_market": {
                "Warehouse Associate": {
                    "London, UK": {
                        "total_job_postings": 65_000,
                        "posting_observation": "benchmark",
                        "posting_provider_id": "jooble_benchmark",
                    }
                }
            },
        },
    )

    assert demand["posting_sources"] == []
    assert demand["total_postings"] == 0
    assert demand["market_evidence"]["metrics"]["job_postings"]["status"] == "estimated"
    assert demand["_meta"]["source_count"] == 0


def test_fusion_never_sums_providers_and_prefers_the_narrower_scope() -> None:
    """Adzuna (whole country) and Jooble (the plan's city) both observed: the
    headline is ONE provider's count -- the city-scoped one -- not 5,000 +
    2,000. The other provider is kept as separate detail."""
    demand = _fuse(
        {
            "Warehouse Associate": {
                "posting_count": 5_000,
                "source": "Adzuna",
                "posting_observation": "observed",
                "posting_provider_id": "adzuna",
                "posting_observed_at": "2026-10-15",
                "posting_scope": "national",
                "posting_scope_label": "United Kingdom",
            }
        },
        {
            "source": "Jooble API",
            "provider_id": "jooble",
            "job_market": {
                "Warehouse Associate": {
                    "London, UK": {
                        "total_job_postings": 2_000,
                        "posting_observation": "observed",
                        "posting_provider_id": "jooble",
                        "posting_observed_at": "2026-10-14",
                    }
                }
            },
        },
    )

    assert demand["total_postings"] == 2_000
    assert demand["posting_sources"] == ["Jooble API"]
    assert demand["posting_primary"] == {
        "count": 2_000,
        "provider": "Jooble API",
        "provider_id": "jooble",
        "observed_at": "2026-10-14",
        "observed_dates": ["2026-10-14"],
        "scope": "location",
        "scope_label": "London, UK",
    }
    assert (
        demand["posting_summary"]
        == "2,000 jobs \u00b7 Jooble API \u00b7 London, UK \u00b7 2026-10-14"
    )
    assert demand["posting_other_providers"] == [
        {
            "count": 5_000,
            "provider": "Adzuna",
            "provider_id": "adzuna",
            "observed_at": "2026-10-15",
            "observed_dates": ["2026-10-15"],
            "scope": "national",
            "scope_label": "United Kingdom",
        }
    ]
    # The metric evidence describes the figure shown: one provider, its date.
    job_postings = demand["market_evidence"]["metrics"]["job_postings"]
    assert job_postings["provider_ids"] == ["jooble"]
    assert job_postings["observed_at"] == ["2026-10-14"]
    assert demand["posting_source_details"] == [
        {
            "source": "Jooble API",
            "provider_id": "jooble",
            "observed_at": "2026-10-14",
            "count": 2_000,
            "scope": "location",
            "scope_label": "London, UK",
        }
    ]


def test_fusion_scope_tie_takes_the_larger_count() -> None:
    demand = _fuse(
        {
            "Warehouse Associate": {
                "posting_count": 5_000,
                "source": "Adzuna",
                "posting_observation": "observed",
                "posting_provider_id": "adzuna",
                "posting_observed_at": "2026-10-15",
                "posting_scope": "location",
                "posting_scope_label": "London",
            }
        },
        {
            "source": "Jooble API",
            "provider_id": "jooble",
            "job_market": {
                "Warehouse Associate": {
                    "London, UK": {
                        "total_job_postings": 2_000,
                        "posting_observation": "observed",
                        "posting_provider_id": "jooble",
                        "posting_observed_at": "2026-10-14",
                    }
                }
            },
        },
    )

    assert demand["total_postings"] == 5_000
    assert demand["posting_primary"]["provider"] == "Adzuna"
    assert (
        demand["posting_summary"]
        == "5,000 jobs \u00b7 Adzuna \u00b7 London \u00b7 2026-10-15"
    )
    assert [o["provider"] for o in demand["posting_other_providers"]] == ["Jooble API"]


def test_fusion_nationwide_scope_is_stated_in_the_summary() -> None:
    demand = _fuse(
        {
            "Warehouse Associate": {
                "posting_count": 7_000,
                "source": "Adzuna",
                "posting_observation": "observed",
                "posting_provider_id": "adzuna",
                "posting_observed_at": "2026-10-15",
                "posting_scope": "national",
                "posting_scope_label": "United Kingdom",
            }
        },
        {},
    )

    assert demand["total_postings"] == 7_000
    assert (
        demand["posting_summary"]
        == "7,000 jobs \u00b7 Adzuna \u00b7 nationwide \u00b7 2026-10-15"
    )
    assert demand["posting_other_providers"] == []


def test_fusion_without_observed_postings_has_no_headline() -> None:
    demand = _fuse({}, {})
    assert demand["total_postings"] == 0
    assert demand["posting_primary"] is None
    assert demand["posting_summary"] == ""
    assert demand["posting_other_providers"] == []


def test_fusion_keeps_only_observed_portion_of_mixed_provider_response() -> None:
    demand = _fuse(
        {
            "Warehouse Associate": {
                "posting_count": 5_000,
                "source": "Adzuna",
                "posting_observation": "observed",
                "posting_provider_id": "adzuna",
                "posting_observed_at": "2026-10-15",
            }
        },
        {
            "source": "Jooble API",
            "provider_id": "jooble",
            "job_market": {
                "Warehouse Associate": {
                    "London, UK": {
                        "total_job_postings": 2_000,
                        "posting_observation": "observed",
                        "posting_provider_id": "jooble",
                        "posting_observed_at": "2026-10-14",
                    },
                    "Manchester, UK": {
                        "total_job_postings": 65_000,
                        "posting_observation": "benchmark",
                    },
                }
            },
        },
    )

    # The 65,000 benchmark never counts; Jooble's observed London count
    # (city scope) is the headline over Adzuna's scope-less 5,000.
    assert demand["total_postings"] == 2_000
    assert demand["posting_sources"] == ["Jooble API"]
    assert [o["provider"] for o in demand["posting_other_providers"]] == ["Adzuna"]


def test_fallback_location_has_no_postings_or_market_activity() -> None:
    demand = data_synthesizer.fuse_job_market_demand(
        {
            "jooble_data": {
                "source": "Jooble API",
                "provider_id": "jooble",
                "job_market": {
                    "Warehouse Associate": {
                        "London, UK": {
                            "total_job_postings": 2_000,
                            "market_activity": "Medium",
                            "posting_observation": "observed",
                            "posting_provider_id": "jooble",
                            "posting_observed_at": "2026-10-14",
                        },
                        "Manchester, UK": {
                            "total_job_postings": 65_000,
                            "market_activity": "High",
                            "posting_observation": "benchmark",
                        },
                    }
                },
            }
        },
        {},
        {
            "roles": ["Warehouse Associate"],
            "locations": ["London, UK", "Manchester, UK"],
        },
    )["Warehouse Associate"]

    assert demand["by_location"]["London, UK"] == {
        "postings": 2_000,
        "market_activity": "Medium",
        "sources": ["Jooble API"],
    }
    assert demand["by_location"]["Manchester, UK"] == {
        "postings": 0,
        "market_activity": "",
        "sources": [],
    }
    assert demand["_meta"]["source_count"] == 1


def test_fallback_payload_is_not_counted_as_a_live_enrichment_success() -> None:
    result, status, metadata = api_enrichment._safe_call(
        lambda: {
            "Warehouse Associate": {
                "posting_count": 95_000,
                "posting_observation": "benchmark",
            }
        },
        "Adzuna",
    )

    assert result
    assert status == "estimated"
    assert metadata["source"] == "benchmark"
    assert metadata["success"] is False


def test_unknown_collector_payload_is_not_labelled_observed_in_workbook() -> None:
    data = {
        "client_name": "Provenance Test",
        "industry": "logistics",
        "roles": ["Warehouse Associate"],
        "target_roles": [{"title": "Warehouse Associate", "count": 20}],
        "locations": ["London, UK"],
        "budget": "150000",
        "_enriched": {
            "enrichment_summary": {
                "apis_called": ["UnknownFallback"],
                "apis_succeeded": ["UnknownFallback"],
                "apis_estimated": [],
                "apis_not_applicable": [],
            }
        },
        "_synthesized": {},
    }
    xlsx = excel_v2.generate_excel_v2(dict(data))
    if isinstance(xlsx, tuple):
        xlsx = xlsx[0]
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx))
    values = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert not any("Observed API sources" in value for value in values)
    assert any(
        "Collectors returned payloads (1/1 applicable)" in value for value in values
    )


def test_workbook_provenance_separates_estimated_collectors_from_observed_apis() -> (
    None
):
    """An offline payload must not create a blanket "Live market APIs" claim."""
    data = {
        "client_name": "Provenance Test",
        "industry": "logistics",
        "roles": ["Warehouse Associate"],
        "target_roles": [{"title": "Warehouse Associate", "count": 20}],
        "locations": ["London, UK"],
        "budget": "150000",
        "_enriched": {
            "enrichment_summary": {
                "apis_called": ["Adzuna", "Jooble"],
                "apis_succeeded": [],
                "apis_estimated": ["Adzuna", "Jooble"],
                "apis_not_applicable": [],
            }
        },
        "_synthesized": {},
    }
    xlsx = excel_v2.generate_excel_v2(dict(data))
    if isinstance(xlsx, tuple):
        xlsx = xlsx[0]
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx))
    values = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert not any("Live market APIs" in value for value in values)
    assert any("Estimated collectors (2; not observed)" in value for value in values)


def _deck_texts(job_market_demand: dict) -> list[str]:
    data = {
        "client_name": "Provenance Test",
        "industry": "logistics",
        "industry_label": "Logistics",
        "budget": "$150,000",
        "budget_period": "campaign",
        "campaign_duration": "3 months",
        "campaign_start_month": 9,
        "hire_volume": "35 hires",
        "work_environment": "onsite",
        # A US market: on a non-US plan slide 5's content-derived row cap drops
        # the 5th benchmark row (this one) before the provenance wording is
        # reachable, so the provenance assertions would test nothing there.
        "locations": [{"city": "Columbus", "state": "OH", "country": "United States"}],
        "roles": [{"title": "Warehouse Associate", "count": 20, "tier": "mid"}],
        "_synthesized": {"job_market_demand": job_market_demand},
    }
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(data)))
    return [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]


def test_deck_uses_collector_observation_date_not_render_date() -> None:
    texts = _deck_texts(
        {
            "Warehouse Associate": {
                "total_postings": 7_000,
                "posting_sources": ["Adzuna", "Jooble API"],
                "posting_source_details": [
                    {
                        "source": "Adzuna",
                        "provider_id": "adzuna",
                        "observed_at": "2026-10-15",
                    },
                    {
                        "source": "Jooble API",
                        "provider_id": "jooble",
                        "observed_at": "2026-10-14",
                    },
                ],
                "market_temperature": "warm",
                "market_evidence": {
                    "contract": "market_evidence_v1",
                    "metrics": {
                        "job_postings": {
                            "status": "observed",
                            "provider_ids": ["adzuna", "jooble"],
                            "observed_at": ["2026-10-14", "2026-10-15"],
                        },
                        "market_temperature": _observed_metric("adzuna", "2026-10-15"),
                    },
                },
            }
        }
    )

    postings = [text for text in texts if "7,000 active jobs" in text]
    assert postings
    assert "Adzuna (2026-10-15)" in postings[0]
    assert "Jooble API (2026-10-14)" in postings[0]
    assert "as of Oct 16" not in postings[0]


def _observed_metric(provider_id: str, observed_at: str) -> dict:
    return {
        "status": "observed",
        "provider_ids": [provider_id],
        "observed_at": [observed_at],
    }


def test_benchmark_search_and_audience_stay_hidden_through_every_consumer(
    monkeypatch,
) -> None:
    """Curated Google/LinkedIn values are useful internally, never market proof."""
    for key in (
        "GOOGLE_ADS_DEVELOPER_TOKEN",
        "GOOGLE_ADS_REFRESH_TOKEN",
        "GOOGLE_ADS_CLIENT_ID",
        "GOOGLE_ADS_CLIENT_SECRET",
        "GOOGLE_ADS_CUSTOMER_ID",
        "LINKEDIN_ACCESS_TOKEN",
        "LINKEDIN_AD_ACCOUNT_ID",
    ):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setattr(api_enrichment, "_get_cached", lambda _key: None)
    monkeypatch.setattr(api_enrichment, "_set_cached", lambda *_args: None)

    google = api_enrichment.fetch_google_ads_data(
        ["Warehouse Associate"], ["London, UK"]
    )
    linkedin = api_enrichment.fetch_linkedin_ads_data(
        ["Warehouse Associate"], ["London, UK"]
    )
    for label, payload in (("Google Ads", google), ("LinkedIn Ads", linkedin)):
        returned, status, metadata = api_enrichment._safe_call(
            lambda payload=payload: payload, label
        )
        assert returned is payload
        assert status == "estimated"
        assert metadata["success"] is False
    demand = data_synthesizer.fuse_job_market_demand(
        {"google_ads_data": google, "linkedin_ads_data": linkedin},
        {},
        {"roles": ["Warehouse Associate"], "locations": ["London, UK"]},
    )["Warehouse Associate"]

    metrics = demand["market_evidence"]["metrics"]
    assert metrics["search_volume"]["status"] == "estimated"
    assert metrics["talent_pool"]["status"] == "estimated"
    assert demand["total_postings"] == 0
    assert demand["search_volume_monthly"] == 0
    assert demand["talent_pool_estimate"] == 0
    assert demand["market_temperature"] == ""
    assert demand["trend_direction"] == ""
    assert demand["_meta"]["source_count"] == 0

    texts = _deck_texts({"Warehouse Associate": demand})
    rendered = "\n".join(texts).lower()
    assert "market temp" not in rendered
    assert "talent market" not in rendered
    assert "buyer's market" not in rendered
    assert "low-demand" not in rendered

    plan = {
        "client_name": "Provenance Test",
        "industry": "logistics",
        "roles": ["Warehouse Associate"],
        "target_roles": [{"title": "Warehouse Associate", "count": 20}],
        "locations": ["London, UK"],
        "budget": "150000",
        "_synthesized": {"job_market_demand": {"Warehouse Associate": demand}},
    }
    xlsx = excel_v2.generate_excel_v2(dict(plan))
    if isinstance(xlsx, tuple):
        xlsx = xlsx[0]
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx))
    rows = list(workbook["Market Intelligence"].iter_rows(values_only=True))
    role_rows = [
        row for row in rows if len(row) > 2 and row[1] == "Warehouse Associate"
    ]
    assert role_rows
    assert all(
        "5,500,000" not in str(row) and "8,500,000" not in str(row) for row in role_rows
    )
    assert any(row[2:8] == ("Data not available",) * 6 for row in role_rows)

    unsupported = dict(demand)
    unsupported["market_temperature"] = "hot"
    unsupported["total_postings"] = 0
    validated = {
        "_synthesized": {"job_market_demand": {"Warehouse Associate": unsupported}}
    }
    findings = plan_validator._check_demand_vs_temperature(validated)
    repaired = validated["_synthesized"]["job_market_demand"]["Warehouse Associate"]
    assert repaired["market_temperature"] == ""
    assert findings and findings[0]["new_value"] == ""


def test_one_observed_jooble_provider_across_locations_counts_once() -> None:
    demand = data_synthesizer.fuse_job_market_demand(
        {
            "jooble_data": {
                "source": "Jooble API",
                "provider_id": "jooble",
                "job_market": {
                    "Warehouse Associate": {
                        "London, UK": {
                            "total_job_postings": 2_000,
                            "posting_observation": "observed",
                            "posting_provider_id": "jooble",
                            "posting_observed_at": "2026-10-14",
                        },
                        "Manchester, UK": {
                            "total_job_postings": 3_000,
                            "posting_observation": "observed",
                            "posting_provider_id": "jooble",
                            "posting_observed_at": "2026-10-15",
                        },
                    }
                },
            }
        },
        {},
        {
            "roles": ["Warehouse Associate"],
            "locations": ["London, UK", "Manchester, UK"],
        },
    )["Warehouse Associate"]

    assert demand["total_postings"] == 5_000
    assert demand["posting_sources"] == ["Jooble API"]
    assert demand["market_evidence"]["metrics"]["job_postings"] == {
        "status": "observed",
        "provider_ids": ["jooble"],
        "observed_at": ["2026-10-14", "2026-10-15"],
    }
    assert demand["_meta"]["source_count"] == 1


def test_mixed_observations_keep_values_dates_and_unique_provider_confidence() -> None:
    demand = data_synthesizer.fuse_job_market_demand(
        {
            "job_market": {
                "Warehouse Associate": {
                    "posting_count": 5_000,
                    "source": "Adzuna",
                    "posting_observation": "observed",
                    "posting_provider_id": "adzuna",
                    "posting_observed_at": "2026-10-15",
                }
            },
            "google_ads_data": {
                "keywords": {
                    "Warehouse Associate": {
                        "avg_monthly_searches": 1_200,
                        "search_volume_observation": "observed",
                        "search_volume_provider_id": "google_ads",
                        "search_volume_observed_at": "2026-10-13",
                    }
                }
            },
            "linkedin_ads_data": {
                "roles": {
                    "Warehouse Associate": {
                        "estimated_audience": "250K",
                        "audience_observation": "observed",
                        "audience_provider_id": "linkedin",
                        "audience_observed_at": "2026-10-12",
                    }
                }
            },
        },
        {},
        {"roles": ["Warehouse Associate"], "locations": ["London, UK"]},
    )["Warehouse Associate"]

    assert demand["total_postings"] == 5_000
    assert demand["search_volume_monthly"] == 1_200
    assert demand["talent_pool_estimate"] == 250_000
    assert demand["market_temperature"] == "warm"
    assert demand["market_evidence"]["status"] == "observed"
    assert demand["market_evidence"]["observed_provider_ids"] == [
        "adzuna",
        "google_ads",
        "linkedin",
    ]
    assert demand["_meta"]["source_count"] == 3


def test_cached_observation_keeps_its_actual_date(monkeypatch) -> None:
    cached = {
        "posting_count": 4_200,
        "source": "Adzuna",
        "posting_observation": "observed",
        "posting_provider_id": "adzuna",
        "posting_observed_at": "2026-10-11",
    }
    monkeypatch.setenv("ADZUNA_APP_ID", "test-id")
    monkeypatch.setenv("ADZUNA_APP_KEY", "test-key")
    monkeypatch.setattr(api_enrichment, "_get_cached", lambda _key: dict(cached))

    result = api_enrichment.fetch_job_market(["Warehouse Associate"], ["London, UK"])

    assert result["Warehouse Associate"] == cached


def test_legacy_provider_names_without_markers_are_unavailable() -> None:
    demand = _fuse(
        {"Warehouse Associate": {"posting_count": 5_000, "source": "Adzuna"}},
        {
            "source": "Jooble API",
            "job_market": {
                "Warehouse Associate": {"London, UK": {"total_job_postings": 2_000}}
            },
        },
    )

    assert demand["total_postings"] == 0
    assert demand["market_evidence"]["metrics"]["job_postings"]["status"] != "observed"
    assert demand["_meta"]["source_count"] == 0


def test_fusion_rejects_noncanonical_and_invalid_raw_observations() -> None:
    demand = data_synthesizer.fuse_job_market_demand(
        {
            "job_market": {
                "Warehouse Associate": {
                    "posting_count": 5_000,
                    "source": "Adzuna",
                    "posting_observation": "observed",
                    "posting_provider_id": "display-name-is-not-provenance",
                    "posting_observed_at": "2026-10-15",
                }
            },
            "google_ads_data": {
                "keywords": {
                    "Warehouse Associate": {
                        "avg_monthly_searches": 1_200,
                        "search_volume_observation": "observed",
                        "search_volume_provider_id": "google_ads",
                        "search_volume_observed_at": "2099-01-01",
                    }
                }
            },
            "linkedin_ads_data": {
                "roles": {
                    "Warehouse Associate": {
                        "estimated_audience": "250K",
                        "audience_observation": "observed",
                        "audience_provider_id": "linkedin",
                        "audience_observed_at": "not-a-date",
                    }
                }
            },
            "search_trends": {
                "Warehouse Associate": {
                    "avg_interest": 50,
                    "latest_interest": 60,
                    "trend_observation": "observed",
                    "trend_provider_id": "google_trends",
                    "trend_observed_at": "2026-02-30",
                }
            },
        },
        {},
        {"roles": ["Warehouse Associate"], "locations": ["London, UK"]},
    )["Warehouse Associate"]

    assert demand["total_postings"] == 0
    assert demand["search_volume_monthly"] == 0
    assert demand["talent_pool_estimate"] == 0
    assert demand["trend_direction"] == ""
    assert demand["market_temperature"] == ""
    assert demand["market_evidence"]["observed_provider_ids"] == []
    assert demand["_meta"]["source_count"] == 0


@pytest.mark.parametrize(
    "metric, provider_ids, observed_at",
    [
        ("job_postings", ["Adzuna"], ["2026-10-15"]),
        ("job_postings", ["adzuna"], ["2099-01-01"]),
        ("job_postings", ["adzuna"], ["2026-02-30"]),
        ("job_postings", [], ["2026-10-15"]),
        ("job_postings", ["adzuna"], []),
        ("job_postings", [{"id": "adzuna"}], ["2026-10-15"]),
        ("job_postings", ["adzuna"], [["2026-10-15"]]),
        ("search_volume", ["linkedin"], ["2026-10-15"]),
    ],
)
def test_strict_market_contract_rejects_forged_metric_evidence(
    metric: str, provider_ids: list[object], observed_at: list[object]
) -> None:
    evidence_module = importlib.import_module("market_evidence")
    role_data = {
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                metric: {
                    "status": "observed",
                    "provider_ids": provider_ids,
                    "observed_at": observed_at,
                }
            },
        }
    }

    assert evidence_module.market_metric_is_observed(role_data, metric) is False


def test_temperature_requires_valid_posting_and_talent_components() -> None:
    evidence_module = importlib.import_module("market_evidence")
    role_data = {
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "market_temperature": {
                    "status": "observed",
                    "provider_ids": ["adzuna", "linkedin"],
                    "observed_at": ["2026-10-15"],
                }
            },
        }
    }

    assert (
        evidence_module.market_metric_is_observed(role_data, "market_temperature")
        is False
    )


def test_temperature_rejects_unhashable_contract_entries_without_error() -> None:
    evidence_module = importlib.import_module("market_evidence")
    role_data = {
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": ["adzuna"],
                    "observed_at": ["2026-10-15"],
                },
                "talent_pool": {
                    "status": "observed",
                    "provider_ids": ["linkedin"],
                    "observed_at": ["2026-10-15"],
                },
                "market_temperature": {
                    "status": "observed",
                    "provider_ids": [["adzuna"], "linkedin"],
                    "observed_at": ["2026-10-15"],
                },
            },
        }
    }

    assert (
        evidence_module.market_metric_is_observed(role_data, "market_temperature")
        is False
    )


def test_all_market_paths_share_one_evidence_validator() -> None:
    evidence_module = importlib.import_module("market_evidence")

    assert (
        ppt_generator.market_metric_is_observed
        is evidence_module.market_metric_is_observed
    )
    assert (
        ppt_generator.validated_market_source_details
        is evidence_module.validated_market_source_details
    )
    assert (
        excel_v2.market_metric_is_observed is evidence_module.market_metric_is_observed
    )
    assert (
        plan_validator.market_metric_is_observed
        is evidence_module.market_metric_is_observed
    )
    assert data_synthesizer.observation_is_valid is evidence_module.observation_is_valid


def test_every_consumer_rejects_forged_observed_contract() -> None:
    forged = {
        "total_postings": 9_999,
        "posting_sources": ["Adzuna"],
        "posting_source_details": [
            {
                "source": "Adzuna",
                "provider_id": "not-adzuna",
                "observed_at": "2026-10-15",
            }
        ],
        "talent_pool_estimate": 100_000,
        "competition_index": 0.1,
        "market_temperature": "hot",
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": [{"id": "adzuna"}],
                    "observed_at": ["2026-10-15"],
                },
                "talent_pool": {
                    "status": "observed",
                    "provider_ids": ["linkedin"],
                    "observed_at": [["2026-10-15"]],
                },
                "market_temperature": {
                    "status": "observed",
                    "provider_ids": [["adzuna"], "linkedin"],
                    "observed_at": ["2026-10-15"],
                },
            },
        },
    }

    rendered = "\n".join(_deck_texts({"Warehouse Associate": forged})).lower()
    assert "active jobs" not in rendered
    assert "market temp" not in rendered
    assert "talent market" not in rendered

    data = {
        "client_name": "Provenance Test",
        "industry": "logistics",
        "roles": ["Warehouse Associate"],
        "target_roles": [{"title": "Warehouse Associate", "count": 20}],
        "locations": ["London, UK"],
        "budget": "150000",
        "_synthesized": {"job_market_demand": {"Warehouse Associate": forged}},
    }
    workbook = openpyxl.load_workbook(
        io.BytesIO(excel_v2.generate_excel_v2(dict(data)))
    )
    rows = list(workbook["Market Intelligence"].iter_rows(values_only=True))
    role_rows = [
        row for row in rows if len(row) > 2 and row[1] == "Warehouse Associate"
    ]
    assert any(row[2:7] == ("Data not available",) * 5 for row in role_rows)

    findings = plan_validator._check_demand_vs_temperature(data)
    assert forged["market_temperature"] == ""
    assert findings and findings[0]["new_value"] == ""


def test_live_posting_deck_row_requires_valid_dated_source_detail() -> None:
    role_data = {
        "total_postings": 4_200,
        "posting_sources": ["Adzuna"],
        "posting_source_details": [{"source": "Adzuna", "provider_id": "adzuna"}],
        "market_temperature": "",
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": ["adzuna"],
                    "observed_at": ["2026-10-15"],
                }
            },
        },
    }

    rendered = "\n".join(_deck_texts({"Warehouse Associate": role_data})).lower()
    assert "active jobs" not in rendered


def test_live_posting_deck_row_requires_complete_source_detail_coverage() -> None:
    role_data = {
        "total_postings": 6_300,
        "posting_source_details": [
            {
                "source": "Adzuna",
                "provider_id": "adzuna",
                "observed_at": "2026-10-15",
            },
            {
                "source": "Jooble API",
                "provider_id": "jooble",
                "observed_at": "not-a-date",
            },
        ],
        "market_temperature": "",
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": ["adzuna", "jooble"],
                    "observed_at": ["2026-10-14", "2026-10-15"],
                }
            },
        },
    }

    rendered = "\n".join(_deck_texts({"Warehouse Associate": role_data})).lower()
    assert "active jobs" not in rendered


def test_live_posting_deck_row_summarizes_all_observed_dates() -> None:
    role_data = {
        "total_postings": 6_300,
        "posting_source_details": [
            {
                "source": "Adzuna",
                "provider_id": "adzuna",
                "observed_at": observed_at,
            }
            for observed_at in ("2026-10-13", "2026-10-14", "2026-10-15")
        ],
        "market_temperature": "",
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": ["adzuna"],
                    "observed_at": [
                        "2026-10-13",
                        "2026-10-14",
                        "2026-10-15",
                    ],
                }
            },
        },
    }

    rendered = "\n".join(_deck_texts({"Warehouse Associate": role_data}))
    postings = [text for text in rendered.splitlines() if "6,300 active jobs" in text]
    assert postings
    assert "2026-10-13–2026-10-15; 3 dates" in postings[0]


def test_cached_jooble_observation_uses_canonical_live_source_label() -> None:
    result = data_synthesizer.fuse_job_market_demand(
        {
            "jooble_data": {
                "source": "Jooble Market Benchmarks",
                "provider_id": "jooble",
                "job_market": {
                    "Warehouse Associate": {
                        "London, UK": {
                            "total_job_postings": 2_300,
                            "posting_observation": "observed",
                            "posting_provider_id": "jooble",
                            "posting_observed_at": "2026-10-14",
                        }
                    }
                },
            }
        },
        {},
        {"roles": ["Warehouse Associate"], "locations": ["London, UK"]},
    )["Warehouse Associate"]

    assert result["posting_sources"] == ["Jooble API"]
    assert result["posting_source_details"] == [
        {
            "source": "Jooble API",
            "provider_id": "jooble",
            "observed_at": "2026-10-14",
            "count": 2_300,
            "scope": "location",
            "scope_label": "London, UK",
        }
    ]
    assert result["by_location"]["London, UK"]["sources"] == ["Jooble API"]


def test_live_posting_deck_derives_source_label_from_provider_id() -> None:
    role_data = {
        "total_postings": 2_300,
        "posting_source_details": [
            {
                "source": "Industry Benchmark",
                "provider_id": "adzuna",
                "observed_at": "2026-10-15",
            }
        ],
        "market_temperature": "",
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": ["adzuna"],
                    "observed_at": ["2026-10-15"],
                }
            },
        },
    }

    rendered = "\n".join(_deck_texts({"Warehouse Associate": role_data}))
    postings = [text for text in rendered.splitlines() if "2,300 active jobs" in text]
    assert postings
    assert "Adzuna (2026-10-15)" in postings[0]
    assert "Industry Benchmark" not in postings[0]


def test_zero_observed_market_providers_force_zero_confidence_despite_kb() -> None:
    result = data_synthesizer.fuse_job_market_demand(
        {},
        {
            "industry_specific_benchmarks": {
                "healthcare": {"hiring_strength": "Strong US demand"}
            }
        },
        {
            "roles": ["Registered Nurse"],
            "locations": ["Dallas, TX"],
            "industry": "healthcare",
        },
    )

    role_data = result["Registered Nurse"]
    assert role_data["market_evidence"]["observed_provider_ids"] == []
    assert role_data["_meta"] == {"source_count": 0, "kb_validated": False}
    assert data_synthesizer._score_section(result) == 0.0
