"""Boards and supply honesty on non-US plans (international audit 2026-10-18).

What these pin, each of which failed before the fix:

* #1  a local board nothing in the supply data covers for the country carried
      NO marker ("unknown" was treated as "fine"); "Pole Emploi (France
      Travail)" and the repository's "France Travail France" were two
      unrelated names, so a VONQ-only board led the French list unmarked;
      the dated record behind a board was not exposed; no surface said how
      much of the in-network list a plan shows.
* #5  the footers said "supply network 2026-10-17" (the newest single-board
      edit) and the deck footer dropped the repository's own date.
* #10 the workbook had no Local Job Boards table for 4 of 23 industries
      (general_entry_level, legal_services, military_recruitment, rideshare)
      while deck slide 5 still named the boards.

Every expected date and count is read from the data files in the test, never
typed: a bumped file moves both sides, a made-up value fails.
"""

from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Set

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402
from pptx import Presentation  # noqa: E402

import excel_v2  # noqa: E402
import plan_trust  # noqa: E402
import ppt_generator  # noqa: E402
import publisher_filter as pf  # noqa: E402
import research  # noqa: E402

DATA_DIR = PROJECT_ROOT / "data"
REPO = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
GROUPS = json.loads((DATA_DIR / "joveo_network_groups.json").read_text())
INTL = json.loads((DATA_DIR / "international_benchmarks_2026.json").read_text())
IN_NETWORK = "In Joveo Network"


# ---------------------------------------------------------------------------
# helpers (independent of publisher_filter's own matching on purpose)
# ---------------------------------------------------------------------------


def _repo_rows() -> List[Dict[str, Any]]:
    return [p for r in REPO["by_region"].values() for p in r["publishers"]]


def _row_countries(row: Dict[str, Any]) -> Set[str]:
    raw = str(row.get("countries") or "")
    if raw.strip().lower() in ("", "nan"):
        return set()
    return {c.strip() for c in raw.split(",") if c.strip()}


def _repo_row(name: str) -> Dict[str, Any]:
    return next(p for p in _repo_rows() if p["name"] == name)


def _plan(industry: str = "general_entry_level", loc: str = "Hanoi, Vietnam") -> dict:
    plan = {
        "client_name": "Fixture Co",
        "company_name": "Fixture Co",
        "industry": industry,
        "budget": "$100,000",
        "locations": [loc],
        "roles": ["Agent"],
        "target_roles": ["Agent"],
        "campaign_duration": "3 months",
        "hire_volume": "50",
        "work_environment": "onsite",
        "_enriched": {},
        "_synthesized": {},
        "_budget_allocation": {},
        # what app.py hands the generators for a non-US market (adds the
        # Intl Benchmarks sheet)
        "_intl_benchmarks": INTL,
    }
    return plan


def _workbook(plan: dict):
    raw = excel_v2.generate_excel_v2(dict(plan), research_mod=research)
    return openpyxl.load_workbook(io.BytesIO(raw))


def _cells(ws) -> List[str]:
    return [
        str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None
    ]


def _deck_texts(plan: dict) -> List[str]:
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(plan))))
    out: List[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return out


# ---------------------------------------------------------------------------
# #1 (a): a named board with no supply record is marked, not assumed
# ---------------------------------------------------------------------------


def test_board_nothing_covers_is_unverified_everywhere():
    name, country = "Zzyzx Careers Portal", "Vietnam"
    record = pf.board_supply_record(name, country)
    assert record["state"] == pf.SUPPLY_NO_RECORD
    assert record["as_of"] == "" and record["status"] == ""
    assert pf.relationship_unverified(name, country) is True
    assert pf.annotate([name], country) == ([f"{name} {pf.RELATIONSHIP_MARK}"], True)


@pytest.mark.parametrize(
    "board,country,key",
    [
        ("TopCV.vn", "Vietnam", "topcv"),
        ("GulfTalent", "United Arab Emirates", "gulftalent"),
        ("InfoJobs Brazil", "Brazil", "infojobs"),
        ("MyCareersFuture (MCF)", "Singapore", "mycareersfuture"),
    ],
)
def test_the_audits_unrecorded_boards_are_marked(board, country, key):
    # guard: the repository really has no row for the board in that country
    covering = [
        p
        for p in _repo_rows()
        if re.sub(r"[^a-z0-9]", "", p["name"].lower()).startswith(key)
        and (country in _row_countries(p) or "Global" in _row_countries(p))
    ]
    assert not covering, covering
    assert pf.relationship_unverified(board, country) is True
    assert pf.annotate([board], country)[1] is True


def test_in_network_board_is_never_marked():
    assert pf.relationship_unverified("VietnamWorks", "Vietnam") is False
    assert pf.relationship_unverified("Totaljobs", "United Kingdom") is False
    assert pf.annotate(["Reed.co.uk"], "United Kingdom") == (["Reed.co.uk"], False)


def test_parenthesised_compound_is_in_network_when_one_named_board_is():
    # "Recruit (Rikunabi / Indeed Japan)": Indeed is Global / In Joveo Network
    assert (
        pf.relationship_unverified("Recruit (Rikunabi / Indeed Japan)", "Japan")
        is False
    )


# ---------------------------------------------------------------------------
# #1 (b): a rename resolves to ONE record
# ---------------------------------------------------------------------------


def test_france_travail_and_pole_emploi_are_one_record():
    row = _repo_row("France Travail France")
    assert row["status"] != IN_NETWORK  # the data says so; the plan must too
    spellings = [
        "Pole Emploi (France Travail)",
        "Pôle Emploi",
        "France Travail",
        "France Travail France",
    ]
    records = [pf.board_supply_record(s, "France") for s in spellings]
    assert {r["state"] for r in records} == {pf.SUPPLY_NOT_IN_NETWORK}
    assert {r["status"] for r in records} == {row["status"]}
    assert pf.names_match("Pole Emploi (France Travail)", "France Travail France")
    assert pf.relationship_unverified("Pole Emploi (France Travail)", "France") is True


@pytest.mark.parametrize(
    "name,country,row_name",
    [
        ("Mynavi", "Japan", "Mynavi Corporation"),
        ("Bundesagentur fuer Arbeit (Jobboerse)", "Germany", "Arbeitsagentur Germany"),
    ],
)
def test_name_variants_with_a_repository_row_resolve_to_it(name, country, row_name):
    row = _repo_row(row_name)
    assert country in _row_countries(row) and row["status"] == IN_NETWORK
    record = pf.board_supply_record(name, country)
    assert record["state"] == pf.SUPPLY_IN_NETWORK
    assert pf.relationship_unverified(name, country) is False


def test_alias_groups_stay_sourced():
    raw = json.loads((DATA_DIR / "publisher_availability.json").read_text())
    for group in raw["alias_groups"]:
        assert str(group.get("source") or "").strip(), group


# ---------------------------------------------------------------------------
# #1 (c): the dated record is exposed, nothing invented
# ---------------------------------------------------------------------------


def test_record_carries_only_dates_the_data_holds():
    # a confirmed group: the group's own as_of
    group_as_of = next(
        g["as_of"] for g in GROUPS["groups"] if g["id"] == "stepstone_group"
    )
    confirmed = pf.board_supply_record("Totaljobs", "United Kingdom")
    assert confirmed["as_of"] == group_as_of
    assert pf.availability_label(confirmed) == f"availability as of {group_as_of}"
    # a repository row: the repository's own last_updated
    plain = pf.board_supply_record("VietnamWorks", "Vietnam")
    assert plain["as_of"] == REPO["last_updated"] == pf.repository_date()
    # no record: no date to show
    nothing = pf.board_supply_record("Zzyzx Careers Portal", "Vietnam")
    assert nothing["as_of"] == "" and pf.availability_label(nothing) == ""


def test_status_phrase_states_state_and_date():
    assert pf.supply_status_phrase(
        pf.board_supply_record("VietnamWorks", "Vietnam")
    ) == (f"VietnamWorks: in network, availability as of {REPO['last_updated']}")
    assert pf.supply_status_phrase(
        pf.board_supply_record("Zzyzx Careers Portal", "Vietnam")
    ) == ("Zzyzx Careers Portal: no supply record")
    apec = pf.supply_status_phrase(pf.board_supply_record("APEC", "France"))
    assert apec.startswith("APEC: listed as ") and "not in network" in apec


def test_market_band_exposes_the_record_of_every_board_shown():
    (band,) = pf.plan_market_boards({"locations": ["Paris, France"]})
    assert len(band["records"]) == len(band["boards"])
    for record, shown in zip(band["records"], band["boards"]):
        assert shown.rstrip(f" {pf.RELATIONSHIP_MARK}") == record["name"]
        assert shown.endswith(pf.RELATIONSHIP_MARK) == (
            record["state"] != pf.SUPPLY_IN_NETWORK
        )


# ---------------------------------------------------------------------------
# #1 (d): confirmed boards first among peers; no leader is hidden
# ---------------------------------------------------------------------------


def test_confirmed_boards_precede_equal_rank_boards_without_a_record():
    # France: the market's #1 board (France Travail, 22%) is VONQ-only; Indeed
    # (20%) and LinkedIn (18%) are peers and in network -> they come first
    ordered = pf.supply_ordered(
        ["Pole Emploi (France Travail)", "Indeed France", "LinkedIn France"], "France"
    )
    assert ordered == [
        "Indeed France",
        "LinkedIn France",
        "Pole Emploi (France Travail)",
    ]


def test_market_leading_board_without_a_record_stays_and_is_marked():
    (band,) = pf.plan_market_boards({"locations": ["Paris, France"]})
    assert "Pole Emploi (France Travail) †" in band["boards"]
    assert band["boards"][0] == "Indeed France"
    # Vietnam: TopCV.vn (20%) is NOT a peer of the leader (30%): it keeps its
    # rank position, marked
    (vn,) = pf.plan_market_boards({"locations": ["Hanoi, Vietnam"]})
    assert vn["boards"][:2] == ["VietnamWorks", "TopCV.vn †"]


def test_ordering_never_adds_or_removes_a_board():
    names = ["MyCareersFuture (MCF)", "LinkedIn Singapore", "JobStreet Singapore"]
    assert sorted(pf.supply_ordered(names, "Singapore")) == sorted(names)
    assert pf.supply_ordered(names, "") == names  # no country -> untouched
    assert pf.supply_ordered(["One"], "Singapore") == ["One"]


# ---------------------------------------------------------------------------
# supply-depth line: counted from the repository or skipped
# ---------------------------------------------------------------------------


def _depth_expected(country: str, shown_keys: List[str]) -> Dict[str, int]:
    keys: Set[str] = set()
    for p in _repo_rows():
        if p["status"] != IN_NETWORK:
            continue
        cs = _row_countries(p)
        if country in cs or "Global" in cs:
            keys.add(pf.normalize(p["name"]))
    return {
        "total": len(keys),
        "shown": sum(1 for k in shown_keys if k in keys),
    }


def test_depth_is_counted_straight_from_the_repository():
    shown = ["VietnamWorks", "TopCV.vn", "LinkedIn Vietnam"]
    depth = pf.supply_depth("Vietnam", shown)
    want = _depth_expected("Vietnam", [pf.normalize(n) for n in shown])
    assert depth == {"country": "Vietnam", **want}
    assert 0 < depth["shown"] < depth["total"]
    line = pf.supply_depth_line([{"depth": depth}])
    assert line == (
        f"{want['shown']} of {want['total']} in-network boards for Vietnam shown; "
        "the rest are not role-matched."
    )


def test_depth_is_skipped_when_it_cannot_be_derived():
    assert pf.supply_depth("Atlantis", ["Anything"]) is None
    assert pf.supply_depth("", ["Anything"]) is None
    assert pf.supply_depth_line([{"depth": None}]) == ""
    assert pf.supply_depth_line([]) == ""
    two = pf.supply_depth_line(
        [
            {"depth": {"country": "France", "shown": 2, "total": 9}},
            {"depth": {"country": "Germany", "shown": 3, "total": 12}},
        ]
    )
    assert "France 2 of 9" in two and "Germany 3 of 12" in two
    assert two.endswith("The rest are not role-matched.")


# ---------------------------------------------------------------------------
# every surface: deck slide 5, workbook, Intl Benchmarks, Nova
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def vn_surfaces():
    plan = _plan("general_entry_level", "Hanoi, Vietnam")
    return {"wb": _workbook(plan), "deck": _deck_texts(plan)}


def test_deck_slide_five_marks_names_and_counts(vn_surfaces):
    deck = vn_surfaces["deck"]
    assert any(t.startswith("Vietnam:") and "TopCV.vn †" in t for t in deck)
    assert pf.RELATIONSHIP_NOTE in deck
    depth = pf.supply_depth(
        "Vietnam",
        [b.rstrip(f" {pf.RELATIONSHIP_MARK}") for b in _vn_deck_boards(deck)],
    )
    assert pf.supply_depth_line([{"depth": depth}]) in deck


def _vn_deck_boards(deck: List[str]) -> List[str]:
    line = next(t for t in deck if t.startswith("Vietnam:"))
    return [b.strip() for b in line.split(":", 1)[1].split(",")]


def test_workbook_local_boards_and_intl_benchmarks_carry_the_marker(vn_surfaces):
    wb = vn_surfaces["wb"]
    channels = _cells(wb["Channels & Strategy"])
    assert any("TopCV.vn †" in c for c in channels)
    assert pf.RELATIONSHIP_NOTE in channels
    intl = _cells(wb["Intl Benchmarks"])
    assert any("TopCV.vn †" in c for c in intl)
    assert any(pf.RELATIONSHIP_NOTE in c for c in intl)


def test_workbook_shows_availability_as_of_the_date_the_data_holds(vn_surfaces):
    channels = _cells(vn_surfaces["wb"]["Channels & Strategy"])
    status = next(c for c in channels if c.startswith("VietnamWorks: in network"))
    assert f"availability as of {REPO['last_updated']}" in status
    assert "TopCV.vn: no supply record" in status
    # no invented date on the board without a record
    assert not re.search(r"TopCV\.vn[^;.]*as of", status)


def test_workbook_depth_line_matches_the_deck_method(vn_surfaces):
    channels = _cells(vn_surfaces["wb"]["Channels & Strategy"])
    depth_cells = [c for c in channels if "in-network boards for Vietnam shown" in c]
    assert len(depth_cells) == 1
    assert depth_cells[0].endswith("the rest are not role-matched.")


def test_nova_ad_platform_tool_carries_the_same_verdict():
    from nova import Nova

    nova = Nova.__new__(Nova)
    nova._data_cache = {
        "international_benchmarks": json.loads(
            (DATA_DIR / "international_benchmarks_2026.json").read_text()
        )
    }
    result = nova._query_ad_platform(
        {"role_type": "professional", "country": "Vietnam"}
    )
    status = {row["name"]: row for row in result["supply_status"]}
    assert status["TopCV.vn"]["state"] == pf.SUPPLY_NO_RECORD
    assert status["TopCV.vn"]["as_of"] == ""
    assert status["VietnamWorks"]["state"] == pf.SUPPLY_IN_NETWORK
    assert status["VietnamWorks"]["as_of"] == REPO["last_updated"]
    assert result["supply_note"] == pf.RELATIONSHIP_NOTE
    # nothing is re-ranked or hidden by the annotation
    assert result["top_platforms"][0] == result["platforms_ranked"][0]["name"]


def _nova() -> Any:
    from nova import Nova

    nova = Nova.__new__(Nova)
    nova._data_cache = {"international_benchmarks": INTL}
    return nova


def test_every_nova_tool_that_reads_the_local_board_list_carries_the_verdict():
    nova = _nova()
    benchmarks = nova._query_recruitment_benchmarks(
        {"industry": "retail", "country": "Vietnam"}
    )
    collar = nova._query_collar_strategy(
        {"role": "Call Centre Agent", "country": "Vietnam"}
    )
    for result in (benchmarks, collar):
        states = {r["name"]: r["state"] for r in result["supply_status"]}
        assert states["TopCV.vn"] == pf.SUPPLY_NO_RECORD, result.keys()
        assert states["VietnamWorks"] == pf.SUPPLY_IN_NETWORK
        assert result["supply_note"] == pf.RELATIONSHIP_NOTE


def test_a_us_query_gets_no_supply_annotation():
    result = _nova()._query_ad_platform({"role_type": "professional", "country": ""})
    assert "supply_status" not in result and "supply_note" not in result


# ---------------------------------------------------------------------------
# new copy stays free of US-only vocabulary on every market
# ---------------------------------------------------------------------------


def test_new_copy_carries_no_us_only_marker():
    from us_only_markers import has_us_only_marker

    info = plan_trust.data_asof({})
    records = [
        pf.board_supply_record(n, "France")
        for n in ("Indeed France", "APEC", "Zzyzx Careers Portal")
    ]
    copy = [
        info["headline"],
        info["footer_headline"],
        pf.RELATIONSHIP_NOTE,
        pf.supply_depth_line([{"depth": pf.supply_depth("France", ["Indeed France"])}]),
        *[pf.supply_status_phrase(r) for r in records],
        *[pf.availability_label(r) for r in records],
    ]
    copy = [t for t in copy if t]  # a board with no record has no date label
    assert len(copy) >= 8
    assert not [t for t in copy if has_us_only_marker(t)]


# ---------------------------------------------------------------------------
# #10: the workbook lists local boards for every industry
# ---------------------------------------------------------------------------

_NO_NICHE_LIST = [
    "general_entry_level",
    "legal_services",
    "military_recruitment",
    "rideshare",
]


@pytest.mark.parametrize("industry", _NO_NICHE_LIST)
def test_workbook_lists_local_boards_for_industries_without_a_niche_list(industry):
    assert not excel_v2.INDUSTRY_NICHE_CHANNELS.get(industry)
    wb = _workbook(_plan(industry, "Paris, France"))
    ws = wb["Channels & Strategy"]
    cells = _cells(ws)
    boards = next(c for c in cells if "Indeed France" in c and "," in c)
    assert "Pole Emploi (France Travail) †" in boards
    assert any(c.upper() == "LOCAL JOB BOARDS" for c in cells)
    # no niche section is invented for an industry that has none
    assert not any("NICHE CHANNELS" in c.upper() for c in cells)
    assert not any("US-domiciled niche boards" in c for c in cells)


def test_industry_with_a_niche_list_keeps_its_niche_section_and_note():
    assert excel_v2.INDUSTRY_NICHE_CHANNELS.get("healthcare_medical")
    cells = _cells(
        _workbook(_plan("healthcare_medical", "Paris, France"))["Channels & Strategy"]
    )
    assert any(c.upper().startswith("NICHE CHANNELS: ") for c in cells)
    assert any("US-domiciled niche boards" in c for c in cells)
    assert any("Indeed France" in c and "," in c for c in cells)


def test_workbook_names_every_board_the_deck_names():
    plan = _plan("rideshare", "Hanoi, Vietnam")
    deck_boards = {
        b.rstrip(f" {pf.RELATIONSHIP_MARK}") for b in _vn_deck_boards(_deck_texts(plan))
    }
    row = next(
        c
        for c in _cells(_workbook(plan)["Channels & Strategy"])
        if "VietnamWorks" in c and "," in c
    )
    for board in deck_boards:
        assert board in row, (board, row)


def test_us_plan_gets_no_local_boards_block():
    cells = _cells(
        _workbook(_plan("general_entry_level", "Dallas, TX"))["Channels & Strategy"]
    )
    assert not any(c.upper() == "LOCAL JOB BOARDS" for c in cells)
    assert not any("in-network boards for" in c for c in cells)
    assert pf.plan_market_boards({"locations": ["Dallas, TX"]}) == []
