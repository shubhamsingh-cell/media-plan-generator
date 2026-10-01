"""Regression: deck slide 5 and the workbook print the SAME cost-per-hire
benchmark, labelled for what it is (design-judge panel, 2026-10-01, items
2 and 3).

Pre-fix (branch at 8424448):
* Hershey (food & beverage -- an industry with no range of its own in
  budget_engine.INDUSTRY_CPH_RANGES) printed the cross-industry DEFAULT
  "$4,000-$8,000 (avg $6,000)" under "Food & Beverage Range" on slide 5,
  while the workbook's KB row said "$3,000-$5,000 (food manufacturing)".
* The figure was labelled "avg" although it is the midpoint of a range,
  and a local figure ("₹35,000-₹80,000 (median ₹55,000)") invented a
  median the source does not state, wrapped onto two lines, and the
  sources line did not name the local source.
* The workbook's KB row for a matched industry (retail "$2,700", finance
  "$4,500-$8,000") disagreed with the engine range on slide 5.
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import budget_engine as be  # noqa: E402
import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402


def _brief(name, industry, budget, locations, roles, hire_volume="Not specified"):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": hire_volume,
        "work_environment": "onsite",
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
    }


def _slide(data, n):
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    return [
        sh.text_frame.text
        for sh in prs.slides[n - 1].shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]


def _exec_summary_rows(data):
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    if isinstance(raw, tuple):
        raw = raw[0]
    ws = load_workbook(io.BytesIO(raw))["Executive Summary"]
    return [[c.value for c in r if c.value not in (None, "")] for r in ws.iter_rows()]


def _value_after(texts, label):
    i = texts.index(label)
    return texts[i + 1]


@pytest.fixture(scope="module")
def hershey():
    return T.build_plan_data(_brief(
        "Probe Foods", "food_beverage", "$45,000",
        ["Hershey, PA", "Reading, PA"], ["Machine Operator", "Sanitation"],
        "100-500 hires"))


@pytest.fixture(scope="module")
def hospital():
    return T.build_plan_data(_brief(
        "Probe Health", "Healthcare & Medical", "$250,000", ["Dallas, TX"],
        ["Registered Nurse", "Pharmacist"], "100 hires"))


@pytest.fixture(scope="module")
def india_it():
    return T.build_plan_data(_brief(
        "Probe Infotech", "tech_engineering", "₹25,000,000",
        ["Bengaluru, Karnataka, India", "Pune, Maharashtra, India"],
        ["Java Developer", "QA Engineer"], "500+ hires"))


class TestNoCrossIndustryDefaultUnderAnIndustryName:
    def test_slide5_shows_the_industrys_own_kb_row(self, hershey):
        assert hershey["_budget_allocation"]["metadata"]["industry_cph"][
            "industry_matched"
        ] is False
        texts = _slide(hershey, 5)
        val = _value_after(texts, "Industry Cost-per-Hire")
        assert "4,000" not in val and "6,000" not in val, val
        assert "$3,000-$5,000" in val, val

    def test_workbook_states_the_same_row_and_names_the_default(self, hershey):
        flat = [str(v) for r in _exec_summary_rows(hershey) for v in r]
        assert any("$3,000-$5,000 (food manufacturing)" in v for v in flat)
        # round 2 (item 3): the range's low end is the SHOWN row's midpoint
        rng = [v for v in flat if v.startswith("Projected hires:")]
        assert rng and (
            "$4,000, the midpoint of the industry row "
            "($3,000-$5,000, food manufacturing)" in rng[0]
        ), rng


class TestMidpointLabelAndAgreement:
    def test_us_row_says_midpoint_and_matches_workbook(self, hospital):
        val = _value_after(_slide(hospital, 5), "Industry Cost-per-Hire")
        assert "midpoint" in val and "avg" not in val, val
        rows = _exec_summary_rows(hospital)
        row = next(r for r in rows if r and r[0] == "Industry Cost-per-Hire")
        assert row[1] == "$9,000–$12,000 (midpoint $10,500)", row

    def test_local_row_is_cited_range_midpoint_one_line_and_sourced(self, india_it):
        texts = _slide(india_it, 5)
        val = _value_after(texts, "Local Cost-per-Hire")
        assert "median" not in val and "midpoint" in val, val
        assert ppt_generator._estimate_lines(val, 2.3, 10.0) == 1, val
        assert any(
            t.startswith("Sources:") and "Shework 2026 (India)" in t for t in texts
        ), texts
        rows = _exec_summary_rows(india_it)
        row = next(r for r in rows if r and r[0] == "Local Cost-per-Hire")
        assert row[1] == "₹35,000–₹80,000 (midpoint ₹57,500)", row


def test_display_marks_the_default_and_prefers_kb():
    d = be.industry_cph_display(
        be.resolve_industry_cph("food_beverage"), "USD"
    )
    assert d["prefer_kb_row"] is True
    assert "cross-industry default" in d["label"]
