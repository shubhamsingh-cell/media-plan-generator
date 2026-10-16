"""Mixed plans and empty labels (review of the hires-suppression rule).

Hires / CPH are suppressed only when NO market in the plan has a cost-per-hire
benchmark in the plan's currency. A USD plan with at least one US market keeps
hires (whole-plan, US basis) and says so; a London-only GBP plan stays suppressed.
"""

from __future__ import annotations

import io
import logging
import os
import re
import sys
import types

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
for _n in ("anthropic", "openai", "supabase", "redis", "qdrant_client", "sentence_transformers", "posthog"):
    sys.modules.setdefault(_n, types.ModuleType(_n))
logging.disable(logging.CRITICAL)

import plan_projection  # noqa: E402

LABEL = "basis: US benchmark; non-US markets not separately projected"
NOTE = plan_projection.NO_LOCAL_HIRES_NOTE


def _project(budget_text, locs, role="Registered Nurse", ind="healthcare_medical", budget=60000):
    return plan_projection.project_plan(
        {"budget": budget, "roles": [role], "locations": locs, "industry": ind, "budget_text": budget_text}
    )


def test_dallas_plus_london_in_usd_keeps_hires_with_the_us_basis_label():
    p = _project("$60,000", ["Dallas, TX", "London, UK"])
    assert p["assumptions"]["hires_suppressed"] is False
    assert p["hires"]["low"] and p["hires"]["high"] > p["hires"]["low"]
    assert p["cph"]["low"] and p["cph"]["high"]
    assert p["assumptions"]["mixed_markets"] is True
    assert LABEL in p["text"]["hires_short"] and LABEL in p["text"]["cph_short"]


def test_london_only_gbp_stays_suppressed():
    p = _project("£60,000", ["London, UK"])
    assert p["assumptions"]["hires_suppressed"] is True and p["hires"]["low"] is None
    assert p["text"]["hires_note"] == NOTE


def test_london_only_usd_stays_suppressed():
    assert _project("$60,000", ["London, UK"])["assumptions"]["hires_suppressed"] is True


def test_us_only_is_unchanged_and_unlabelled():
    p = _project("$60,000", ["Dallas, TX"])
    assert p["assumptions"]["hires_suppressed"] is False and p["assumptions"]["mixed_markets"] is False
    assert LABEL not in p["text"]["hires_short"]
    assert (p["text"]["hires"], p["text"]["cph"]) == ("5–11", "$5,250–$10,500")


def test_mixed_plan_agrees_across_wizard_quick_plan_and_nova():
    import app
    import quick_plan
    from nova import Nova

    locs = ["Dallas, TX", "London, UK"]
    wiz = app._compute_plan_estimate(
        {"budget": "$60,000", "industry": "healthcare_medical", "roles": ["Registered Nurse"],
         "locations": locs, "client_name": "Acme"}
    )["projection"]
    q = quick_plan.generate_quick_plan("Registered Nurse", "Dallas, TX", "$60,000", "healthcare_medical")  # single-location tool
    n = Nova.__new__(Nova)
    n._data_cache = {}
    nova = n._query_budget_projection(
        {"budget": 60000, "roles": ["Registered Nurse"], "locations": locs, "industry": "healthcare_medical"}
    )["projection"]
    assert wiz["assumptions"]["hires_suppressed"] is False
    assert {m: (wiz[m]["low"], wiz[m]["high"]) for m in ("hires", "cph")} == {
        m: (nova[m]["low"], nova[m]["high"]) for m in ("hires", "cph")
    }
    assert LABEL in wiz["text"]["hires_short"] and q["success"]


def _bundle(budget, locs):
    import excel_v2
    import ppt_generator
    import tools_regen_bundles as T
    from kb_loader import load_knowledge_base
    from openpyxl import load_workbook
    from pptx import Presentation

    data = T.build_plan_data(
        {"client_name": "Probe", "industry": "Healthcare & Medical", "budget": budget, "campaign_duration": "3 months",
         "hire_volume": "Not specified", "work_environment": "onsite", "locations": locs,
         "roles": ["Registered Nurse"], "target_roles": ["Registered Nurse"]}
    )
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    wb = load_workbook(io.BytesIO(raw))
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    return wb, prs


def _deck(prs):
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    yield i, p.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        yield i, c.text


def test_mixed_bundle_keeps_hires_and_labels_the_basis():
    wb, prs = _bundle("$150,000", ["Dallas, TX", "London, UK"])
    cells = [str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value is not None]
    assert any(v.startswith("Projected hires") for v in cells), "mixed plan lost its Projected hires line"
    assert not any(NOTE in v for v in cells)
    assert any(LABEL in v for v in cells)
    deck = " || ".join(t for _, t in _deck(prs))
    assert NOTE not in deck and "This plan projects" in deck


# ---- no empty labels left by the suppression --------------------------------
EMPTY_LABEL = re.compile(r"^\s*(?:Impact|Hires|Expected hires|Projected hires|Mitigation)\s*:\s*$", re.I)


@pytest.mark.parametrize("loc", ["London, UK"])
def test_no_label_is_left_empty_in_a_suppressed_bundle(loc):
    wb, prs = _bundle("£60,000", [loc])
    empties = [(i, t) for i, t in _deck(prs) if EMPTY_LABEL.match(t)]
    assert not empties, empties
    # the risk slide: every card still has a non-empty impact statement
    risk = [t for i, t in _deck(prs) if i == 9]
    impacts = [t for t in risk if t.startswith("Impact:")]
    assert len(impacts) >= 3 and all(len(t) > len("Impact: ") + 8 for t in impacts), risk
    assert not any(re.search(r"\bhires?\b", t, re.I) for t in impacts), impacts
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and EMPTY_LABEL.match(c.value):
                    nxt = ws.cell(row=c.row, column=c.column + 1).value
                    assert nxt not in (None, ""), (ws.title, c.coordinate, c.value)
