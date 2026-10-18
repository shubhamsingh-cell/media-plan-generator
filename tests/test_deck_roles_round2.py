"""Slide 5 / slide 8 trust and layout fixes (design-judge round 2).

DJ1  the hardcoded ad-platform table is never captioned or coloured "live"
DJ2  slide 5 keeps OBSERVED evidence rows and evicts curated rows first, and
     the "Live Postings" value is one line built from structured fields
DJ4  slide 8's industry column is labelled a benchmark, not a projection
DJ5  slide 8 withholds hires like slides 2 and 6 do on a no-local-CPH market
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

import pytest
from pptx import Presentation
from pptx.util import Inches

import data_synthesizer
import ppt_generator
import tools_regen_bundles as T
from kb_loader import load_knowledge_base

logging.disable(logging.CRITICAL)

TODAY = date.today().isoformat()

# The deck receives the LEGACY industry key from app.py; that key (not the
# normalised "healthcare_medical") is what routes slide 5 through the generic
# ad-platform range layer that used to stamp confidence "live_api".
US_TWO_ROLE = {
    "client_name": "Lone Star Health Network",
    "industry": "healthcare",
    "budget": "$250,000",
    "campaign_duration": "3 months",
    "hire_volume": "100 hires",
    "work_environment": "onsite",
    "locations": ["Houston, TX", "Phoenix, AZ"],
    "roles": ["Registered Nurse", "Travel Nurse"],
    "target_roles": ["Registered Nurse", "Travel Nurse"],
}
UK_WAREHOUSE = {
    "client_name": "Northern Fulfilment Group",
    "industry": "logistics",
    "budget": "£400,000",
    "campaign_duration": "3 months",
    "hire_volume": "100 hires",
    "work_environment": "onsite",
    "locations": ["Manchester, United Kingdom"],
    "roles": ["Warehouse Associate"],
    "target_roles": ["Warehouse Associate"],
}
US_ONE_MARKET = {
    "client_name": "Buckeye Logistics",
    "industry": "logistics",
    "budget": "$150,000",
    "campaign_duration": "3 months",
    "hire_volume": "20 hires",
    "work_environment": "onsite",
    "locations": ["Columbus, OH"],
    "roles": ["Warehouse Associate"],
    "target_roles": ["Warehouse Associate"],
}


def _plan(brief: Dict[str, Any], observed: bool, extra_jm: Optional[dict] = None):
    data = T.build_plan_data(brief)
    # legacy industry key / label, as app.py hands them to the deck
    legacy = {"healthcare_medical": "healthcare", "logistics_supply_chain": "logistics"}
    data["industry"] = legacy.get(data["industry"], data["industry"])
    data["industry_label"] = data["industry"].title()
    roles = brief["roles"]
    if observed:
        jm = {
            role: {
                "posting_count": 7000,
                "avg_salary": 52000,
                "competition": "medium",
                "source": "Adzuna",
                "posting_observation": "observed",
                "posting_provider_id": "adzuna",
                "posting_observed_at": TODAY,
                # the collector's own scope (api_enrichment.fetch_job_market)
                "posting_scope": "national",
                "posting_scope_label": "",
                **(extra_jm or {}),
            }
            for role in roles
        }
    else:
        import api_enrichment

        jm = api_enrichment._adzuna_benchmark_fallback(roles, brief["locations"])
    enriched = {"job_market": jm}
    data["_enriched"] = enriched
    data["_synthesized"] = data_synthesizer.synthesize(
        enriched, data.get("_knowledge_base") or load_knowledge_base(), data
    )
    return data


def _deck(data: dict) -> Presentation:
    return Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))


def _slide(prs: Presentation, title: str):
    for slide in prs.slides:
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        if title in texts:
            return slide
    raise AssertionError(f"no slide titled {title!r}")


def _slide5(prs):
    return _slide(prs, "Channel Strategy & Investment")


def _slide8(prs):
    return _slide(prs, "Plan Comparison & Implementation")


def _shapes_with_text(slide) -> List[Any]:
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def _run_rgb(shape) -> Optional[Any]:
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                return run.font.color.rgb
            except (AttributeError, TypeError):
                return None
    return None


def _inch(shape, attr: str) -> float:
    return getattr(shape, attr) / 914400.0


# ---------------------------------------------------------------------- DJ1


def test_platform_benchmark_table_is_never_confidence_live_api() -> None:
    data = _plan(US_TWO_ROLE, observed=False)
    bench = ppt_generator._get_benchmarks(data["industry"], data)
    assert bench["cpa_is_generic_platform_range"], "fixture must hit the platform layer"
    assert bench["confidence"] != "live_api"
    assert bench["confidence"] == ppt_generator.PLATFORM_BENCHMARK_CONFIDENCE


def test_platform_benchmark_confidence_is_curated_not_live_styled() -> None:
    colour, label = ppt_generator._confidence_color(
        ppt_generator.PLATFORM_BENCHMARK_CONFIDENCE
    )
    assert "live" not in label.lower()
    assert colour != ppt_generator.GREEN
    # the genuine live label/colour is still reachable for an observed source
    assert ppt_generator._confidence_color("live_api") == (
        ppt_generator.GREEN,
        "Live Data",
    )


@pytest.mark.parametrize("brief", [US_TWO_ROLE, UK_WAREHOUSE], ids=["us", "uk"])
@pytest.mark.parametrize("observed", [False, True], ids=["nokeys", "observed"])
def test_slide5_does_not_caption_the_hardcoded_table_as_live(brief, observed) -> None:
    slide = _slide5(_deck(_plan(brief, observed)))
    texts = [s.text_frame.text for s in _shapes_with_text(slide)]
    assert "Live Data" not in texts
    # no benchmark cell (CPA/CPC/CPH/apply rate) is painted live-green
    labels = {"Industry CPA", "Industry CPC", "Industry Cost-per-Hire", "Apply Rate"}
    shapes = _shapes_with_text(slide)
    for i, shape in enumerate(shapes):
        if any(shape.text_frame.text.startswith(lbl) for lbl in labels):
            value = shapes[i + 1]
            assert _run_rgb(value) != ppt_generator.GREEN, (
                shape.text_frame.text,
                value.text_frame.text,
            )


# ---------------------------------------------------------------------- DJ2

# the collector scopes Adzuna counts to the whole country, so the cell says so
# ("US nationwide" on a US-only plan) on a short first line; the provider and
# the observation date follow
_POSTINGS_RE = re.compile(
    r"^7,000 jobs \u00b7 (US )?nationwide\nAdzuna \u00b7 \d{4}-\d{2}-\d{2}$"
)


def _live_postings(slide) -> List[Tuple[Any, Any]]:
    """[(label_shape, value_shape)] for every Live Postings row."""
    shapes = _shapes_with_text(slide)
    return [
        (shape, shapes[i + 1])
        for i, shape in enumerate(shapes)
        if shape.text_frame.text.startswith("Live Postings: ")
    ]


@pytest.mark.parametrize(
    "brief", [US_TWO_ROLE, UK_WAREHOUSE], ids=["us-two-role", "uk"]
)
def test_slide5_keeps_the_observed_live_postings_row(brief) -> None:
    slide = _slide5(_deck(_plan(brief, observed=True)))
    rows = _live_postings(slide)
    assert rows, "observed postings row was evicted from slide 5"
    label, value = rows[0]
    text = value.text_frame.text
    assert _POSTINGS_RE.match(text), text
    assert _run_rgb(value) == ppt_generator.GREEN  # observed => live green
    # a stacked cell is real paragraphs: a raw newline inside one run is plain
    # whitespace to PowerPoint and the two lines would run together
    for para in value.text_frame.paragraphs:
        assert all("\n" not in run.text for run in para.runs)
    assert len(value.text_frame.paragraphs) == len(text.split("\n"))
    # every line is fitted to the value column (no unplanned third wrap)
    pt = value.text_frame.paragraphs[0].runs[0].font.size.pt
    for line in text.split("\n"):
        assert ppt_generator._fits_one_line(line, 2.3, pt, bold=True), (line, pt)
    # old nested-bracket string is gone
    assert "active jobs" not in text and "(Adzuna (" not in text


@pytest.mark.parametrize(
    "brief", [US_TWO_ROLE, UK_WAREHOUSE], ids=["us-two-role", "uk"]
)
def test_slide5_orders_observed_rows_right_after_cpa_cpc_cph(brief) -> None:
    slide = _slide5(_deck(_plan(brief, observed=True)))
    labels = [
        s.text_frame.text
        for s in _shapes_with_text(slide)
        if _inch(s, "left") < 6.0 or _inch(s, "left") > 6.9
    ]
    rows = [
        t
        for t in labels
        if t.startswith(
            ("Industry CPA", "Industry CPC", "Industry Cost", "Live Postings")
        )
        or t == "Apply Rate"
    ]
    assert rows[:3] and all(r.startswith("Industry") for r in rows[:3]), rows
    assert rows[3].startswith("Live Postings"), rows
    if "Apply Rate" in rows:
        assert rows.index("Apply Rate") > rows.index(rows[3])


def test_slide5_evicts_curated_rows_before_observed_rows() -> None:
    """Space runs out on the two-role US table: BOTH observed Live Postings
    rows (each says what the count covers, so takes two short lines) are kept
    and the apply-rate benchmark / ad-platform rows are what give way."""
    slide = _slide5(_deck(_plan(US_TWO_ROLE, observed=True)))
    texts = [s.text_frame.text for s in _shapes_with_text(slide)]
    assert any(t.startswith("Live Postings: Registered Nurse") for t in texts)
    assert any(t.startswith("Live Postings: Travel Nurse") for t in texts)
    assert "Apply Rate" not in texts
    assert not any(t.startswith("Google Ads") for t in texts)


def test_slide5_nokeys_deck_has_no_live_postings_row() -> None:
    """No observation -> no row: the benchmark fallback is never a live count."""
    slide = _slide5(_deck(_plan(UK_WAREHOUSE, observed=False)))
    assert not _live_postings(slide)


def test_slide5_table_closes_the_band_above_the_attribution_header() -> None:
    slide = _slide5(_deck(_plan(US_TWO_ROLE, observed=True)))
    shapes = _shapes_with_text(slide)
    src = next(
        s
        for s in shapes
        if s.text_frame.text.startswith(("Sources: ", "Industry ranges"))
    )
    header = next(
        s for s in shapes if s.text_frame.text == "CHANNEL CATEGORY ATTRIBUTION"
    )
    gap = _inch(header, "top") - (_inch(src, "top") + _inch(src, "height"))
    # pre-fix the table stopped at ~4.4in and left a ~0.48in gap below the
    # sources line (observed rows evicted, body rows at the 0.38in default)
    assert gap < 0.3, gap


def _jm(primary: dict, details_dates=None) -> dict:
    """A role's job_market_demand block in the data layer's real shape."""
    dates = details_dates or primary.get("observed_dates") or [primary["observed_at"]]
    return {
        "total_postings": primary["count"],
        "posting_primary": primary,
        "posting_source_details": [
            {
                "source": primary["provider"],
                "provider_id": primary["provider_id"],
                "observed_at": d,
            }
            for d in dates
        ],
        "market_evidence": {
            "contract": "market_evidence_v1",
            "metrics": {
                "job_postings": {
                    "status": "observed",
                    "provider_ids": [primary["provider_id"]],
                    "observed_at": sorted(set(dates)),
                }
            },
        },
    }


_ADZUNA = {
    "count": 7000,
    "provider": "Adzuna",
    "provider_id": "adzuna",
    "observed_at": "2026-10-15",
    "observed_dates": ["2026-10-15"],
    "scope": "unspecified",
    "scope_label": "",
}


def test_observed_postings_cell_formats_posting_primary() -> None:
    cell = ppt_generator._observed_postings_cell
    # no stated scope: nothing is guessed, one line
    text, pt, lines = cell(_jm(_ADZUNA), 2.3)
    assert (text, lines) == ("7,000 jobs \u00b7 Adzuna \u00b7 2026-10-15", 1)
    # a nationwide count says so; one line cannot hold it -> two short lines
    national = dict(_ADZUNA, scope="national")
    text, pt, lines = cell(_jm(national), 2.3)
    assert text == "7,000 jobs \u00b7 nationwide\nAdzuna \u00b7 2026-10-15"
    assert lines == 2
    # a US-only plan says "US nationwide"
    text, _, _ = cell(_jm(national), 2.3, {"locations": ["Columbus, OH"]})
    assert text.split("\n")[0] == "7,000 jobs \u00b7 US nationwide"
    # a location-scoped count names its place
    london = dict(
        _ADZUNA,
        count=2000,
        provider="Jooble API",
        provider_id="jooble",
        observed_at="2026-10-14",
        observed_dates=["2026-10-14"],
        scope="location",
        scope_label="London, UK",
    )
    text, _, _ = cell(_jm(london), 2.3)
    assert text.split("\n")[0] == "2,000 jobs \u00b7 London, UK"
    assert "Jooble API \u00b7 2026-10-14" in text
    # several observation dates stack, none dropped
    multi = dict(
        _ADZUNA,
        count=6300,
        scope="national",
        observed_dates=["2026-10-13", "2026-10-14", "2026-10-15"],
    )
    text, _, lines = cell(_jm(multi), 2.3)
    assert lines == 2
    assert text.endswith("Adzuna \u00b7 2026-10-13\u20132026-10-15")


def test_observed_postings_cell_suppressed_without_backed_primary() -> None:
    cell = ppt_generator._observed_postings_cell
    assert cell({}, 2.3) is None  # no posting_primary
    assert cell(dict(_jm(_ADZUNA), posting_primary=None), 2.3) is None
    # a date the market evidence does not back is not printed
    forged = _jm(_ADZUNA, details_dates=["2026-10-01"])
    assert cell(forged, 2.3) is None
    # a provider the evidence does not name is not printed
    other = _jm(dict(_ADZUNA, provider_id="jooble", provider="Jooble API"))
    other["posting_primary"] = dict(_ADZUNA)  # claims Adzuna; evidence says Jooble
    assert cell(other, 2.3) is None


def test_slide5_says_nationwide_when_the_collector_scopes_the_count_national() -> None:
    slide = _slide5(_deck(_plan(US_TWO_ROLE, observed=True)))
    rows = _live_postings(slide)
    assert rows
    assert "nationwide" in rows[0][1].text_frame.text


def test_slide5_names_the_place_for_a_location_scoped_count() -> None:
    plan = _plan(
        US_TWO_ROLE,
        observed=True,
        extra_jm={"posting_scope": "location", "posting_scope_label": "Houston, TX"},
    )
    slide = _slide5(_deck(plan))
    rows = _live_postings(slide)
    assert rows
    assert "Houston, TX" in rows[0][1].text_frame.text


# ---------------------------------------------------------------------- DJ4


@pytest.mark.parametrize("brief", [US_TWO_ROLE, UK_WAREHOUSE], ids=["us", "uk"])
def test_slide8_labels_the_industry_cpa_a_nova_estimate(brief) -> None:
    slide = _slide8(_deck(_plan(brief, observed=False)))
    right = [
        s.text_frame.text
        for s in _shapes_with_text(slide)
        if _inch(s, "left") > 6.5 and _inch(s, "top") > 1.5 and _inch(s, "top") < 5.0
    ]
    assert "Industry CPA (Nova estimate)" in right
    assert "Projected CPA" not in right
    left = [
        s.text_frame.text for s in _shapes_with_text(slide) if _inch(s, "left") < 6.5
    ]
    # the plan's own figure keeps its label (D1: "(plan efficiency)" is added
    # where slide 2 prints the same label over a range)
    assert any(t.startswith("Projected CPA") for t in left)


# ---------------------------------------------------------------------- DJ5


def _value_after(slide, label: str, *, left_panel: bool = True) -> Optional[str]:
    shapes = _shapes_with_text(slide)
    for i, shape in enumerate(shapes):
        on_left = _inch(shape, "left") < 6.5
        if shape.text_frame.text == label and on_left == left_panel:
            return shapes[i + 1].text_frame.text
    return None


def test_slide8_withholds_hires_on_a_market_with_no_local_cph() -> None:
    prs = _deck(_plan(UK_WAREHOUSE, observed=False))
    slide8 = _slide8(prs)
    shown = _value_after(slide8, "Projected Hires")
    assert shown is not None, "fixture must render the Projected Hires row"
    assert not re.search(r"\d", shown), shown
    assert shown == "n/a"
    # slides 2 and 6 say the same thing
    # looked up by title: origin/main inserted a Client Story slide, so
    # positional indexes silently point at a different slide
    all2 = " ".join(
        s.text_frame.text for s in _shapes_with_text(_slide(prs, "Executive Summary"))
    )
    assert "n/a" in all2
    all6 = " ".join(
        s.text_frame.text
        for s in _shapes_with_text(_slide(prs, "Budget Allocation & Projections"))
    )
    assert "hires not projected" in all6
    # and no client-goal band computed from a count that is not projected
    assert "this plan projects" not in " ".join(
        s.text_frame.text for s in _shapes_with_text(slide8)
    )


def test_slide8_still_projects_hires_when_a_local_benchmark_exists() -> None:
    slide8 = _slide8(_deck(_plan(US_ONE_MARKET, observed=False)))
    shown = _value_after(slide8, "Projected Hires")
    assert shown is not None
    assert re.fullmatch(r"[\d,]+", shown.strip()), shown


# ----------------------------------------------------- final round: D1, D2, D3


def _text_of(slide) -> str:
    return "\n".join(s.text_frame.text for s in _shapes_with_text(slide))


def _kpi_value_before(slide, label: str) -> Optional[str]:
    """The KPI value printed directly above ``label`` on the slide-2 bar."""
    shapes = _shapes_with_text(slide)
    for i, shape in enumerate(shapes):
        if shape.text_frame.text == label and i:
            return shapes[i - 1].text_frame.text
    return None


def test_d1_slides_2_and_9_never_carry_one_cpa_label_over_two_numbers() -> None:
    prs = _deck(_plan(US_TWO_ROLE, observed=False))
    kpi = _kpi_value_before(_slide(prs, "Executive Summary"), "Projected CPA")
    assert kpi is not None and "\u2013" in kpi, kpi  # slide 2 prints a range
    cmp_slide = _slide8(prs)
    # slide 9's own figure is the plan-efficiency point of that range, and its
    # label says so; the bare slide-2 label does not appear over a lone point
    assert _value_after(cmp_slide, "Projected CPA") is None
    point = _value_after(cmp_slide, "Projected CPA (plan efficiency)")
    assert point is not None
    point_num = float(re.search(r"[\d,.]+", point).group(0).replace(",", ""))
    low, high = (
        float(x.lstrip("$\u00a3\u20ac").replace(",", "").strip())
        for x in kpi.split("\u2013")
    )
    # The point is a point OF the range. It used to be the range's low end (an
    # engine CPA below the KB band was the hull's low); after the healthcare
    # apply-rate re-level (Appcast 2026 healthcare apply rate 3.22%, CPA-realism
    # C1) the engine CPA sits inside the KB band, so the range low is the band's
    # low ($35) and the plan-efficiency point lies within it ($36.80). Slide 2
    # then reads $35-$58 beside slide 9's point under a different label.
    assert low <= point_num <= high, (kpi, point)


def test_d1_slide_9_keeps_the_plain_label_when_slide_2_prints_a_point() -> None:
    prs = _deck(_plan(UK_WAREHOUSE, observed=False))
    kpi = _kpi_value_before(_slide(prs, "Executive Summary"), "Projected CPA")
    assert kpi is not None and "\u2013" not in kpi, kpi
    point = _value_after(_slide8(prs), "Projected CPA")
    assert point is not None and point.split()[0] == kpi


def test_d2_uk_slide_2_keeps_the_sourced_ons_vacancies_line() -> None:
    # _autofit_textframe measured at the unfloored 7pt size, chose scale 0.92,
    # then floored the runs to 8pt, overflowed and trimmed the cited ONS line
    # (~1in of the RESOLUTION card was left empty).
    text = _text_of(
        _slide(_deck(_plan(UK_WAREHOUSE, observed=False)), "Executive Summary")
    )
    assert "(ONS, 2026)" in text, text
    assert "705,000" in text


def test_d3_national_count_names_its_country_when_the_collector_did() -> None:
    cell = ppt_generator._observed_postings_cell
    uk = dict(_ADZUNA, scope="national", scope_label="United Kingdom")
    text, _, _ = cell(_jm(uk), 3.6, {"locations": ["London, UK", "Leeds, UK"]})
    assert "United Kingdom \u00b7 nationwide" in text, text
    # the US keeps its short form, whatever the plan's other markets
    us = dict(_ADZUNA, scope="national", scope_label="United States")
    text, _, _ = cell(_jm(us), 3.6, {"locations": ["Columbus, OH", "London, UK"]})
    assert "US nationwide" in text and "United States" not in text, text
    # no label: unchanged fallbacks
    bare = dict(_ADZUNA, scope="national", scope_label="")
    assert ppt_generator._posting_scope_phrase(bare, None) == "nationwide"
    assert (
        ppt_generator._posting_scope_phrase(bare, {"locations": ["Columbus, OH"]})
        == "US nationwide"
    )
