"""Regression: the hiring-goal gap never prints per-hire or top-up arithmetic
it cannot stand behind (design-judge panel, 2026-10-01, item 5).

Pre-fix (branch at 6e2aaf5):
* A Brazil plan -- no local cost-per-hire benchmark, so the hero card shows
  "--" -- still printed "Closing the ~363-hire gap at this plan's
  R$8,759/hire would need roughly R$3.2M of additional budget" in the
  Executive Summary gap row, and the strategic summary repeated the top-up
  and "At a blended cost of R$8,759.12 per hire". That per-hire cost is
  bounded by an FX-translated US figure the plan must never print.
* A $3,000 plan projecting 0 hires against a goal of 1 was told to add
  $3,350 (total ~$6,350): the full cost of one hire on TOP of the budget
  already committed, i.e. (goal - projected) x cost per hire, with a cost
  the plan never achieves.
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import display_format  # noqa: E402
import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402


def _brief(name, industry, budget, locations, roles, hire_volume):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": hire_volume,
        "work_environment": "onsite",
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
    }


def _exec_texts(data):
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    if isinstance(raw, tuple):
        raw = raw[0]
    ws = load_workbook(io.BytesIO(raw))["Executive Summary"]
    return [str(c.value) for r in ws.iter_rows() for c in r if c.value not in (None, "")]


def _gap_row(texts):
    rows = [t for t in texts if t.startswith("Hiring-goal gap")]
    assert len(rows) == 1, rows
    return rows[0]


@pytest.fixture(scope="module")
def brazil():
    return T.build_plan_data(_brief(
        "Probe Contact BR", "general_entry_level", "R$1,200,000",
        ["São Paulo, Brazil", "Rio de Janeiro, Brazil"],
        ["Customer Service Representative", "Telemarketing Agent"], "500+ hires"))


@pytest.fixture(scope="module")
def tiny():
    return T.build_plan_data(_brief(
        "Probe Food Bank", "general_entry_level", "$3,000", ["Sacramento, CA"],
        ["Volunteer Coordinator", "Warehouse Associate"], "1-10 hires"))


class TestGoalGapHelper:
    def test_fewer_than_one_hire_has_no_top_up(self):
        gap = display_format.goal_gap(0, 1, 3350.0, budget=3000.0)
        assert gap["additional_budget"] is None
        assert gap["cost_per_hire"] == 3350.0

    def test_top_up_is_goal_cost_minus_committed_budget(self):
        gap = display_format.goal_gap(60, 100, 5000.0, budget=300_000.0)
        assert gap["additional_budget"] == 100 * 5000.0 - 300_000.0


class TestSuppressedMarket:
    def test_gap_row_has_no_per_hire_or_top_up(self, brazil):
        cph = brazil["_budget_allocation"]["metadata"]["industry_cph"]
        assert cph["claim_suppressed"] is True
        row = _gap_row(_exec_texts(brazil))
        assert "/hire" not in row and "additional budget" not in row, row
        assert "no local cost-per-hire benchmark" in row, row

    def test_summary_cites_no_per_hire_cost(self, brazil):
        texts = _exec_texts(brazil)
        summary = [t for t in texts if t.startswith("This recruitment media plan")]
        assert summary, "no strategic summary"
        assert "per hire" not in summary[0], summary[0]
        assert "additional budget" not in summary[0], summary[0]

    def test_deck_states_no_top_up(self, brazil):
        prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(brazil))))
        blob = "\n".join(
            sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame
        )
        assert "would close the gap" not in blob, blob


class TestFewerThanOneHire:
    def test_gap_row_states_the_cost_of_one_hire_not_an_add(self, tiny):
        assert tiny["_budget_allocation"]["total_projected"]["hires"] == 0
        row = _gap_row(_exec_texts(tiny))
        assert "fewer than 1 hire" in row, row
        assert "indicative budget for one hire ≈ $3,350" in row, row
        assert "additional budget" not in row and "total ~" not in row, row
