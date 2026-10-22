"""ONE CPA benchmark resolver for the deck and the workbook (2026-10-23).

Review findings (scores 95 / 88): the deck took its CPA range from
``ppt_generator._get_benchmarks`` (which falls back to the cross-industry
ad-platform span when the knowledge base has no row for the industry) while the
verdict looked the knowledge base up by the raw industry key and a non-existent
``general_entry_level`` row -- so a CPA far below the printed "$12 - $56" got a
green "Ahead of Nova estimate" badge and no note. And for a plan with no US
market the verdict fell through to the US industry figure, so a UK + Germany
plan in USD was told its CPA was "below the indicative industry range ($12.00)"
on a slide that prints no range.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

import pytest
from pptx import Presentation

import excel_v2
import plan_validator
import ppt_generator

from tests.test_deck_iterate_1022 import _us_plan
from tests.test_plan_integrity import _german_plan, _workbook_texts

_ROOT = Path(__file__).resolve().parent.parent

# the industries the wizard offers (templates/partials/index/body_content.html)
WIZARD_INDUSTRIES = sorted(
    set(
        re.findall(
            r'data-industry="([a-z_]+)"',
            (_ROOT / "templates/partials/index/body_content.html").read_text(),
        )
    )
)

# every platform's CPA range, as data_synthesizer hands it to the generators
_PLATFORMS = {
    "google": {"avg_cpa": 12.4},
    "indeed": {"avg_cpa": 28.0},
    "linkedin": {"avg_cpa": 55.6},
}

_BASE = None


def _base_us_plan():
    global _BASE
    if _BASE is None:
        _BASE = _us_plan(10.0)
    return copy.deepcopy(_BASE)


def _plan(industry, cpa, platforms=True):
    data = _base_us_plan()
    data["industry"] = industry
    if platforms:
        data["_synthesized"] = {"ad_platform_analysis": copy.deepcopy(_PLATFORMS)}
    data["_budget_allocation"]["total_projected"]["cost_per_application"] = cpa
    return data


def _printed_range(industry, data):
    """(low, high) of the CPA range the deck prints for this plan."""
    nums = [
        float(n.replace(",", ""))
        for n in re.findall(r"\d[\d,]*(?:\.\d+)?", ppt_generator._get_benchmarks(industry, data)["cpa"])
    ]
    return min(nums), max(nums)


def _slide(data):
    prs = Presentation()
    ppt_generator._build_slide_comparison_timeline(prs, copy.deepcopy(data))
    return prs.slides[0]


def _cpa_value_text(data, cpa):
    """The text of the 'Projected CPA' value cell (figure + badge)."""
    shown = (f"${cpa:,.0f}", f"${cpa:,.2f}")  # whole dollars print without cents
    for sh in _slide(data).shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(shown):
            return sh.text_frame.text
    raise AssertionError(f"no CPA value cell for ${cpa:,.2f}")


def _texts(data):
    return [sh.text_frame.text for sh in _slide(data).shapes if sh.has_text_frame]


def test_wizard_industry_list_is_found():
    assert "general_entry_level" in WIZARD_INDUSTRIES and len(WIZARD_INDUSTRIES) >= 20


# ---------------------------------------------------------------------------
# (1) every wizard industry: below the printed low -> never green, note on both
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("industry", WIZARD_INDUSTRIES)
def test_cpa_below_the_printed_low_is_never_green_and_both_surfaces_say_so(industry):
    low, _ = _printed_range(industry, _plan(industry, 10.0))
    cpa = round(low * 0.4, 2)
    data = _plan(industry, cpa)
    cell = _cpa_value_text(data, cpa)
    assert "▲" not in cell, f"{industry}: green badge on {cell!r}"
    assert "!" in cell, f"{industry}: no amber badge on {cell!r}"

    gap = plan_validator.plan_cpa_gap(data)
    assert gap and gap["status"] == "below", f"{industry}: verdict never fires"
    long_msg = plan_validator.benchmark_gap_message(gap)
    short_msg = plan_validator.benchmark_gap_message(gap, short=True)
    fallback = plan_validator.benchmark_gap_fallback(gap)
    assert long_msg and short_msg
    # the deck prints the short sentence (or its one-line fallback) ...
    assert any(short_msg in t or fallback in t for t in _texts(data)), industry
    # ... the workbook the long one, both from the one gap, naming the same CPA
    # and the same range low
    wb = _workbook_texts(data)
    assert long_msg in " ".join(wb["Executive Summary"]), industry
    assert excel_v2._plan_integrity_notes(data)[1] == long_msg
    assert f"${cpa:,.2f}" in long_msg and f"${cpa:,.2f}" in short_msg
    rng = plan_validator._usd_range_text(gap["low_usd"], gap["high_usd"])
    assert rng in short_msg
    assert rng in long_msg or f"${gap['low_usd']:,.2f}" in long_msg


@pytest.mark.parametrize("industry", ["general_entry_level", "tech_engineering", "logistics_supply_chain"])
def test_the_range_the_verdict_tests_is_the_range_the_deck_prints(industry):
    data = _plan(industry, 10.0)
    printed = _printed_range(industry, data)
    bench = plan_validator.resolve_cpa_benchmark(data)
    assert (bench["low"], bench["high"]) == printed
    # the comparison row prints the resolver's own text
    assert any(bench["text"] == t for t in _texts(data)), (bench["text"], _texts(data))


def test_cross_industry_fallback_is_labelled_as_such():
    data = _plan("general_entry_level", 3.91)
    bench = plan_validator.resolve_cpa_benchmark(data)
    assert bench["scope"] == "cross_industry_platforms"
    assert "cross-industry" in bench["basis_label"]
    # the row prints the range the verdict tests (label pinned by DJ4 in
    # test_deck_roles_round2), and the note names it cross-industry
    assert any(bench["text"] == t for t in _texts(data))
    assert "cross-industry" in plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(data))


# ---------------------------------------------------------------------------
# (2) a plan with no US market never gets the US industry verdict
# ---------------------------------------------------------------------------


def _multi_country_usd_plan(cpa=7.48):
    data = _us_plan(cpa)
    data["industry"] = "logistics_supply_chain"
    data["locations"] = ["London, United Kingdom", "Berlin, Germany"]
    data["country"] = ""
    return data


def test_multi_country_non_us_plan_gets_no_us_industry_verdict():
    data = _multi_country_usd_plan()
    assert plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(data)) == ""
    assert plan_validator.plan_cpa_gap(data) is None
    assert plan_validator.industry_cpa_gap(data) is None
    assert plan_validator.resolve_cpa_benchmark(data) is None
    # deck: no note, no legend line, no CPA comparison row (the slide prints no range)
    texts = _texts(data)
    assert not any("optimistic until" in t for t in texts), texts
    assert not any("Industry CPA" in t for t in texts), texts
    # workbook: no verdict anywhere
    wb = _workbook_texts(data)
    joined = " ".join(t for texts in wb.values() for t in texts)
    assert "Model vs. benchmark" not in joined
    assert "indicative industry range" not in joined
    assert excel_v2._plan_integrity_notes(data)[1] == ""
    assert excel_v2._model_vs_benchmark_note(data, {"cpa": {"range": "$12-$35", "median": "$20"}}) is None


# ---------------------------------------------------------------------------
# (3) single-country non-US: unchanged
# ---------------------------------------------------------------------------


def test_single_country_market_range_is_unchanged():
    inside = _german_plan(cpa_eur=14.0, confidence=0.9)
    bench = plan_validator.resolve_cpa_benchmark(inside)
    assert bench["scope"] == "market_industry" and bench["country"] == "Germany"
    assert plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(inside)) == ""
    assert excel_v2._plan_integrity_notes(inside)[1] == ""
    below = _german_plan(cpa_eur=5.0, confidence=0.9)
    gap = plan_validator.plan_cpa_gap(below)
    assert gap["status"] == "below" and gap["basis"] == "industry"
    assert "this market's industry benchmark range (US$8–25 per application)" in (
        plan_validator.benchmark_gap_message(gap)
    )
    # the US-industry check never answers for it
    assert plan_validator.industry_cpa_gap(below) is None


# ---------------------------------------------------------------------------
# (4) no benchmark at all -> nothing, no crash
# ---------------------------------------------------------------------------


def test_no_benchmark_at_all_says_nothing():
    data = _plan("general_entry_level", 3.0, platforms=False)
    assert plan_validator.resolve_cpa_benchmark(data) is None
    assert plan_validator.plan_cpa_gap(data) is None
    cell = _cpa_value_text(data, 3.0)  # the slide still builds
    assert "▲" not in cell and "!" not in cell and "▼" not in cell
    assert not any("optimistic until" in t for t in _texts(data))
    assert excel_v2._plan_integrity_notes(data)[1] == ""
    assert excel_v2._model_vs_benchmark_note(data, {}) is None


def test_resolver_tolerates_junk():
    for junk in (None, {}, {"industry": None}, {"_synthesized": {"ad_platform_analysis": {"x": 1, "_m": {}}}}):
        assert plan_validator.resolve_cpa_benchmark(junk) is None


# ---------------------------------------------------------------------------
# (5) above the high -> behind, no note;  inside -> neutral, no note
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("industry", ["general_entry_level", "logistics_supply_chain"])
def test_above_the_high_is_behind_without_a_note(industry):
    _, high = _printed_range(industry, _plan(industry, 10.0))
    cpa = round(high * 1.5, 2)
    data = _plan(industry, cpa)
    cell = _cpa_value_text(data, cpa)
    assert "▼" in cell and "▲" not in cell and "!" not in cell
    assert plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(data)) == ""
    assert not any("optimistic until" in t for t in _texts(data))
    assert excel_v2._plan_integrity_notes(data)[1] == ""


@pytest.mark.parametrize("industry", ["general_entry_level", "logistics_supply_chain"])
def test_inside_the_range_is_neutral_without_a_note(industry):
    low, high = _printed_range(industry, _plan(industry, 10.0))
    cpa = round((low + high) / 2, 2)
    data = _plan(industry, cpa)
    cell = _cpa_value_text(data, cpa)
    assert "▲" not in cell and "▼" not in cell and "!" not in cell
    assert plan_validator.plan_cpa_gap(data)["status"] == "within"
    assert excel_v2._plan_integrity_notes(data)[1] == ""
