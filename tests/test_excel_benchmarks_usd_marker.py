"""Regression: US-calibrated KB benchmark rows carry "(USD)" on non-USD plans.

The Executive Summary's "Industry Benchmarks" block prints the knowledge
base's recruitment benchmarks verbatim -- US-dollar constants (CPA/CPC/CPH
ranges, "$289,000/year per hospital" turnover cost). On a GBP plan those
rows printed a bare "$" beside the plan's own £ figures, and the delivery
gate (bundle_qa currency_symbol_mixing) raised two criticals on a London
healthcare plan (2026-09-16). Longer rows escaped only because the gate
skips text over its length cutoff -- they were equally unmarked.

Contract: on a non-USD plan, every row in that block whose value carries a
"$" figure gets a "(USD)" label (the convention the gate and the Intl
Benchmarks sheet already use); rows with no money figure keep a plain label;
USD plans are unchanged.
"""

from __future__ import annotations

import io
import logging
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

import bundle_qa
import tools_regen_bundles as regen


def _london_brief():
    b = dict(regen.ATRIA_BRIEF)
    b.update(
        {
            "client_name": "London NHS Trust",
            "budget": "£2,000,000",
            "locations": ["London, UK"],
            "roles": ["Registered Nurse"],
            "target_roles": [
                {"title": "Registered Nurse", "count": 40, "tier": "Hourly"}
            ],
            "hire_volume": "40 hires",
            "notes": "Registered nurses for London hospitals.",
            "competitors": [],
        }
    )
    return b


def _bundle(brief):
    logging.disable(logging.CRITICAL)
    try:
        return regen.generate_bundle(dict(brief), Path(tempfile.mkdtemp()), "t")
    finally:
        logging.disable(logging.NOTSET)


def _benchmark_rows(xlsx_bytes):
    ws = load_workbook(io.BytesIO(xlsx_bytes))["Executive Summary"]
    rows, inside = [], False
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        value = ws.cell(r, 4).value
        if isinstance(label, str) and label.startswith("Industry Benchmarks"):
            inside = True
            continue
        if inside:
            if not label:
                break
            rows.append((str(label), str(value or "")))
    return rows


@pytest.fixture(scope="module")
def london():
    return _bundle(_london_brief())


@pytest.fixture(scope="module")
def atria():
    return _bundle(regen.ATRIA_BRIEF)


def test_gbp_plan_marks_every_dollar_benchmark_row_usd(london):
    rows = _benchmark_rows(london["xlsx_bytes"])
    money = [(lbl, v) for lbl, v in rows if "$" in v]
    assert len(money) >= 3, rows  # CPA / CPC / CPH at minimum
    for lbl, v in money:
        assert lbl.endswith("(USD)"), (lbl, v)
    for lbl, v in rows:
        if "$" not in v:
            assert "(USD)" not in lbl, (lbl, v)


def test_gbp_plan_passes_currency_gate(london):
    findings = bundle_qa.run_bundle_qa(
        london["pptx_bytes"], london["xlsx_bytes"], london["data"]
    )
    mixing = [f for f in findings if f.get("code") == "currency_symbol_mixing"]
    assert mixing == [], mixing


def test_usd_plan_labels_unchanged(atria):
    rows = _benchmark_rows(atria["xlsx_bytes"])
    assert rows, "benchmark block missing"
    assert all("(USD)" not in lbl for lbl, _ in rows), rows


# ---------------------------------------------------------------------------
# The "(USD)" row labels must not blind the currency gate elsewhere.
#
# bundle_qa used to treat ANY "(USD)" cell as sanctioning its whole column on
# the sheet, so the benchmark row label "CPA (USD)" in Executive Summary
# column B exempted every other column-B cell -- the budget tiles, channel
# labels, risk factors and recommendations -- on every non-USD plan. A
# labelled data row now clears its own row only; a header clears only the
# table body beneath it.
# ---------------------------------------------------------------------------
def _qa_with_injected(london, sheet, coord, text):
    wb = load_workbook(io.BytesIO(london["xlsx_bytes"]))
    wb[sheet][coord] = text
    buf = io.BytesIO()
    wb.save(buf)
    findings = bundle_qa.run_bundle_qa(
        london["pptx_bytes"], buf.getvalue(), london["data"]
    )
    return [
        f
        for f in findings
        if f.get("code") == "currency_symbol_mixing"
        and f.get("location") == f"{sheet}!{coord}"
    ]


def _column_b_targets(london):
    """Column-B cells OUTSIDE the benchmark rows: first budget tile, the
    first text row after the benchmark block, and the last text row."""
    ws = load_workbook(io.BytesIO(london["xlsx_bytes"]))["Executive Summary"]
    col_b = [
        r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(r, 2).value, str) and ws.cell(r, 2).value.strip()
    ]
    marked = [r for r in col_b if "(USD)" in str(ws.cell(r, 2).value)]
    assert marked, "fixture no longer has a (USD) benchmark row"
    after = [
        r for r in col_b if r > max(marked) and "(USD)" not in str(ws.cell(r, 2).value)
    ]
    return sorted({col_b[0], after[0], col_b[-1]})


def test_usd_row_label_does_not_blind_column_b(london):
    for r in _column_b_targets(london):
        hits = _qa_with_injected(london, "Executive Summary", f"B{r}", "Budget: $2.0M")
        assert hits, f"bare '$' injected at Executive Summary!B{r} went unflagged"


def test_value_column_still_checked(london):
    ws = load_workbook(io.BytesIO(london["xlsx_bytes"]))["Executive Summary"]
    # The campaign-overview tile row: budget in B ("£2.0M"), duration in D.
    r = next(
        r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(r, 2).value, str)
        and ws.cell(r, 2).value.startswith("£")
        and isinstance(ws.cell(r, 4).value, str)
    )
    assert _qa_with_injected(london, "Executive Summary", f"D{r}", "$2.0M")
