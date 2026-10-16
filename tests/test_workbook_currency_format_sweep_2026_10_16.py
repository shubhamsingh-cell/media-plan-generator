"""Regression tests for the workbook currency / format defects of the
2026-10-16 bundle sweep (ledger B_sweep FINDINGS #8, #9, #12, #13).

1. Ambiguous "$": a CAD plan printed CPC "$0.31" beside "C$" amounts, and a
   GBP plan's Executive Summary read "Range: $11-$25" where slide 5 reads
   "US$11-US$25". Root causes: ``excel_v2._cpc_number_format``'s fixed-USD
   format was a BARE "$", and ``intl_benchmark_lookup`` did not read
   "Toronto, ON" as Canada, so a CAD plan never got a local CPC basis and
   priced C$ budgets with unconverted US CPCs.
2. "150K GBP" / "300000 EUR" with US locations priced in USD: plan_currency
   read only symbols, never an ISO code.
3. Intl Benchmarks printed "$0.00-$0.00" for the free government boards.
4. "Toronto, ON" printed "ON" in a Country column.
5. Every "=SUM" total had no cached value (blank in preview apps).
6. "1 hires" wherever a count was templated.
7. "Remote" counted / priced / phrased as a place.

Every test here fails on the pre-fix parent (6d90dbb).
"""

from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

import openpyxl
import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import excel_v2  # noqa: E402
import gold_standard  # noqa: E402
import intl_benchmark_lookup  # noqa: E402
import plan_currency  # noqa: E402

# A "$" not preceded by a letter: "$0.31" / '"$"#,##0.00' -- not "US$"/"C$".
_BARE_DOLLAR_FMT_RE = re.compile(r'(?:^|[^A-Za-z"])\$|"\$"')
_BARE_DOLLAR_TXT_RE = re.compile(r"(?<![A-Za-z])\$\d")


# ── fixtures ────────────────────────────────────────────────────────────


def _plan(locations, budget_text, *, industry="hospitality_travel", role="Line Cook"):
    """A plan dict the way /api/generate builds it for the generators: the
    real budget engine, fed the raw locations and the typed budget so it
    resolves the plan currency and the locale CPC basis itself."""
    import budget_engine
    import data_synthesizer

    from tests.geo_audit_helpers import canonical_plan_locations

    display, sidecar = canonical_plan_locations(list(locations))
    roles = [{"title": role, "count": 10, "tier": "mid"}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=250_000,
        roles=roles,
        locations=display,
        industry=industry,
        channel_percentages={
            "programmatic_dsp": 35,
            "global_boards": 30,
            "regional_boards": 15,
            "niche_boards": 10,
            "social_media": 10,
        },
        locations_raw=display,
        budget_text=budget_text,
    )
    data = {
        "client_name": "Sweep Fixture Co",
        "industry": industry,
        "budget": budget_text,
        "budget_range": budget_text,
        "campaign_duration": "3 months",
        "campaign_weeks": 13,
        "roles": [role],
        "target_roles": roles,
        "locations": display,
        "_location_resolution": sidecar,
        "_budget_allocation": alloc,
        "_enriched": {},
    }
    data["_synthesized"] = {
        "location_profiles": data_synthesizer.fuse_location_profiles({}, {}, data)
    }
    return data


def _xlsx(data, data_only=False):
    import research
    from kb_loader import load_knowledge_base

    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    # the same knowledge-base loader app.py hands the generator
    blob = excel_v2.generate_excel_v2(
        data, research_mod=research, load_kb_fn=load_knowledge_base
    )
    return openpyxl.load_workbook(io.BytesIO(blob), data_only=data_only), blob


@pytest.fixture(scope="module")
def cad_workbook():
    data = _plan(["Toronto, ON", "Vancouver, BC", "Calgary, AB"], "C$250,000")
    wb, blob = _xlsx(data)
    return data, wb, blob


@pytest.fixture(scope="module")
def remote_workbook():
    data = _plan(
        ["Remote", "San Francisco, CA", "Austin, TX", "New York, NY"],
        "$1,200,000",
        industry="tech_engineering",
        role="Software Engineer",
    )
    wb, blob = _xlsx(data)
    return data, wb, blob


@pytest.fixture
def cad_active():
    excel_v2._set_active_currency({"currency_code": "CAD"})
    yield
    excel_v2._set_active_currency({"currency_code": "USD"})


def _header_columns(ws, header):
    """(row, col) of every cell whose value is ``header``."""
    return [(c.row, c.column) for r in ws.iter_rows() for c in r if c.value == header]


def _column_below(ws, row, col):
    out = []
    r = row + 1
    while ws.cell(row=r, column=2).value not in (None, "Total"):
        out.append(ws.cell(row=r, column=col))
        r += 1
    return out


# ── 1. CPC / CPA formats never show a bare "$" on a non-USD plan ────────


def test_unconverted_cpc_format_is_marked_us_dollar(cad_active):
    fmt = excel_v2._cpc_number_format({"cpc_source": "trend_engine"})
    assert "US$" in fmt
    assert not _BARE_DOLLAR_FMT_RE.search(fmt), fmt


def test_canadian_province_locations_get_the_canada_cpc_basis():
    basis = intl_benchmark_lookup.get_locale_cpc_basis(
        ["Toronto, ON", "Calgary, AB"], "CAD"
    )
    assert basis is not None and basis["source"] == "intl_local:canada"
    assert intl_benchmark_lookup._normalize_country("Vancouver, BC") == "canada"
    # a US state is still never read as a country
    assert intl_benchmark_lookup._normalize_country_2026("Springfield, IL") is None


def test_cad_plan_cpc_cells_are_in_the_plans_currency(cad_workbook):
    data, wb, _blob = cad_workbook
    sources = {
        k: v.get("cpc_source")
        for k, v in data["_budget_allocation"]["channel_allocations"].items()
    }
    assert all(
        str(s).startswith("intl_local") or "->CAD" in str(s) for s in sources.values()
    ), sources
    checked = 0
    for sheet in ("Executive Summary", "Channels & Strategy", "Channel Recommendations"):
        ws = wb[sheet]
        for row, col in _header_columns(ws, "CPC"):
            for cell in _column_below(ws, row, col):
                if isinstance(cell.value, (int, float)):
                    checked += 1
                    assert not _BARE_DOLLAR_FMT_RE.search(cell.number_format), (
                        f"{sheet}!{cell.coordinate} {cell.number_format}"
                    )
                    assert "C$" in cell.number_format, f"{sheet}!{cell.coordinate}"
    assert checked >= 5


def test_usd_benchmark_text_is_marked_on_a_non_usd_plan(cad_workbook):
    _data, wb, _blob = cad_workbook
    ws = wb["Executive Summary"]
    rows = [
        (r[1].value, r[3].value)
        for r in ws.iter_rows()
        if isinstance(r[1].value, str) and r[1].value.endswith("(USD)")
    ]
    assert rows, "no (USD)-labelled benchmark rows on the CAD Executive Summary"
    for label, text in rows:
        assert isinstance(text, str) and "US$" in text, (label, text)
        assert not _BARE_DOLLAR_TXT_RE.search(text), (label, text)


# ── 2. an ISO code typed in the budget is a declaration ─────────────────


@pytest.mark.parametrize(
    "budget,locations,code",
    [
        ("150K GBP", ["Austin, TX"], "GBP"),
        ("300000 EUR", ["Austin, TX"], "EUR"),
        ("£150k", ["Austin, TX"], "GBP"),
        ("150K GBP", ["London, UK"], "GBP"),
        ("150k gbp", ["Austin, TX"], "GBP"),
        ("GBP 150,000", ["Austin, TX"], "GBP"),
        ("150,000 USD", ["London, UK"], "USD"),
        ("$150k CAD", ["Austin, TX"], "CAD"),
        ("£150,000 (~190,000 USD)", ["Austin, TX"], "GBP"),
    ],
)
def test_typed_iso_code_or_symbol_beats_location(budget, locations, code):
    assert plan_currency.currency_for_plan_with_basis(
        {"budget_range": budget, "locations": locations}
    ) == (code, "declared")


def test_a_word_is_not_a_currency_code():
    # "per" / "and" are not codes; a lower-case token away from the amount is
    # never read; nothing typed -> the market decides, as before
    for budget in ("150k per month", "10k and 15k", "150,000"):
        assert plan_currency.currency_codes_from_symbol(budget) == ()
    assert plan_currency.currency_for_plan_with_basis(
        {"budget_range": "150k per month", "locations": ["Austin, TX"]}
    ) == ("USD", "market")


@pytest.mark.parametrize(
    "raw",
    ["150K GBP", "300000 EUR", "INR 500000", "$1.5M USD", "150k NOK", "$150k CAD",
     "AUD$100k", "USD$90,000", "CHF 50,000", "£75,000", "C$250,000", "150,000"],
)
def test_canonical_budget_prefix_declares_what_was_typed(raw):
    """app rewrites the budget to "<declared_currency_symbol><amount>"; that
    string must declare exactly what the raw text did."""
    prefix = plan_currency.declared_currency_symbol(raw)
    assert plan_currency.currency_codes_from_symbol(
        f"{prefix}1,000"
    ) == plan_currency.currency_codes_from_symbol(raw), (raw, prefix)


def test_request_normalisation_and_all_three_resolvers_agree():
    import app
    import bundle_qa
    import ppt_generator

    for budget, locations, code in (
        ("150K GBP", ["Austin, TX"], "GBP"),
        ("300000 EUR", ["Austin, TX"], "EUR"),
        ("150K GBP", ["London, UK"], "GBP"),
    ):
        data = {
            "budget_range": budget,
            "locations": locations,
            "budget_period": "campaign",
            "campaign_duration": "3 months",
        }
        app._normalize_request_budget(data, log=False)
        reading = data["_budget_resolution"]["plan_currency"]
        assert reading["code"] == code and reading["note"] == "", (budget, reading)
        assert ppt_generator._plan_currency_code(dict(data)) == code
        assert excel_v2._plan_currency_code(dict(data)) == code
        assert bundle_qa._resolve_plan_currency(dict(data))[0] == code


# ── 3. Intl Benchmarks: a missing figure is "n/a", never 0.00 ──────────


def test_intl_benchmarks_print_na_not_zero_range():
    raw = json.loads(
        (PROJECT_ROOT / "data" / "international_benchmarks_2026.json").read_text(
            encoding="utf-8"
        )
    )
    countries = {k: raw["countries"][k] for k in ("austria", "belgium", "sweden", "uk")}
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_international_benchmarks(
        ws, {"_budget_allocation": {}}, {"countries": countries, "regions": {}}
    )
    text = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
    assert not any("$0.00" in t or "$0-$0" in t for t in text), [
        t for t in text if "$0" in t
    ]
    assert any(t.startswith("n/a") or "; n/a" in t for t in text)
    assert any("n/a = no CPC/CPA figure" in t for t in text)
    # each range sits under its own platform: Austria's #2 board (AMS) is free
    austria = next(
        [c.value for c in r] for r in ws.iter_rows() if r[1].value == "Austria"
    )
    platforms = austria[3].split(", ")
    assert platforms[1].startswith("AMS")
    assert austria[4].split("; ")[1] == "n/a"  # CPC
    assert austria[5].split("; ")[1] == "n/a"  # CPA: AMS has none either


def test_intl_usd_range_helper():
    assert excel_v2._intl_usd_range({"min": 0, "max": 0}, 2) is None
    assert excel_v2._intl_usd_range({"min": None, "max": None}, 2) is None
    assert excel_v2._intl_usd_range({"min": 1.09, "max": 5.45}, 2) == "$1.09-$5.45"
    assert excel_v2._intl_usd_range({"min": 0.3, "max": 2.4}, 0) == "$0.30-$2.40"


# ── 4. a province is not a country ──────────────────────────────────────


def test_country_column_reads_canada_for_province_locations(cad_workbook):
    _data, wb, _blob = cad_workbook
    ws = wb["Market Intelligence"]
    seen = {}
    for row, col in _header_columns(ws, "Country"):
        if ws.cell(row=row, column=col - 1).value != "Location":
            continue
        r = row + 1
        while ws.cell(row=r, column=col - 1).value:
            seen[ws.cell(row=r, column=col - 1).value] = ws.cell(row=r, column=col).value
            r += 1
    assert seen, "no Location/Country table on Market Intelligence"
    for loc, country in seen.items():
        assert country == "Canada", (loc, country)


# ── 5. formula totals carry a cached value ──────────────────────────────


def test_sum_totals_have_cached_values(cad_workbook):
    _data, wb, blob = cad_workbook
    values = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    formulas = 0
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("=SUM("):
                    formulas += 1
                    cached = values[ws.title][c.coordinate].value
                    assert isinstance(cached, (int, float)), f"{ws.title}!{c.coordinate}"
                    col, lo, hi = re.match(r"=SUM\(([A-Z]+)(\d+):[A-Z]+(\d+)\)", c.value).groups()
                    want = sum(
                        v
                        for v in (ws[f"{col}{i}"].value for i in range(int(lo), int(hi) + 1))
                        if isinstance(v, (int, float))
                    )
                    assert cached == pytest.approx(want), f"{ws.title}!{c.coordinate}"
    assert formulas >= 5


# ── 6. counts agree with their nouns ────────────────────────────────────


def test_count_noun():
    assert gold_standard.count_noun(1, "hire") == "1 hire"
    assert gold_standard.count_noun(0, "hire") == "0 hires"
    assert gold_standard.count_noun(1250, "hire") == "1,250 hires"
    assert gold_standard.count_noun(1, "city", "cities") == "1 city"


def test_one_hire_is_singular_in_workbook_copy():
    note = excel_v2._client_history_note_text(
        {"historical_performance": {"total_applications": 1, "total_hires": 1}}
    )
    assert "1 application" in note and "1 applications" not in note
    assert "1 hire" in note and "1 hires" not in note
    summary = excel_v2._build_deterministic_executive_summary(
        {"client_name": "Riverbend", "budget_num": 3000, "duration": "3 months", "header_hires": 1}
    )
    assert "1 hires" not in summary and "deliver 1 hire" in summary, summary


# ── 7. "Remote" is a work mode, not a place ─────────────────────────────


def test_remote_is_not_a_city_in_gold_standard(remote_workbook):
    data, _wb, _blob = remote_workbook
    cities = gold_standard.enrich_city_level_data(data)
    assert not any(str(k).strip().lower() == "remote" for k in cities), list(cities)
    assert gold_standard.is_remote_location("Remote")
    assert gold_standard.is_remote_location("Work From Home")
    assert not gold_standard.is_remote_location("Austin, TX")


def test_remote_is_not_counted_or_phrased_as_a_place(remote_workbook):
    _data, wb, _blob = remote_workbook
    es = wb["Executive Summary"]
    card = next(
        es.cell(row=c.row - 1, column=c.column).value
        for r in es.iter_rows()
        for c in r
        if c.value == "Locations"
    )
    assert card == "3"  # the deck's Locations tile: physical sites only
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            first = next((c.value for c in r if c.value is not None), None)
            if ws.title == "Quality Intelligence":
                assert str(first).strip() != "Remote", f"{ws.title} row {r[0].row}"
            for c in r:
                if isinstance(c.value, str):
                    assert not re.search(r"\bin Remote\b", c.value), (
                        f"{ws.title}!{c.coordinate}: {c.value[:80]}"
                    )
    mi = wb["Market Intelligence"]
    rows = [[c.value for c in r] for r in mi.iter_rows()]
    start = next(
        i
        for i, r in enumerate(rows)
        if any(str(v).strip().lower() == "geographic cost variance" for v in r if v)
    )
    markets = []
    for r in rows[start + 2 :]:
        if not r[1]:
            break
        markets.append(r[1])
    assert markets and "Remote" not in markets, markets


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
