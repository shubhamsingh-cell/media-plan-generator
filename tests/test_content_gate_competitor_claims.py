"""Competitor claims need evidence -- and the client is not a competitor.

Prod (all 4 Hershey runs, 2026-09-24): bundle_qa raised a critical
``unsourced_competitor_claim`` on the LLM executive summary (Executive
Summary!B58) after every competitor search tier failed (Tavily 432, Jina
401, DDG timeouts). Two shapes, both verbatim from the Render log:
  * "...this risk is compounded by named competitors (Nestle Purina,
    Campbell's, Land O'Lakes, Treehouse Foods) drawing from the same..."
    -- model knowledge asserted as fact about client-typed competitors;
  * "...labor market where Hershey is drawing from the same electrical,
    HVAC, and industrial maintenance talent" -- the CLIENT itself, which the
    gate misreported as competitor 'Hershey' (and, one finding per cell,
    that false positive masked the real competitor claim).
Slide 7 cards also printed template claims ("X is a major employer in this
industry and a likely competitor...") for competitors with no evidence
record at all.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from unittest import mock

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import competitor_claims  # noqa: E402
import excel_v2  # noqa: E402

CLIENT_SENTENCE = (
    "Against a competitive skilled-trades labor market where Hershey is "
    "drawing from the same electrical, HVAC, and industrial maintenance "
    "talent as other regional manufacturers, this plan concentrates spend "
    "on programmatic and job-board channels near each plant."
)
COMPETITOR_SENTENCE = (
    "The key risk to monitor is qualified-to-interview conversion, and this "
    "risk is compounded by named competitors Nestle Purina, Campbell's, Land "
    "O'Lakes, Treehouse Foods drawing from the same maintenance and "
    "sanitation talent pools."
)
TYPED = ["Nestle Purina", "Campbell's", "Land O'Lakes", "Treehouse Foods"]


def _narrative_data(**overrides) -> dict:
    data = {
        "client_name": "The Hershey Company",
        "company_name": "The Hershey Company",
        "industry": "food_beverage",
        "budget": "$150,000",
        "locations": ["Hershey, PA"],
        "roles": ["Machine Operator"],
        "target_roles": ["Machine Operator"],
        "campaign_duration": "3 months",
        "hire_volume": "50",
        "work_environment": "onsite",
        "competitors": list(TYPED),
        "_enriched": {},
        "_synthesized": {},
        "_budget_allocation": {},
    }
    data.update(overrides)
    return data


def _exec_summary_text(xlsx_bytes: bytes) -> str:
    """The Executive Strategic Summary narrative cell (the competitor NAMES
    legitimately appear elsewhere -- the client-typed competitor tables)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Executive Summary"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    for i, c in enumerate(cells):
        if c.value.strip().upper() == "EXECUTIVE STRATEGIC SUMMARY":
            return cells[i + 1].value
    raise AssertionError("no Executive Strategic Summary section")


def test_llm_narrative_competitor_claims_are_removed_before_render():
    """The prod narrative shape through excel_v2's real narrative path (LLM
    client faked the same way tests/test_excel_narrative_reliability.py
    does). The grounded figures survive; the competitor sentence does not;
    the client sentence stays."""
    narrative = (
        f"{CLIENT_SENTENCE} Programmatic (DSP) is the core engine of the mix. "
        f"{COMPETITOR_SENTENCE} We recommend reviewing channel performance "
        "after the first month."
    )
    prompts: list = []

    def _fake_call_llm(**kwargs):
        prompts.append(kwargs["messages"][0]["content"])
        return {"text": narrative, "provider": "deepseek", "model": "m", "attempts": []}

    data = _narrative_data()
    with mock.patch("llm_router.call_llm", side_effect=_fake_call_llm):
        text = _exec_summary_text(excel_v2.generate_excel_v2(data))
    assert "Treehouse" not in text and "Nestle Purina" not in text
    assert "where Hershey is drawing from the same" in text
    assert data["_narrative_status"]["status"] == "llm_grounded"
    assert data["_narrative_status"]["competitor_sentences_removed"] == 1
    # the prompt now tells the model it has no competitor evidence
    assert "NO verified data about any competitor" in prompts[0]


def test_narrative_that_is_mostly_competitor_claims_falls_back_to_template():
    narrative = f"{COMPETITOR_SENTENCE} Campbell's is drawing from the same pool too."

    def _fake_call_llm(**kwargs):
        return {"text": narrative, "provider": "deepseek", "model": "m", "attempts": []}

    data = _narrative_data()
    with mock.patch("llm_router.call_llm", side_effect=_fake_call_llm):
        text = _exec_summary_text(excel_v2.generate_excel_v2(data))
    assert "drawing from the same" not in text
    assert data["_narrative_status"]["status"] == "llm_rejected_fabrication"


def _plan(competitors, **extra):
    from tools_regen_bundles import build_plan_data

    data = build_plan_data(
        {
            "client_name": "The Hershey Company",
            "industry": "food_beverage",
            "budget": "$150,000",
            "campaign_duration": "6 months",
            "hire_volume": "50-100 hires",
            "locations": ["Hershey, PA", "Lancaster, PA"],
            "roles": ["Machine Operator"],
            "target_roles": ["Machine Operator"],
            "competitors": competitors,
        }
    )
    data.update(extra)
    return data


def _slide7_texts(data) -> list:
    import ppt_generator
    from pptx import Presentation

    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(data)))
    for slide in prs.slides:
        texts = [
            sh.text_frame.text
            for sh in slide.shapes
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()
        ]
        if "COMPETITOR LANDSCAPE" in texts:
            return texts
    raise AssertionError("no Competitive Landscape slide")


def _slide7_why_lines(competitors) -> list:
    return [t for t in _slide7_texts(_plan(competitors)) if t.startswith("Why:")]


LOW_CI_CONFIDENCE = {
    "_synthesized": {
        "confidence_scores": {"per_section": {"competitive_intelligence": 0.2}}
    }
}


def test_no_evidence_competitors_render_as_one_group_block_not_cards():
    """Design review 2026-10-01: three cards each repeating the same "no
    verified data" line read as a broken template. With NO evidence record,
    every typed name is listed once (beyond the 3-card cap too) with one
    neutral sentence, and nothing says "inferred from industry
    classification" about names the client typed -- even when the
    competitive-intelligence confidence score is low."""
    texts = _slide7_texts(_plan(list(TYPED), **LOW_CI_CONFIDENCE))
    assert not [t for t in texts if t.startswith(("Why:", "Counter:"))], texts
    group = [t for t in texts if competitor_claims.BRIEF_GROUP_LABEL in t]
    assert len(group) == 1, texts
    assert all(n in group[0] for n in TYPED)  # all 4, not just 3
    assert group[0].count(competitor_claims.GROUP_NO_EVIDENCE_SENTENCE) == 1
    assert not [t for t in texts if "inferred from industry classification" in t]


def test_evidence_backed_competitor_keeps_its_card_and_the_rest_share_one_line():
    sourced = {
        "name": "Mars Wrigley",
        "description": (
            "Confectionery manufacturer with plants in Hackettstown, NJ and "
            "Elizabethtown, PA (careers page lists 40 open maintenance roles "
            "across its Pennsylvania sites this quarter)"
        ),
        "source_url": "https://www.example.com/mars-careers",
    }
    texts = _slide7_texts(_plan([sourced, "Campbell's", "Land O'Lakes"]))
    whys = [t for t in texts if t.startswith("Why:")]
    assert len(whys) == 1 and "Confectionery manufacturer with plants" in whys[0]
    # attributable + cut on a boundary, never a dangling "(careers ..."
    assert whys[0].endswith("(source: example.com)"), whys[0]
    assert "(careers" not in whys[0]
    assert len([t for t in texts if t.startswith("Counter:")]) == 1
    rest = [t for t in texts if t.startswith("Also named in your brief")]
    assert len(rest) == 1 and "Campbell's" in rest[0] and "Land O'Lakes" in rest[0]
    assert not [t for t in texts if "inferred from industry classification" in t]


_CLAIM_WORDS = (
    "presence",
    "known name",
    "well-known",
    "reputation",
    "recognition",
    "top employer",
    "major employer",
    "plausible competitor",
    "likely competitor",
    "is a staffing agency",
    "is a direct employer",
)


def test_claim_free_counter_bank_makes_no_claim_about_the_competitor():
    import insight_composer

    seen = set()
    for i in range(12):
        text = insight_composer.compose_counter_strategy(
            "Flagger Force",
            {"role": "Flagger", "city": "Wheeling", "ordinal": i, "has_evidence": False},
        )
        seen.add(text)
        low = text.lower()
        assert not [w for w in _CLAIM_WORDS if w in low], text
        assert not competitor_claims.find_asserted_claims(text, "AWP")
    assert len(seen) >= 10  # ordinal rows never repeat


def _workbook_cells(data) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(excel_v2.generate_excel_v2(data)))
    return {
        f"{ws.title}!{c.coordinate}": c.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str)
    }


def test_workbook_competitor_sheets_have_no_claims_for_unevidenced_names():
    """Cell-by-cell scan of EVERY sheet (design review 2026-10-01): typed
    competitors with no evidence are listed under "Named in brief", the
    neutral sentence appears once per sheet that lists them, and no cell
    pairs one of their names with a presence / known-name / reputation /
    top-employer claim."""
    from tools_regen_bundles import build_plan_data

    names = ["Flagger Force", "Traffic Management Inc.", "RoadSafe Traffic Systems", "Altus Traffic"]
    data = build_plan_data(
        {
            "client_name": "AWP",
            "industry": "construction_real_estate",
            "budget": "$495,000",
            "campaign_duration": "6-12 months",
            "hire_volume": "100-500 hires",
            "locations": ["Wheeling, WV", "Nashville, TN", "Raleigh, NC", "Columbus, OH"],
            "roles": ["Traffic Control Flagger"],
            "target_roles": ["Traffic Control Flagger"],
            "competitors": names,
        }
    )
    cells = _workbook_cells(data)
    offenders = {
        k: v
        for k, v in cells.items()
        if any(n in v for n in names) and any(w in v.lower() for w in _CLAIM_WORDS)
    }
    assert not offenders, offenders
    headers = {k: v for k, v in cells.items() if v in ("Top Employers", competitor_claims.NAMED_IN_BRIEF_HEADER)}
    assert headers and set(headers.values()) == {competitor_claims.NAMED_IN_BRIEF_HEADER}, headers
    by_sheet = {}
    for k, v in cells.items():
        if competitor_claims.GROUP_NO_EVIDENCE_SENTENCE in v:
            by_sheet[k.split("!")[0]] = by_sheet.get(k.split("!")[0], 0) + 1
    assert by_sheet.get("Quality Intelligence") == 1, by_sheet
    assert by_sheet.get("Market Intelligence") == 1, by_sheet
    assert all(n == 1 for n in by_sheet.values()), by_sheet
    # the inferred padding is not presented as the client's "Named in brief"
    qi_lists = [v for k, v in cells.items() if k.startswith("Quality Intelligence!C") and "Flagger Force" in v]
    assert qi_lists and all(
        set(x.strip() for x in v.split(",")) <= set(names) for v in qi_lists
    ), qi_lists


def test_workbook_inferred_roster_is_not_called_top_employers():
    """A brief with NO competitors gets gold_standard's static per-industry
    roster (e.g. Amazon/Walmart/UPS for a food bank). Nothing about those
    employers was observed, so the column is "Inferred competitors" -- not
    "Top Employers" -- and the Why column stays market-level ("Active but
    not dominant" characterises employers nobody looked at)."""
    from tools_regen_bundles import build_plan_data

    data = build_plan_data(
        {
            "client_name": "Riverbend Community Food Bank",
            "industry": "general_entry_level",
            "budget": "$3,000",
            "campaign_duration": "1 month",
            "hire_volume": "1-10 hires",
            "locations": ["Sacramento, CA"],
            "roles": ["Volunteer Coordinator", "Warehouse Associate"],
            "target_roles": ["Volunteer Coordinator", "Warehouse Associate"],
            "competitors": [],
        }
    )
    qi = {
        k: v
        for k, v in _workbook_cells(data).items()
        if k.startswith("Quality Intelligence!")
    }
    assert competitor_claims.INFERRED_HEADER in qi.values(), sorted(set(qi.values()))[:40]
    assert "Top Employers" not in qi.values()
    assert not [v for v in qi.values() if "Active but not dominant" in v]
    assert [v for v in qi.values() if v.startswith("Moderate hiring competition in ")]


# ---------------------------------------------------------------------------
# Verifier follow-ups (2026-10-01)
# ---------------------------------------------------------------------------
def test_a_client_name_first_does_not_hide_the_competitor_after_it():
    """'Hershey and Mars are drawing from the same...' -- the leftmost name
    is the client, the claim is still about Mars (base flagged it; the first
    client-exclusion version reported it clean)."""
    text = "Hershey and Mars are drawing from the same maintenance talent pool."
    hits = competitor_claims.find_asserted_claims(text, "The Hershey Company")
    assert [h["name"] for h in hits] == ["Mars"]
    clean, removed = competitor_claims.strip_claim_sentences(text, "The Hershey Company")
    assert removed and clean == ""


COMMON_WORD_SENTENCES = [
    ("Target", "The hiring target is ambitious for this market."),
    ("Target", "Target CPA sits below the benchmark across channels."),
    ("Gap", "Closing the gap requires a faster offer cycle."),
    ("Gap", "Gap analysis shows the weekend shift is understaffed."),
    ("Apple", "An apple-a-day wellness perk is part of the benefits pitch."),
    ("Delta", "Delta in applications versus last quarter is small."),
    ("Visa", "Applicants needing a visa get a sponsorship contact."),
    ("Shell", "The shell of the career site needs mobile fixes."),
]


@pytest.mark.parametrize("typed,sentence", COMMON_WORD_SENTENCES)
def test_common_words_are_never_treated_as_competitor_mentions(typed, sentence):
    clean, removed = competitor_claims.strip_claim_sentences(sentence, "Acme Health", [typed])
    assert removed == [] and clean == sentence


@pytest.mark.parametrize("typed", ["Target", "Gap", "Apple", "Delta", "Visa", "Shell"])
def test_the_typed_company_as_a_proper_noun_is_still_removed(typed):
    sentence = f"Candidates also weigh offers from {typed} and other large employers."
    _clean, removed = competitor_claims.strip_claim_sentences(sentence, "Acme Health", [typed])
    assert removed == [sentence]
    upper = f"Candidates also weigh offers from {typed.upper()} stores nearby."
    assert competitor_claims.strip_claim_sentences(upper, "Acme Health", [typed])[1] == [upper]


def test_narrative_keeps_common_word_sentences_for_a_plan_that_typed_target():
    narrative = (
        "The hiring target is ambitious for this market. "
        "Target CPA sits below the benchmark across channels. "
        "Programmatic (DSP) is the core engine of the mix. "
        "Candidates also weigh offers from Target and other large employers."
    )

    def _fake_call_llm(**kwargs):
        return {"text": narrative, "provider": "deepseek", "model": "m", "attempts": []}

    data = _narrative_data(
        competitors=["Target", "Gap"], client_name="Acme Health", company_name="Acme Health"
    )
    with mock.patch("llm_router.call_llm", side_effect=_fake_call_llm):
        text = _exec_summary_text(excel_v2.generate_excel_v2(data))
    assert "The hiring target is ambitious for this market." in text
    assert "Target CPA sits below the benchmark across channels." in text
    assert "weigh offers from Target" not in text
