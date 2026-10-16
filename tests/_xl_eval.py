"""A deliberately small Excel formula evaluator for tests.

Supports exactly what the generated workbooks use: numbers, strings, cell and
range references (optionally sheet-qualified), + - * /, comparisons, and
SUM / IF / ROUND / ABS / COUNT / COUNTA / COUNTIF / AND / OR.  IF is lazy.
It exists so a test can prove a cached value equals what Excel would compute,
without LibreOffice (not available on the dev machines).
"""

from __future__ import annotations

import re
from typing import Any, Callable, List

from openpyxl.utils import column_index_from_string

_TOKEN = re.compile(
    r"""\s*(?:
    (?P<num>\d+(?:\.\d+)?)
    |(?P<str>"(?:[^"]|"")*")
    |(?P<ref>(?:(?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
    |(?P<func>[A-Z]+)\(
    |(?P<op><>|<=|>=|[-+*/=<>(),&])
    )""",
    re.X,
)


class _Blank:
    """An empty cell: 0 in arithmetic, equal to "" and 0."""

    def __eq__(self, o: Any) -> bool:
        return o is None or o == "" or o == 0 or isinstance(o, _Blank)

    def __ne__(self, o: Any) -> bool:
        return not self.__eq__(o)

    def __hash__(self) -> int:
        return 0


BLANK = _Blank()


def _n(v: Any) -> float:
    if isinstance(v, _Blank) or v is None or v == "":
        return 0.0
    return float(v)


class Evaluator:
    def __init__(self, wb: Any) -> None:
        self.wb = wb
        self._memo: dict = {}

    # -- cell access ------------------------------------------------------
    def cell(self, sheet: str, ref: str) -> Any:
        key = (sheet, ref.replace("$", ""))
        if key in self._memo:
            return self._memo[key]
        raw = self.wb[sheet][key[1]].value
        if isinstance(raw, str) and raw.startswith("="):
            val = self.evaluate(raw, sheet)
        elif raw is None:
            val = BLANK
        else:
            val = raw
        self._memo[key] = val
        return val

    def rng(self, sheet: str, a: str, b: str) -> List[Any]:
        a, b = a.replace("$", ""), b.replace("$", "")
        ma, mb = re.match(r"([A-Z]+)(\d+)", a), re.match(r"([A-Z]+)(\d+)", b)
        out = []
        for r in range(int(ma.group(2)), int(mb.group(2)) + 1):
            for c in range(column_index_from_string(ma.group(1)), column_index_from_string(mb.group(1)) + 1):
                from openpyxl.utils import get_column_letter

                out.append(self.cell(sheet, f"{get_column_letter(c)}{r}"))
        return out

    # -- parsing ----------------------------------------------------------
    def evaluate(self, formula: str, sheet: str) -> Any:
        toks = []
        pos, text = 0, formula[1:]
        while pos < len(text):
            m = _TOKEN.match(text, pos)
            if not m or m.end() == pos:
                raise ValueError(f"cannot tokenize {formula!r} at {pos}")
            pos = m.end()
            kind = m.lastgroup
            toks.append((kind, m.group(kind)))
        self._t, self._i, self._sheet = toks, 0, sheet
        # parse into a closure tree, then run it
        tree = self._expr()
        if self._i != len(toks):
            raise ValueError(f"trailing tokens in {formula!r}")
        return tree()

    def _peek(self) -> Any:
        return self._t[self._i] if self._i < len(self._t) else (None, None)

    def _take(self) -> Any:
        tok = self._t[self._i]
        self._i += 1
        return tok

    def _expr(self) -> Callable[[], Any]:
        left = self._add()
        while self._peek()[0] == "op" and self._peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self._take()[1]
            right = self._add()
            left = self._cmp(op, left, right)
        return left

    @staticmethod
    def _cmp(op: str, a: Callable, b: Callable) -> Callable:
        def run() -> bool:
            x, y = a(), b()
            if op == "=":
                return x == y or (_is_num(x) and _is_num(y) and _n(x) == _n(y))
            if op == "<>":
                return not (x == y or (_is_num(x) and _is_num(y) and _n(x) == _n(y)))
            if isinstance(x, str) or isinstance(y, str):
                x, y = str(x), str(y)
            else:
                x, y = _n(x), _n(y)
            return {"<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y}[op]

        return run

    def _add(self) -> Callable[[], Any]:
        left = self._mul()
        while self._peek()[0] == "op" and self._peek()[1] in ("+", "-"):
            op = self._take()[1]
            right = self._mul()
            left = (lambda a, b, o: lambda: _n(a()) + _n(b()) if o == "+" else _n(a()) - _n(b()))(left, right, op)
        return left

    def _mul(self) -> Callable[[], Any]:
        left = self._unary()
        while self._peek()[0] == "op" and self._peek()[1] in ("*", "/"):
            op = self._take()[1]
            right = self._unary()
            left = (lambda a, b, o: lambda: _n(a()) * _n(b()) if o == "*" else _n(a()) / _n(b()))(left, right, op)
        return left

    def _unary(self) -> Callable[[], Any]:
        if self._peek() == ("op", "-"):
            self._take()
            inner = self._unary()
            return lambda: -_n(inner())
        return self._atom()

    def _atom(self) -> Callable[[], Any]:
        kind, val = self._take()
        sheet = self._sheet
        if kind == "num":
            return lambda: float(val)
        if kind == "str":
            s = val[1:-1].replace('""', '"')
            return lambda: s
        if kind == "ref":
            sh = sheet
            if "!" in val:
                sh, val = val.split("!", 1)
                sh = sh.strip("'")
            if ":" in val:
                a, b = val.split(":")
                return lambda: self.rng(sh, a, b)
            return lambda: self.cell(sh, val)
        if kind == "op" and val == "(":
            inner = self._expr()
            assert self._take() == ("op", ")")
            return inner
        if kind == "func":
            args: List[Callable] = []
            if self._peek() != ("op", ")"):
                args.append(self._expr())
                while self._peek() == ("op", ","):
                    self._take()
                    args.append(self._expr())
            assert self._take() == ("op", ")")
            return self._func(val, args)
        raise ValueError(f"unexpected token {kind} {val}")

    def _func(self, name: str, args: List[Callable]) -> Callable[[], Any]:
        def flat(vals: List[Any]) -> List[Any]:
            out: List[Any] = []
            for v in vals:
                out.extend(v if isinstance(v, list) else [v])
            return out

        if name == "IF":
            return lambda: args[1]() if args[0]() else (args[2]() if len(args) > 2 else False)
        if name == "SUM":
            return lambda: sum(float(v) for v in flat([a() for a in args]) if _is_num(v) and not isinstance(v, _Blank))
        if name == "COUNT":
            return lambda: float(sum(1 for v in flat([a() for a in args]) if _is_num(v) and not isinstance(v, _Blank)))
        if name == "COUNTA":
            return lambda: float(sum(1 for v in flat([a() for a in args]) if not isinstance(v, _Blank) and v != ""))
        if name == "COUNTIF":
            return lambda: float(sum(1 for v in args[0]() if v == args[1]()))
        if name == "ROUND":
            return lambda: _round(_n(args[0]()), int(_n(args[1]())))
        if name == "ABS":
            return lambda: abs(_n(args[0]()))
        if name == "AND":
            return lambda: all(a() for a in args)
        if name == "OR":
            return lambda: any(a() for a in args)
        raise ValueError(f"unsupported function {name}")


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _round(x: float, digits: int) -> float:
    from decimal import ROUND_HALF_UP, Decimal

    q = Decimal(1).scaleb(-digits)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))
