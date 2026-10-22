"""Design-panel iteration on the client deck (2026-10-22): the draft strip, the
funded-boards line, short timelines, the employer card, and ONE benchmark
verdict for the deck and the workbook."""

from __future__ import annotations

import copy

import pytest
from pptx import Presentation
from pptx.util import Inches

import excel_v2
import plan_validator
import ppt_generator

from tests.test_plan_integrity import (
    _deck_slide_texts,
    _german_plan,
    _workbook_texts,
)

EMU = 914400


def _in(v) -> float:
    return v / EMU


# ---------------------------------------------------------------------------
# 2. cover draft strip
# ---------------------------------------------------------------------------

_STAMPS = {
    1: "Internal draft: data confidence 47%. Review before sharing with a client.",
    2: "Internal draft: job market demand, salary and location data missing (0%). "
    "Review before sharing with a client. Do not forward outside Joveo until the "
    "missing sections are filled and re-checked by the supply team.",
    3: "Internal draft: job market demand, salary and location data missing (0%). "
    "Review before sharing with a client. Do not forward outside Joveo until the "
    "missing sections are filled and re-checked by the supply team. "
    + "Every figure in this deck is a working estimate until a campaign test "
    "confirms it against live publisher pricing and the client's own history. " * 2,
}


def _cover(monkeypatch, stamp):
    monkeypatch.setattr(plan_validator, "internal_draft_stamp", lambda d: stamp)
    prs = Presentation()
    ppt_generator._build_slide_cover(prs, {"client_name": "Acme", "industry_label": "X"})
    return prs.slides[0]


@pytest.mark.parametrize("n_lines", [1, 2, 3])
def test_cover_strip_sizes_to_its_text_and_stays_clear_of_logo_and_bars(
    monkeypatch, n_lines
):
    stamp = _STAMPS[n_lines]
    slide = _cover(monkeypatch, stamp)
    box = next(
        sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text == stamp
    )
    left, top, w, h = (_in(box.left), _in(box.top), _in(box.width), _in(box.height))
    # anchored on the left content margin, full content width
    assert left == pytest.approx(0.62, abs=0.01)
    assert w == pytest.approx(13.333 - 2 * 0.62, abs=0.05)
    # height follows the wrapped text: more lines, taller strip
    lines = ppt_generator._measure_lines(stamp, (w - 0.2) * 0.93, 10.0, bold=True)
    assert h >= lines * 10.0 * 1.35 / 72.0 or box.text_frame.paragraphs[0].runs[0].font.size.pt < 10
    # clear of the wordmark (0.62..4.62 x 0.5..1.05) and the title (1.95)
    assert top >= 1.05
    assert top + h <= 1.95
    # clear of the top-right bars (x >= 10.0, y <= 0.33)
    assert top > 0.33
    # a brand status colour, not the old orange
    fills = [
        sh
        for sh in slide.shapes
        if not sh.has_text_frame or not sh.text_frame.text
        if sh.shape_type is not None and abs(_in(sh.left) - 0.62) < 0.01
        and abs(_in(sh.top) - top) < 0.01
    ]
    assert fills, "strip background shape missing"
    rgb = str(fills[0].fill.fore_color.rgb)
    assert rgb.lower() != "b45309"


def test_cover_strip_grows_with_the_text(monkeypatch):
    heights = []
    for n in (1, 3):
        stamp = _STAMPS[n]
        slide = _cover(monkeypatch, stamp)
        box = next(
            sh
            for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text == stamp
        )
        heights.append(box.height)
    assert heights[1] > heights[0]


# ---------------------------------------------------------------------------
# 3. short timelines fill the row
# ---------------------------------------------------------------------------


def _timeline_cards(weeks):
    data = {
        "client_name": "Acme",
        "industry": "logistics_supply_chain",
        "locations": ["Austin, TX"],
        "campaign_weeks": weeks,
        "_budget_allocation": {
            "metadata": {"total_budget": 100000},
            "channel_allocations": {
                "programmatic_dsp": {
                    "label": "Programmatic DSP",
                    "dollar_amount": 100000,
                    "projected_hires": 10,
                    "projected_applications": 1000,
                    "cpa": 100,
                }
            },
            "total_projected": {
                "applications": 1000,
                "hires": 10,
                "cost_per_application": 100,
            },
        },
    }
    prs = Presentation()
    ppt_generator._build_slide_comparison_timeline(prs, data)
    bars = [
        sh
        for sh in prs.slides[0].shapes
        if abs(sh.height - Inches(0.05)) < 2000 and sh.width > Inches(3.0)
    ]
    return sorted(bars, key=lambda s: s.left)


@pytest.mark.parametrize("weeks,n", [(1, 1), (2, 2), (12, 3)])
def test_timeline_cards_fill_the_row(weeks, n):
    cards = _timeline_cards(weeks)
    assert len(cards) == n
    right = max(_in(c.left + c.width) for c in cards)
    assert right >= 12.5, f"{n} card(s) end at {right:.2f}in: dead right third"
    if n == 3:
        assert all(_in(c.width) == pytest.approx(3.85, abs=0.01) for c in cards)
    else:
        assert len({round(_in(c.width), 2) for c in cards}) == 1


# ---------------------------------------------------------------------------
# 4. employer card
# ---------------------------------------------------------------------------


def _employer_slide(names, match_h_in=0.0, width_in=6.3):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = {
        "label": "Largest employers in Germany (Automotive & Manufacturing)",
        "names": names,
        "note": "Largest employers in this country and industry from public "
        "sources; a size list, not measured competition for this exact role.",
        "source": "Source: public company reports and pages (a.de), as of FY2024",
    }
    bottom = ppt_generator._render_large_employers_card(
        slide, Inches(6.5), Inches(2.0), Inches(width_in), group, match_h_in=match_h_in
    )
    paras = [
        p.text
        for sh in slide.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
    ]
    return paras, bottom


def test_employer_names_are_never_broken_across_lines():
    names = [
        "Volkswagen Group",
        "ZF / Friedrichshafen",
        "Robert Bosch GmbH / Stuttgart",
        "Mercedes-Benz Group",
        "BMW Group",
        "Schaeffler Technologies / Herzogenaurach",
    ]
    paras, _ = _employer_slide(names, width_in=4.2)
    for name in names:
        assert any(name in p for p in paras), f"{name!r} was split: {paras}"
    # no line ends or begins on a bare separator
    for p in paras:
        assert not p.strip().endswith("·") and not p.strip().startswith("·")


def test_fewer_than_three_employers_read_as_one_sentence():
    paras, _ = _employer_slide(["Royal Mail Group"])
    assert (
        "Largest employer on record for Germany (Automotive & Manufacturing): "
        "Royal\u00a0Mail\u00a0Group"
    ) in paras
    assert not any(p.strip() == "Royal Mail Group" for p in paras)
    paras2, _ = _employer_slide(["A Co", "B Co"])
    assert any(
        p.endswith(": A\u00a0Co and B\u00a0Co") and p.startswith("Largest employers on record")
        for p in paras2
    )
    # the not-measured note and the source line stay
    assert any("not measured competition" in p for p in paras)
    assert any(p.startswith("Source: ") for p in paras)


def test_employer_card_matches_its_neighbour_height_without_dead_space():
    _, bottom_plain = _employer_slide(["Royal Mail Group"])
    _, bottom_match = _employer_slide(["Royal Mail Group"], match_h_in=2.2)
    assert bottom_match - 2.0 == pytest.approx(2.2, abs=0.01)
    # without a taller neighbour the card is only as tall as its own text
    assert bottom_plain - 2.0 < 1.4


# ---------------------------------------------------------------------------
# 1 + V2. slide 5 funded-boards line
# ---------------------------------------------------------------------------


def _push_pull_texts(data):
    deck = {
        "push_meets_pull": {
            "summary": "Joveo combines outreach with brand pull.",
            "push": {"name": "Push", "detail": "Targeted ads and alerts."},
            "pull": {"name": "Pull", "detail": "Content and career pages."},
        }
    }
    prs = Presentation()
    ppt_generator._build_slide_push_meets_pull(prs, copy.deepcopy(data), deck)
    return [
        (sh.top, sh.height, sh.text_frame.text)
        for sh in prs.slides[0].shapes
        if sh.has_text_frame and sh.text_frame.text
    ]


def _us_plan(cpa=6.47):
    data = _german_plan(cpa_eur=cpa, confidence=0.9)
    data["industry"] = "logistics_supply_chain"
    data["locations"] = ["Austin, TX"]
    data["country"] = "United States"
    data["budget"] = "$250,000"
    data["_budget_allocation"]["metadata"]["intl_cpc_basis"] = {"currency": "USD"}
    data.pop("_intl_benchmarks", None)
    return data


def test_germany_note_does_not_count_boards_the_funded_line_lists():
    rows = _push_pull_texts(_german_plan(cpa_eur=14.0, confidence=0.9))
    texts = [t for _, _, t in rows]
    assert any(t.startswith("Regional & aggregator budget funds (Germany)") for t in texts)
    assert not any("boards shown" in t for t in texts), texts
    funded = next(t for t in texts if t.startswith("Regional & aggregator"))
    named = next((t for t in texts if t.startswith("Germany")), "")
    for board in ("XING", "StepStone"):
        if board in funded:
            assert board not in named, f"{board} repeated: {named!r} / {funded!r}"


def test_funded_band_lines_never_overlap_the_note():
    rows = sorted(_push_pull_texts(_german_plan(cpa_eur=14.0, confidence=0.9)))
    band = [r for r in rows if Inches(5.0) < r[0] < Inches(7.0)]
    for (t1, h1, _), (t2, _, _) in zip(band, band[1:]):
        assert t1 + h1 <= t2 + Inches(0.02), (t1 + h1, t2)


@pytest.mark.parametrize("plan", [_german_plan, _us_plan])
def test_funded_boards_line_always_carries_its_basis_note(plan):
    data = plan() if plan is _us_plan else plan(cpa_eur=14.0, confidence=0.9)
    texts = [t for _, _, t in _push_pull_texts(data)]
    if not any(t.startswith("Regional & aggregator budget funds") for t in texts):
        pytest.skip("plan funds no named regional boards")
    note = next((t for t in texts if "ranked by local relevance" in t), "")
    assert note, texts
    assert "equal scores receive equal amounts" in note


# ---------------------------------------------------------------------------
# 5. ONE benchmark verdict for the deck and the workbook
# ---------------------------------------------------------------------------


def _comparison_texts(data):
    prs = Presentation()
    ppt_generator._build_slide_comparison_timeline(prs, copy.deepcopy(data))
    return [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]


def test_us_plan_below_the_industry_range_gets_the_note_on_both_surfaces():
    data = _us_plan(6.47)
    gap = plan_validator.plan_cpa_gap(data)
    assert gap and gap["status"] == "below"
    long_msg = plan_validator.benchmark_gap_message(gap)
    short_msg = plan_validator.benchmark_gap_message(gap, short=True)
    assert "$6.47" in long_msg and "$6.47" in short_msg
    assert "Nova working estimate" in long_msg and "Nova working estimate" in short_msg
    rng = plan_validator._usd_range_text(gap["low_usd"], gap["high_usd"])
    assert rng in short_msg
    # deck
    deck_texts = _comparison_texts(data)
    # the legend note: the short sentence, or (when it would need a second
    # line on a dense slide) the shorter legend wording -- never nothing
    fallback = plan_validator.benchmark_gap_fallback(gap)
    assert any(short_msg in t or fallback in t for t in deck_texts), deck_texts
    # workbook: the Executive Summary note and the integrity line are the same string
    wb = _workbook_texts(data)
    assert long_msg in wb["Executive Summary"]
    assert excel_v2._plan_integrity_notes(data)[1] == long_msg


def test_us_cpa_below_the_printed_range_never_gets_the_green_badge():
    data = _us_plan(6.47)
    prs = Presentation()
    ppt_generator._build_slide_comparison_timeline(prs, copy.deepcopy(data))
    value = next(
        sh
        for sh in prs.slides[0].shapes
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("$6.47")
    )
    assert "▲" not in value.text_frame.text
    assert "!" in value.text_frame.text


def test_us_cpa_inside_the_range_has_no_note_anywhere():
    data = _us_plan(20.0)
    assert plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(data)) == ""
    assert excel_v2._plan_integrity_notes(data)[1] == ""
    assert not any("Nova working estimate industry" in t for t in _comparison_texts(data))


def test_median_only_row_fires_below_half_the_median_only():
    ref = plan_validator.industry_cpa_reference({"median": "$20"})
    assert ref == {"low": None, "high": None, "median": 20.0}
    data = _us_plan(9.0)
    below = plan_validator.industry_cpa_gap(data, {"median": "$20"})
    assert below["status"] == "below" and below["threshold_usd"] == 10.0
    data = _us_plan(11.0)
    assert plan_validator.industry_cpa_gap(data, {"median": "$20"})["status"] == "within"
    msg = plan_validator.benchmark_gap_message(below)
    assert "industry median ($20.00)" in msg and "$9.00" in msg


def test_germany_deck_and_workbook_print_the_same_verdict():
    below = _german_plan(cpa_eur=5.0, confidence=0.9)
    gap = plan_validator.plan_cpa_gap(below)
    long_msg = plan_validator.benchmark_gap_message(gap)
    short_msg = plan_validator.benchmark_gap_message(gap, short=True)
    assert long_msg and short_msg
    assert excel_v2._plan_integrity_notes(below)[1] == long_msg
    assert any(short_msg in t for t in _comparison_texts(below))
    inside = _german_plan(cpa_eur=14.0, confidence=0.9)
    assert plan_validator.benchmark_gap_message(plan_validator.plan_cpa_gap(inside)) == ""
    assert excel_v2._plan_integrity_notes(inside)[1] == ""


def test_non_usd_plan_is_never_compared_to_the_us_industry_figure():
    data = _us_plan(6.47)
    data["_budget_allocation"]["metadata"]["intl_cpc_basis"] = {"currency": "GBP"}
    assert plan_validator.industry_cpa_gap(data) is None


# ---------------------------------------------------------------------------
# V1. the salary figure names its source (and its statistic)
# ---------------------------------------------------------------------------


def _salary_plan(raw, source="BLS OES"):
    return {
        "_enriched": {"salary_data": {"Warehouse Associate": raw}},
        "_gold_standard": {
            "city_level_data": {
                "Dallas": {
                    "per_role_salary": {
                        "Warehouse Associate": {
                            "median": 43459,
                            "min": 33772,
                            "max": 59551,
                            "source": source,
                            "currency": "USD",
                        }
                    }
                }
            }
        },
    }


def test_salary_basis_names_the_bls_period_soc_and_statistic():
    raw = {
        "source": "BLS OES",
        "median": 40240,
        "mean": 42260,
        "soc_code": "53-7062",
        "reference_period": "May 2025",
        "series_ids": {"median": "OEUN000000000000053706213"},
    }
    word, text = ppt_generator._salary_basis(_salary_plan(raw), "Warehouse Associate")
    assert word == "median"
    assert text.startswith("Source: BLS OEWS May 2025, SOC 53-7062 national median")


def test_salary_basis_says_mean_when_bls_published_no_median():
    raw = {"source": "BLS OES", "mean": 42260, "soc_code": "53-7062"}
    word, text = ppt_generator._salary_basis(_salary_plan(raw), "Warehouse Associate")
    assert word == "mean" and "national mean" in text


def test_salary_basis_is_a_nova_estimate_without_a_bls_row():
    word, text = ppt_generator._salary_basis(
        _salary_plan({}, source="Industry Benchmark"), "Warehouse Associate"
    )
    assert text == "Source: Nova estimate"


def test_role_breakdown_difficulty_column_is_labelled_modeled():
    # Source-level guard: the slide needs >=4 roles plus a difficulty
    # framework to render, so both header variants (with and without the
    # salary column) are checked in the builder itself.
    import inspect

    src = inspect.getsource(ppt_generator._build_slide_role_breakdown)
    assert src.count('"Difficulty (modeled)"') == 2
    assert '"Difficulty",' not in src
