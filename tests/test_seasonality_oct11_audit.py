"""Regression tests for the 2026-10-11 seasonality / campaign-timing audit
(RTI International: healthcare, "2 months", start month 10; The Hershey
Company: food & beverage manufacturer, "6 months", start month 10).

One test group per audited defect (S1-S9). Each builds the plan the way the
headless regen harness does (``tools_regen_bundles.build_plan_data``), then
generates the real workbook and deck and asserts on what a client reads.
"""

from __future__ import annotations

import datetime
import functools
import io
import json
import re
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from pptx import Presentation  # noqa: E402

import display_format  # noqa: E402
import excel_v2  # noqa: E402
import gold_standard  # noqa: E402
import kb_loader  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles  # noqa: E402

_MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]

_BRIEFS = {
    "rti": {
        "client_name": "RTI International",
        "industry": "healthcare_medical",
        "locations": ["Philadelphia, PA", "Washington, DC", "Jacksonville, FL"],
        "target_roles": ["Coordinator", "Worker", "RN"],
        "budget": "$200,000",
        "campaign_duration": "2 months",
        "campaign_start_month": 10,
        "hire_volume": "200 hires",
        "competitors": ["Westat", "Abt Associates"],
    },
    "hershey": {
        "client_name": "The Hershey Company",
        "industry": "food_beverage",
        "locations": ["Hershey, PA", "Hazleton, PA", "Lancaster, PA"],
        "target_roles": [
            "Production Operator",
            "Maintenance Technician",
            "Food Scientist",
            "Warehouse Associate",
        ],
        "budget": "$90,000",
        "campaign_duration": "6 months",
        "campaign_start_month": 10,
        "hire_volume": "150 hires",
        "competitors": ["Mars", "Mondelez International", "Ferrero"],
    },
}


@functools.lru_cache(maxsize=None)
def _bundle(slug: str):
    """(plan data, workbook cell strings by sheet, deck slide strings)."""
    brief = dict(_BRIEFS[slug])
    data = tools_regen_bundles.build_plan_data(brief)
    # The harness's own weeks ladder differs from the app's: use the app's.
    data["campaign_weeks"] = display_format.resolve_campaign_weeks(
        brief["campaign_duration"]
    )
    data.pop("_gold_standard", None)
    gold_standard.apply_all_quality_gates(data)
    xlsx = excel_v2.generate_excel_v2(dict(data), load_kb_fn=kb_loader.load_knowledge_base)
    if isinstance(xlsx, tuple):
        xlsx = xlsx[0]
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True)
    sheets: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip():
                    cells.append(c.value)
        sheets[ws.title] = cells
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    slides: list[str] = []
    for slide in prs.slides:
        parts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text)
        slides.append("\n".join(parts))
    return data, sheets, slides


def _timeline(slug: str) -> list[dict]:
    return _bundle(slug)[0]["_gold_standard"]["activation_calendar"]["timeline"]


def _all_text(slug: str) -> str:
    _data, sheets, slides = _bundle(slug)
    return "\n".join([c for cells in sheets.values() for c in cells] + slides)


def _month_list(text: str) -> list[str]:
    return [m for m in _MONTHS if m in text]


# ---------------------------------------------------------------------------
# S1 -- ONE seasonality source: summary, timing sentence, deck and calendar
# ---------------------------------------------------------------------------


def test_s1_hershey_every_named_month_agrees_with_calendar_labels():
    labels = {m["month_name"]: m["hiring_intensity"] for m in _timeline("hershey")}
    text = _all_text("hershey")
    checked = 0
    for kind, ok in (
        ("Peak Months", {"high", "very_high"}),
        ("Trough Months", {"low"}),
    ):
        for m in re.finditer(rf"{kind}: ([A-Za-z, ]+?)(?:;|\.|$)", text):
            for month in _month_list(m.group(1)):
                checked += 1
                assert month in labels, (kind, month, "outside the campaign window")
                assert labels[month] in ok, (kind, month, labels[month])
    assert checked, "no Peak/Trough Months text found to check"
    # The deck's timing risk names only high-intensity calendar months.
    deck = "\n".join(_bundle("hershey")[2])
    named = re.search(r"Impact: ([A-Za-z, ]+?) runs? high hiring intensity", deck)
    assert named, "deck timing risk not found"
    for month in _month_list(named.group(1)):
        assert labels[month] in {"high", "very_high"}, (month, labels[month])


def test_s1_hershey_summary_no_longer_calls_calendar_high_months_a_trough():
    labels = {m["month_name"]: m["hiring_intensity"] for m in _timeline("hershey")}
    seasonal = next(
        c
        for c in _bundle("hershey")[1]["Executive Summary"]
        if "Peak Months" in c or "Trough Months" in c
    )
    for m in re.finditer(r"Trough Months: ([A-Za-z, ]+?)(?:;|$)", seasonal):
        for month in _month_list(m.group(1)):
            assert labels[month] == "low", (month, labels[month], seasonal)


def test_s1_food_manufacturing_seasonality_matches_the_repo_food_benchmark():
    deep = json.loads((PROJECT_ROOT / "data/recruitment_benchmarks_deep.json").read_text())
    fb = deep["industry_benchmarks"]["food_beverage"]["seasonal_patterns"]
    trends = json.loads((PROJECT_ROOT / "data/seasonal_hiring_trends.json").read_text())
    entry = trends["seasonal_patterns"]["food_manufacturing"]
    assert [_MONTHS[m - 1] for m in entry["peak_months"]] == fb["peak_months"]
    assert [_MONTHS[m - 1] for m in entry["low_months"]] == fb["trough_months"]
    assert gold_standard._SEASONAL_PATTERNS_KEY_MAP.get("food_manufacturing") in (
        None,
        "food_manufacturing",
    )


# ---------------------------------------------------------------------------
# S2 -- the calendar covers only the campaign's months
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "slug,expected",
    [
        ("rti", ["October", "November"]),
        ("hershey", ["October", "November", "December", "January", "February", "March"]),
    ],
)
def test_s2_calendar_lists_only_campaign_months(slug, expected):
    assert [m["month_name"] for m in _timeline(slug)] == expected
    table_months = [
        c
        for c in _bundle(slug)[1]["Quality Intelligence"]
        if c in _MONTHS
    ]
    assert table_months == expected


def test_s2_unknown_length_keeps_the_full_year_and_long_plans_cap_at_12():
    base = {"industry": "healthcare", "campaign_start_month": 3}
    assert len(gold_standard.build_activation_calendar(dict(base))["timeline"]) == 12
    longer = dict(base, campaign_duration="18 months", campaign_weeks=78)
    assert len(gold_standard.build_activation_calendar(longer)["timeline"]) == 12
    weeks_only = dict(base, campaign_weeks=8)
    assert len(gold_standard.build_activation_calendar(weeks_only)["timeline"]) == 2


# ---------------------------------------------------------------------------
# S3 -- deck timing risk: months inside the window, strongest weight first
# ---------------------------------------------------------------------------


def test_s3_rti_deck_timing_risk_never_names_months_after_the_campaign():
    deck = "\n".join(_bundle("rti")[2])
    risk = re.search(r"Impact: ([^\n]*hiring intensity[^\n]*)", deck)
    assert risk, "deck timing risk not found"
    assert "January" not in risk.group(1) and "February" not in risk.group(1)
    for month in _month_list(risk.group(1).split("(")[0]):
        assert month in ("October", "November")


def test_s3_deck_timing_risk_ranks_by_budget_weight_and_includes_march():
    timeline = _timeline("hershey")
    top = max(timeline, key=lambda m: m["budget_weight"])
    assert top["month_name"] == "March"
    deck = "\n".join(_bundle("hershey")[2])
    named = re.search(r"Impact: ([A-Za-z, ]+?) runs? high hiring intensity", deck)
    assert named
    order = [m.strip() for m in named.group(1).split(",")]
    assert order[0] == "March", order


# ---------------------------------------------------------------------------
# S4 -- no comparative / other-industry copy that contradicts the weights
# ---------------------------------------------------------------------------


def test_s4_no_superlative_or_other_industry_recommendations():
    for slug in ("rti", "hershey"):
        text = _all_text(slug)
        assert "Retail: maximum seasonal spend" not in text, slug
        assert "highest hiring season" not in text, slug
        assert "Second biggest hiring wave" not in text, slug
    for m in range(1, 13):
        rec = gold_standard._HIRING_EVENTS_CALENDAR[m]["recommendation"].lower()
        assert not re.search(r"highest|biggest|lowest|maximum", rec), (m, rec)


def test_s4_generic_fallback_drops_other_industry_events():
    # An industry with no month-by-month events of its own falls back to the
    # generic list, which must not show another industry's shows.
    cal = gold_standard.build_activation_calendar(
        {"industry": "pharma", "campaign_start_month": 1, "campaign_duration": "12 months"}
    )
    events = " ".join(e for m in cal["timeline"] for e in m["key_events"])
    for foreign in ("SXSW (tech)", "Dreamforce", "Tax season (finance)"):
        assert foreign not in events


# ---------------------------------------------------------------------------
# S5 -- "2 months" forecast ends on the last day of month 2
# ---------------------------------------------------------------------------


def test_s5_two_month_forecast_window_ends_nov_30_with_no_december_spend():
    _data, sheets, _slides = _bundle("rti")
    title = next(t for t in sheets if t.endswith("Forecast"))
    cells = sheets[title]
    period = next(c for c in cells if c.startswith("Oct 01, "))
    assert period.endswith("Nov 30, " + period[-4:]), period
    assert not any(c.startswith("December ") for c in cells), [
        c for c in cells if c.startswith("December ")
    ]
    note = next(c for c in cells if c.startswith("This forecast phases budget"))
    assert "December" not in note, note


def test_s5_campaign_window_end_helper():
    start = datetime.date(2026, 10, 1)
    assert gold_standard.campaign_window_end({"campaign_duration": "2 months"}, start) == (
        datetime.date(2026, 11, 30)
    )
    assert gold_standard.campaign_window_end({"campaign_duration": "3 months"}, start) == (
        datetime.date(2026, 12, 31)
    )
    assert gold_standard.campaign_window_end({"campaign_duration": "9 weeks"}, start) is None
    # A range or a months figure that disagrees with the carried weeks is not
    # a stated whole-month duration.
    assert gold_standard.campaign_window_end({"campaign_duration": "1-3 months"}, start) is None
    conflict = {"campaign_duration": "2 months", "campaign_weeks": 12}
    assert gold_standard.campaign_window_end(conflict, start) is None


# ---------------------------------------------------------------------------
# S6 -- calendar weights printed to 2 decimals
# ---------------------------------------------------------------------------


def test_s6_calendar_weight_uses_two_decimals_so_moderate_and_low_differ():
    cells = _bundle("hershey")[1]["Quality Intelligence"]
    start = cells.index("Activation Event Calendar")
    weights = [c for c in cells[start:] if re.fullmatch(r"\d\.\d+x", c)][:6]
    assert weights, "no calendar weight cells found"
    assert all(re.fullmatch(r"\d\.\d\dx", w) for w in weights), weights
    # Direct: 0.85 (Moderate) and 0.75 (Low) must not collapse to one string.
    assert excel_v2._fmt_budget_weight(0.85, "calendar") != excel_v2._fmt_budget_weight(
        0.75, "calendar"
    )


# ---------------------------------------------------------------------------
# S7 -- event text agrees with the month's label
# ---------------------------------------------------------------------------


def test_s7_no_peak_event_in_a_moderate_or_low_month():
    for slug in ("rti", "hershey"):
        for m in _timeline(slug):
            if m["hiring_intensity"] in ("moderate", "low"):
                for e in m["key_events"]:
                    assert not re.search(r"\bpeak\b", e, re.I), (slug, m["month_name"], e)
    cal = gold_standard.build_activation_calendar(
        {"industry": "food_beverage", "campaign_start_month": 10, "campaign_duration": "6 months"}
    )
    assert not any("Holiday production peak" in e for e in cal["industry_events"])


# ---------------------------------------------------------------------------
# S8 -- phasing note derived from the window's weights
# ---------------------------------------------------------------------------


def test_s8_phasing_note_names_the_windows_top_months_not_a_fixed_rule():
    cal = _bundle("hershey")[0]["_gold_standard"]["activation_calendar"]
    note = cal["budget_phasing_note"]
    assert "first 2 months" not in note
    assert "March" in note and "October" in note
    assert note.index("March") < note.index("October")  # highest weight first
    rti = _bundle("rti")[0]["_gold_standard"]["activation_calendar"]["budget_phasing_note"]
    assert "first 2 months" not in rti and "October" in rti


# ---------------------------------------------------------------------------
# S9 -- unsourced statistic removed
# ---------------------------------------------------------------------------


def test_s9_unsourced_january_search_statistic_is_gone():
    for m in range(1, 13):
        assert "25-30%" not in gold_standard._HIRING_EVENTS_CALENDAR[m]["recommendation"]
    for slug in ("rti", "hershey"):
        assert "25-30% more job searches" not in _all_text(slug)
