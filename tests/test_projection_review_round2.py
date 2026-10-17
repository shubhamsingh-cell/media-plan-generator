"""Review round 2 of the projection work (F1-F9 + US-benchmark labelling).

Each test names the defect it pins; all fail on the previous branch HEAD.
"""

from __future__ import annotations

import io
import logging
import os
import sys
import types

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
for _name in ("anthropic", "openai", "supabase", "redis", "qdrant_client", "sentence_transformers", "posthog"):
    sys.modules.setdefault(_name, types.ModuleType(_name))
logging.disable(logging.CRITICAL)

import plan_projection  # noqa: E402

DASH = "–"
_METRICS = ("applications", "hires", "cpa", "cph")


def _read(rel):
    with open(os.path.join(PROJECT_ROOT, rel), encoding="utf-8") as fh:
        return fh.read()


def _band(proj):
    return {m: (proj[m]["low"], proj[m]["high"]) for m in _METRICS}


def _wizard(role, loc, ind, budget_text):
    import app

    return app._compute_plan_estimate(
        {"budget": budget_text, "industry": ind, "roles": [role], "locations": [loc],
         "campaign_duration": "3 months", "client_name": "Acme"}
    )


def _nova(role, loc, ind, budget):
    from nova import Nova

    n = Nova.__new__(Nova)
    n._data_cache = {}
    return n._query_budget_projection({"budget": budget, "roles": [role], "locations": [loc], "industry": ind})


# ---- F1: non-US plans agree across surfaces --------------------------------
def test_f1_london_plan_is_the_same_on_wizard_quick_plan_and_nova():
    import quick_plan

    role, loc, ind = "Warehouse Associate", "London, UK", "logistics_supply_chain"
    # all three priced in the SAME money: a typed "$60,000" is dollars in London
    # (declare, never convert), and Nova's tool budget is US dollars by its schema
    wiz = _wizard(role, loc, ind, "$60,000")["projection"]
    quick = quick_plan.generate_quick_plan(role, loc, "$60,000", ind)["total_projected"]["projection"]
    nova = _nova(role, loc, ind, 60000)["projection"]
    assert _band(quick) == _band(wiz), (quick["text"], wiz["text"])
    assert _band(nova) == _band(wiz), (nova["text"], wiz["text"])
    assert wiz["currency"]["symbol"] == "$"
    # and a typed pound budget is pounds on the wizard and Quick Plan
    gbp = _wizard(role, loc, ind, "£60,000")["projection"]
    assert gbp["currency"]["symbol"] == "£" and gbp["text"]["cpa"].startswith("£")


# ---- F2: Quick Plan speaks ONE currency ------------------------------------
def test_f2_quick_plan_currency_is_the_plans_everywhere():
    import quick_plan

    plan = quick_plan.generate_quick_plan("Warehouse Associate", "London, UK", "£60,000", "logistics_supply_chain")
    assert plan["currency"] == {"code": "GBP", "symbol": "£"}
    # no local cost-per-hire benchmark -> no per-hire grading (and no stray "$")
    text = plan["budget_assessment"]["explanation"]
    assert plan["budget_assessment"]["status"] == "unknown" and "$" not in text, text
    assert plan["total_projected"]["projection"]["text"]["cpa"].startswith("£")
    page = _read("templates/quick-plan.html")
    assert 'var CUR = "$"' in page and "CUR = (data.currency" in page
    assert 'return "$" + (n / 1000000)' not in page  # no hardcoded dollar left in fmtDollar


def test_f2_a_typed_dollar_budget_stays_dollars_in_london():
    import quick_plan

    plan = quick_plan.generate_quick_plan("Warehouse Associate", "London, UK", "$60,000", "logistics_supply_chain")
    assert plan["currency"]["symbol"] == "$"


# ---- F3: wizard banner words the same ranges as the tiles ------------------
def test_f3_banner_and_insights_read_the_projection_not_a_typed_table():
    prev = _read("templates/partials/index/body_preview_js.html")
    app_js = _read("templates/partials/index/body_app_js.html")
    assert "highVolume: 450" not in prev and "BENCH = {" not in prev
    assert "txt.hires" in prev and "txt.cph" in prev
    assert "f.ranged" in app_js


# ---- F4: continuous CPA range ----------------------------------------------
def test_f4_cpa_range_is_continuous_and_never_sub_cent():
    kb = {"low": 14.78, "high": 89.0, "scope": "industry", "source": "x"}
    prev = None
    sweep = [round(7 + 0.01 * i, 2) for i in range(0, 9500)]
    for e in sweep:
        lo, hi, _ = plan_projection._cpa_range(e, kb)
        assert lo <= e <= hi or e > kb["high"] or e < kb["low"]
        assert hi - lo >= 0.5 or lo == hi == e, (e, lo, hi)  # never a sub-cent sliver
        assert hi >= kb["high"] and lo <= min(e, kb["low"])
        if prev is not None:
            assert abs(hi - prev[1]) <= 0.011 and abs(lo - prev[0]) <= 0.011, (e, prev, (lo, hi))
        prev = (lo, hi)
    # the reviewer's pair: $14.77 and $14.79 against a $14.78 benchmark low
    a = plan_projection._cpa_range(14.77, kb)
    b = plan_projection._cpa_range(14.79, kb)
    assert a[1] == b[1] == 89.0


def test_f4_customer_service_is_not_a_cent_wide_range():
    p = plan_projection.project_plan(
        {"budget": 60000, "roles": ["Customer Service Representative"], "locations": ["United States"],
         "industry": "general_entry_level"}
    )
    assert p["cpa"]["high"] - p["cpa"]["low"] > 5
    assert p["applications"]["low"] < p["applications"]["high"] - 1000


# ---- F5: below one hire ----------------------------------------------------
def test_f5_budget_below_one_hire_has_no_cost_per_hire():
    p = plan_projection.project_plan(
        {"budget": 1000, "roles": ["Registered Nurse"], "locations": ["Columbus, OH"], "industry": "healthcare_medical"}
    )
    assert p["hires"] == {"low": 0, "mid": 0, "high": 0}
    assert p["cph"]["low"] is None and p["text"]["cph"] == "n/a"
    notes = " ".join(p["notes"])
    assert "below one hire" in notes
    assert "in its own currency" not in notes
    assert "below one hire" in p["text"]["hires_label"]
    assert "n/a" in p["text"]["cph_label"]


# ---- F6: CPH stays inside its stated basis ---------------------------------
@pytest.mark.parametrize("budget", [5_000, 8_000, 15_000, 20_000, 33_000, 60_000, 250_000])
def test_f6_cph_range_is_the_basis_and_hires_are_budget_over_it(budget):
    p = plan_projection.project_plan(
        {"budget": budget, "roles": ["Registered Nurse"], "locations": ["Columbus, OH"], "industry": "healthcare_medical"}
    )
    a = p["assumptions"]
    avg, floor = a["industry_cph_avg"], a["industry_cph_floor"]  # 10,500 / 5,250
    assert (avg, floor) == (10500.0, 5250.0)
    if p["hires"]["high"] == 0:
        assert p["cph"]["low"] is None
        return
    assert p["hires"]["high"] == int(budget / floor)
    assert p["hires"]["low"] == min(int(budget / avg), p["hires"]["mid"])
    assert p["cph"]["low"] >= floor - 0.01 and p["cph"]["high"] >= avg - 0.01
    assert p["cph"]["low"] <= p["engine"]["cost_per_hire"] <= p["cph"]["high"] + 0.01


def test_f6_fifteen_thousand_rn_reads_the_basis_not_7500_to_15000():
    p = plan_projection.project_plan(
        {"budget": 15000, "roles": ["Registered Nurse"], "locations": ["Columbus, OH"], "industry": "healthcare_medical"}
    )
    assert (p["cph"]["low"], p["cph"]["high"]) == (5250.0, 10500.0)
    assert (p["hires"]["low"], p["hires"]["high"]) == (1, 2)


# ---- F7 (superseded 2026-10-17): no local CPH -> hires / CPH not projected --
def test_f7_uk_plan_projects_no_hires_or_cph_but_keeps_applications_and_cpa():
    p = plan_projection.project_plan(
        {"budget": 60000, "roles": ["Warehouse Associate"], "locations": ["London, UK"],
         "industry": "logistics_supply_chain", "budget_text": "£60,000"}
    )
    assert p["hires"]["low"] is None and p["cph"]["low"] is None
    assert p["assumptions"]["hires_suppressed"] is True
    assert p["notes"].count(plan_projection.NO_LOCAL_HIRES_NOTE) == 1
    assert p["applications"]["low"] and p["cpa"]["low"]
    assert p["text"]["cpa"].startswith("£")


# ---- F8: deck slide 6 ------------------------------------------------------
@pytest.fixture(scope="module")
def deck_us():
    import ppt_generator
    import tools_regen_bundles as T
    from pptx import Presentation

    data = T.build_plan_data(
        {"client_name": "Probe Logistics", "industry": "Logistics & Supply Chain", "budget": "$250,000",
         "campaign_duration": "6 months", "hire_volume": "100 hires", "work_environment": "onsite",
         "locations": ["Dallas, TX"], "roles": ["Warehouse Associate", "CDL Driver"],
         "target_roles": ["Warehouse Associate", "CDL Driver"]}
    )
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    return [sh.text_frame.text for sh in prs.slides[6].shapes if sh.has_text_frame]


def test_f8_slide6_funnel_totals_and_takeaway_do_not_lead_with_the_point(deck_us):
    joined = " || ".join(deck_us)
    # (labelled by what it is -- the engine's channel scenario -- not "upper")
    assert "engine scenario" in joined
    funnel = next(t for t in deck_us if t.startswith("Funnel:"))
    assert funnel.endswith("(engine scenario)")
    take = next(t for t in deck_us if "Budget engine projects" in t)
    assert DASH in take and "estimated cost per hire" in take and "estimated" in take
    assert "average CPA" not in take


# ---- F9: headline ranges ignore live enrichment ----------------------------
def test_f9_bundle_headline_ranges_equal_the_estimates_even_with_synthesized_data():
    import app
    import budget_engine
    import ppt_generator
    from kb_loader import load_knowledge_base

    role, loc, ind = "Registered Nurse", "Columbus, OH", "healthcare_medical"
    est = _wizard(role, loc, ind, "$60,000")
    plan = {"budget": "$60,000", "industry": ind, "client_name": "Acme", "locations": [loc],
            "target_roles": [role],
            "campaign_start_month": 0, "hire_volume": ""}
    synth = {"ad_platform_analysis": {"google": {"avg_cpc": 4.0}, "meta": {"avg_cpc": 4.0}, "linkedin": {"avg_cpc": 4.0}}}
    kb = load_knowledge_base()
    roles = [{"title": role, "count": 1, "tier": "Professional"}]
    locs = [plan_projection.split_location(loc)]
    pcts = plan_projection.default_channel_percentages(ind)
    enriched = budget_engine.calculate_budget_allocation(
        total_budget=60000.0, roles=roles, locations=locs, industry=ind, channel_percentages=pcts,
        synthesized_data=synth, knowledge_base=kb, plan_currency="USD", locations_raw=[loc],
        budget_text="$60,000",
    )
    # precondition: live data really does move the engine's own numbers
    assert enriched["total_projected"]["applications"] != est["est_applications"]
    app._attach_headline_projection(
        enriched, plan, budget=60000.0, roles=roles, locations=locs, channel_pcts=pcts,
        vendor_availability=None, kb=kb,
    )
    head = enriched["headline_projection"]
    assert _band(head) == _band(est["projection"])
    data = {"_budget_allocation": enriched, "locations": [loc]}
    assert _band(ppt_generator._deck_projection(data)) == _band(est["projection"])
    import excel_v2

    line = excel_v2._projection_range_line(enriched, 60000.0)
    assert est["projection"]["text"]["applications"] in line


# ---- roi_projector headline ------------------------------------------------
def test_roi_projector_headline_is_the_plan_projection():
    import roi_projector

    res = roi_projector.project_roi("Indeed", 60000, "healthcare", "Registered Nurse", ["Columbus, OH"])
    head = plan_projection.project_plan(
        {"budget": 60000, "roles": ["Registered Nurse"], "locations": ["Columbus, OH"], "industry": "healthcare"}
    )
    assert res["headline_source"] == "plan_projection.project_plan"
    assert res["projections"]["pessimistic"]["hires"] == head["hires"]["low"]
    assert res["projections"]["optimistic"]["hires"] == head["hires"]["high"]
    cph = res["roi_metrics"]["cost_per_hire"]
    assert cph["low"] == round(head["cph"]["low"]) and cph["high"] == round(head["cph"]["high"])
    assert "channel_detail" in res


# ---- US benchmark labelling on non-US plans --------------------------------
def _bundle_texts(brief):
    import excel_v2
    import ppt_generator
    import tools_regen_bundles as T
    from kb_loader import load_knowledge_base
    from openpyxl import load_workbook
    from pptx import Presentation

    data = T.build_plan_data(brief)
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    deck = {i: [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame] for i, sl in enumerate(prs.slides, 1)}
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    ws = load_workbook(io.BytesIO(raw))["Executive Summary"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value not in (None, "")]
    return deck, cells


def _brief(industry, budget, loc, role):
    return {"client_name": "Probe", "industry": industry, "budget": budget, "campaign_duration": "3 months",
            "hire_volume": "Not specified", "work_environment": "onsite", "locations": [loc],
            "roles": [role], "target_roles": [role]}


TAG = "US benchmark — no local data"


@pytest.mark.parametrize(
    "brief",
    [
        _brief("Logistics & Supply Chain", "£60,000", "London, UK", "Warehouse Associate"),
        _brief("tech_engineering", "₹25,000,000", "Bengaluru, Karnataka, India", "Java Developer"),
        _brief("Healthcare & Medical", "$150,000", "London, UK", "Registered Nurse"),
    ],
    ids=["uk-gbp", "india-inr", "uk-usd-typed"],
)
def test_non_us_plans_label_the_us_calibrated_rows(brief):
    deck, cells = _bundle_texts(brief)
    assert any("Apply Rate" in t and TAG in t for t in deck[2]), deck[2]
    assert any(TAG in t for t in deck[6]), deck[6]
    rows = [c for c in cells if c.startswith("Apply Rate")]
    assert all(TAG in c for c in rows), rows  # every US-calibrated row is labelled
    if "India" not in brief["locations"][0]:  # (the India sheet carries no such row)
        assert rows


def test_uk_usd_plan_does_not_present_the_us_midpoint_as_the_markets():
    deck, cells = _bundle_texts(_brief("Healthcare & Medical", "$150,000", "London, UK", "Registered Nurse"))
    foot = " ".join(deck[2])
    # no hires are projected, so there is no US-midpoint "plan efficiency" note
    assert "midpoint" not in foot and "not projected" in foot
    assert any("not projected" in c for c in cells)


def test_us_plan_is_untouched_by_the_labelling():
    deck, cells = _bundle_texts(_brief("Healthcare & Medical", "$150,000", "Dallas, TX", "Registered Nurse"))
    assert not any(TAG in t for ts in deck.values() for t in ts)
    assert not any(TAG in c for c in cells)
