"""Trust pack: "data as of" freshness (T1) and the review-before-sending
checklist (T2), in plan_trust and in the workbook and deck that use it.

Dates are asserted against the data files' OWN metadata, read independently
in the test, so a bumped file date moves both sides and a made-up date fails.
"""

import io
import json
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402
from pptx import Presentation  # noqa: E402

import excel_v2  # noqa: E402
import plan_trust  # noqa: E402
import ppt_generator  # noqa: E402

DATA_DIR = PROJECT_ROOT / "data"


def _plan(**overrides) -> dict:
    data = {
        "client_name": "Acme Health",
        "company_name": "Acme Health",
        "industry": "healthcare_medical",
        "budget": "$100,000",
        "locations": ["Dallas, TX"],
        "roles": ["Registered Nurse"],
        "target_roles": ["Registered Nurse"],
        "campaign_duration": "3 months",
        "hire_volume": "50",
        "work_environment": "onsite",
        "_enriched": {},
        "_synthesized": {},
        "_budget_allocation": {},
    }
    data.update(overrides)
    return data


def _sheet_text(ws) -> str:
    return "\n".join(
        str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None
    )


def _fam(info, label_start):
    return next(f for f in info["families"] if f["label"].startswith(label_start))


# ---------------------------------------------------------------------------
# T1: read_as_of / data_asof
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "doc,expected",
    [
        ({"_metadata": {"last_updated": "2026-06-13"}}, "2026-06-13"),
        ({"metadata": {"last_updated": "2026-04-24"}}, "2026-04-24"),
        ({"_meta": {"last_updated": "2025-Q1"}}, "2025-Q1"),
        ({"last_updated": "2026-03-31"}, "2026-03-31"),
        ({"_metadata": {"_last_verified": "2026-03-08"}}, "2026-03-08"),
        ({"_refreshed_iso": "2026-04-03T13:38:36+00:00"}, "2026-04-03"),
        ({"fetched_at": "2026-04-05T07:55:39+00:00"}, "2026-04-05"),
        ({"_metadata": {"researched_date": "2026-03-26"}}, "2026-03-26"),
        # last_updated beats the original research date
        (
            {"_metadata": {"researched_date": "2026-03-26", "last_updated": "2026-06-13"}},
            "2026-06-13",
        ),
        # metadata block beats a stray top-level key
        (
            {"_metadata": {"last_updated": "2026-06-13"}, "last_updated": "2020-01-01"},
            "2026-06-13",
        ),
        # nothing readable
        ({"_metadata": {"last_updated": "recently"}}, None),
        ({"_metadata": {"last_updated": "2026-13-45"}}, None),
        ({"entries": []}, None),
        ([], None),
    ],
)
def test_read_as_of_shapes(tmp_path, doc, expected):
    path = tmp_path / "f.json"
    path.write_text(json.dumps(doc))
    assert plan_trust.read_as_of(str(path)) == expected


def test_read_as_of_epoch_only(tmp_path):
    path = tmp_path / "f.json"
    path.write_text(json.dumps({"_refreshed_at": 1775223516.2}))
    assert plan_trust.read_as_of(str(path)) == "2026-04-03"


def test_read_as_of_missing_or_corrupt_file(tmp_path):
    assert plan_trust.read_as_of(str(tmp_path / "nope.json")) is None
    bad = tmp_path / "bad.json"
    bad.write_text("{not json")
    assert plan_trust.read_as_of(str(bad)) is None


def test_source_without_metadata_reads_date_not_recorded(tmp_path):
    (tmp_path / "dated.json").write_text(json.dumps({"last_updated": "2026-02-01"}))
    (tmp_path / "undated.json").write_text(json.dumps({"entries": [1, 2]}))
    info = plan_trust.data_asof(
        {},
        data_dir=str(tmp_path),
        families=[
            ("Dated", ("dated.json",)),
            ("Undated", ("undated.json",)),
            ("Mixed", ("dated.json", "undated.json")),
            ("Absent file", ("missing.json",)),
        ],
    )
    assert _fam(info, "Dated")["as_of"] == "2026-02-01"
    assert _fam(info, "Undated")["as_of"] == plan_trust.NOT_RECORDED
    assert _fam(info, "Absent")["as_of"] == plan_trust.NOT_RECORDED
    mixed = _fam(info, "Mixed")["as_of"]
    assert mixed.startswith("2026-02-01") and plan_trust.NOT_RECORDED in mixed


def test_family_range_when_files_differ(tmp_path):
    (tmp_path / "a.json").write_text(json.dumps({"last_updated": "2026-01-05"}))
    (tmp_path / "b.json").write_text(json.dumps({"last_updated": "2026-05-09"}))
    info = plan_trust.data_asof(
        {}, data_dir=str(tmp_path), families=[("Both", ("a.json", "b.json"))]
    )
    assert info["families"][0]["as_of"] == "2026-01-05 to 2026-05-09"


def test_repo_families_match_the_files_own_metadata():
    info = plan_trust.data_asof({})
    repo = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
    assert _fam(info, "Supply repository")["as_of"] == repo["last_updated"]
    seasonal = json.loads((DATA_DIR / "seasonal_hiring_trends.json").read_text())
    assert seasonal["metadata"]["last_updated"] in _fam(info, "Seasonal")["as_of"]
    adzuna = json.loads((DATA_DIR / "adzuna_benchmarks.json").read_text())
    assert adzuna["_refreshed_iso"][:10] in _fam(info, "Scheduled refresh")["as_of"]
    assert repo["last_updated"] in info["headline"]
    for fam in info["families"]:
        assert fam["files"], fam
        for name in fam["files"]:
            assert (DATA_DIR / name).exists(), f"{name} missing from data/"


def _nested_dates(path, key):
    out = []

    def walk(o):
        if isinstance(o, dict):
            for k, v in o.items():
                if k == key and isinstance(v, str):
                    out.append(v[:10])
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)

    walk(json.loads(Path(path).read_text()))
    return out


def test_supply_network_family_reports_the_latest_real_date():
    info = plan_trust.data_asof({})
    network = _fam(info, "Supply network")
    groups = max(_nested_dates(DATA_DIR / "joveo_network_groups.json", "as_of"))
    avail = max(_nested_dates(DATA_DIR / "publisher_availability.json", "as_of"))
    rows = max(
        _nested_dates(DATA_DIR / "joveo_global_supply_repository.json", "network_as_of")
    )
    newest = max(groups, avail, rows)
    assert network["as_of"] == f"updated {newest}"
    # the full listing keeps its own, older, top-level date
    repo = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
    assert _fam(info, "Supply repository")["as_of"] == repo["last_updated"]
    assert newest > repo["last_updated"]
    assert f"supply network {newest}" in info["headline"]
    assert f"repository {repo['last_updated']}" in info["headline"]


def test_nested_as_of_readers_and_latest_mode(tmp_path):
    (tmp_path / "groups.json").write_text(
        json.dumps({"groups": [{"as_of": "2026-10-01"}, {"as_of": "2026-10-09"}]})
    )
    (tmp_path / "repo.json").write_text(
        json.dumps(
            {
                "last_updated": "2026-03-31",
                "rows": [{"network_as_of": "2026-10-12"}, {"x": 1}],
            }
        )
    )
    assert plan_trust.read_as_of_nested(str(tmp_path / "groups.json"), ("as_of",)) == "2026-10-09"
    assert plan_trust.read_as_of_nested(str(tmp_path / "repo.json"), ("network_as_of",)) == "2026-10-12"
    assert plan_trust.read_as_of_nested(str(tmp_path / "repo.json"), ("as_of",)) is None
    info = plan_trust.data_asof(
        {},
        data_dir=str(tmp_path),
        families=[
            (
                "Supply network status",
                (("groups.json", "as_of"), ("repo.json", "network_as_of")),
                "latest",
            ),
            ("Supply repository (full listing)", ("repo.json",)),
        ],
    )
    assert info["families"][0]["as_of"] == "updated 2026-10-12"
    assert info["families"][1]["as_of"] == "2026-03-31"


def test_live_api_line_counts_real_enrichment_and_stamps_generation_time():
    none = plan_trust.data_asof({"_enriched": {}})
    assert "not recorded" in none["live_line"]
    zero = plan_trust.data_asof({"_enriched": {"enrichment_summary": {"apis_succeeded": []}}})
    assert "none returned data" in zero["live_line"]
    some = plan_trust.data_asof(
        {"_enriched": {"enrichment_summary": {"apis_succeeded": ["BLS", "FRED"]}}}
    )
    assert "2 sources returned data" in some["live_line"]
    assert some["generated_label"] in some["live_line"]


def test_version_stamp_follows_deploy_source(monkeypatch):
    monkeypatch.setitem(
        sys.modules,
        "app",
        type(sys)("app_stub"),
    )
    sys.modules["app"]._DEPLOY_VERSION = "4.0.0-abcd1234"
    sys.modules["app"]._DEPLOY_GIT_COMMIT = "abcd1234ffff"
    version, sha = plan_trust.deploy_version()
    assert (version, sha) == ("4.0.0-abcd1234", "abcd1234ffff")
    assert "Nova 4.0.0-abcd1234" in plan_trust.provenance_stamp(
        plan_trust.data_asof({})
    )


# ---------------------------------------------------------------------------
# T1: workbook
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def workbook():
    raw = excel_v2.generate_excel_v2(_plan(excluded_publishers=["Monster"]))
    return openpyxl.load_workbook(io.BytesIO(raw))


def test_workbook_executive_summary_has_asof_line(workbook):
    ws = workbook["Executive Summary"]
    text = _sheet_text(ws)
    repo = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
    info = plan_trust.data_asof({})
    assert info["headline"] in text
    assert "(repository " + repo["last_updated"] + ")" in text
    assert "Nova " in text and "UTC" in text


def test_workbook_sources_sheet_lists_every_family_with_metadata_dates(workbook):
    text = _sheet_text(workbook["Sources & Confidence"])
    assert "DATA AS OF (LAST REFRESHED)" in text
    info = plan_trust.data_asof({})
    for fam in info["families"]:
        assert fam["label"] in text
        assert fam["as_of"] in text
    assert "Live API pulls" in text and "Workbook generated" in text


def test_workbook_properties_carry_generation_stamp(workbook):
    desc = workbook.properties.description
    assert "Generated 20" in desc and "UTC | Nova " in desc


# ---------------------------------------------------------------------------
# T1: deck
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def deck():
    raw = ppt_generator.generate_pptx(_plan(excluded_publishers=["Monster"]))
    return Presentation(io.BytesIO(raw))


def test_deck_slide_two_footer_states_data_as_of(deck):
    repo = json.loads((DATA_DIR / "joveo_global_supply_repository.json").read_text())
    texts = [
        sh.text_frame.text
        for sh in deck.slides[1].shapes
        if sh.has_text_frame and sh.text_frame.text
    ]
    assert plan_trust.data_asof({})["footer_headline"] in texts
    assert repo["last_updated"]  # the full repository date lives in the notes
    assert len(plan_trust.data_asof({})["footer_headline"]) <= (
        ppt_generator._ASOF_FOOTER_MAX_CHARS
    )


def test_deck_slide_two_notes_list_every_family(deck):
    notes = deck.slides[1].notes_slide.notes_text_frame.text
    for fam in plan_trust.data_asof({})["families"]:
        assert f"{fam['label']}: {fam['as_of']}" in notes


def test_deck_properties_carry_generation_stamp(deck):
    assert "UTC | Nova " in deck.core_properties.comments


def test_deck_footer_asof_falls_back_when_dates_are_long():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ppt_generator._add_footer(slide, "October 17, 2026")
    long_headline = (
        "Data as of: supply date not recorded | benchmarks date not recorded"
    )
    assert len(long_headline) > ppt_generator._ASOF_FOOTER_MAX_CHARS
    ppt_generator._stamp_data_asof_footer(slide, long_headline, "October 17, 2026")
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert "Data as of: see workbook, Sources sheet" in texts
    assert long_headline not in texts


# ---------------------------------------------------------------------------
# T2: review items (flag-driven)
# ---------------------------------------------------------------------------


def _items(**plan_overrides):
    unverified = plan_overrides.pop("unverified", 0)
    return plan_trust.review_items(_plan(**plan_overrides), unverified_boards=unverified)


def _has(items, needle):
    return any(needle in i for i in items)


def test_projection_item_only_without_client_history():
    assert _has(_items(), "no client hire history was provided")
    assert not _has(
        _items(historical_performance={"platform_count": 2}),
        "no client hire history",
    )


def test_unverified_boards_item_uses_the_count():
    assert not _has(_items(), "relationship unverified")
    one = _items(unverified=1)
    assert _has(one, "1 board is marked † relationship unverified")
    many = _items(unverified=4)
    assert _has(many, "4 boards are marked † relationship unverified")


@pytest.mark.parametrize(
    "tier,label,flagged",
    [
        ("impossible", "UNREALISTIC", True),
        ("severely_underfunded", "SEVERELY UNDERFUNDED", True),
        ("underfunded", "UNDERFUNDED", True),
        ("tight", "TIGHT BUT FEASIBLE", True),
        ("adequate", "ADEQUATE", False),
        ("generous", "WELL-FUNDED", False),
        ("not_assessed", "NOT ASSESSED", False),
    ],
)
def test_budget_check_item_only_for_flagged_tiers(tier, label, flagged):
    alloc = {
        "sufficiency": {
            "budget_reality_check": {
                "feasibility_tier": tier,
                "feasibility_label": label,
            }
        }
    }
    items = _items(_budget_allocation=alloc)
    assert _has(items, "Budget reality check flagged") is flagged
    if flagged:
        assert _has(items, label.title())


def test_markets_without_benchmarks_item():
    assert not _has(_items(locations=["Dallas, TX"]), "Markets without local")
    items = _items(locations=["Dallas, TX", "Madrid, Spain"])
    assert _has(items, "Markets without local benchmarks: Spain")


def test_suppressed_cph_names_the_non_us_markets():
    alloc = {"metadata": {"industry_cph": {"claim_suppressed": True}}}
    items = _items(locations=["Paris, France"], _budget_allocation=alloc)
    assert _has(items, "Markets without local benchmarks: France")


def test_us_costed_non_us_market_is_listed():
    alloc = {
        "metadata": {
            "industry_cph": {
                "basis": "us_benchmark",
                "us_benchmark_non_us_market": True,
                "claim_suppressed": False,
            }
        }
    }
    items = _items(
        locations=["Lagos, Nigeria", "Dallas, TX"], _budget_allocation=alloc
    )
    assert _has(items, "Markets without local benchmarks: Nigeria")
    assert not _has(items, "Dallas")
    # a US-only plan carrying the same flag has no non-US market to list
    assert not _has(
        _items(locations=["Dallas, TX"], _budget_allocation=alloc),
        "Markets without local",
    )
    # flag false -> nothing to say
    off = {"metadata": {"industry_cph": {"us_benchmark_non_us_market": False}}}
    assert not _has(
        _items(locations=["Lagos, Nigeria"], _budget_allocation=off),
        "Markets without local",
    )


def test_projection_copy_claims_only_what_main_shows():
    first = _items()[0]
    assert "ranges" not in first.lower()
    assert "industry benchmarks" not in first  # the Nova table is not a published benchmark
    assert (
        "estimates from Nova's assumed industry cost-per-hire table "
        "(an all-in figure, not a published benchmark); media budget divided by that figure"
    ) in first
    from us_only_markers import _US_ONLY_MARKERS

    assert not [m for m in _US_ONLY_MARKERS if m.lower() in first.lower()]
    assert "no client hire history was provided" in first
    assert "indicative" in first
    guide = (PROJECT_ROOT / "templates" / "help-media-plans.html").read_text()
    assert "Projections are ranges" not in guide
    assert "Ranges and assumptions" not in guide
    assert "Estimates and assumptions" in guide


def test_exclusions_item_names_the_boards():
    assert not _has(_items(), "Client exclusions applied")
    items = _items(excluded_publishers=["Monster", "Glassdoor"])
    assert _has(
        items,
        "Client exclusions applied. Excluded at client request: Monster, Glassdoor",
    )


def test_standing_items_always_present_and_total_in_range():
    base = _items()
    assert _has(base, "Confirm the channel mix and budget split with supply")
    assert _has(base, "Data as of: supply network")
    assert _has(base, "responsible for what you send")
    everything = _items(
        unverified=2,
        excluded_publishers=["Monster"],
        locations=["Dallas, TX", "Madrid, Spain"],
        _budget_allocation={
            "sufficiency": {
                "budget_reality_check": {
                    "feasibility_tier": "underfunded",
                    "feasibility_label": "UNDERFUNDED",
                }
            }
        },
    )
    assert _has(base, "against the brief you were given")
    assert len(base) == 5  # 4 standing + the no-client-history item
    assert len(everything) == 9  # every conditional item on top of the 4


def test_count_unverified_boards():
    cells = [
        "Pole Emploi, Indeed France, APEC †",
        "APEC †",  # same board again -> counted once
        "Seek †, Jora †, Totaljobs",
        "† relationship unverified — confirm before quoting",  # legend
        "no mark here",
    ]
    assert plan_trust.count_unverified_boards(cells) == 3
    assert plan_trust.count_unverified_boards([]) == 0


# ---------------------------------------------------------------------------
# T2: workbook sheet + deck notes
# ---------------------------------------------------------------------------


def test_review_sheet_is_last_and_reflects_this_plan(workbook):
    assert workbook.sheetnames[-1] == "Review before sending"
    text = _sheet_text(workbook["Review before sending"])
    # Monster was excluded by the client: the sheet names it (the board sweep
    # runs before this sheet is added, on purpose)
    assert "Client exclusions applied. Excluded at client request: Monster" in text
    assert "no client hire history was provided" in text
    assert "Budget reality check flagged" not in text  # no alloc -> omitted


def test_review_sheet_follows_flags_end_to_end(monkeypatch):
    monkeypatch.setattr(excel_v2, "_workbook_unverified_board_count", lambda wb: 2)
    raw = excel_v2.generate_excel_v2(
        _plan(
            historical_performance={"platform_count": 1},
            _budget_allocation={
                "sufficiency": {
                    "budget_reality_check": {
                        "feasibility_tier": "severely_underfunded",
                        "feasibility_label": "SEVERELY UNDERFUNDED",
                    }
                }
            },
        )
    )
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    text = _sheet_text(wb["Review before sending"])
    assert "2 boards are marked † relationship unverified" in text
    assert "Budget reality check flagged: Severely Underfunded" in text
    assert "no client hire history" not in text
    assert "Client exclusions applied" not in text


def test_workbook_dagger_count_reads_the_rendered_cells():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Pole Emploi, APEC †"
    wb.create_sheet()["B2"] = "Seek †"
    assert excel_v2._workbook_unverified_board_count(wb) == 2


def test_deck_first_slide_notes_carry_the_checklist(deck):
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert notes.startswith("REVIEW BEFORE SENDING")
    assert "Client exclusions applied. Excluded at client request: Monster" in notes
    assert "no client hire history was provided" in notes
    assert "responsible for what you send" in notes


def test_deck_notes_omit_unflagged_items():
    raw = ppt_generator.generate_pptx(
        _plan(historical_performance={"platform_count": 1})
    )
    notes = Presentation(io.BytesIO(raw)).slides[0].notes_slide.notes_text_frame.text
    assert "no client hire history" not in notes
    assert "Client exclusions applied" not in notes
