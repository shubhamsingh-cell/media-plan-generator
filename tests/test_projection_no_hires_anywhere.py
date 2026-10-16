"""A market with no local cost-per-hire benchmark projects NO hire count or cost
per hire anywhere: every workbook sheet, every deck slide (text boxes, tables,
notes) and Nova's tool result (per-channel included). A US plan is unchanged."""

from __future__ import annotations

import io
import logging
import os
import re
import sys
import types

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
for _n in ("anthropic", "openai", "supabase", "redis", "qdrant_client", "sentence_transformers", "posthog"):
    sys.modules.setdefault(_n, types.ModuleType(_n))
logging.disable(logging.CRITICAL)

# a STATED projected hire count: "25 hires", "17 projected hires", "$2,400/hire"
COUNT = re.compile(r"(?<![\w,.])[1-9][\d,]*\s+(?:projected\s+)?hires?\b|[£$€₹]\s?[\d,.]+[KkMm]?\s*/\s*hire|projected to deliver", re.I)
HDR = re.compile(r"^(?:total\s+)?(?:proj\.?\s*|projected\s+|est\.?\s*|avg\.?\s*)?(?:cost\s*(?:/|per)\s*hire|hires?)\s*(?:\([^)]*\))?$", re.I)


def _bundle(budget, loc):
    import excel_v2
    import ppt_generator
    import tools_regen_bundles as T
    from kb_loader import load_knowledge_base
    from openpyxl import load_workbook
    from pptx import Presentation

    data = T.build_plan_data(
        {"client_name": "Probe", "industry": "Logistics & Supply Chain", "budget": budget,
         "campaign_duration": "3 months", "hire_volume": "Not specified", "work_environment": "onsite",
         "locations": [loc], "roles": ["Warehouse Associate"], "target_roles": ["Warehouse Associate"]}
    )
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    wb = load_workbook(io.BytesIO(raw))
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    return wb, prs


def _sheet_strings(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    yield ws.title, c.coordinate, c.value


def _numeric_under_hire_headers(wb):
    bad = []
    num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)  # noqa: E731
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and len(c.value) < 40 and HDR.match(c.value.strip()):
                    below = [ws.cell(row=c.row + k, column=c.column).value for k in range(1, 12)]
                    right = [ws.cell(row=c.row, column=c.column + k).value for k in range(1, 10)]
                    # (a numeric run directly under / beside a hires header)
                    for seq in (below, right):
                        run = []
                        for v in seq:
                            if v is None:
                                break
                            run.append(v)
                        if any(num(v) for v in run):
                            bad.append((ws.title, c.coordinate, c.value, run[:4]))
    return bad


def _deck_strings(prs):
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame:
                yield i, sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        yield i, c.text
        if sl.has_notes_slide:
            yield i, sl.notes_slide.notes_text_frame.text


@pytest.mark.parametrize("budget", ["£60,000", "$60,000"], ids=["gbp", "usd-in-london"])
def test_london_bundle_states_no_projected_hire_count_anywhere(budget):
    wb, prs = _bundle(budget, "London, UK")
    leaks = [(s, c, v[:90]) for s, c, v in _sheet_strings(wb) if COUNT.search(v)]
    assert not leaks, leaks
    assert not _numeric_under_hire_headers(wb), _numeric_under_hire_headers(wb)
    deck = [(i, t[:90]) for i, t in _deck_strings(prs) if COUNT.search(t)]
    assert not deck, deck
    # but the one note that says why is still there
    assert any("not projected" in v for _, _, v in _sheet_strings(wb))
    assert any("not projected" in t for _, t in _deck_strings(prs))


def test_us_bundle_is_unchanged():
    wb, prs = _bundle("$60,000", "Dallas, TX")
    assert _numeric_under_hire_headers(wb), "a US workbook still carries its hires columns"
    assert any(COUNT.search(t) for _, t in _deck_strings(prs))
    assert not any("not projected" in v for _, _, v in _sheet_strings(wb))


def _all_keys(obj, out):
    if isinstance(obj, dict):
        for k, v in obj.items():
            out.add(k)
            _all_keys(v, out)
    elif isinstance(obj, list):
        for v in obj:
            _all_keys(v, out)
    return out


def test_nova_result_has_no_per_channel_hires_or_cost_per_hire():
    from nova import Nova

    n = Nova.__new__(Nova)
    n._data_cache = {}
    res = n._query_budget_projection(
        {"budget": 60000, "roles": ["Warehouse Associate"], "locations": ["London, UK"], "industry": "logistics_supply_chain"}
    )
    keys = _all_keys(res.get("channel_allocations"), set()) | set(res.get("total_projected") or {})
    assert not ({"projected_hires", "cost_per_hire", "hires", "hires_low", "hires_high"} & keys), keys
    assert not any("hire" in r.lower() for r in res.get("recommendations") or [] if isinstance(r, str))
    assert res["hires_note"]


def test_nova_us_result_keeps_per_channel_hires():
    from nova import Nova

    n = Nova.__new__(Nova)
    n._data_cache = {}
    res = n._query_budget_projection(
        {"budget": 60000, "roles": ["Registered Nurse"], "locations": ["Columbus, OH"], "industry": "healthcare_medical"}
    )
    assert any("projected_hires" in ch for ch in res["channel_allocations"].values())


def test_strip_hire_claims_keeps_the_declaration_and_the_clients_goal():
    from plan_projection import strip_hire_claims as s

    assert s("Hires: no local cost-per-hire benchmark for this market — not projected.").endswith("not projected.")
    assert s("CLIENT GOAL:  100 hires target  —  this plan projects 25 (25%)") == "CLIENT GOAL:  100 hires target"
    assert s("If CPA rises 20%, projected hires drop from 25 to 20 (5 fewer hires)") == ""
    assert s("Timing: Peak Months: October. The plan is projected to deliver 25 hires.") == "Timing: Peak Months: October."
