"""2026-10-16 job 2617d01b: client "COHERENT CORP." shipped with
``client_name_wrong_casing@Market Intelligence!D30`` ("bundle delivered
anyway").

Root causes pinned here:

1. excel_v2's Market Intelligence "Company Profile" block wrote the SEC
   EDGAR legal name (always ALL CAPS: "COHERENT CORP.") raw into its
   "Company Name" row, bypassing ``display_format.client_display_name`` that
   every other sheet used; the delivery repair pass never covered casing
   (REPAIRABLE_CODES has no casing entry), so nothing fixed the cell.
2. bundle_qa's casing gate wrapped the canonical name in ``\\b...\\b``; for a
   canonical name ending in "." (Corp., Inc., Co.) ``\\b`` after the dot needs
   a following word character, so "COHERENT CORP." at the end of a cell (the
   usual place) never matched and the gate was blind to the dotted spelling.

The canonicaliser table at the bottom pins the casing rules themselves.
"""

from __future__ import annotations

import copy
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import display_format as fmt  # noqa: E402
import excel_v2  # noqa: E402
from bundle_qa import _check_client_name_casing, _TextUnit, run_bundle_qa  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _gate(raw_client_name: str, text: str) -> list[dict]:
    findings: list[dict] = []
    units = [_TextUnit(text, "Market Intelligence!D19", 0, 0, 0)]
    _check_client_name_casing(units, {"client_name": raw_client_name}, findings)
    return findings


def _plan(client_name: str) -> dict:
    return {
        "client_name": client_name,
        "industry": "technology",
        "budget": "$90,000",
        "campaign_duration": "6 months",
        "locations": ["Pittsburgh, PA"],
        "roles": ["Optical Engineer"],
        "_enriched": {
            "sec_data": {
                "name": "COHERENT CORP.",
                "cik": "0000820318",
                "sic": "3674",
                "state": "PA",
            }
        },
    }


def _all_strings(raw_xlsx: bytes):
    wb = load_workbook(io.BytesIO(raw_xlsx))
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    yield ws.title, cell.coordinate, cell.value


class TestGateSeesDottedNames:
    def test_all_caps_dotted_name_at_end_of_cell_is_flagged(self):
        findings = _gate("COHERENT CORP.", "Company Name: COHERENT CORP.")
        assert len(findings) == 1, findings
        assert findings[0]["code"] == "client_name_wrong_casing"

    def test_canonical_dotted_name_passes(self):
        assert _gate("COHERENT CORP.", "Company Name: Coherent Corp.") == []
        assert _gate("COHERENT CORP.", "Coherent Corp. is hiring.") == []


@pytest.mark.parametrize("typed", ["COHERENT CORP.", "COHERENT CORP"])
class TestWorkbookWritesTheDisplayName:
    def test_sec_legal_name_is_not_written_shouted(self, typed):
        raw = excel_v2.generate_excel_v2(copy.deepcopy(_plan(typed)))
        shouted = [(s, c, v) for s, c, v in _all_strings(raw) if "COHERENT" in v]
        assert not shouted, shouted

    def test_bundle_qa_reports_no_casing_finding(self, typed):
        plan = _plan(typed)
        raw = excel_v2.generate_excel_v2(copy.deepcopy(plan))
        casing = [
            f
            for f in run_bundle_qa(None, raw, plan)
            if f.get("code") == "client_name_wrong_casing"
        ]
        assert not casing, casing

    def test_company_name_row_carries_the_display_name(self, typed):
        raw = excel_v2.generate_excel_v2(copy.deepcopy(_plan(typed)))
        mi = [
            v
            for s, _c, v in _all_strings(raw)
            if s == "Market Intelligence" and v.startswith("Coherent Corp")
        ]
        assert mi, "Company Name row missing from Market Intelligence"


class TestCanonicaliserOutsideTheBrandTable:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("COHERENT CORP.", "Coherent Corp."),
            ("COHERENT CORP", "Coherent Corp"),
            ("THE HERSHEY COMPANY", "The Hershey Company"),
            ("HERSHEY", "Hershey"),
            ("COHERENT", "Coherent"),
            ("ACME CORP", "Acme Corp"),
            # Brand-table / acronym names stay shouted.
            ("ADP", "ADP"),
            ("3M", "3M"),
            ("IBM", "IBM"),
            ("UPS", "UPS"),
            ("IBM CORP.", "IBM Corp."),
            # Mixed-case typed names are the client's own spelling.
            ("Coherent Corp.", "Coherent Corp."),
            ("McKesson Corp", "McKesson Corp"),
            ("eBay", "eBay"),
        ],
    )
    def test_case(self, raw, expected):
        assert fmt.client_display_name(raw) == expected
        assert fmt.client_display_name(expected) == expected  # idempotent

    def test_workbook_and_deck_resolvers_agree(self):
        import ppt_generator

        for raw in ("COHERENT CORP.", "THE HERSHEY COMPANY", "ADP", "3M"):
            assert (
                excel_v2._proper_client_name(raw)
                == ppt_generator._proper_client_name(raw)
                == fmt.client_display_name(raw)
            )
