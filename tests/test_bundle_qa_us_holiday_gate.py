"""bundle_qa's US-holiday rule follows the generator's per-market gate
(2026-10-16).

The deliverables branch (gold_standard._us_holiday_mode) prints US holidays
by the plan's OWN markets: bare on an all-US plan, attributed "(US markets)"
on a plan that mixes US and non-US markets, and not at all when the plan has
no US market. ``bundle_qa`` still decided "is this a US plan" with
``plan_geo.is_us_plan`` (False for a New York + London plan), so it flagged
every legitimately labelled holiday on a mixed plan as
``us_data_on_non_us_plan`` critical -- six false criticals on the real
bundle, and the Slack post would have read "QA: 6 critical".

Fixture strings are the exact cells / deck sentence the branch (287cd78)
generated for hospitality + ["New York, NY", "London, UK"]; the 8 per-event
shapes are every US-holiday event the branch's calendar can emit across all
13 industries (labelled when the plan mixes markets, bare for an all-US plan).

Each "does not fire" test below was run against the pre-fix bundle_qa; the
"still fires" tests are positive controls that the fix did not turn the rule
off.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import bundle_qa  # noqa: E402

CODE = "us_data_on_non_us_plan"

MIXED = {"locations": ["New York, NY", "London, UK"]}
LONDON = {"locations": ["London, UK"]}
LONDON_DEFAULT_REGION = {"locations": ["London, UK"], "target_region": "us_only"}
US_ONLY = {"locations": ["Dallas, TX"]}
US_ON_GLOBAL_REGION = {
    "locations": ["Dallas, TX", "Austin, TX"],
    "target_region": "global",
}

# Every US-holiday event the branch calendar emits on a mixed plan ("label").
LABELLED_EVENTS = [
    "Memorial Day prep (US markets)",
    "Memorial Day staffing (US markets)",
    "Memorial Day weekend lull (US markets)",
    "Spring break peak hiring (Mar, US markets)",
    "Spring break peak hiring (US markets)",
    "Spring break staffing prep (US markets)",
    "Summer plant shutdown / maintenance week (often around July 4 in US markets)",
    "Thanksgiving staffing surge (US markets)",
]
# The same events on an all-US plan ("keep"): no attribution.
BARE_EVENTS = [
    "Memorial Day prep",
    "Memorial Day staffing",
    "Memorial Day weekend lull",
    "Spring break peak hiring",
    "Spring break peak hiring (Mar)",
    "Spring break staffing prep",
    "Thanksgiving staffing surge",
]
# Cells / sentence exactly as the generated mixed-plan bundle rendered them.
WORKBOOK_CELLS = [
    (
        "Quality Intelligence!F87",
        "Spring break staffing prep (US markets); Valentine's Day event surge",
    ),
    (
        "Quality Intelligence!F88",
        "Spring break peak hiring (US markets); Convention season begins",
    ),
    (
        "Quality Intelligence!F90",
        "Memorial Day staffing (US markets); Peak summer hiring push",
    ),
    ("Quality Intelligence!F96", "Thanksgiving staffing surge (US markets)"),
    (
        "Quality Intelligence!D99",
        (
            "Spring break peak hiring (Mar, US markets); NRA Show (Jun); "
            "Summer tourism peak (Jun-Aug); Holiday event season (Nov-Dec)"
        ),
    ),
]
DECK_SENTENCE = (
    "Impact: March, April, May run high hiring intensity for this plan "
    "(Spring break peak hiring (US markets), Summer season recruitment ramp, "
    "Memorial Day staffing (US markets))"
)


def _findings(plan: dict, *texts: str) -> list[dict]:
    units = [bundle_qa._TextUnit(t, f"Sheet!A{i + 1}") for i, t in enumerate(texts)]
    out: list[dict] = []
    bundle_qa._check_us_data_on_non_us_plan(units, None, plan, out)
    return [f for f in out if f["code"] == CODE]


# ---------------------------------------------------------------------------
# Mixed US + non-US plan: labelled holidays pass, bare ones do not
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("event", LABELLED_EVENTS)
def test_labelled_us_holiday_is_clean_on_a_mixed_plan(event):
    assert _findings(MIXED, event) == []


@pytest.mark.parametrize("location,text", WORKBOOK_CELLS)
def test_real_workbook_cells_are_clean_on_a_mixed_plan(location, text):
    assert _findings(MIXED, text) == []


def test_real_deck_sentence_is_clean_on_a_mixed_plan():
    assert _findings(MIXED, DECK_SENTENCE) == []


@pytest.mark.parametrize("event", BARE_EVENTS)
def test_bare_us_holiday_on_a_mixed_plan_is_critical(event):
    found = _findings(MIXED, event)
    assert len(found) == 1 and found[0]["severity"] == "critical", found
    assert "not attributed to US markets" in found[0]["message"]


def test_a_neighbouring_events_label_does_not_vouch_for_a_bare_one():
    """The deck lists events in one sentence: only Memorial Day is
    labelled, so Spring break must still be reported."""
    sentence = (
        "Impact: March, April, May run high hiring intensity for this plan "
        "(Spring break peak hiring, Summer season recruitment ramp, "
        "Memorial Day staffing (US markets))"
    )
    found = _findings(MIXED, sentence)
    assert len(found) == 1 and "Spring break" in found[0]["message"], found

    cell = "Spring break peak hiring; Memorial Day staffing (US markets)"
    found = _findings(MIXED, cell)
    assert len(found) == 1 and "Spring break" in found[0]["message"], found


@pytest.mark.parametrize(
    "text,labelled",
    [
        # The aside the generator builds for a holiday inside a parenthetical.
        ("Summer shutdown week (often around Memorial Day in US markets)", True),
        ("Summer shutdown week (often around Memorial Day)", False),
        # The generator's trailing-aside form.
        ("Spring break peak hiring (Mar, US markets)", True),
        ("Spring break peak hiring (Mar)", False),
        ("Spring break peak hiring (Mar, US markets); NRA Show (Jun)", True),
        # The label belongs to the NEXT event, past a separator.
        ("Spring break peak hiring; Convention season (US markets)", False),
        ("Spring break peak hiring, Convention season (US markets)", False),
        ("Spring break peak hiring\nConvention season (US markets)", False),
        # A deck sentence: the closing paren of the list ends the event.
        ("(Spring break peak hiring, Summer ramp (US markets))", False),
    ],
)
def test_event_label_is_read_inside_the_events_own_span(text, labelled):
    start = text.lower().index(
        next(m for m in ("memorial day", "spring break") if m in text.lower())
    )
    assert bundle_qa._us_holiday_event_is_labelled(text, start) is labelled


def test_one_bare_occurrence_in_a_cell_is_enough():
    cell = "Memorial Day staffing (US markets); Thanksgiving staffing surge"
    found = _findings(MIXED, cell)
    assert len(found) == 1 and "Thanksgiving" in found[0]["message"], found


# ---------------------------------------------------------------------------
# No US market: a US holiday is a defect, labelled or not
# ---------------------------------------------------------------------------
# "July 4" is not a bundle_qa marker (the shared list holds Thanksgiving,
# Memorial Day and Spring break; widening it would also catch Independence
# Day on an India / Mexico plan, where it is the plan's own holiday), so the
# plant-shutdown aside is not detectable by this rule on any plan.
_DETECTABLE_LABELLED = [e for e in LABELLED_EVENTS if "July 4" not in e]


@pytest.mark.parametrize(
    "event", ["Thanksgiving staffing surge", *_DETECTABLE_LABELLED]
)
def test_us_holiday_on_a_plan_with_no_us_market_is_critical(event):
    found = _findings(LONDON, event)
    assert len(found) == 1 and found[0]["severity"] == "critical", found


def test_london_plan_left_on_the_default_us_region_still_flags_us_holidays():
    """The wizard's region select defaults to "US Only", which is_us_plan
    honours over the locations; the generator drops US holidays for this
    plan, so one that appears is a defect."""
    assert bundle_qa.plan_geo.is_us_plan(LONDON_DEFAULT_REGION) is True
    found = _findings(LONDON_DEFAULT_REGION, "Thanksgiving staffing surge")
    assert len(found) == 1 and found[0]["severity"] == "critical", found


# ---------------------------------------------------------------------------
# All-US markets: bare holidays are right
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("event", BARE_EVENTS)
def test_bare_us_holiday_is_clean_on_a_us_only_plan(event):
    assert _findings(US_ONLY, event) == []


@pytest.mark.parametrize("event", BARE_EVENTS)
def test_bare_us_holiday_is_clean_on_us_locations_under_a_global_region(event):
    """is_us_plan is False here (region "global") yet every market is US:
    the generator prints the holidays bare."""
    assert bundle_qa.plan_geo.is_us_plan(US_ON_GLOBAL_REGION) is False
    assert _findings(US_ON_GLOBAL_REGION, event) == []


# ---------------------------------------------------------------------------
# Scope: only holidays moved; US-sourced data keeps the plan-wide gate
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "text",
    [
        "Fed Funds Rate: 3.64%",
        "BLS JOLTS openings (US markets)",
        "Expected Starting Salary: US$101,500 (US markets)",
    ],
)
def test_non_holiday_us_markers_are_still_critical_on_a_mixed_plan(text):
    found = _findings(MIXED, text)
    assert len(found) == 1 and found[0]["severity"] == "critical", found


def test_non_holiday_us_markers_are_still_clean_on_a_us_plan():
    assert _findings(US_ONLY, "Fed Funds Rate: 3.64%") == []


def test_a_labelled_holiday_does_not_hide_a_us_data_marker_in_the_same_cell():
    found = _findings(
        MIXED, "Memorial Day staffing (US markets); Fed Funds Rate: 3.64%"
    )
    assert len(found) == 1 and "Fed Funds" in found[0]["message"], found


def test_holiday_markers_are_a_subset_of_the_shared_marker_list():
    shared = {m.lower() for m in bundle_qa._US_ONLY_MARKERS}
    assert bundle_qa._US_HOLIDAY_MARKERS <= shared


# ---------------------------------------------------------------------------
# Mirror of the generator's gate
# ---------------------------------------------------------------------------
_MODE_CASES = [
    (["Dallas, TX", "Remote"], {}, "keep"),
    (["Dallas, TX", "Toronto, ON, Canada"], {}, "label"),
    (["New York, NY", "London, UK"], {}, "label"),
    (["Dallas, TX", "Munich, Germany"], {"target_region": "us_only"}, "label"),
    (["London, UK"], {}, "drop"),
    (["London, UK"], {"target_region": "us_only"}, "drop"),
    (["Toronto, ON, Canada", "Vancouver, BC, Canada"], {}, "drop"),
    (["Dallas, TX", "Austin, TX"], {"target_region": "global"}, "keep"),
]


@pytest.mark.parametrize("locations,extra,mode", _MODE_CASES)
def test_holiday_mode_follows_the_plans_own_markets(locations, extra, mode):
    assert bundle_qa._us_holiday_mode({"locations": locations, **extra}) == mode


@pytest.mark.parametrize("locations,extra,_mode", _MODE_CASES)
def test_market_presence_agrees_with_the_generators_gate(locations, extra, _mode):
    """gold_standard.plan_market_presence / _us_holiday_mode (deliverables
    branch 287cd78) and this module's mirror must never drift apart. Skips
    until that branch is on main."""
    import gold_standard

    if not hasattr(gold_standard, "plan_market_presence"):
        pytest.skip("generator holiday gate not on this branch yet")
    data = {"locations": locations, **extra}
    assert bundle_qa._plan_market_presence(data) == gold_standard.plan_market_presence(
        data
    )
    assert bundle_qa._us_holiday_mode(data) == gold_standard._us_holiday_mode(data)


def test_market_presence_fails_open_on_garbage():
    assert bundle_qa._plan_market_presence(None) == (True, False)
    assert bundle_qa._us_holiday_mode(None) == "keep"


# ---------------------------------------------------------------------------
# End to end through run_bundle_qa, workbook and deck parsed for real
# ---------------------------------------------------------------------------
def _bundle(cells, sentence):
    openpyxl = pytest.importorskip("openpyxl")
    pptx_mod = pytest.importorskip("pptx")
    from pptx.util import Inches

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quality Intelligence"
    for i, (_loc, text) in enumerate(cells, start=1):
        ws.cell(row=i, column=6, value=text)
    xbuf = io.BytesIO()
    wb.save(xbuf)

    prs = pptx_mod.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = sentence
    pbuf = io.BytesIO()
    prs.save(pbuf)
    return pbuf.getvalue(), xbuf.getvalue()


def _run(plan, cells, sentence):
    pptx_bytes, xlsx_bytes = _bundle(cells, sentence)
    findings = bundle_qa.run_bundle_qa(pptx_bytes, xlsx_bytes, plan)
    return [f for f in findings if f["code"] == CODE]


def test_a_mixed_plan_bundle_with_labelled_holidays_has_no_critical():
    assert _run(MIXED, WORKBOOK_CELLS, DECK_SENTENCE) == []


def test_a_mixed_plan_bundle_with_one_bare_holiday_reports_exactly_it():
    cells = list(WORKBOOK_CELLS)
    cells[3] = ("Quality Intelligence!F96", "Thanksgiving staffing surge")
    found = _run(MIXED, cells, DECK_SENTENCE)
    assert len(found) == 1, found
    assert found[0]["severity"] == "critical"
    assert "Thanksgiving" in found[0]["message"]
    assert found[0]["location"] == "Quality Intelligence!F4"


def test_a_london_bundle_carrying_us_holidays_is_still_critical():
    found = _run(LONDON, WORKBOOK_CELLS, DECK_SENTENCE)
    assert len(found) == len(WORKBOOK_CELLS) + 1, found
    assert all(f["severity"] == "critical" for f in found)


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
