"""Regression tests for the 2026-09-24 deck-defect audit (findings #11, #13,
#21, #22 -- see docs/deck audit closure trail; #10/#12 were already closed
by efdab5e and re-verified, not re-covered here).

Runs under pytest, or standalone: ``python3 tests/test_ppt_deck_defects_2026_09_24.py``.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import ppt_generator as ppt  # noqa: E402
from pptx import Presentation  # noqa: E402

EMU_PER_IN = 914400


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            try:
                yield from _iter_shapes(sh.shapes)
            except Exception:
                pass


def _text_shapes(slide):
    out = []
    for sh in _iter_shapes(slide.shapes):
        try:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh)
        except Exception:
            pass
    return out


def _slide_containing(prs: Presentation, needle: str):
    for slide in prs.slides:
        for sh in _text_shapes(slide):
            if needle in sh.text_frame.text:
                return slide
    return None


def _shape_with_text(slide, needle: str, exact: bool = False):
    for sh in _text_shapes(slide):
        t = sh.text_frame.text
        if (t.strip() == needle) if exact else (needle in t):
            return sh
    return None


def _run_font_pt(shape, default: float = 10.0) -> float:
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                return r.font.size.pt
    return default


def _v_overlap(a, b) -> float:
    """Vertical overlap in inches between two shapes' declared boxes."""
    a_top, a_bot = a.top / EMU_PER_IN, (a.top + a.height) / EMU_PER_IN
    b_top, b_bot = b.top / EMU_PER_IN, (b.top + b.height) / EMU_PER_IN
    return max(0.0, min(a_bot, b_bot) - max(a_top, b_top))


# ---------------------------------------------------------------------------
# #11 -- slide 6 (Budget Allocation) hero KPI value must fit on one line for
# a low-denomination currency total, same as the structurally identical
# slide-2 exec-summary hero already does via _fit_font_single_line.
# ---------------------------------------------------------------------------
def _idr_budget_plan() -> Dict[str, Any]:
    return {
        "client_name": "Acme",
        "industry": "hospitality_travel",
        "industry_label": "Hospitality & Travel",
        "locations": ["Jakarta, Indonesia"],
        "roles": ["Commercial Cab Driver"],
        "budget": "Rp10,000,000,000",
        "work_environment": "hybrid",
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "social_media": True,
        },
        "_budget_allocation": {
            "metadata": {"total_budget": 10000000000},
            "total_projected": {"applications": 50000, "hires": 500},
            "channel_allocations": {
                "programmatic_dsp": {"projected_hires": 200},
                "global_boards": {"projected_hires": 150},
                "social_media": {"projected_hires": 150},
            },
        },
    }


def test_budget_slide_hero_value_fits_one_line_for_low_denomination_currency():
    prs = Presentation(__import__("io").BytesIO(ppt.generate_pptx(_idr_budget_plan())))
    slide = _slide_containing(prs, "Total Investment")
    assert slide is not None, "Budget Allocation slide (hero KPI band) not found"

    value_shape = _shape_with_text(slide, "10,000,000,000")
    label_shape = _shape_with_text(slide, "Total Investment", exact=True)
    assert value_shape is not None
    assert label_shape is not None

    font_pt = _run_font_pt(value_shape)
    n_lines = ppt._measure_lines(
        value_shape.text_frame.text, value_shape.width / EMU_PER_IN, font_pt, bold=True
    )
    assert n_lines == 1, (
        f"IDR hero value wrapped to {n_lines} lines at {font_pt}pt in a "
        f"{value_shape.width / EMU_PER_IN:.2f}in box -- no fit-to-width applied"
    )
    # The wrapped 2nd line is what physically overprinted the caption below
    # it -- assert the two declared boxes don't vertically overlap either.
    assert _v_overlap(value_shape, label_shape) <= 0.02, (
        "hero value box overlaps 'Total Investment' caption box"
    )


# ---------------------------------------------------------------------------
# #13 -- slide 5 (Channel Strategy) action subhead must not collide with the
# CHANNEL MIX / INDUSTRY BENCHMARKS section headers when it needs 3 lines.
# ---------------------------------------------------------------------------
def _long_subhead_plan() -> Dict[str, Any]:
    return {
        "client_name": (
            "The International Consolidated Hospitality Leisure and Gaming "
            "Holdings Group PLC & Co KGaA"
        ),
        "industry": "hospitality_travel",
        "industry_label": (
            "Hospitality, Travel, Leisure, Gaming and Entertainment Services (Global)"
        ),
        "locations": ["London, United Kingdom"],
        "roles": ["Commercial Cab Driver"],
        "budget": "£20,000,000",
        "work_environment": "hybrid",
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "niche_boards": True,
            "social_media": True,
            "regional_boards": True,
            "employer_branding": True,
        },
        "_budget_allocation": {
            "metadata": {"total_budget": 20000000},
            "total_projected": {"applications": 50000, "hires": 5000},
            "channel_allocations": {
                "programmatic_dsp": {"projected_hires": 1000},
                "global_boards": {"projected_hires": 1000},
                "niche_boards": {"projected_hires": 1000},
                "social_media": {"projected_hires": 1000},
                "regional_boards": {"projected_hires": 500},
                "employer_branding": {"projected_hires": 500},
            },
        },
    }


def test_channel_strategy_subhead_never_overlaps_section_headers():
    prs = Presentation(__import__("io").BytesIO(ppt.generate_pptx(_long_subhead_plan())))
    slide = _slide_containing(prs, "CHANNEL MIX")
    assert slide is not None, "Channel Strategy slide not found"

    subhead = _shape_with_text(slide, "channel strategy allocates")
    assert subhead is not None, "action subhead shape not found"
    font_pt = _run_font_pt(subhead)
    n_lines = ppt._measure_lines(
        subhead.text_frame.text, subhead.width / EMU_PER_IN, font_pt, bold=True
    )
    assert n_lines >= 2, "fixture should force a multi-line subhead"

    # The declared box height is not a reliable collision signal here -- the
    # pre-fix box was a fixed Inches(0.5) regardless of content, so an
    # overflowing 3rd line still reports a "0.5in" shape even though the
    # RENDERED glyphs print past it. Compute the subhead's actual measured
    # ink bottom (same estimate _autofit_textframe/collide.py use: line
    # count * point size * a calibrated line-height factor) and assert it
    # clears the header start, rather than comparing declared boxes.
    subhead_top_in = subhead.top / EMU_PER_IN
    measured_bottom_in = subhead_top_in + n_lines * (font_pt * 1.35) / 72.0

    channel_mix_header = _shape_with_text(slide, "CHANNEL MIX", exact=True)
    industry_bm_header = _shape_with_text(slide, "INDUSTRY BENCHMARKS", exact=True)
    assert channel_mix_header is not None
    assert industry_bm_header is not None

    header_top_in = channel_mix_header.top / EMU_PER_IN
    assert channel_mix_header.top == industry_bm_header.top, (
        "left/right section headers expected at the same cascaded y"
    )
    assert measured_bottom_in <= header_top_in + 0.02, (
        f"multi-line subhead's measured ink (bottom={measured_bottom_in:.2f}in) "
        f"overlaps the CHANNEL MIX / INDUSTRY BENCHMARKS headers "
        f"(top={header_top_in:.2f}in)"
    )


# ---------------------------------------------------------------------------
# #21 -- a USD plan whose CPH benchmark row is the plan's OWN derived
# cost-per-hire (not a real external industry constant) must carry the
# "(this plan)" qualifier just like a non-USD plan already does.
# ---------------------------------------------------------------------------
def _usd_plan_derived_cph() -> Dict[str, Any]:
    return {
        "client_name": "Acme",
        "industry": "general_entry_level",
        "industry_label": "General / Entry Level",
        "locations": ["Austin, TX"],
        "roles": ["Clerk"],
        "budget": "$500,000",
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "social_media": True,
        },
        # Force _kb_recruitment_industry_benchmark to miss (no real KB CPH
        # to prefer) so the plan-derived budget-engine figure is what's
        # actually shown -- this is the exact scenario the fix targets.
        "_knowledge_base": {},
        "_budget_allocation": {
            "metadata": {"total_budget": 500000},
            "total_projected": {
                "applications": 5000,
                "hires": 200,
                "cost_per_hire": 2500,
            },
            "channel_allocations": {
                "programmatic_dsp": {"projected_hires": 100},
                "global_boards": {"projected_hires": 50},
                "social_media": {"projected_hires": 50},
            },
        },
    }


def test_usd_plan_derived_cph_is_labeled_this_plan_not_industry():
    data = _usd_plan_derived_cph()
    prs = Presentation(__import__("io").BytesIO(ppt.generate_pptx(data)))
    slide = _slide_containing(prs, "INDUSTRY BENCHMARKS")
    assert slide is not None, "Channel Strategy slide (benchmark table) not found"

    bare_label = _shape_with_text(slide, "Est. Cost-per-Hire", exact=True)
    qualified_label = _shape_with_text(
        slide, "Est. Cost-per-Hire (this plan)", exact=True
    )
    assert bare_label is None, (
        "USD plan's own derived CPH rendered under the bare "
        "'Est. Cost-per-Hire' label -- reads as an external industry "
        "benchmark, not the plan's own number"
    )
    assert qualified_label is not None, (
        "expected 'Est. Cost-per-Hire (this plan)' label for a USD plan "
        "whose CPH value is its own budget-engine projection"
    )


# ---------------------------------------------------------------------------
# #22 -- the "<Industry> Average" comparison row for Channels Selected must
# stay the coded INDUSTRY_BENCHMARKS_COMPARISON value, never the client's
# own funded-channel count.
# ---------------------------------------------------------------------------
def _six_channel_healthcare_plan() -> Dict[str, Any]:
    return {
        "client_name": "Acme Health",
        "industry": "healthcare_medical",
        "industry_label": "Healthcare & Medical",
        "locations": ["Dallas, TX"],
        "roles": ["Registered Nurse"],
        "budget": "$1,000,000",
        "channel_categories": {
            "programmatic_dsp": True,
            "global_boards": True,
            "niche_boards": True,
            "social_media": True,
            "regional_boards": True,
            "employer_branding": True,
        },
        "_budget_allocation": {
            "metadata": {"total_budget": 1000000},
            "total_projected": {"applications": 8000, "hires": 300},
            "channel_allocations": {
                "programmatic_dsp": {"percentage": 25, "projected_hires": 60},
                "global_boards": {"percentage": 20, "projected_hires": 50},
                "niche_boards": {"percentage": 15, "projected_hires": 40},
                "social_media": {"percentage": 15, "projected_hires": 50},
                "regional_boards": {"percentage": 15, "projected_hires": 50},
                "employer_branding": {"percentage": 10, "projected_hires": 50},
            },
        },
    }


def test_industry_comparison_avg_channels_not_overwritten_by_client_count():
    # Unit-level: the coded constant must survive the overlay unchanged even
    # though the plan funds 6 channels (avg_channels for healthcare_medical
    # is hardcoded to 4 in INDUSTRY_BENCHMARKS_COMPARISON).
    data = _six_channel_healthcare_plan()
    ind = ppt._get_industry_comparison("healthcare_medical", data)
    coded_avg = ppt.INDUSTRY_BENCHMARKS_COMPARISON["healthcare_medical"]["avg_channels"]
    assert ind["avg_channels"] == coded_avg, (
        f"avg_channels overlay returned {ind['avg_channels']!r}, expected the "
        f"coded industry constant {coded_avg!r} (not the client's own "
        f"6-channel count)"
    )

    # Slide-level: the rendered "Channels Selected" row must show the SAME
    # coded value in its industry column, not the client's own count.
    prs = Presentation(__import__("io").BytesIO(ppt.generate_pptx(data)))
    slide = _slide_containing(prs, "Channels Selected")
    assert slide is not None, "Plan Comparison slide not found"
    coded_shape = _shape_with_text(slide, str(coded_avg), exact=True)
    assert coded_shape is not None, (
        f"expected the industry column to render the coded avg_channels "
        f"value {coded_avg!r} somewhere on the comparison slide"
    )


if __name__ == "__main__":
    import traceback

    tests = [
        test_budget_slide_hero_value_fits_one_line_for_low_denomination_currency,
        test_channel_strategy_subhead_never_overlaps_section_headers,
        test_usd_plan_derived_cph_is_labeled_this_plan_not_industry,
        test_industry_comparison_avg_channels_not_overwritten_by_client_count,
    ]
    failures = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except Exception:  # noqa: BLE001
            failures += 1
            print(f"FAIL {t.__name__}")
            traceback.print_exc()
    sys.exit(1 if failures else 0)
