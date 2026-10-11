"""Plan-output audit (2026-10-11) -- competitor / claim honesty (C1-C6).

``hiring_intensity`` in the gold-standard competitor map is a static
hiring-difficulty proxy (gold_standard.py), not an observed count of
employers, so no surface may phrase it as "Fortune 500+ companies are
actively hiring" or "high hiring volume".
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402

import excel_v2  # noqa: E402


def _sheet_text(ws) -> str:
    return "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value not in (None, "")
    )


def _data_with_high_intensity() -> dict:
    return {
        "client_name": "Hershey Co",
        "industry": "food_beverage",
        "industry_label": "Food & Beverage",
        "locations": ["Hershey, PA", "Chicago, IL"],
        "roles": ["Production Operator"],
        "target_roles": ["Production Operator"],
        "budget": "$200,000",
        "competitors": [],
        "_synthesized": {},
        "_enriched": {},
        "_budget_allocation": {},
        "_gold_standard": {
            "competitor_mapping": {
                "Hershey": {
                    "top_employers": ["Acme Foods"],
                    "hiring_intensity": "high",
                },
                "Chicago": {
                    "top_employers": ["Acme Foods"],
                    "hiring_intensity": "very_high",
                },
                "Dallas": {"top_employers": [], "hiring_intensity": "moderate"},
            }
        },
    }


# -- C1 ---------------------------------------------------------------------


def _risk_texts_excel(data: dict) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_executive_summary(ws, data)
    return _sheet_text(ws)


def test_c1_excel_risk_row_does_not_claim_fortune_500_hiring():
    text = _risk_texts_excel(_data_with_high_intensity())
    assert "Competitive Pressure" in text, "risk row did not render (vacuous test)"
    assert "Fortune 500" not in text
    assert "actively hiring" not in text
    assert "2 markets score high hiring difficulty" in text


def test_c1_deck_risk_slide_does_not_claim_fortune_500_hiring():
    from pptx import Presentation

    import ppt_generator

    prs = Presentation()
    ppt_generator._build_slide_risk_analysis(prs, _data_with_high_intensity())
    texts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    texts.extend(c.text for c in r.cells)
    blob = "\n".join(texts)
    assert "COMPETITION" in blob or "High Competition" in blob, "risk row missing"
    assert "Fortune 500" not in blob
    assert "2 markets score high on modeled hiring difficulty" in blob


# -- C2 ---------------------------------------------------------------------


def _quality_intel_text(data: dict) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_quality_intelligence(ws, data, data["_gold_standard"])
    return _sheet_text(ws)


def test_c2_quality_intelligence_says_hiring_difficulty_not_hiring_volume():
    for brief in ([], ["Acme Foods"]):
        data = _data_with_high_intensity()
        data["competitors"] = brief
        text = _quality_intel_text(data)
        assert "Why They Matter" in text, "competitive table missing (vacuous)"
        assert "High hiring volume" not in text
        assert "Hiring Intensity" not in text
        assert "Hiring Difficulty" in text
        assert "Hard-to-fill market in Hershey (hiring-difficulty score)" in text


# -- C3 / C5 ----------------------------------------------------------------


def _client_only_enrichment() -> dict:
    """Everything the pipeline typically has: the CLIENT's own Wikipedia,
    Clearbit, SEC and BLS data -- and not one byte about any competitor."""
    return {
        "company_info": {
            "description": "Hershey is a confectionery manufacturer.",
            "summary": "Hershey makes chocolate.",
            "url": "https://en.wikipedia.org/wiki/Hershey",
        },
        "company_metadata": {"domain": "thehersheycompany.com", "logo": "x.png"},
        "sec_data": {"ticker": "HSY", "cik": "0000047111", "recent_filings": [{}]},
        "industry_employment": {"employment": 1000, "avg_weekly_wage": 1200},
    }


def test_c3_competitive_intelligence_not_graded_on_the_clients_own_sources():
    import data_synthesizer

    ci = data_synthesizer.fuse_competitive_intelligence(
        _client_only_enrichment(),
        {},
        {"company_name": "Hershey Co", "industry": "food_beverage"},
    )
    scores = data_synthesizer.compute_confidence_scores({"competitive_intelligence": ci})
    score = scores["per_section"]["competitive_intelligence"]
    assert score == 0.0, f"graded {score} on client-only sources"
    meta = ci["_meta"]
    assert meta["competitor_evidence_count"] == 0
    assert "no competitor evidence" in meta["note"].lower()


def test_c3_competitive_intelligence_still_earns_credit_for_real_competitor_evidence():
    import data_synthesizer

    enriched = _client_only_enrichment()
    enriched["competitor_logos"] = {
        "Acme Foods": {
            "domain": "acmefoods.com",
            "logo": "a.png",
            "description": "Acme Foods operates 4 plants in PA per its 10-K.",
            "source_url": "https://example.com/10k",
        },
        "Beta Co": {
            "domain": "beta.com",
            "evidence": "careers page lists 12 open production roles",
        },
    }
    ci = data_synthesizer.fuse_competitive_intelligence(
        enriched, {}, {"company_name": "Hershey Co", "industry": "food_beverage"}
    )
    assert ci["_meta"]["competitor_evidence_count"] == 2
    scores = data_synthesizer.compute_confidence_scores({"competitive_intelligence": ci})
    assert scores["per_section"]["competitive_intelligence"] >= 0.8


def test_c3_sources_sheet_carries_the_no_competitor_evidence_note():
    import data_synthesizer

    ci = data_synthesizer.fuse_competitive_intelligence(
        _client_only_enrichment(),
        {},
        {"company_name": "Hershey Co", "industry": "food_beverage"},
    )
    synthesized = {
        "competitive_intelligence": ci,
        "confidence_scores": data_synthesizer.compute_confidence_scores(
            {"competitive_intelligence": ci}
        ),
    }
    data = _data_with_high_intensity()
    data["_synthesized"] = synthesized
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_sources(ws, data)
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[1] == "Competitive Intelligence"]
    assert rows, _sheet_text(ws)
    assert rows[0][3] == "F"
    assert "no competitor evidence" in str(rows[0][4]).lower()


def test_c5_unknown_public_status_is_unknown_not_no():
    import data_synthesizer

    ci = data_synthesizer.fuse_competitive_intelligence(
        {"company_info": {"description": "A confectionery manufacturer."}},
        {},
        {"company_name": "Hershey Co", "industry": "food_beverage"},
    )
    mp = ci["market_positioning"]
    assert mp["is_public_company"] is None
    assert mp["has_sec_filings"] is None

    data = _data_with_high_intensity()
    data["_synthesized"] = {"competitive_intelligence": ci}
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_market_intelligence(ws, data)
    text = _sheet_text(ws)
    assert "Is Public Company: Unknown" in text
    assert "Has Sec Filings: Unknown" in text
    assert "Is Public Company: No" not in text


def test_c5_known_public_company_still_reads_yes():
    import data_synthesizer

    ci = data_synthesizer.fuse_competitive_intelligence(
        _client_only_enrichment(), {}, {"company_name": "Hershey Co", "industry": "food_beverage"}
    )
    mp = ci["market_positioning"]
    assert mp["is_public_company"] is True
    assert mp["has_sec_filings"] is True


# -- C4 ---------------------------------------------------------------------


def _landscape_slide_text(synth: dict) -> str:
    from pptx import Presentation

    import ppt_generator

    prs = Presentation()
    data = _data_with_high_intensity()
    data["_synthesized"] = {"competitive_intelligence": synth}
    ppt_generator._build_slide_competitive_landscape(prs, data)
    assert len(prs.slides) == 1
    texts = []
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame:
            texts.append(sh.text_frame.text)
    return "\n".join(texts)


def test_c4_deck_source_footer_lists_only_sources_actually_present():
    import data_synthesizer

    bare = data_synthesizer.fuse_competitive_intelligence(
        {"company_info": {"description": "x"}}, {}, {"company_name": "Hershey Co"}
    )
    text = _landscape_slide_text(bare)
    assert "Public company filings" not in text
    assert "industry employment data" not in text
    assert "curated knowledge base" not in text


def test_c4_deck_source_footer_names_filings_and_bls_when_present():
    import data_synthesizer

    full = data_synthesizer.fuse_competitive_intelligence(
        _client_only_enrichment(), {}, {"company_name": "Hershey Co", "industry": "food_beverage"}
    )
    text = _landscape_slide_text(full)
    source_line = next(t for t in text.splitlines() if t.startswith("Sources"))
    assert "SEC filings" in source_line
    assert "BLS" in source_line
    assert "no competitor-level sources" in source_line
