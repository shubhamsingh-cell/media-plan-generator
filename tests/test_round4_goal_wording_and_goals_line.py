"""Regression: design-judge round 4 (2026-10-01) on mpg-projection-honesty.

Pre-fix (branch at da3c62f):
F1 India's break-even read three ways: slide 9 "met only at ₹35K/hire" (the
   plan's own cost per hire, not the break-even), B50 "only at plan
   efficiency", B17 "goal holds only if hires cost at most ₹50,000". One
   phrasing now, from budget ÷ goal: "goal met if hires cost ≤ ₹50K each".
F2 Protecting the "Client goal:" line made the card's autofit trim the
   client's own goals bullet by bullet: Hershey showed "Client Goals:
   • Direct Hiring" (Talent Pipeline gone), hospital and Brazil lost the
   whole block. The goals are now ONE protected inline line.
F3 Slide 2's top-up "+~$281.9K" did not say which cost per hire it assumes.
F4 _KPI_SUBLABEL_LINE_IN was 0.14 (clearance held only in the preview
   renderer's 1.2x line model).
F5 Hershey's workbook cost-per-hire cell held three bases (media-only,
   total, a sister sector's $1,070).
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import display_format  # noqa: E402
import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

_BE_INDIA = "goal met if hires cost ≤ ₹50K each"


def _brief(name, industry, budget, locations, roles, hire_volume, goals, duration):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": duration,
        "hire_volume": hire_volume,
        "work_environment": ["On-site"],
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
        "campaign_goals": goals,
    }


class _Bundle:
    def __init__(self, brief):
        self.data = T.build_plan_data(brief)
        self.prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(self.data))))
        raw = excel_v2.generate_excel_v2(dict(self.data), load_kb_fn=load_knowledge_base)
        raw = raw[0] if isinstance(raw, tuple) else raw
        self.wb = load_workbook(io.BytesIO(raw))

    def deck_paras(self):
        return [
            p.text
            for s in self.prs.slides
            for sh in s.shapes
            if sh.has_text_frame
            for p in sh.text_frame.paragraphs
            if p.text.strip()
        ]

    def slide2(self):
        return [
            p.text
            for sh in self.prs.slides[1].shapes
            if sh.has_text_frame
            for p in sh.text_frame.paragraphs
            if p.text.strip()
        ]

    def exec_texts(self):
        return [
            str(c.value)
            for r in self.wb["Executive Summary"].iter_rows()
            for c in r
            if isinstance(c.value, str)
        ]


# Mirrors of the design-judge sweep briefs (B_sweep/briefs) -- enough
# content to overfill the slide-2 RESOLUTION card.
@pytest.fixture(scope="module")
def india():
    return _Bundle(_brief(
        "Nimbus Infotech Services", "tech_engineering", "₹25,000,000",
        ["Bengaluru, Karnataka, India", "Hyderabad, Telangana, India",
         "Pune, Maharashtra, India"],
        ["Software Engineer (Fresher)", "Java Developer", "QA Engineer",
         "Senior Data Engineer (Lateral)", "DevOps Engineer", "Business Analyst"],
        "500+ hires", ["Direct Hiring", "Talent Pipeline"], "6-12 months"))


@pytest.fixture(scope="module")
def hershey():
    return _Bundle(_brief(
        "THE HERSHEY COMPANY", "food_beverage", "$45,000",
        ["Hershey, PA", "Hazleton, PA", "Reading, PA", "Palmyra, PA",
         "Lancaster, PA", "Stuarts Draft, VA", "Memphis, TN", "Columbus, OH",
         "Dayton, OH", "Houston, TX"],
        ["Machine Operator", "Maintenance Technician", "Sanitation",
         "Quality Tech", "Research Associate"],
        "100-500 hires", ["Direct Hiring", "Talent Pipeline"], "3-6 months"))


@pytest.fixture(scope="module")
def hospital():
    return _Bundle(_brief(
        "Lakeside Regional Health System", "healthcare_medical", "$250,000",
        ["Houston, TX", "Dallas, TX", "Phoenix, AZ", "Tampa, FL",
         "Charlotte, NC", "Columbus, OH"],
        ["Registered Nurse", "Certified Nursing Assistant",
         "Clinical Pharmacist", "ICU Nurse"],
        "100-500 hires", ["Direct Hiring"], "3 months"))


@pytest.fixture(scope="module")
def brazil():
    return _Bundle(_brief(
        "Conecta Contact Center Brasil", "general_entry_level", "R$1,200,000",
        ["São Paulo, Brazil", "Rio de Janeiro, Brazil"],
        ["Customer Service Representative", "Telemarketing Agent",
         "Bilingual Support Agent", "Team Leader"],
        "500+ hires", ["Direct Hiring", "Lead Generation"], "3-6 months"))


class TestOneBreakEvenPhrasing:
    def test_helper(self):
        assert display_format.goal_break_even_phrase(
            25_000_000, 500, lambda v: f"₹{v / 1000:.0f}K"
        ) == _BE_INDIA
        assert display_format.goal_break_even_phrase(0, 500, str) == ""

    def test_every_surface_says_it_the_same_way(self, india):
        deck = india.deck_paras()
        assert any(_BE_INDIA in p for p in india.slide2()), india.slide2()
        assert any(p.startswith("CLIENT GOAL") and _BE_INDIA in p for p in deck)
        xl = india.exec_texts()
        assert any(t.startswith("Hiring goal:") and _BE_INDIA in t for t in xl)
        assert any(
            t.startswith("This recruitment media plan") and _BE_INDIA in t for t in xl
        ), [t for t in xl if t.startswith("This recruitment")]
        everything = " ".join(deck + xl)
        assert "met only at" not in everything
        assert "goal holds only if" not in everything
        assert "only at plan efficiency" not in everything


class TestClientGoalsNeverPartial:
    @pytest.mark.parametrize(
        "name,line",
        [
            ("india", "Client goals: Direct Hiring · Talent Pipeline"),
            ("hershey", "Client goals: Direct Hiring · Talent Pipeline"),
            ("hospital", "Client goals: Direct Hiring"),
            ("brazil", "Client goals: Direct Hiring · Lead Generation"),
        ],
    )
    def test_goals_line_is_complete(self, name, line, request):
        b = request.getfixturevalue(name)
        assert line in b.slide2(), b.slide2()
        assert not any(p.startswith("•") for p in b.slide2())

    def test_autofit_keeps_the_goals_line(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        tf = box.text_frame
        texts = ["MARKET THESIS"] + [f"filler {i}" for i in range(12)] + [
            "Client goals: Direct Hiring · Talent Pipeline",
            "Client goal: 100 hires — this plan projects 15 (15% of goal)",
            "trailing filler",
        ]
        for i, t in enumerate(texts):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = t
            r.font.size = Pt(9)
        ppt_generator._autofit_textframe(tf, 3.0, 1.2)
        kept = [p.text for p in tf.paragraphs]
        assert "Client goals: Direct Hiring · Talent Pipeline" in kept, kept


class TestTopUpStatesItsBasis:
    def test_hospital_top_up_names_the_cost_per_hire(self, hospital):
        line = next(p for p in hospital.slide2() if p.startswith("Client goal:"))
        assert "at $5.3K/hire would close the gap" in line, line
        assert "at the $10.5K midpoint" in line, line


def test_kpi_sublabel_line_matches_the_generator_line_box():
    # 8pt x 1.42 line box = 0.158in
    assert ppt_generator._KPI_SUBLABEL_LINE_IN >= 0.158


class TestOneIndustryRowInTheWorkbook:
    def test_hershey_row_is_the_shown_row_only(self, hershey):
        rows = [
            [c.value for c in r if c.value]
            for r in hershey.wb["Executive Summary"].iter_rows()
            if any(c.value == "Industry Cost-per-Hire" for c in r)
        ]
        assert rows and rows[0][1] == "$3,000-$5,000 (food manufacturing)", rows
        assert not any("$1,070" in t for t in hershey.exec_texts())
