"""Regression tests for four verified defects in app.py's plan-generation
path (2026-09-24 sweep):

C17 -- the async (wizard job) /api/generate path never set
``data["campaign_weeks"]`` (only the synchronous path's inline block did),
so ppt_generator's deck timeline (``data.get("campaign_weeks", 12)``)
silently fell back to a fixed 12-week phasing regardless of the plan's
real duration, and bundle_qa's ``campaign_duration_incoherence`` rule (which
treats ``campaign_weeks`` as authoritative) then flagged every async plan
whose duration wasn't coincidentally ~12 weeks as a blocking critical.
Fixed by extracting ``app._resolve_campaign_weeks(data)`` -- the single
shared helper -- and calling it from BOTH the sync handler and the async
worker (``_run_async_generate``).

C5 -- the budget engine's channel_percentages were built purely from
INDUSTRY_ALLOC_PROFILES by industry, with no awareness of which channels
the wizard's toggles (``data["channel_categories"]``) turned off, so a
disabled channel still received budget (and projected clicks/applications/
hires) in every deliverable. Fixed by ``app._apply_channel_selection``,
called at both ``calculate_budget_allocation`` call sites in
/api/generate (sync + async).

C8 -- app.py called ``calculate_budget_allocation`` with no
``plan_currency`` at any of its three call sites, so the engine fell back
to its own best-effort location guess for currency instead of the single
shared resolver the deck/workbook/scorecard/gate all use
(``plan_currency.currency_for_plan_with_basis``) -- e.g. a USD-typed
budget for a single UK market could guess GBP CPCs from the location.
Fixed by ``app._resolve_plan_currency``, called (and passed as
``plan_currency=``) at every call site.

C3 -- ``_infer_industry_from_signals``'s Detector 1 (company-name-only
scan) raised a CRITICAL ``industry_client_conflict`` on correct Blue
Collar / Skilled Trades plans whenever the contractor's own name happened
to contain a keyword from an ADJACENT occupational sector's profile
("Electric" -> Energy & Utilities, "Industrial" -> Manufacturing,
"Construction" -> Construction & Real Estate, "Freight" -> Transportation
& Logistics) -- because Trades is a cross-cutting occupational category
whose own keyword list ("welder", "electrician", "plumber", ...) almost
never appears in a real trade contractor's name. Fixed by
``app._BLUE_COLLAR_ADJACENT_LEGACY_KEYS``: a name hit that falls entirely
within that adjacency set now AGREES with a blue_collar_trades selection
instead of conflicting with it. A name that implies a genuinely unrelated
industry (a hospital name, a hotel chain) still conflicts as before.

Also covers the currency-symbol-preservation fix in budget-period
normalisation (monthly/quarterly/annual -> campaign total): that block
used to hardcode ``f"${scaled:,.0f}"`` regardless of the currency symbol
the client actually typed, so a GBP/EUR budget came out of normalisation
relabelled as USD. Fixed by ``app._budget_currency_prefix``.

Root-cause gate: every test below that asserts NEW behaviour was proven to
fail against the pre-fix code (throwaway ``git worktree add`` at the
parent commit, e6b1e34d) -- see the task report for the fail-then-pass
output. app.py's request handler is ~26k lines deep inside a single
method (not a standalone testable unit -- see test_app_duration_wiring.py's
own note on this), so the wiring itself (are both /api/generate call
sites actually calling the new shared helper, not a re-implemented local
copy) is verified by source inspection, the same approach
test_app_duration_wiring.py and test_app_vendor_gate_wiring.py already use
for this file.

Runs under pytest, or standalone:
``python3 tests/test_app_async_wiring_and_trades_adjacency_2026_09_24.py``.
"""

from __future__ import annotations

import logging
import sys
from pathlib import Path

import openpyxl
import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import app  # noqa: E402
import bundle_qa  # noqa: E402
import display_format  # noqa: E402


@pytest.fixture(autouse=True)
def _quiet():
    logging.disable(logging.CRITICAL)
    yield
    logging.disable(logging.NOTSET)


def _app_source() -> str:
    return (PROJECT_ROOT / "app.py").read_text()


def _async_worker_source(src: str) -> str:
    """Slice out just the ``_run_async_generate`` closure body (from its
    ``def`` through the outer ``t = threading.Thread(`` that starts it) so
    wiring assertions can't accidentally match the sync handler's copy."""
    start = src.index("def _run_async_generate(jid, gen_data, rid):")
    end = src.index("t = threading.Thread(\n", start)
    return src[start:end]


def _sync_handler_source(src: str, async_end: int) -> str:
    """The sync /api/generate handler body runs after the async branch's
    early ``return`` (the request either goes async and returns, or falls
    through to this code) -- slice from there through the next top-level
    handler definition."""
    start = src.index("# ── Process uploaded files", async_end)
    end = src.index("\n    def do_HEAD(self)", start)
    return src[start:end]


# ---------------------------------------------------------------------------
# C17: campaign_weeks -- single shared resolver, both paths
# ---------------------------------------------------------------------------
class TestCampaignWeeksSharedResolver:
    def test_resolve_campaign_weeks_one_month(self):
        data = {"campaign_duration": "1 month"}
        weeks = app._resolve_campaign_weeks(data)
        assert weeks == 4 == display_format.resolve_campaign_weeks("1 month")
        assert data["campaign_weeks"] == 4

    def test_resolve_campaign_weeks_two_weeks(self):
        data = {"campaign_duration": "2 weeks"}
        weeks = app._resolve_campaign_weeks(data)
        assert weeks == 2 == display_format.resolve_campaign_weeks("2 weeks")
        assert data["campaign_weeks"] == 2

    def test_resolve_campaign_weeks_falls_back_to_timeline_field(self):
        # The async worker's gen_data may carry the raw request field as
        # "timeline" rather than "campaign_duration" depending on caller;
        # the sync path's own pre-existing normalization already handles
        # this fallback (data.get("campaign_duration") or data.get("timeline")),
        # so the shared helper must too.
        data = {"timeline": "1 month"}
        assert app._resolve_campaign_weeks(data) == 4

    def test_deck_timeline_reads_the_resolved_weeks_not_the_12_week_default(self):
        """Directly exercises the exact lookup ppt_generator's deck-timeline
        builder makes (``data.get("campaign_weeks", 12)`` at
        ppt_generator.py ~7561): pre-fix, an async plan's gen_data never got
        campaign_weeks set at all, so this lookup silently returned the
        12-week default regardless of the plan's real duration. Post-fix,
        calling the shared helper first (exactly what both /api/generate
        paths now do) makes the deck timeline agree with the real duration.
        """
        # Pre-fix-equivalent gen_data: campaign_duration set, campaign_weeks
        # never populated (this IS the bug -- the async worker's gen_data
        # before this fix).
        buggy_gen_data = {"campaign_duration": "1 month"}
        assert buggy_gen_data.get("campaign_weeks", 12) == 12

        # Post-fix: the shared helper runs first, as both call sites now do.
        fixed_gen_data = {"campaign_duration": "1 month"}
        app._resolve_campaign_weeks(fixed_gen_data)
        assert fixed_gen_data.get("campaign_weeks", 12) == 4

    def test_both_generate_paths_call_the_shared_resolver_not_a_local_copy(self):
        src = _app_source()
        async_src = _async_worker_source(src)
        assert "_resolve_campaign_weeks(gen_data)" in async_src, (
            "the async worker (_run_async_generate) must call the shared "
            "app._resolve_campaign_weeks helper -- it used to skip "
            "campaign_weeks resolution entirely"
        )
        assert "display_format.resolve_campaign_weeks(" not in async_src, (
            "the async worker must not re-implement its own duration->weeks "
            "ladder inline -- that reintroduces the drift class this fix "
            "closes (single shared resolver only)"
        )

        async_end = src.index("t = threading.Thread(\n", src.index(
            "def _run_async_generate(jid, gen_data, rid):"
        ))
        sync_src = _sync_handler_source(src, async_end)
        assert "campaign_weeks = _resolve_campaign_weeks(data)" in sync_src, (
            "the sync /api/generate handler must call the shared "
            "app._resolve_campaign_weeks helper, not a re-implemented "
            "inline ladder"
        )


# ---------------------------------------------------------------------------
# C17 (gate level): bundle_qa's campaign_duration_incoherence rule
# ---------------------------------------------------------------------------
def _exec_summary_workbook(duration_label: str) -> openpyxl.Workbook:
    """Minimal workbook matching the exact KV layout
    bundle_qa._check_campaign_duration_incoherence reads: the "Duration"
    label cell, with its VALUE one row above it in the same column (see
    excel_v2's own Executive Summary KV block, which this mirrors)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    ws["B60"] = duration_label
    ws["B61"] = "Duration"
    return wb


class TestCampaignDurationIncoherenceGate:
    def test_missing_campaign_weeks_trips_the_incoherence_gate(self):
        """Pre-fix reproduction: campaign_weeks never resolved -> deck
        timeline phases default to a fixed 12-week span (last phase 'Weeks
        7-12') while the workbook's Duration cell correctly says '1 month'
        -- an 8-week spread, over the 3-week tolerance."""
        wb = _exec_summary_workbook("1 month")
        units = [
            bundle_qa._TextUnit("Weeks 1-2", "Deck!slide6", top=1, left=1),
            bundle_qa._TextUnit("Weeks 3-6", "Deck!slide6", top=1, left=2),
            bundle_qa._TextUnit("Weeks 7-12", "Deck!slide6", top=1, left=3),
        ]
        findings: list = []
        bundle_qa._check_campaign_duration_incoherence(units, wb, findings)
        codes = [f["code"] for f in findings]
        assert "campaign_duration_incoherence" in codes, findings

    def test_resolved_campaign_weeks_closes_the_incoherence_gate(self):
        """Post-fix: campaign_weeks resolved to 4 for '1 month' (both
        paths, via the shared helper) -> the deck timeline's last phase is
        'Weeks 4-4' (ppt_generator's own cw<=12 branch: p3_start=min(7,cw),
        p3_end=cw), which agrees with the workbook's '1 month' Duration
        cell -- zero spread."""
        data = {"campaign_duration": "1 month"}
        cw = app._resolve_campaign_weeks(data)
        assert cw == 4
        p3_start = min(7, cw)
        p3_end = cw

        wb = _exec_summary_workbook("1 month")
        units = [
            bundle_qa._TextUnit("Weeks 1-2", "Deck!slide6", top=1, left=1),
            bundle_qa._TextUnit(f"Weeks 3-{min(6, cw)}", "Deck!slide6", top=1, left=2),
            bundle_qa._TextUnit(
                f"Weeks {p3_start}-{p3_end}", "Deck!slide6", top=1, left=3
            ),
        ]
        findings: list = []
        bundle_qa._check_campaign_duration_incoherence(units, wb, findings)
        codes = [f["code"] for f in findings]
        assert "campaign_duration_incoherence" not in codes, findings

    def test_two_weeks_duration_end_to_end(self):
        data = {"campaign_duration": "2 weeks"}
        cw = app._resolve_campaign_weeks(data)
        assert cw == 2
        p3_start = min(7, cw)
        p3_end = cw
        assert (p3_start, p3_end) == (2, 2)

        wb = _exec_summary_workbook("2 weeks")
        units = [
            bundle_qa._TextUnit(f"Weeks {p3_start}-{p3_end}", "Deck!slide6", top=1, left=1),
        ]
        findings: list = []
        bundle_qa._check_campaign_duration_incoherence(units, wb, findings)
        assert not any(
            f["code"] == "campaign_duration_incoherence" for f in findings
        ), findings


# ---------------------------------------------------------------------------
# C5: channel_categories honoured before calculate_budget_allocation
# ---------------------------------------------------------------------------
_ALLOC = {
    "programmatic_dsp": 35,
    "global_boards": 20,
    "niche_boards": 15,
    "social_media": 12,
    "regional_boards": 8,
    "employer_branding": 5,
    "apac_regional": 3,
    "emea_regional": 2,
}


class TestChannelSelectionHonoured:
    def test_disabled_channels_are_zeroed_and_the_rest_renormalised(self):
        data = {
            "channel_categories": {
                "niche_boards": False,
                "social_media": False,
                "programmatic_dsp": True,
            }
        }
        out = app._apply_channel_selection(dict(_ALLOC), data)
        assert "niche_boards" not in out
        assert "social_media" not in out
        assert set(out) == set(_ALLOC) - {"niche_boards", "social_media"}
        assert sum(out.values()) == pytest.approx(100.0)
        # Proportions among the surviving channels are preserved.
        assert out["programmatic_dsp"] > out["global_boards"] > out["regional_boards"]

    def test_absent_channel_categories_is_unchanged_behaviour(self):
        out = app._apply_channel_selection(dict(_ALLOC), {})
        assert out == _ALLOC
        out2 = app._apply_channel_selection(dict(_ALLOC), {"channel_categories": {}})
        assert out2 == _ALLOC
        out3 = app._apply_channel_selection(dict(_ALLOC), {"channel_categories": None})
        assert out3 == _ALLOC

    def test_disabling_every_category_falls_back_to_original_allocation(self):
        """Guard: a plan that (incorrectly, or via a legacy client) disables
        every category must not hand calculate_budget_allocation an
        empty/zero-sum channel split."""
        data = {"channel_categories": {k: False for k in _ALLOC}}
        out = app._apply_channel_selection(dict(_ALLOC), data)
        assert out == _ALLOC

    def test_both_generate_paths_apply_channel_selection_before_the_call(self):
        src = _app_source()
        async_src = _async_worker_source(src)
        assert "_apply_channel_selection(\n                                        channel_pcts, gen_data\n                                    )" in async_src or (
            "_apply_channel_selection(" in async_src
            and "channel_pcts, gen_data" in async_src
        ), "the async worker must filter channel_pcts through _apply_channel_selection before calculate_budget_allocation"

        async_end = src.index(
            "t = threading.Thread(\n",
            src.index("def _run_async_generate(jid, gen_data, rid):"),
        )
        sync_src = _sync_handler_source(src, async_end)
        assert "_apply_channel_selection(" in sync_src and "channel_pcts, data" in sync_src, (
            "the sync /api/generate handler must filter channel_pcts through "
            "_apply_channel_selection before calculate_budget_allocation"
        )


# ---------------------------------------------------------------------------
# C8: plan_currency resolved once, passed at every call site
# ---------------------------------------------------------------------------
class TestPlanCurrencyResolvedAndPassed:
    def test_explicit_currency_code_wins(self):
        assert app._resolve_plan_currency({"currency_code": "GBP"}) == "GBP"

    def test_market_based_resolution_matches_shared_resolver(self):
        import plan_currency as pc

        data = {"locations": ["London, United Kingdom"]}
        assert app._resolve_plan_currency(data) == pc.currency_for_plan_with_basis(data)[0]

    def test_default_never_raises_and_returns_none_or_usd_basis(self):
        # currency_for_plan_with_basis returns ("USD", "default") when
        # nothing resolves; app._resolve_plan_currency passes that through.
        assert app._resolve_plan_currency({}) == "USD"

    def test_every_calculate_budget_allocation_call_site_passes_plan_currency(self):
        src = _app_source()
        call_sites = [
            m.start() for m in __import__("re").finditer(r"calculate_budget_allocation\(", src)
        ]
        assert len(call_sites) == 3, (
            f"expected exactly 3 calculate_budget_allocation call sites in "
            f"app.py, found {len(call_sites)} -- update this test's site "
            f"count (and check the new site passes plan_currency=) if a "
            f"site was legitimately added/removed"
        )
        for start in call_sites:
            # The call's own closing paren is well within a few hundred
            # characters for all three sites (checked by hand); slicing a
            # generous window avoids needing a full paren-matcher.
            window = src[start : start + 1200]
            assert "plan_currency=" in window, (
                "a calculate_budget_allocation call site is missing "
                f"plan_currency=...: {window[:200]}..."
            )


# ---------------------------------------------------------------------------
# C3: Blue Collar / Skilled Trades adjacency
# ---------------------------------------------------------------------------
class TestBlueCollarTradesAdjacency:
    @pytest.mark.parametrize(
        "company,roles",
        [
            ("Sparks Electric", ["Electrician"]),
            ("Mister Sparky Electric", ["Electrician"]),
            ("Summit Industrial Group", ["Welder", "Machinist"]),
            ("RoadRunner Freight Lines", ["CDL-A Truck Driver"]),
            ("Allied Construction Services", ["Carpenter"]),
            ("Sunrun Solar", ["Solar Installer"]),
        ],
    )
    def test_adjacent_sector_names_agree_with_blue_collar_trades(self, company, roles):
        result = app._infer_industry_from_signals(company, roles, "blue_collar_trades")
        assert result is None, (
            f"{company!r} on blue_collar_trades should AGREE (adjacent "
            f"occupational sector), not conflict: {result}"
        )

    def test_unrelated_industry_name_still_conflicts_on_blue_collar_trades(self):
        """The real cross-category conflict case: a hospital's name
        implies an unrelated industry even when the plan is (incorrectly)
        marked blue_collar_trades -- this must still be CRITICAL."""
        result = app._infer_industry_from_signals(
            "Mercy General Hospital", ["Warehouse Associate"], "blue_collar_trades"
        )
        assert result is not None
        assert result.get("legacy_key") == "healthcare_medical"
        assert result.get("_inferred_signal") == "client_name"

    def test_uber_cab_driver_still_conflicts_on_hospitality(self):
        """Unrelated to the blue-collar adjacency fix: the original Uber
        incident (rideshare-implying name/role on a Hospitality & Travel
        selection) must remain unaffected."""
        result = app._infer_industry_from_signals(
            "Uber", ["commercial cab driver"], "hospitality_travel"
        )
        assert result is not None
        assert result.get("_inferred_signal") == "client_name"

    def test_adjacency_is_scoped_to_blue_collar_trades_selection_only(self):
        """The same 'Sparks Electric' name against a DIFFERENT, non-adjacent
        selection must still conflict -- the adjacency set only widens
        agreement when blue_collar_trades itself is the selection."""
        result = app._infer_industry_from_signals(
            "Sparks Electric", ["Electrician"], "hospitality_travel"
        )
        assert result is not None
        assert result.get("_inferred_signal") == "client_name"


# ---------------------------------------------------------------------------
# Budget-period normalisation: currency symbol must survive
# ---------------------------------------------------------------------------
class TestBudgetPeriodCurrencySymbolSurvives:
    @pytest.mark.parametrize(
        "raw,expected_prefix",
        [
            ("£10,000", "£"),
            ("£10,000 per month", "£"),
            ("€5,000", "€"),
            ("€5.000 monthly", "€"),
            ("A$400,000", "A$"),
            ("10,000", "$"),
            ("$10,000", "$"),
        ],
    )
    def test_budget_currency_prefix(self, raw, expected_prefix):
        assert app._budget_currency_prefix(raw) == expected_prefix

    def test_gbp_monthly_budget_scales_without_becoming_usd(self):
        # £10,000/month over a 3-month campaign -> £30,000, not $30,000.
        prefix = app._budget_currency_prefix("£10,000")
        scaled = app.parse_budget("£10,000") * 3
        assert f"{prefix}{scaled:,.0f}" == "£30,000"

    def test_eur_comma_budget_scales_without_becoming_usd(self):
        prefix = app._budget_currency_prefix("€5,000")
        scaled = app.parse_budget("€5,000") * 3
        assert f"{prefix}{scaled:,.0f}" == "€15,000"

    def test_budget_period_block_uses_the_shared_prefix_helper(self):
        src = _app_source()
        start = src.index(
            "# ── Budget period normalization (monthly/quarterly/annual"
        )
        end = src.index("# ── Gold Standard: Campaign start month validation", start)
        block = src[start:end]
        assert "_budget_currency_prefix(" in block
        assert 'f"${_scaled:,.0f}"' not in block, (
            "the budget-period block must not hardcode a '$' prefix when "
            "formatting the scaled budget"
        )


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
