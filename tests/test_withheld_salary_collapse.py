"""Design-judge 2026-10-01: withheld salaries must not look like a broken
template.

Shipped defect: a Japan plan's Role Breakdown slide showed a column of
"Local salary data n/a" under a subtitle promising a "salary band", and the
workbook's Quality Intelligence sheet carried ~50-110 "n/a" cells. When
EVERY salary row of a table is withheld, the salary column/table is dropped
and one plain sentence says why; the slide-7 subtitle never promises a
band that isn't there. A partially withheld table keeps its n/a cells (and
its header, which then has data under it).
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from pptx import Presentation

import data_synthesizer
import excel_v2
import gold_standard as gs
import plan_geo
import ppt_generator
import research as research_mod
import tools_regen_bundles as T


def _bundle(roles, locations, budget, industry, synthesize=False):
    brief = {
        "client_name": "Hoshino Electronics",
        "budget": budget,
        "campaign_duration": "3-6 months",
        "industry": industry,
        "locations": list(locations),
        "roles": list(roles),
        "target_roles": [{"title": r, "count": 5} for r in roles],
    }
    data = T.build_plan_data(brief)
    if synthesize:
        data["_synthesized"] = data_synthesizer.synthesize({}, {}, data)
        data["_gold_standard"] = gs.apply_all_quality_gates(data)
    xlsx = excel_v2.generate_excel_v2(dict(data), research_mod=research_mod)
    pptx = ppt_generator.generate_pptx(dict(data))
    xlsx = xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx
    pptx = pptx.getvalue() if hasattr(pptx, "getvalue") else pptx
    return {"xlsx": xlsx, "pptx": pptx}


_JP_ROLES = ["Embedded Software Engineer", "Firmware Engineer", "Hardware Design Engineer", "Product Manager"]
# Partially withheld: two roles with an audited Irish band, two without.
_IE_ROLES = ["Clinical Nurse Specialist", "Software Engineer", "QA Engineer", "Business Analyst"]
_US_ROLES = ["Registered Nurse", "Certified Nursing Assistant", "Pharmacist", "ICU Nurse"]


@pytest.fixture(scope="module")
def japan():
    return _bundle(_JP_ROLES, ["Tokyo, Japan", "Osaka, Japan"], "¥45,000,000", "technology", True)


@pytest.fixture(scope="module")
def ireland():
    return _bundle(_IE_ROLES, ["Dublin, Ireland"], "€400,000", "technology")


@pytest.fixture(scope="module")
def us_plan():
    return _bundle(_US_ROLES, ["Houston, TX", "Dallas, TX"], "$250,000", "healthcare")


def _role_slide(pptx: bytes) -> list[str]:
    for slide in Presentation(io.BytesIO(pptx)).slides:
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        if "Role Breakdown" in texts:
            return texts
    raise AssertionError("Role Breakdown slide not rendered")


def _sheet_text(xlsx: bytes, name: str) -> list[str]:
    ws = load_workbook(io.BytesIO(xlsx))[name]
    return [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]


def test_market_names():
    data = {"locations": ["Bengaluru, Karnataka, India", "London, UK", "Columbus, OH", "Pune, India"]}
    assert plan_geo.non_us_market_names(data) == ["India", "UK"]
    assert plan_geo.join_market_names(["India", "UK", "Japan"]) == "India, UK and Japan"
    assert plan_geo.non_us_market_names({"locations": ["new zealand", "uk"]}) == [
        "New Zealand",
        "UK",
    ]


def test_all_withheld_slide7_drops_salary_column_and_says_so(japan):
    texts = _role_slide(japan["pptx"])
    assert "Est. Median Salary" not in texts
    assert "Local salary data n/a" not in texts
    assert not any("salary band" in t for t in texts)
    assert any(t.startswith("Tier, difficulty, and channel emphasis") for t in texts)
    assert (
        "Role-level salary benchmarks are not available for Japan in our sourced "
        "data; US salary figures are not used for non-US markets."
    ) in texts


def _slide2(pptx: bytes) -> str:
    slide = Presentation(io.BytesIO(pptx)).slides[1]
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_slide2_omits_an_na_salary_line(japan):
    """Design-judge round 2: "Salary Range: Local salary data n/a" said
    nothing -- the item is omitted; the card's other items flow up."""
    text = _slide2(japan["pptx"])
    assert "Salary Range" not in text
    assert "Local salary data n/a" not in text
    assert "Budget:" in text and "Target Roles:" in text


def test_slide2_keeps_a_sourced_salary_line(ireland):
    assert "Salary Range:" in _slide2(ireland["pptx"])


def test_partially_withheld_slide7_keeps_the_column(ireland):
    texts = _role_slide(ireland["pptx"])
    assert "Salary Benchmark" in texts
    assert "€55K-€65K band" in texts and "€51K median" in texts
    assert "Local salary data n/a" in texts
    assert any(t.startswith("Tier, salary band, and channel emphasis") for t in texts)


def test_all_withheld_workbook_tables_collapse(japan):
    qi = _sheet_text(japan["xlsx"], "Quality Intelligence")
    assert "n/a" not in qi and "Local salary data n/a" not in qi
    assert "Estimated Salary" not in qi and "Salary Range" not in qi
    assert "Salary Multiplier" not in qi and "Median Salary" not in qi
    assert any(t.startswith("Role-level salary benchmarks are not available for Japan") for t in qi)
    mi = _sheet_text(japan["xlsx"], "Market Intelligence")
    assert "Local salary data n/a" not in mi and "Median" not in mi
    assert any(t.startswith("Role-level salary benchmarks are not available for Japan") for t in mi)


def test_partially_withheld_workbook_keeps_role_rows(ireland):
    qi = _sheet_text(ireland["xlsx"], "Quality Intelligence")
    assert "Median Salary" in qi  # header with data under it
    assert "n/a" in qi  # the withheld roles' cells
    # City salary columns: every city's estimate is withheld -> dropped.
    assert "Estimated Salary" not in qi


def test_us_plan_tables_unchanged(us_plan):
    texts = _role_slide(us_plan["pptx"])
    assert "Est. Median Salary" in texts
    assert any(t.startswith("Tier, salary band, and channel emphasis") for t in texts)
    qi = _sheet_text(us_plan["xlsx"], "Quality Intelligence")
    for header in ("Salary Multiplier", "Estimated Salary", "Salary Range", "Median Salary"):
        assert header in qi
    assert not any(t.startswith("Role-level salary benchmarks") for t in qi)
