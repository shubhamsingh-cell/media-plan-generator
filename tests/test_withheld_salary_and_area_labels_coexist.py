"""Two workbook features that touch the same tables must both survive:

* mpg-salary-geo-correctness: a non-US market's salary is withheld; a table
  whose every salary is withheld drops its salary columns / rows and says so
  in one sentence (tests/test_withheld_salary_collapse.py).
* mpg-data-sources: area-level figures are labelled -- Quality Intelligence
  rows that fell back to their state carry " [PA state-level estimate]", the
  Location Intelligence population reads "(country-level, ...; not city)",
  and a wrapped merged footnote gets an explicit row height
  (tests/test_location_labelling.py).

They met in one conflict (the Quality Intelligence city row); this pins the
resolution: the area label lives on the Market cell, which the salary-column
collapse always keeps, and every collapse sentence goes through
``_write_footnote`` so it gets the same row height as any other footnote.
"""

from __future__ import annotations

import io
from typing import Any, Dict

import openpyxl
import pytest

import api_enrichment
import budget_engine
import data_synthesizer
import excel_v2
import gold_standard
import public_data_sources as pds
from tests.public_data_fakes import FakeNet

_WITHHELD = "salary benchmarks are not available for"


@pytest.fixture
def net(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(api_enrichment, "_get_cached", lambda key: None)
    monkeypatch.setattr(api_enrichment, "_set_cached", lambda key, data: None)
    monkeypatch.delenv("CENSUS_API_KEY", raising=False)
    pds.clear_caches()
    fake = FakeNet().install(monkeypatch)
    yield fake
    pds.clear_caches()


def _workbook(locations: list) -> openpyxl.Workbook:
    roles = [{"title": "Financial Analyst", "count": 4, "tier": "mid"}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=150_000,
        roles=roles,
        locations=locations,
        industry="finance_banking",
        channel_percentages={"programmatic_dsp": 40, "global_boards": 40, "social_media": 20},
    )
    enriched: Dict[str, Any] = {
        "location_demographics": api_enrichment.fetch_location_demographics(locations)
    }
    data: Dict[str, Any] = {
        "client_name": "Coexist Fixture Bank",
        "industry": "finance_banking",
        "budget": "150000",
        "campaign_duration": "3 months",
        "campaign_weeks": 13,
        "roles": [r["title"] for r in roles],
        "target_roles": roles,
        "locations": locations,
        "_budget_allocation": alloc,
        "_enriched": enriched,
    }
    # Full synthesis (as production): location_profiles carry the labelled
    # area figures, salary_intelligence feeds the Market Intelligence table.
    data["_synthesized"] = data_synthesizer.synthesize(enriched, {}, data)
    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    return openpyxl.load_workbook(io.BytesIO(excel_v2.generate_excel_v2(data)))


def _texts(ws) -> list[str]:
    return [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]


def test_mixed_plan_keeps_salary_columns_and_the_state_level_label(net: FakeNet) -> None:
    qi = _texts(_workbook(["Hershey, PA", "London, UK"])["Quality Intelligence"])
    # Not every market is withheld -> the salary columns stay ...
    assert "Estimated Salary" in qi and "Salary Range" in qi
    # ... the US town's state fallback is still labelled, in both city tables ...
    labelled = [t for t in qi if t.startswith("Hershey [PA state-level estimate]")]
    assert len(labelled) >= 2, labelled
    assert any("share identical figures" in t for t in qi)
    # ... and the non-US market's salary is still withheld, not a US figure.
    assert "n/a" in qi and "Local salary data n/a" in qi


def test_all_withheld_plan_collapses_salary_and_keeps_area_labels(net: FakeNet) -> None:
    wb = _workbook(["London, UK", "Manchester, UK"])
    qi_ws, mi_ws = wb["Quality Intelligence"], wb["Market Intelligence"]
    qi, mi = _texts(qi_ws), _texts(mi_ws)

    # Salary side: columns/tables collapsed to one sentence each.
    for header in ("Estimated Salary", "Salary Range", "Salary Multiplier", "Median Salary"):
        assert header not in qi
    assert any(_WITHHELD in t for t in qi) and any(_WITHHELD in t for t in mi)

    # Data-sources side: population is labelled country-level, never the city's.
    population_cells = [t for t in mi if "country-level" in t and "not city" in t]
    assert len(population_cells) >= 2, population_cells

    # Every collapse sentence (and the withheld city-table footnote) is a
    # merged B..H footnote with the shared helper's row height.
    for ws in (qi_ws, mi_ws):
        for r in ws.iter_rows():
            for c in r:
                if not (isinstance(c.value, str) and "not available for" in c.value):
                    continue
                assert any(
                    m.min_row == c.row == m.max_row and m.min_col == 2 and m.max_col == 8
                    for m in ws.merged_cells.ranges
                ), (ws.title, c.coordinate)
                expected = excel_v2._footnote_row_height(ws, c.value)
                if expected is not None:
                    assert (ws.row_dimensions[c.row].height or 0) >= expected, (
                        ws.title,
                        c.coordinate,
                        len(c.value),
                    )
