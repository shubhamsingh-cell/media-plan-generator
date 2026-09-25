"""Regression tests (Hershey follow-up #3): the workbook forecast must agree
with the SAME plan's Activation Event Calendar at EVERY campaign duration,
not just the 3-month 90-Day frame.

Client-reported defect (Hershey Company): the forecast narrative claimed the
heaviest-spend period was a month the same workbook's Activation Event
Calendar calls "Low". Two earlier fixes (b802d68, 01f8d6c) repaired only the
3-month 90-Day path. Short plans (<=13 weeks) ran through a different path:
they took the 3-month split for the start month AND THE TWO FOLLOWING
MONTHS, stretched those three phases across the campaign's weeks with
display_format.scale_week_phases, then re-bucketed the weeks into columns
labelled with calendar-month names. On 01f8d6c that meant:

* manufacturing, June start, 5 weeks: the "July 2026" column (really weeks
  4-5, mostly June days, priced at July's weight) got 63% "(heaviest planned
  spend)"; June (Moderate, 1.0x) got 37% "(lightest)" -- while claiming to
  match the Activation Event Calendar, where July is Low (0.7x);
* manufacturing, November start, 8 weeks: December (Low) crowned heaviest;
* a 4-week plan (always inside ONE calendar month) was priced partly at the
  weights of the NEXT two months -- a flat calendar still rendered 33/17/17/33
  with heaviest/lightest labels;
* a malformed timeline (entries without a ``month`` key) fell back to generic
  math but still printed "matching this plan's Activation Event Calendar",
  because the gate checked only "the timeline list is non-empty".

The fix (excel_v2._short_plan_forecast_split) weights every campaign DAY by
the budget weight of the calendar month it falls in, and makes each
month-framed column hold exactly that month's campaign days. A column's
share therefore reflects BOTH its month's pacing and its day count, so for
short plans the invariant that must hold is on DAILY PACING (share / campaign
days), and the narrative ranks pacing (never raw share) and prints each
column's day count and calendar weight. The calendar claim is gated on an
explicit ``used_real_calendar`` returned by the function that computed the
split.

The sweep below is black-box: it renders the real forecast sheet with
excel_v2._build_sheet_rolling_forecast and judges only what a reader sees
(column headers, the Spend row, the "Forecast Period" row, the footnote),
so it runs unchanged against any commit.
"""

from __future__ import annotations

import calendar as _calendar
import datetime
import io
import re
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402

import budget_engine  # noqa: E402
import excel_v2  # noqa: E402
import gold_standard  # noqa: E402

_INDUSTRIES = [
    "manufacturing",
    "retail_consumer",
    "healthcare_medical",
    "tech_engineering",
    "hospitality_travel",
    "construction_real_estate",
    "finance_banking",
    "logistics_supply_chain",
    "education",
    "government",
    "food_beverage",
    "automotive",
    "insurance",
    "telecommunications",
    "restaurant",
]
# Every short-plan duration (weekly frame 1-4, 2-month frame 5-8, 3/4-month
# frame 9-13) plus two 90-Day-frame durations.
_DURATIONS = list(range(1, 14)) + [26, 78]

_CAL_CLAIM = "Activation Event Calendar"


# ---------------------------------------------------------------------------
# Harness
# ---------------------------------------------------------------------------


def _calendar_for(industry: str, start_month: int, **extra) -> list[dict]:
    data = {
        "client_name": "Sweep Co",
        "industry": industry,
        "campaign_start_month": start_month,
        "roles": ["Generic Role"],
    }
    data.update(extra)
    return gold_standard.build_activation_calendar(data)["timeline"]


def _render(industry: str, start_month: int, weeks: int, timeline) -> "openpyxl.worksheet.worksheet.Worksheet":
    data = {
        "client_name": "Sweep Co",
        "industry": industry,
        "campaign_start_month": start_month,
        "campaign_weeks": weeks,
        "campaign_duration": f"{weeks} weeks",
        "roles": ["Generic Role"],
        "_budget_allocation": {
            "metadata": {"total_budget": 100_000.0},
            "total_projected": {"applications": 1000, "hires": 50},
            "channel_allocations": {},
        },
        "_gold_standard": {"activation_calendar": {"timeline": timeline}},
    }
    ws = openpyxl.Workbook().active
    excel_v2._build_sheet_rolling_forecast(ws, data)
    return ws


def _parse(ws) -> dict:
    rows = list(ws.iter_rows(min_col=2, values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Metric")
    hdr = [v for v in rows[hi] if v is not None]
    cols = [str(c) for c in hdr[1:-2]]
    spend = next(r for r in rows[hi + 1 :] if r and r[0] == "Spend")
    vals = [v for v in spend[1:] if isinstance(v, (int, float))]
    total = vals[len(cols)]
    shares = [v / total for v in vals[: len(cols)]]
    period = next(r for r in rows if r and r[0] == "Forecast Period")
    pv = next(v for v in period[1:] if isinstance(v, str))
    start, end = (
        datetime.datetime.strptime(x.strip(), "%b %d, %Y").date()
        for x in pv.split(" - ")
    )
    note = next(
        c.value
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and "phases budget" in c.value
    )
    return {
        "title": ws.title,
        "cols": cols,
        "shares": shares,
        "start": start,
        "end": end,
        "note": note,
    }


def _column_days(label: str, start: datetime.date, end: datetime.date) -> list:
    """The campaign days a reader would take a column to cover: "Week k" is
    days 7(k-1)..7k-1 of the Forecast Period; "June 2026" is June's days
    inside the Forecast Period."""
    m = re.fullmatch(r"Week (\d+)", label)
    if m:
        k = int(m.group(1))
        d0 = start + datetime.timedelta(days=7 * (k - 1))
        return [d0 + datetime.timedelta(days=i) for i in range(7)]
    first = datetime.datetime.strptime(label, "%B %Y").date()
    last = first.replace(day=_calendar.monthrange(first.year, first.month)[1])
    lo, hi = max(start, first), min(end, last)
    return [lo + datetime.timedelta(days=i) for i in range((hi - lo).days + 1)]


def _note_labels(note: str, cols: list[str]) -> dict[int, str]:
    """Column index -> the parenthetical the note attaches to it. Handles
    both "NN% Month k (...)" (90-Day frame / old short frame) and
    "NN% June 2026 (...)" / "NN% Week k (...)"."""
    out: dict[int, str] = {}
    for m in re.finditer(
        r"(\d+)% (Month \d+|Week \d+|[A-Z][a-z]+ \d{4}) \(([^)]*)\)", note
    ):
        col, text = m.group(2), m.group(3).lower()
        mm = re.fullmatch(r"Month (\d+)", col)
        if mm:
            out[int(mm.group(1)) - 1] = text
        elif col in cols:
            out[cols.index(col)] = text
    return out


def _violations(parsed: dict, timeline: list[dict]) -> list[tuple]:
    """Reader-grounded invariants against this plan's own calendar."""
    w = {e["month"]: e["budget_weight"] for e in timeline}
    inten = {e["month"]: e["hiring_intensity"] for e in timeline}
    cols, shares, note = parsed["cols"], parsed["shares"], parsed["note"]
    is90 = parsed["title"] == "90-Day Forecast"
    days = [_column_days(c, parsed["start"], parsed["end"]) for c in cols]
    out: list[tuple] = []
    if any(not ds for ds in days):
        return [("column_with_no_campaign_days", cols)]
    n = len(cols)
    cw = [sum(w[d.month] for d in ds) / len(ds) for ds in days]
    # 90-Day frame: three full months, shares proportional to weight.
    # Short frames: columns hold different day counts -> judge daily pacing.
    metric = shares if is90 else [s / len(ds) for s, ds in zip(shares, days)]
    for i in range(n):
        for j in range(n):
            if cw[i] > cw[j] + 1e-9 and metric[i] < metric[j] - 1e-9:
                out.append(("pacing_inversion", cols[i], cw[i], cols[j], cw[j]))
    if not is90:
        # Every share must be exactly the calendar-weighted day count of the
        # column's own days -- no month outside the window can leak in.
        tot = sum(w[d.month] for ds in days for d in ds)
        for c, s, ds in zip(cols, shares, days):
            exp = sum(w[d.month] for d in ds) / tot
            if abs(exp - s) > 0.0006:
                out.append(("share_not_from_own_days", c, round(s, 4), round(exp, 4)))

    # Labels: short-plan notes print the calendar-table-formatted weight
    # ("0.7x") beside each label, so they are judged on that displayed
    # value; 90-Day notes print no weight and rank share, so raw weight.
    def shown(i: int) -> float:
        months = {d.month for d in days[i]}
        if not is90 and len(months) == 1:
            return float(f"{w[months.pop()]:.1f}")
        return round(cw[i], 6)

    dw = [shown(i) for i in range(n)]
    for i, text in _note_labels(note, cols).items():
        heavy = "heaviest" in text or "heavier" in text
        light = "lightest" in text or "lighter" in text
        for j in range(n):
            if heavy and dw[j] > dw[i]:
                out.append(("heavy_label_on_lower_weight", cols[i], cols[j]))
            if light and dw[i] > dw[j]:
                out.append(("light_label_on_higher_weight", cols[i], cols[j]))
        if (
            heavy
            and {inten[d.month] for d in days[i]} == {"low"}
            and any(inten[d.month] != "low" for ds in days for d in ds)
        ):
            out.append(("calendar_low_month_labeled_heavy", cols[i]))
    if max(dw) == min(dw) and re.search(r"heavi|light", note, re.I):
        out.append(("ranking_claim_without_weight_difference", note))
    if not is90 and "heaviest planned spend" in note:
        # Short-plan columns differ in length: a total-share "heaviest"
        # claim is not a seasonal statement and must never be made.
        out.append(("share_rank_claim_on_unequal_columns", note))
    return out


# ---------------------------------------------------------------------------
# 1. The exhaustive sweep: 15 industries x 12 start months x 15 durations.
# ---------------------------------------------------------------------------


def test_forecast_agrees_with_calendar_at_every_duration():
    failures = []
    per_duration: dict[int, int] = {}
    for weeks in _DURATIONS:
        for industry in _INDUSTRIES:
            for start_month in range(1, 13):
                timeline = _calendar_for(industry, start_month)
                parsed = _parse(_render(industry, start_month, weeks, timeline))
                v = _violations(parsed, timeline)
                if _CAL_CLAIM not in parsed["note"]:
                    v.append(("calendar_claim_missing_despite_full_calendar",))
                if v:
                    per_duration[weeks] = per_duration.get(weeks, 0) + 1
                    failures.append((industry, start_month, weeks, v[:3]))
    assert not failures, (
        f"{len(failures)} of {len(_DURATIONS) * len(_INDUSTRIES) * 12} cases "
        f"contradict the calendar; failing cases per duration: {per_duration}; "
        f"first: {failures[:3]}"
    )


def test_subvertical_calendar_short_plans_agree_with_calendar():
    """A sub-vertical override (propane/heating-fuel inside logistics) has the
    OPPOSITE seasonality of its parent industry -- the forecast must follow
    the plan's own (overridden) calendar at every short duration too."""
    extra = {"client_name": "Propane Unlimited Co", "roles": ["Propane CDL Driver"]}
    failures = []
    for weeks in _DURATIONS:
        for start_month in range(1, 13):
            timeline = _calendar_for("logistics_supply_chain", start_month, **extra)
            assert all(
                str(e["seasonal_phase"]).startswith("subvertical:") for e in timeline
            ), "propane sub-vertical override did not trigger -- test is vacuous"
            parsed = _parse(
                _render("logistics_supply_chain", start_month, weeks, timeline)
            )
            v = _violations(parsed, timeline)
            if v:
                failures.append((start_month, weeks, v[:2]))
    assert not failures, failures[:5]


# ---------------------------------------------------------------------------
# 2. The verifier's named failure cases.
# ---------------------------------------------------------------------------


def test_manufacturing_june_5_weeks_july_low_is_not_heaviest():
    timeline = _calendar_for("manufacturing", 6)
    w = {e["month"]: e["budget_weight"] for e in timeline}
    inten = {e["month"]: e["hiring_intensity"] for e in timeline}
    # S50 fix (seasonal key-mismatch): "manufacturing" now correctly matches
    # its own seasonal_hiring_trends.json pattern (previously it resolved to
    # _get_industry_key() bucket "blue_collar_trades", which never matched
    # any JSON key, so July's word came only from the generic
    # _HIRING_EVENTS_CALENDAR and happened to read "low"). July isn't a
    # peak/low month in the JSON manufacturing pattern, so its budget_weight
    # is now blended toward the neutral 1.0 seasonal multiplier and the word
    # is "moderate" -- the repro's actual invariant (July must stay lighter
    # than June, never the heaviest month) still holds on the weights.
    assert inten[7] == "moderate" and w[7] < w[6], (w[6], w[7])  # repro is real

    p = _parse(_render("manufacturing", 6, 5, timeline))
    assert [c.split()[0] for c in p["cols"]] == ["June", "July"], p["cols"]
    # June 1-30 + July 1-5: each column holds its own month's days only.
    exp_june = 30 * w[6] / (30 * w[6] + 5 * w[7])
    assert abs(p["shares"][0] - exp_june) < 0.0006, (p["shares"], exp_june)
    assert p["shares"][0] > p["shares"][1]
    note = p["note"]
    assert "heaviest planned spend" not in note, note
    # 0.7x -> 0.8x: July's budget_weight is now 0.85 (was a generic 0.7 pre-
    # fix), since it's blended against the real manufacturing seasonal
    # pattern's neutral 1.0x instead of never being blended at all.
    assert re.search(r"July \d{4} \(5 campaign days; 0\.8x, lighter daily pacing\)", note), note
    assert re.search(r"June \d{4} \(30 campaign days; 1\.0x, heavier daily pacing\)", note), note
    assert "matching this plan's Activation Event Calendar" in note
    assert _violations(p, timeline) == []


def test_manufacturing_november_8_weeks_december_low_is_not_heaviest():
    timeline = _calendar_for("manufacturing", 11)
    inten = {e["month"]: e["hiring_intensity"] for e in timeline}
    assert inten[12] == "low"

    p = _parse(_render("manufacturing", 11, 8, timeline))
    assert [c.split()[0] for c in p["cols"]] == ["November", "December"]
    labels = _note_labels(p["note"], p["cols"])
    assert "lighter daily pacing" in labels[1], p["note"]
    assert "heavier daily pacing" in labels[0], p["note"]
    assert "heaviest" not in labels[1] and "heavier" not in labels[1]
    assert _violations(p, timeline) == []


def test_months_after_the_campaign_ends_cannot_move_the_split():
    """A 4-week June plan runs June 1-28 only. Before this fix it was priced
    partly at July's and August's weights (the stretched 3-month split).
    Making July and August extreme must change nothing."""
    base = _calendar_for("manufacturing", 6)
    skewed = [
        dict(e, budget_weight=(5.0 if e["month"] in (7, 8) else e["budget_weight"]))
        for e in base
    ]
    a = _parse(_render("manufacturing", 6, 4, base))
    b = _parse(_render("manufacturing", 6, 4, skewed))
    assert a["shares"] == b["shares"], (a["shares"], b["shares"])
    assert all(abs(s - 0.25) < 0.0002 for s in b["shares"]), b["shares"]
    # And a 5-week plan (June 1 - July 5) must NOT feel August at all.
    c = _parse(_render("manufacturing", 6, 5, base))
    only_aug = [
        dict(e, budget_weight=(5.0 if e["month"] == 8 else e["budget_weight"]))
        for e in base
    ]
    d = _parse(_render("manufacturing", 6, 5, only_aug))
    assert c["shares"] == d["shares"], (c["shares"], d["shares"])


# ---------------------------------------------------------------------------
# 3. Flat calendar (no real seasonality): no ranking may be claimed.
# ---------------------------------------------------------------------------


def _flat(timeline: list[dict]) -> list[dict]:
    return [dict(e, budget_weight=1.0, hiring_intensity="moderate") for e in timeline]


def test_flat_calendar_4_weeks_is_an_even_split_with_no_ranking():
    timeline = _flat(_calendar_for("manufacturing", 6))
    p = _parse(_render("manufacturing", 6, 4, timeline))
    assert p["cols"] == ["Week 1", "Week 2", "Week 3", "Week 4"]
    assert all(abs(s - 0.25) < 0.0002 for s in p["shares"]), p["shares"]
    for word in ("heaviest", "lightest", "heavier", "lighter"):
        assert word not in p["note"].lower(), p["note"]
    assert "comparable" in p["note"]


def test_flat_calendar_every_short_duration_makes_no_ranking_claim():
    offenders = []
    for start_month in range(1, 13):
        timeline = _flat(_calendar_for("manufacturing", start_month))
        for weeks in range(1, 14):
            p = _parse(_render("manufacturing", start_month, weeks, timeline))
            if re.search(r"heavi|light", p["note"], re.I):
                offenders.append((start_month, weeks, p["note"]))
            days = [_column_days(c, p["start"], p["end"]) for c in p["cols"]]
            paces = [s / len(ds) for s, ds in zip(p["shares"], days)]
            if max(paces) - min(paces) > 1e-4:
                offenders.append((start_month, weeks, "uneven pacing", paces))
    assert not offenders, offenders[:3]


# ---------------------------------------------------------------------------
# 4. The calendar claim prints only when real calendar values were used.
# ---------------------------------------------------------------------------


def test_malformed_timeline_without_month_key_never_claims_calendar_match():
    timeline = [
        {k: v for k, v in e.items() if k != "month"}
        for e in _calendar_for("manufacturing", 10)
    ]
    for weeks in (1, 4, 5, 8, 9, 13, 26, 78):
        note = _parse(_render("manufacturing", 10, weeks, timeline))["note"]
        assert _CAL_CLAIM not in note, (weeks, note)


def test_timeline_without_budget_weights_never_claims_calendar_match():
    timeline = [
        {k: v for k, v in e.items() if k != "budget_weight"}
        for e in _calendar_for("manufacturing", 10)
    ]
    for weeks in (4, 5, 9, 78):
        note = _parse(_render("manufacturing", 10, weeks, timeline))["note"]
        assert _CAL_CLAIM not in note, (weeks, note)


def test_calendar_missing_a_plotted_month_never_claims_calendar_match():
    full = _calendar_for("manufacturing", 6)
    partial = [e for e in full if e["month"] != 7]  # July missing
    for weeks in (5, 8, 78):  # every one of these plots July
        note = _parse(_render("manufacturing", 6, weeks, partial))["note"]
        assert _CAL_CLAIM not in note, (weeks, note)
    # ...but a plan that never touches July may still claim it truthfully.
    note = _parse(_render("manufacturing", 6, 4, partial))["note"]
    assert "matching this plan's Activation Event Calendar" in note, note


def test_used_real_calendar_flag_is_explicit_and_exact():
    full = _calendar_for("manufacturing", 10)
    no_month = [{k: v for k, v in e.items() if k != "month"} for e in full]
    start = datetime.date(2026, 10, 1)

    d = excel_v2._seasonal_monthly_phasing_detail("manufacturing", 10, full)
    assert d["used_real_calendar"] is True and d["source"] == "calendar"
    d = excel_v2._seasonal_monthly_phasing_detail("manufacturing", 10, no_month)
    assert d["used_real_calendar"] is False and d["source"] != "calendar"
    d = excel_v2._seasonal_monthly_phasing_detail("manufacturing", 10, [])
    assert d["used_real_calendar"] is False

    s = excel_v2._short_plan_forecast_split("manufacturing", start, 5, False, full)
    assert s["used_real_calendar"] is True
    assert s["labels"] == ["October 2026", "November 2026"]
    assert s["days"] == [31, 4]
    s = excel_v2._short_plan_forecast_split("manufacturing", start, 5, False, no_month)
    assert s["used_real_calendar"] is False


# ---------------------------------------------------------------------------
# 5. Full workbook: the note's weights are the calendar table's weights.
# ---------------------------------------------------------------------------


def _hershey_data(weeks: int, start_month: int) -> dict:
    roles = [
        {"title": "Production Associate", "count": 30, "tier": "entry"},
        {"title": "Packaging Technician", "count": 15, "tier": "entry"},
    ]
    locations = [{"city": "Hershey", "state": "PA", "country": "United States"}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=300_000,
        roles=roles,
        locations=locations,
        industry="manufacturing",
        channel_percentages={
            "niche_boards": 15,
            "programmatic_dsp": 25,
            "global_boards": 20,
            "social_media": 10,
            "employer_branding": 10,
            "regional_boards": 20,
        },
        collar_type="blue_collar",
        campaign_start_month=start_month,
    )
    data = {
        "client_name": "Hershey Company",
        "industry": "manufacturing",
        "budget": "$300,000",
        "campaign_duration": f"{weeks} weeks",
        "campaign_weeks": weeks,
        "campaign_start_month": start_month,
        "hire_volume": "45 hires",
        "work_environment": "onsite",
        "roles": [r["title"] for r in roles],
        "target_roles": roles,
        "locations": [f"{loc['city']}, {loc['state']}" for loc in locations],
        "_budget_allocation": alloc,
    }
    data["_gold_standard"] = gold_standard.apply_all_quality_gates(data)
    return data


def _calendar_table(wb) -> dict[str, tuple[str, str]]:
    """Month name -> (Hiring Intensity, Budget Weight) as printed in the
    Quality Intelligence sheet's Activation Event Calendar table."""
    ws = wb["Quality Intelligence"]
    out: dict[str, tuple[str, str]] = {}
    inside = False
    for row in ws.iter_rows(values_only=True):
        if len(row) > 1 and row[1] == "Activation Event Calendar":
            inside = True
            continue
        if inside and len(row) > 4 and row[1] in _calendar.month_name[1:]:
            out[row[1]] = (str(row[3]), str(row[4]))
    return out


def test_generated_workbook_short_plans_quote_the_calendar_tables_weights():
    for weeks, start_month in ((5, 6), (8, 11), (8, 10), (4, 12), (13, 10)):
        data = _hershey_data(weeks, start_month)
        raw = excel_v2.generate_excel_v2(data)
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        table = _calendar_table(wb)
        assert table, "Activation Event Calendar table not found"
        ws = wb[excel_v2._forecast_sheet_title(data)]
        p = _parse(ws)
        note = p["note"]
        timeline = data["_gold_standard"]["activation_calendar"]["timeline"]
        assert _violations(p, timeline) == [], (weeks, start_month, note)
        # Every "<Month YYYY> (<n> campaign days; <w>x, ..." weight quoted in
        # the note is exactly the weight the calendar table prints.
        quoted = re.findall(
            r"([A-Z][a-z]+) \d{4} \(\d+ campaign days?; ([\d.]+x)", note
        )
        if weeks > 4:  # month-framed; these calendars all differ by month
            assert len(quoted) == len(p["cols"]), (weeks, start_month, note)
        for month, wtxt in quoted:
            assert table[month][1] == wtxt, (month, wtxt, table[month], note)
        # A month the calendar table calls "Low" is never the heavier/heaviest.
        for month, _lbl in re.findall(
            r"([A-Z][a-z]+) \d{4} \([^)]*?(heavier|heaviest)[^)]*\)", note
        ):
            assert table[month][0].lower() != "low", (month, table[month], note)


def test_hershey_90_day_repro_still_ranks_october_over_december():
    data = _hershey_data(78, 10)
    data["campaign_duration"] = "18 months"
    raw = excel_v2.generate_excel_v2(data)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    p = _parse(wb["90-Day Forecast"])
    assert [c.split()[0] for c in p["cols"]] == ["October", "November", "December"]
    assert p["shares"][0] > p["shares"][1] > p["shares"][2], p["shares"]
    # S50 fix (seasonal key-mismatch): industry "manufacturing" now correctly
    # matches its own seasonal_hiring_trends.json pattern (low_months
    # include November AND December, not just December), instead of never
    # matching at all (it used to resolve to _get_industry_key() bucket
    # "blue_collar_trades", which has no JSON key of its own). December's
    # share moves from 25% to 28% under the real low_multiplier blend, but
    # the invariant this test is named for -- October ranked heaviest,
    # December lightest -- still holds.
    assert "28% Month 3 (lightest planned spend)" in p["note"], p["note"]
    assert "matching this plan's Activation Event Calendar" in p["note"]
    assert _violations(p, data["_gold_standard"]["activation_calendar"]["timeline"]) == []


if __name__ == "__main__":
    import pytest

    raise SystemExit(pytest.main([__file__, "-q"]))
