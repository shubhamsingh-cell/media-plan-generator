"""Final review of the projection work: headline parity for generated bundles,
labels never stripped, roi_projector honouring suppression, one currency for
the chat tools, truthful scenario labels, ranges (not points) in the narrative,
and no unlabelled US CPA band on a non-US plan."""

from __future__ import annotations

import io
import logging
import os
import re
import sys
import types

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
for _n in ("anthropic", "openai", "supabase", "redis", "qdrant_client", "sentence_transformers", "posthog"):
    sys.modules.setdefault(_n, types.ModuleType(_n))
logging.disable(logging.CRITICAL)

import plan_projection  # noqa: E402

NOTE = plan_projection.NO_LOCAL_HIRES_NOTE
DASH = "–"
_M = ("applications", "hires", "cpa", "cph")
US_SHAPES = [
    ("Warehouse Associate", "Dallas, TX", "logistics_supply_chain"),
    ("Registered Nurse", "Columbus, OH", "healthcare_medical"),
    ("Software Engineer", "Remote, US", "tech_engineering"),
    ("Customer Service Representative", "United States", "general_entry_level"),
]


def _band(p):
    return {m: (p[m]["low"], p[m]["high"]) for m in _M}


def _brief(role, loc, ind, budget="$60,000"):
    return {"client_name": "Acme", "industry": ind, "budget": budget, "campaign_duration": "3 months",
            "hire_volume": "Not specified", "work_environment": "onsite", "locations": [loc],
            "roles": [role], "target_roles": [role]}


def _bundle(brief):
    import excel_v2
    import ppt_generator
    import tools_regen_bundles as T
    from kb_loader import load_knowledge_base
    from openpyxl import load_workbook
    from pptx import Presentation

    data = T.build_plan_data(brief)
    raw = excel_v2.generate_excel_v2(dict(data), load_kb_fn=load_knowledge_base)
    raw = raw[0] if isinstance(raw, tuple) else raw
    wb = load_workbook(io.BytesIO(raw))
    prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))
    return data, wb, prs


def _deck(prs):
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    yield i, p.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        yield i, c.text


def _cells(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    yield ws.title, c.coordinate, c.value


# 1 ------------------------------------------------------------------------
@pytest.mark.parametrize("role,loc,ind", US_SHAPES)
def test_generated_bundle_headline_equals_the_estimate(role, loc, ind):
    import app
    import ppt_generator

    est = app._compute_plan_estimate(
        {"budget": "$60,000", "industry": ind, "roles": [role], "locations": [loc],
         "campaign_duration": "3 months", "client_name": "Acme"}
    )["projection"]
    data, _wb, _prs = _bundle(_brief(role, loc, ind))
    head = data["_budget_allocation"].get("headline_projection")
    assert head, "the bundle carries no headline projection"
    assert _band(head) == _band(est)
    assert _band(ppt_generator._deck_projection(data)) == _band(est)


# 2 ------------------------------------------------------------------------
def test_strip_never_removes_labels():
    s = plan_projection.strip_hire_claims
    for label in ("Projected Hires", "Total Proj. Hires", "Proj. Hires", "Cost / Hire", "Hires"):
        assert s(label) == label, label


def test_london_labels_survive_with_n_a_values():
    _data, wb, prs = _bundle(_brief("Warehouse Associate", "London, UK", "logistics_supply_chain", "£60,000"))
    deck = list(_deck(prs))
    assert any(t.strip() == "Projected Hires" for i, t in deck if i == 2)
    assert any(t.strip() == "n/a" for i, t in deck if i == 2)
    assert any(t.strip() == "Projected Hires" for i, t in deck if i == 7)
    cells = list(_cells(wb))
    assert any(s == "Executive Summary" and v == "Projected Hires" for s, _c, v in cells)
    assert any(s == "ROI Projections" and v == "Total Proj. Hires" for s, _c, v in cells)
    assert any(NOTE in str(v) for _s, _c, v in cells)


# 3 ------------------------------------------------------------------------
def test_roi_projector_projects_no_hires_for_london_and_keeps_them_for_the_us():
    import roi_projector

    ldn = roi_projector.project_roi("Indeed", 60000, "logistics", "Warehouse Associate", ["London, UK"])
    for k in ("pessimistic", "expected", "optimistic"):
        assert ldn["projections"][k]["hires"] is None, ldn["projections"]
    assert ldn["roi_metrics"]["cost_per_hire"] == {"low": None, "expected": None, "high": None}
    assert NOTE in ldn["recommendation"] and not re.search(r"\d\s+hires?", ldn["recommendation"])
    assert "channel_detail" not in ldn and ldn["hires_note"] == NOTE
    us = roi_projector.project_roi("Indeed", 60000, "healthcare", "Registered Nurse", ["Columbus, OH"])
    assert us["projections"]["expected"]["hires"]


# 4 ------------------------------------------------------------------------
def test_chat_tools_price_a_dollar_budget_in_dollars_in_india():
    import roi_projector
    from nova import Nova

    n = Nova.__new__(Nova)
    n._data_cache = {}
    res = n._query_budget_projection(
        {"budget": 60000, "roles": ["Software Engineer"], "locations": ["Bengaluru, India"], "industry": "tech_engineering"}
    )
    assert res["projection"]["currency"]["code"] == "USD"
    assert res["projection"]["text"]["cpa"].startswith("$")
    assert res["projection"]["applications"]["high"] > 1000  # not INR-priced (was 58)
    roi = roi_projector.project_roi("Indeed", 60000, "technology", "Software Engineer", ["Bengaluru, India"])
    assert roi["headline_projection"]["currency"]["code"] == "USD"


# 5 ------------------------------------------------------------------------
def test_slide6_scenario_label_is_truthful():
    for role, loc, ind in US_SHAPES[2:]:
        _d, _w, prs = _bundle(_brief(role, loc, ind))
        s6 = " || ".join(t for i, t in _deck(prs) if i == 7)
        assert "upper scenario" not in s6, s6
        assert "engine scenario" in s6


# 6 ------------------------------------------------------------------------
def test_narrative_tiles_and_risk_quote_ranges_not_the_upper_point():
    data, wb, prs = _bundle(_brief("Registered Nurse", "Columbus, OH", "healthcare_medical"))
    cells = list(_cells(wb))
    flat = [str(v) for _s, _c, v in cells]
    assert not any("projected to deliver 11 hire" in v for v in flat)
    assert any("projected to deliver 5" + DASH + "11 hires" in v for v in flat), [v for v in flat if "deliver" in v]
    ws_vals = [v for s, _c, v in cells if s == "Executive Summary"]
    assert "5" + DASH + "11" in [str(v) for v in ws_vals], "hero Projected Hires tile is not a range"
    assert any(str(v).startswith("$5,250") and DASH in str(v) for v in ws_vals), "hero Cost/Hire tile"
    risk = " || ".join(t for i, t in _deck(prs) if i == 10)
    assert "drop from 11 to" not in risk and "5" + DASH + "11" in risk, risk
    assert not any("drop from 11 to" in v for v in flat)


# 7 ------------------------------------------------------------------------
def test_us_cpa_band_is_not_applied_to_a_non_us_dollar_plan():
    p = plan_projection.project_plan(
        {"budget": 60000, "roles": ["Software Engineer"], "locations": ["Bengaluru, India"],
         "industry": "tech_engineering", "budget_text": "$60,000"}
    )
    assert p["cpa"]["low"] == p["cpa"]["high"]  # no US band widening
    assert "US" in " ".join(p["notes"]) and "single estimate" in " ".join(p["notes"])
    us = plan_projection.project_plan(
        {"budget": 60000, "roles": ["Software Engineer"], "locations": ["Remote, US"],
         "industry": "tech_engineering", "budget_text": "$60,000"}
    )
    assert us["cpa"]["low"] < us["cpa"]["high"]


# ---- headline helper locked directly (not via the harness) ------------------
def _headline_for(plan, caller_mix):
    import budget_engine
    from kb_loader import load_knowledge_base
    import app

    alloc = budget_engine.calculate_budget_allocation(
        total_budget=60000.0,
        roles=[{"title": r, "count": 1, "tier": "Professional"} for r in plan["target_roles"]],
        locations=[plan_projection.split_location(plan["locations"][0])],
        industry=plan["industry"], channel_percentages=caller_mix, synthesized_data=None,
        knowledge_base=load_knowledge_base(), plan_currency="USD", locations_raw=plan["locations"],
    )
    app._attach_headline_projection(
        alloc, plan, budget=60000.0, roles=[], locations=[], channel_pcts=caller_mix,
        vendor_availability={"niche_boards": True}, kb=load_knowledge_base(),
    )
    return alloc["headline_projection"]


@pytest.mark.parametrize(
    "role,loc,raw,region",
    [
        ("Customer Service Representative", "United States", "general_entry_level", "us_only"),
        ("Registered Nurse", "Columbus, OH", "healthcare_medical", "us_only"),
        ("Registered Nurse", "London, UK", "healthcare_medical", "global"),
        ("Warehouse Associate", "Manchester, UK", "logistics_supply_chain", "global"),
    ],
    ids=["us-cs", "us-rn", "uk-global-rn", "uk-global-wh"],
)
def test_headline_helper_ignores_the_callers_mix_key_and_vendor_gate(role, loc, raw, region):
    """Called the way /api/generate calls it: a MAPPED industry key
    (retail_consumer), a different channel mix, a permissive vendor gate and
    `_industry_raw` set. The headline must still be the estimate's."""
    import app

    est = app._compute_plan_estimate(
        {"budget": "$60,000", "industry": raw, "roles": [role], "locations": [loc],
         "target_region": region, "client_name": "Acme"}
    )["projection"]
    plan = {"budget": "$60,000", "industry": "retail_consumer", "_industry_raw": raw,
            "client_name": "Acme", "locations": [loc], "target_roles": [role],
            "target_region": region, "hire_volume": ""}
    other_mix = {"programmatic_dsp": 5, "global_boards": 5, "niche_boards": 80, "social_media": 10}
    assert _band(_headline_for(plan, other_mix)) == _band(est)


def test_headline_helper_applies_the_estimates_vendor_gate(monkeypatch):
    """A non-US plan: the niche-vendor gate must be the estimate's (False), not
    the permissive one the caller passes."""
    import app
    import budget_engine

    seen = []
    real = budget_engine.calculate_budget_allocation

    def spy(*a, **k):
        seen.append(k.get("vendor_availability"))
        return real(*a, **k)

    monkeypatch.setattr(budget_engine, "calculate_budget_allocation", spy)
    role, loc, raw = "Registered Nurse", "London, UK", "healthcare_medical"
    app._compute_plan_estimate(
        {"budget": "$60,000", "industry": raw, "roles": [role], "locations": [loc],
         "target_region": "global", "client_name": "Acme"}
    )
    plan = {"budget": "$60,000", "industry": "retail_consumer", "_industry_raw": raw, "client_name": "Acme",
            "locations": [loc], "target_roles": [role], "target_region": "global", "hire_volume": ""}
    _headline_for(plan, {"programmatic_dsp": 50, "niche_boards": 50})
    est_gate, head_gate = seen[0], seen[-1]
    assert est_gate == head_gate == {"niche_boards": False}, (est_gate, head_gate)


@pytest.mark.parametrize(
    "role,loc,ind,budget",
    [
        ("Registered Nurse", "Columbus, OH", "healthcare_medical", "$150,000"),
        ("Software Engineer", "Remote, US", "tech_engineering", "$250,000"),
        ("Customer Service Representative", "United States", "general_entry_level", "$60,000"),
    ],
)
def test_channel_concentration_risk_quotes_a_range(role, loc, ind, budget):
    import re as _re

    data, wb, prs = _bundle(_brief(role, loc, ind, budget))
    texts = [str(v) for s, _c, v in _cells(wb) if "could impact" in str(v)]
    texts += [t for i, t in _deck(prs) if i == 10 and "could impact" in t]
    assert texts, "fixture must trigger the channel-concentration risk"
    for t in texts:
        assert not _re.search(r"could impact \d[\d,]*\s+(?:projected\s+)?hires?", t), t
        assert DASH in t.split("could impact")[1], t
