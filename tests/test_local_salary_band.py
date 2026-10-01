"""Design-judge 2026-10-01: a local salary is ONE published, role-matched
band -- never a span across different statistics -- and the deck and the
workbook state the same one.

Shipped defect: the deck's slide-2 Salary Range for a UK finance plan read
"£39,000 - £49,983 (GBP) local benchmark" -- the ONS all-occupations median
paired with an Adzuna category average (min/max over every entry of the
vertical) -- while the workbook printed n/a for the same data. India IT
read "₹400,000 - ₹7,500,000" (an entry-level band's low to a senior
top-company band's high).
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation

import cited_data_block
import data_synthesizer
import excel_v2
import gold_standard as gs
import intl_benchmark_lookup as ibl
import ppt_generator
import research as research_mod
import tools_regen_bundles as T

_DATA = json.loads(
    (Path(__file__).resolve().parent.parent / "data" / "intl_role_benchmarks_v1.json")
    .read_text(encoding="utf-8")
)

_INDIA_IT_ROLES = [
    "Software Engineer (Fresher)",
    "Java Developer",
    "QA Engineer",
    "Senior Data Engineer (Lateral)",
]
_INDIA = ["Bengaluru, Karnataka, India", "Pune, Maharashtra, India"]
_UK_FIN_ROLES = ["Financial Analyst", "Compliance Analyst", "Software Developer"]
_UK = ["London, UK", "Manchester, UK"]


def _entries():
    for vertical in _DATA["verticals"].values():
        for slug, block in vertical["by_country"].items():
            for key, entry in (block.get("annual_salary") or {}).items():
                yield slug, key, entry


_SWEEP_ROLES = [
    "Registered Nurse", "Staff Nurse", "Nurse Practitioner", "Software Engineer",
    "Senior Software Engineer", "Software Engineer (Fresher)", "Java Developer",
    "Warehouse Associate", "Accountant", "Chartered Accountant",
    "Financial Analyst", "Compliance Analyst", "Customer Service Representative",
]
_SWEEP_COUNTRIES = [
    "UK", "India", "Germany", "France", "Canada", "Australia", "Netherlands",
    "Spain", "Brazil", "Mexico", "Singapore", "UAE", "Japan", "Ireland",
]


@pytest.mark.parametrize("country", _SWEEP_COUNTRIES)
def test_every_band_is_exactly_one_published_entry(country):
    """low / median / high always come from ONE dataset entry -- never a
    min/max assembled across entries."""
    by_key = {(s, k): e for s, k, e in _entries()}
    slug = ibl._normalize_country(country)
    for role in _SWEEP_ROLES:
        band = ibl.get_local_role_salary_band(country, role)
        if band is None:
            continue
        entry = by_key[(slug, band["statistic"])]
        assert (band["low"], band["median"], band["high"]) == (
            entry["low"], entry["median"], entry["high"],
        ), (country, role, band)
        assert band["source"]


def test_uk_finance_roles_have_no_band():
    for role in _UK_FIN_ROLES:
        assert ibl.get_local_role_salary_band("London, UK", role) is None


def test_india_bands_follow_title_seniority():
    fresher = ibl.get_local_role_salary_band(_INDIA[0], "Software Engineer (Fresher)")
    java = ibl.get_local_role_salary_band(_INDIA[0], "Java Developer")
    assert fresher["statistic"] == "swe_entry_4_6_lpa"
    assert java["statistic"] == "swe_mid_8_15_lpa"
    assert ibl.get_local_role_salary_band(_INDIA[0], "QA Engineer") is None


def test_us_location_never_gets_a_local_band():
    assert ibl.get_local_role_salary_band("Columbus, OH", "Registered Nurse") is None


def test_salary_range_text_never_pairs_unrelated_statistics():
    uk = ppt_generator._local_salary_range_text(
        {"locations": _UK, "roles": _UK_FIN_ROLES, "industry": "finance_banking"}
    )
    assert uk == "Local salary data n/a"
    india = ppt_generator._local_salary_range_text(
        {"locations": _INDIA, "roles": _INDIA_IT_ROLES, "industry": "tech_engineering"}
    )
    # Compact on the deck card: figures, role and source domain. (Outside a
    # render the active currency is USD, so the INR code is declared too.)
    assert india.startswith("₹500K median (₹400K-₹600K)")
    assert india.endswith("- Software Engineer (Fresher) [plugscale.com]")
    assert "7.5M" not in india and "7,500,000" not in india


def test_cited_block_uses_the_same_band():
    uk = cited_data_block.build_cited_2026_block({"locations": _UK, "roles": _UK_FIN_ROLES})
    assert uk["salary_line"] == ""
    india = cited_data_block.build_cited_2026_block(
        {"locations": _INDIA, "roles": _INDIA_IT_ROLES}
    )
    assert "₹500K median (₹400K-₹600K)" in india["salary_line"]
    assert "software engineer, entry level; source: Plugscale" in india["salary_line"]


def test_market_and_quality_intelligence_carry_the_same_band():
    data = {"roles": _INDIA_IT_ROLES, "locations": _INDIA, "industry": "tech_engineering"}
    synth = data_synthesizer.synthesize({}, {}, data)
    mi = synth["salary_intelligence"]["Software Engineer (Fresher)"]
    assert (mi["min"], mi["median"], mi["max"]) == (400000, 500000, 600000)
    assert mi["currency"] == "INR" and mi["p25"] is None and mi["local_band"]
    assert mi["sources"][0] == (
        "Plugscale Software Engineer Compensation Benchmark India 2026 "
        "(software engineer, entry level)"
    )
    data["_synthesized"] = synth
    for path_data in (data, {"roles": _INDIA_IT_ROLES, "locations": _INDIA}):
        for info in gs.enrich_city_level_data(path_data).values():
            row = info["per_role_salary"]["Software Engineer (Fresher)"]
            assert (row["min"], row["median"], row["max"]) == (400000, 500000, 600000)
            assert row["p25"] is None and row["currency"] == "INR"
            assert row["source"] == mi["sources"][0]  # one spelling, both paths
            assert info["per_role_salary"]["QA Engineer"]["local_salary_na"] is True


@pytest.fixture(scope="module")
def india_bundle():
    brief = {
        "client_name": "Nimbus Infotech Services",
        "budget": "₹25,000,000",
        "campaign_duration": "3-6 months",
        "industry": "Information Technology",
        "locations": _INDIA,
        "roles": _INDIA_IT_ROLES,
        "target_roles": [{"title": r, "count": 10} for r in _INDIA_IT_ROLES],
    }
    data = T.build_plan_data(brief)  # no synthesis: the tools_regen path
    xlsx = excel_v2.generate_excel_v2(dict(data), research_mod=research_mod)
    pptx = ppt_generator.generate_pptx(dict(data))
    xlsx = xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx
    pptx = pptx.getvalue() if hasattr(pptx, "getvalue") else pptx
    return {"xlsx": xlsx, "pptx": pptx}


def test_deck_slide2_and_workbook_state_one_decision(india_bundle):
    prs = Presentation(io.BytesIO(india_bundle["pptx"]))
    slide2 = " ".join(
        sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame
    )
    assert "₹500K median (₹400K-₹600K) - Software Engineer (Fresher)" in slide2
    ws = load_workbook(io.BytesIO(india_bundle["xlsx"]))["Quality Intelligence"]
    rows = [
        [c for c in r if c.value is not None]
        for r in ws.iter_rows()
        if len(r) > 2 and r[2].value in ("Software Engineer (Fresher)", "Software Engineer (Fresher) (est.)")
        and r[1].value in ("Bengaluru", "Pune")
    ]
    assert rows, "role row missing on Quality Intelligence"
    for r in rows:
        vals = [c.value for c in r]
        assert vals[2:7] == [400000, "—", 500000, "—", 600000], vals
        assert r[2].number_format == '"₹"#,##0'
        assert str(vals[-1]).startswith("Plugscale")
