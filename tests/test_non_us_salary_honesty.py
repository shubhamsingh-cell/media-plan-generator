"""Audit F 3.4 + K-05b + B-39: non-US plans never print a US number as
their own, and every salary figure is labelled in the currency it is in.

Shipped defects reproduced before this fix (real generate_pptx /
generate_excel_v2 output):
  * Bangalore nurse plan: deck slide 2 "Salary Range: US$78K median" (a DOL
    H-1B US median), slide 7 "₹78K (est., shared band)" (the same US dollars
    with a rupee sign), Quality Intelligence "₹50,700 - ₹180,000 (INR)".
  * London software plan: Quality Intelligence "£97,500 ... £217,500" for US
    H-1B dollars; Market Intelligence "US$150,000" for the same role.
  * New York + London plan budgeted in GBP: New York rows (US dollars)
    printed with "£", the city Salary Range as a bare "$" range on a GBP
    plan (a bundle_qa currency_symbol_mixing critical).
  * B-39: every non-US plan's Market Intelligence sheet carried the US
    "National Economic Snapshot" (~8.0M openings, 62.5% participation).
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from pptx import Presentation

import data_synthesizer
import excel_v2
import gold_standard as gs
import plan_geo
import ppt_generator
import research as research_mod
import tools_regen_bundles as T
from bundle_qa import run_bundle_qa

_ROLES_HC = ["Registered Nurse", "Staff Nurse", "Nurse Practitioner", "Medical Assistant"]
_ROLES_TECH = ["Software Engineer", "Accountant", "Data Scientist", "DevOps Engineer"]
_NA = "Local salary data n/a"


def test_label_constant():
    assert gs.LOCAL_SALARY_NA_LABEL == _NA


def _bundle(roles, locations, budget, industry):
    brief = {
        "client_name": "Northwind Health",
        "budget": budget,
        "campaign_duration": "3-6 months",
        "industry": industry,
        "locations": list(locations),
        "roles": list(roles),
        "target_roles": [{"title": r, "count": 5} for r in roles],
    }
    data = T.build_plan_data(brief)
    data["_synthesized"] = data_synthesizer.synthesize({}, {}, data)
    data["_gold_standard"] = gs.apply_all_quality_gates(data)
    xlsx = excel_v2.generate_excel_v2(dict(data), research_mod=research_mod)
    pptx = ppt_generator.generate_pptx(dict(data))
    xlsx = xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx
    pptx = pptx.getvalue() if hasattr(pptx, "getvalue") else pptx
    return {"data": data, "xlsx": xlsx, "pptx": pptx}


@pytest.fixture(scope="module")
def bangalore():
    return _bundle(_ROLES_HC, ["Bangalore, India"], "₹20,000,000", "healthcare")


@pytest.fixture(scope="module")
def london():
    return _bundle(_ROLES_TECH, ["London, UK"], "£500,000", "technology")


@pytest.fixture(scope="module")
def mixed():
    return _bundle(_ROLES_TECH, ["New York, NY", "London, UK"], "£500,000", "technology")


@pytest.fixture(scope="module")
def columbus():
    return _bundle(
        ["Retail Sales Associate", "Cashier", "Store Manager", "Stock Associate"],
        ["Columbus, OH"],
        "$250,000",
        "retail",
    )


def _slides(pptx: bytes) -> dict[int, list[str]]:
    out: dict[int, list[str]] = {}
    for i, slide in enumerate(Presentation(io.BytesIO(pptx)).slides, 1):
        out[i] = [
            sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text
        ]
    return out


def _sheet_rows(xlsx: bytes, sheet: str):
    ws = load_workbook(io.BytesIO(xlsx))[sheet]
    return ws, [[c for c in row if c.value is not None] for row in ws.iter_rows()]


# --------------------------------------------------------------------------
# Source: data_synthesizer / gold_standard / plan_geo
# --------------------------------------------------------------------------


def test_plan_has_us_market():
    assert plan_geo.plan_has_us_market({"locations": ["Columbus, OH"]}) is True
    assert plan_geo.plan_has_us_market({"locations": ["London, UK"]}) is False
    assert (
        plan_geo.plan_has_us_market(
            {"locations": ["New York, NY", "London, UK"], "target_region": "global"}
        )
        is True
    )


@pytest.mark.parametrize(
    "roles,location",
    [
        (_ROLES_HC, "Bangalore, India"),
        (_ROLES_TECH, "London, UK"),
        (["Retail Sales Associate"], "Sao Paulo, Brazil"),
        (["Warehouse Associate", "Truck Driver"], "Berlin, Germany"),
    ],
)
def test_no_us_salary_survives_on_a_plan_with_no_us_market(roles, location):
    data = {"roles": roles, "locations": [location], "industry": "general"}
    synth = data_synthesizer.synthesize({}, {}, data)
    for role, sal in synth["salary_intelligence"].items():
        assert not (sal.get("median") and sal.get("currency") == "USD"), (role, sal)
    data["_synthesized"] = synth
    info = next(iter(gs.enrich_city_level_data(data).values()))
    assert info["estimated_salary"] == 0
    assert info["salary_range"] == _NA
    for role, row in info["per_role_salary"].items():
        # Either the role's ONE published local band (its own currency,
        # named source -- tests/test_local_salary_band.py) or withheld.
        if row.get("local_band"):
            assert row["currency"] not in ("", "USD"), (role, row)
            assert row["source"] != _NA
            continue
        assert row["local_salary_na"] is True, role
        assert row["median"] == 0
        assert row["source"] == _NA


def test_us_market_rows_are_tagged_usd_and_kept_on_a_mixed_plan():
    data = {"roles": _ROLES_TECH, "locations": ["New York, NY", "London, UK"]}
    data["_synthesized"] = data_synthesizer.synthesize({}, {}, data)
    city = gs.enrich_city_level_data(data)
    for role, row in city["New York"]["per_role_salary"].items():
        assert row["currency"] == "USD" and row["median"] > 0, role
    for role, row in city["London"]["per_role_salary"].items():
        assert row.get("local_salary_na") is True, role


# --------------------------------------------------------------------------
# Rendered deck
# --------------------------------------------------------------------------


def test_bangalore_deck_prints_local_range_not_a_us_salary(bangalore):
    slides = _slides(bangalore["pptx"])
    text = " ".join(t for ts in slides.values() for t in ts)
    assert "US$78K" not in text and "₹78K" not in text
    # One published band (staff nurse, government) -- not a span of entries.
    assert "₹300K median (₹180K-₹480K) - Registered Nurse" in " ".join(slides[2])
    role_breakdown = [i for i, ts in slides.items() if any(t == "Role Breakdown" for t in ts)]
    assert role_breakdown, "Role Breakdown slide not rendered"
    cells = slides[role_breakdown[0]]
    assert cells.count("₹300K") == 2  # Registered Nurse, Staff Nurse
    assert sum("n/a" in t for t in cells) == 2  # Nurse Practitioner, Medical Assistant


def test_mixed_plan_role_breakdown_marks_us_dollars(mixed):
    slides = _slides(mixed["pptx"])
    rb = next(i for i, ts in slides.items() if any(t == "Role Breakdown" for t in ts))
    salaries = [t for t in slides[rb] if "K" in t and "$" in t]
    assert salaries and all(t.startswith("US$") for t in salaries), salaries
    assert not any(t.startswith("£") and "K" in t for t in slides[rb])


def test_columbus_retail_deck_prints_band_salary(columbus):
    slides = _slides(columbus["pptx"])
    text = " ".join(t for ts in slides.values() for t in ts)
    assert "$86K" not in text and "$90K median" not in text
    assert "$35K median ($28K-$42K) - Retail Sales Associate" in " ".join(slides[2])


# --------------------------------------------------------------------------
# Rendered workbook
# --------------------------------------------------------------------------


def _qi_role_rows(xlsx: bytes):
    ws, rows = _sheet_rows(xlsx, "Quality Intelligence")
    return [r for r in rows if len(r) >= 9 and r[8].value in (
        "DOL H-1B/LCA", "Industry Benchmark", "Tier-Scaled Estimate",
        _NA,
    )]


def test_london_quality_intelligence_prints_na_not_us_dollars(london):
    rows = _qi_role_rows(london["xlsx"])
    assert len(rows) == 4
    for r in rows:
        assert [c.value for c in r[2:7]] == ["n/a"] * 5
        assert r[8].value == _NA
    _, all_rows = _sheet_rows(london["xlsx"], "Quality Intelligence")
    city = next(r for r in all_rows if r and r[0].value == "London")
    assert [c.value for c in city][2] == "n/a"
    assert city[-1].value == _NA


def test_london_market_intelligence_withholds_us_salary(london):
    _, rows = _sheet_rows(london["xlsx"], "Market Intelligence")
    flat = [str(c.value) for r in rows for c in r]
    assert not any("US$150,000" in v for v in flat)
    assert flat.count("Local salary data n/a") >= 4


def test_mixed_plan_us_rows_formatted_us_dollars(mixed):
    rows = _qi_role_rows(mixed["xlsx"])
    ny = [r for r in rows if r[0].value == "New York"]
    assert ny
    for r in ny:
        for c in r[2:7]:
            assert c.number_format == '"US$"#,##0', (r[1].value, c.number_format)
    _, all_rows = _sheet_rows(mixed["xlsx"], "Quality Intelligence")
    city = next(r for r in all_rows if r and r[0].value == "New York")
    assert city[2].number_format == '"US$"#,##0'
    assert str(city[-1].value).startswith("US$")


@pytest.mark.parametrize("name", ["bangalore", "london", "mixed", "columbus"])
def test_bundle_qa_has_no_currency_or_us_data_criticals(name, request):
    bundle = request.getfixturevalue(name)
    findings = run_bundle_qa(bundle["pptx"], bundle["xlsx"], bundle["data"])
    bad = [
        f
        for f in findings
        if f.get("severity") == "critical"
        and f.get("code") in ("currency_symbol_mixing", "us_data_on_non_us_plan")
    ]
    assert not bad, bad
