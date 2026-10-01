"""Regression: design-judge + numbers-verifier round 2 (2026-10-01) on the
cost-per-hire disclosure of mpg-projection-honesty.

Pre-fix (branch at 9d1b305):
* Product decision: a single-source LOCAL range ("midpoint of cited range")
  was extrapolated below its own low end -- India IT ₹25M projected 869
  hires at ₹28,769/hire against a cited ₹35K-₹80K range. Floor is now
  max(0.5 x midpoint, cited low end): 714 hires.
* India's footnotes said "CPH: local-market benchmark" and never what the
  plan assumes per hire; "plan efficiency" was undefined; a 500-hire goal
  read as met although at the midpoint cost the plan buys 434.
* Hershey's slide-2 low end ("7 at cross-industry avg") came from the
  cross-industry default ($6,000) while slide 5 showed the food-
  manufacturing row ($3,000-$5,000); B16 also quoted a second range.
* The slide-5 cost-per-hire cell was fitted with a regular-weight estimate
  but renders BOLD: "$2,000–$4,700 (midpoint $3,350)" wrapped.
* B16 printed the floor ($5,250) where the deck prints the plan's own
  $5,319.
* A KSh plan with no exchange rate got a hires range at US$ parity.
* No slide said the plan budget is media spend while industry cost-per-
  hire ranges are all-in hiring costs.
"""

from __future__ import annotations

import io
import os
import sys

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

import budget_engine as be  # noqa: E402
import excel_v2  # noqa: E402
import ppt_generator  # noqa: E402
import tools_regen_bundles as T  # noqa: E402
from kb_loader import load_knowledge_base  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

_MEDIA = (
    "Plan budget covers media spend; industry cost-per-hire ranges include "
    "all hiring costs."
)


def _brief(name, industry, budget, locations, roles, hire_volume="Not specified"):
    return {
        "client_name": name,
        "industry": industry,
        "budget": budget,
        "campaign_duration": "6 months",
        "hire_volume": hire_volume,
        "work_environment": "onsite",
        "locations": locations,
        "roles": roles,
        "target_roles": roles,
    }


class _Bundle:
    def __init__(self, brief):
        self.data = T.build_plan_data(brief)
        self.prs = Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(self.data))))
        raw = excel_v2.generate_excel_v2(dict(self.data), load_kb_fn=load_knowledge_base)
        raw = raw[0] if isinstance(raw, tuple) else raw
        ws = load_workbook(io.BytesIO(raw))["Executive Summary"]
        self.cells = [str(c.value) for r in ws.iter_rows() for c in r if c.value]

    @property
    def tp(self):
        return self.data["_budget_allocation"]["total_projected"]

    def paras(self, n):
        return [
            p.text
            for sh in self.prs.slides[n - 1].shapes
            if sh.has_text_frame
            for p in sh.text_frame.paragraphs
            if p.text.strip()
        ]

    def cell(self, prefix):
        hits = [c for c in self.cells if c.startswith(prefix)]
        return hits[0] if hits else ""


@pytest.fixture(scope="module")
def india():
    return _Bundle(_brief(
        "Probe Infotech", "tech_engineering", "₹25,000,000",
        ["Bengaluru, Karnataka, India", "Pune, Maharashtra, India"],
        ["Java Developer", "QA Engineer"], "500+ hires"))


@pytest.fixture(scope="module")
def hospital():
    return _Bundle(_brief(
        "Probe Health", "Healthcare & Medical", "$250,000", ["Dallas, TX"],
        ["Registered Nurse", "Pharmacist"], "100 hires"))


@pytest.fixture(scope="module")
def hershey():
    return _Bundle(_brief(
        "Probe Foods", "food_beverage", "$45,000", ["Hershey, PA", "Reading, PA"],
        ["Machine Operator", "Sanitation"], "100-500 hires"))


@pytest.fixture(scope="module")
def tiny():
    return _Bundle(_brief(
        "Probe Food Bank", "general_entry_level", "$3,000", ["Sacramento, CA"],
        ["Volunteer Coordinator", "Warehouse Associate"], "1-10 hires"))


class TestCitedLowEndFloor:
    def test_resolver_floor_rule(self):
        local = be._attach_cph_floor({
            "basis": "local_kb", "value_label": "midpoint of cited range",
            "math_value": 57_500.0, "value": 57_500.0, "low": 35_000.0,
        })
        assert (local["floor"], local["floor_rule"]) == (35_000.0, "cited_low_end")
        wide = be._attach_cph_floor({
            "basis": "local_kb", "value_label": "midpoint of cited range",
            "math_value": 15_500.0, "value": 15_500.0, "low": 6_000.0,
        })
        assert (wide["floor"], wide["floor_rule"]) == (7_750.0, "half_midpoint")
        us = be.resolve_industry_cph("healthcare_medical")
        assert (us["floor"], us["floor_rule"]) == (5_250.0, "half_midpoint")

    def test_india_plan_never_below_the_cited_low_end(self, india):
        tp = india.tp
        assert (tp["hires_low"], tp["hires"], tp["hires_high"]) == (434, 714, 714)
        assert tp["cost_per_hire"] >= 35_000

    def test_us_headline_unchanged(self, hospital):
        assert (hospital.tp["hires_low"], hospital.tp["hires"]) == (23, 47)


class TestLocalAssumptionDisclosed:
    def test_slides_2_5_6_state_the_assumption(self, india):
        for n in (2, 5, 6):
            text = " ".join(india.paras(n))
            assert "₹35.0K" in text and "434 hires" in text, (n, text[-600:])
            assert "the low end of the cited ₹35K–₹80K range" in text, n
            assert "plan efficiency" in text.lower(), n

    def test_footnote_paragraphs_are_one_line(self, india):
        notes = [p for p in india.paras(2) + india.paras(5) if "₹35.0K" in p]
        assert notes
        for p in notes:
            assert ppt_generator._fits_one_line(p, 12.0, 9.0), p

    def test_workbook_states_it_too(self, india):
        b16 = india.cell("Plan assumes ₹35.0K per hire")
        assert "at the range midpoint (₹57.5K) this budget buys 434 hires" in b16, b16
        assert "₹35,014/hire" in b16 and "₹35,000/hire, the cited range's low end" in b16

    def test_goal_judged_at_the_conservative_end(self, india):
        b17 = india.cell("Hiring goal:")
        assert (
            "At the midpoint cost (₹57,500/hire) the plan buys 434 hires "
            "against a goal of 500" in b17
        ), b17
        deck = " ".join(india.paras(2))
        assert "at the midpoint cost the plan buys 434 hires against a goal of 500" in deck


class TestOneCostBasis:
    def test_hershey_low_end_is_the_shown_row(self, hershey):
        assert hershey.tp["hires_low"] == 11
        assert "11 at industry-avg cost" in hershey.paras(2)
        b16 = hershey.cell("Projected hires:")
        assert b16.startswith("Projected hires: 11–15."), b16
        assert "$4,000, the midpoint of the industry row" in b16, b16
        assert "range's ends" not in b16 and "$8,000" not in b16, b16

    def test_headline_cph_is_the_plans_own(self, hospital):
        b16 = hospital.cell("Projected hires:")
        assert "Point estimate: 47 hires at this plan's own $5,319/hire" in b16, b16
        # the floor prints only where it is labelled as the floor
        assert "lowest cost per hire ($5,250/hire" in b16, b16
        assert "efficiency floor of" not in b16


class TestBoldCellFit:
    def test_bold_measure_rejects_the_wrapping_form(self):
        assert not ppt_generator._fits_one_line(
            "$2,000–$4,700 (midpoint $3,350)", 2.3, 10.0, bold=True
        )

    def test_nonprofit_row_is_the_compact_form(self, tiny):
        paras = tiny.paras(5)
        i = paras.index("Industry Cost-per-Hire")
        assert paras[i + 1] == "$2K–$4.7K (midpoint $3.35K)", paras[i + 1]


class TestNoRangeWithoutAPlanCurrencyCph:
    def test_ksh_plan_has_no_range(self):
        b = _Bundle(_brief(
            "Probe KE", "general_entry_level", "KSh 100,000", ["Nairobi, Kenya"],
            ["Customer Service Representative"], "20 hires"))
        assert b.tp["hires_low"] is None and b.tp["hires_high"] is None
        assert not any(" at plan efficiency" in p for p in b.paras(2))
        assert not b.cell("Projected hires:")


class TestMediaVsAllIn:
    def test_slides_2_and_5_say_it(self, hospital):
        assert _MEDIA in hospital.paras(2)
        assert _MEDIA in hospital.paras(5)
