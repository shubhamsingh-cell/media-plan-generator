"""Delivery-gate policy: repair what is repairable, re-check, always deliver,
never ship a critical silently.

Prod (A_prod_telemetry 2026-10-01, section 3): bundle_qa raised >=1
CRITICAL on 5 of 5 real runs and every bundle shipped with nothing but
WARNING log lines -- the gate ran AFTER the ZIP was built and only observed.
bundle_qa.gate_bundle (called by app._run_bundle_qa_gate on both the async
job path and the sync /api/generate path, BEFORE packaging):
  1. lints, 2. repairs snake_case_leak / unsourced_competitor_claim /
  confirmed mid_word_truncation in the bytes that ship, re-lints,
  3. delivers regardless, but logs surviving criticals at ERROR with ids,
  audits them, records them on the job/plan result and puts "QA: N
  critical" in the Slack plan notification, 4. fails safe (original bytes,
  no verdict) on any gate error.
"""

from __future__ import annotations

import io
import logging
import sys
import time
import zipfile
from pathlib import Path
from unittest import mock

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TESTS_DIR = Path(__file__).resolve().parent
for _p in (PROJECT_ROOT, TESTS_DIR):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import app as app_module  # noqa: E402
import bundle_qa  # noqa: E402
import competitor_claims  # noqa: E402
import slack_plan_notifier  # noqa: E402

LEAK = "AWP Safety is a company in the construction_real_estate industry."
CLAIM = "Why: Expect Hilton to keep pressure on UK cab-driver hiring all year."
FULL = "The Accuracy International rifle is made by the British company Accuracy International."
MID_WORD_CUT = "The Accuracy International rifle is made by the Brit…"
WORD_CUT = "The Accuracy International rifle is made by the British…"


def _deck(*texts: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, t in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5 + i), Inches(8), Inches(0.8))
        para = box.text_frame.paragraphs[0]
        label, _, value = t.partition("|")
        if value:  # label run + value run, like slide 8's profile rows
            para.add_run().text = label
            para.add_run().text = value
        else:
            para.add_run().text = t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _workbook(cells: dict, with_chart: bool = True) -> bytes:
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Market Intelligence"
    for ref, val in cells.items():
        ws[ref] = val
    if with_chart:
        for r in range(1, 4):
            ws.cell(row=r, column=8, value=r * 10)
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=8, min_row=1, max_row=3))
        chart.title = "Spend by channel"
        ws.add_chart(chart, "J2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cells(xlsx: bytes) -> dict:
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(xlsx))["Market Intelligence"]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}


def _deck_lines(pptx: bytes) -> list:
    from pptx import Presentation

    out = []
    for slide in Presentation(io.BytesIO(pptx)).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out += [p.text for p in sh.text_frame.paragraphs]
    return out


DATA = {"client_name": "AWP Safety", "locations": ["Wheeling, WV"]}


# ---------------------------------------------------------------------------
# 1-2. repair + re-lint
# ---------------------------------------------------------------------------
def test_gate_repairs_snake_case_and_claims_in_the_bytes_that_ship():
    pptx = _deck("Description:  |" + LEAK, CLAIM)
    xlsx = _workbook({"D57": LEAK, "D58": "Hilton keeps pressure on UK hiring."})
    res = bundle_qa.gate_bundle(pptx, xlsx, DATA)
    s = res["summary"]
    assert s["initial_critical_count"] >= 4
    assert s["critical_count"] == 0, s["critical_ids"]
    assert s["repairs"]["snake_case_leak"] >= 2
    assert s["repairs"]["unsourced_competitor_claim"] >= 2

    lines = _deck_lines(res["pptx_bytes"])
    assert "Description:  AWP Safety is a company in the Construction & Real Estate industry." in lines
    assert f"Why: {competitor_claims.NO_EVIDENCE_LINE}" in lines
    cells = _cells(res["xlsx_bytes"])
    assert cells["D57"] == "AWP Safety is a company in the Construction & Real Estate industry."
    assert cells["D58"] == competitor_claims.NO_EVIDENCE_LINE


def test_xlsx_repair_leaves_every_other_package_part_byte_identical():
    xlsx = _workbook({"D57": LEAK})
    res = bundle_qa.gate_bundle(None, xlsx, DATA)
    with zipfile.ZipFile(io.BytesIO(xlsx)) as a, zipfile.ZipFile(io.BytesIO(res["xlsx_bytes"])) as b:
        assert a.namelist() == b.namelist()
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert changed and all(
        n == "xl/sharedStrings.xml" or n.startswith("xl/worksheets/sheet") for n in changed
    ), changed
    assert any(n.startswith("xl/charts/") for n in a.namelist())  # chart survived untouched


def test_mid_word_cut_is_confirmed_and_trimmed_but_a_word_boundary_cut_is_not_flagged():
    xlsx = _workbook({"D10": MID_WORD_CUT, "D11": FULL}, with_chart=False)
    findings = bundle_qa.run_bundle_qa(None, xlsx, DATA)
    assert [f for f in findings if f["code"] == "mid_word_truncation"]
    res = bundle_qa.gate_bundle(None, xlsx, DATA)
    assert _cells(res["xlsx_bytes"])["D10"] == "The Accuracy International rifle is made by the …"
    clean = _workbook({"D10": WORD_CUT, "D11": FULL}, with_chart=False)
    assert not [
        f for f in bundle_qa.run_bundle_qa(None, clean, DATA) if f["code"] == "mid_word_truncation"
    ]


def test_statute_number_is_not_a_raw_float_but_a_float_artifact_is():
    xlsx = _workbook(
        {
            "H6": "Notice: 30d; Argentine Labour Contract Law (Ley 20.744) governs employment",
            "H7": "CPA ratio 0.30000000000000004",
        },
        with_chart=False,
    )
    hits = [f["location"] for f in bundle_qa.run_bundle_qa(None, xlsx, DATA) if f["code"] == "raw_float_precision"]
    assert hits == ["Market Intelligence!H7"]


# ---------------------------------------------------------------------------
# Gate accuracy: the client is not its own competitor (prod Hershey
# 2026-09-24 logged a critical against competitor 'Hershey'; one finding per
# cell let that false positive mask the real competitor claim).
# ---------------------------------------------------------------------------
CLIENT_SENTENCE = (
    "Against a competitive skilled-trades labor market where Hershey is "
    "drawing from the same electrical, HVAC, and industrial maintenance "
    "talent as other regional manufacturers, this plan concentrates spend "
    "on programmatic and job-board channels near each plant."
)
COMPETITOR_SENTENCE = (
    "The key risk to monitor is qualified-to-interview conversion, and this "
    "risk is compounded by named competitors Nestle Purina, Campbell's, Land "
    "O'Lakes, Treehouse Foods drawing from the same maintenance and "
    "sanitation talent pools."
)


def _units(*texts):
    return [bundle_qa._TextUnit(t, f"Executive Summary!B{58 + i}") for i, t in enumerate(texts)]


def test_gate_does_not_flag_the_client_as_its_own_competitor():
    findings: list = []
    bundle_qa._check_unsourced_competitor_claim(
        _units(CLIENT_SENTENCE), findings, "The Hershey Company"
    )
    assert findings == []


def test_gate_still_flags_the_real_competitor_claim_behind_a_client_sentence():
    findings: list = []
    bundle_qa._check_unsourced_competitor_claim(
        _units(f"{CLIENT_SENTENCE} {COMPETITOR_SENTENCE}"), findings, "The Hershey Company"
    )
    assert len(findings) == 1
    assert "Hershey" not in findings[0]["message"].split("competitor", 1)[1].split(":", 1)[0]
    assert "Nestle Purina" in findings[0]["message"]


def test_run_bundle_qa_passes_the_client_name_through():
    """End-to-end through run_bundle_qa on a real workbook cell."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Executive Summary"
    wb.active["B58"] = CLIENT_SENTENCE
    buf = io.BytesIO()
    wb.save(buf)
    findings = bundle_qa.run_bundle_qa(None, buf.getvalue(), {"client_name": "THE HERSHEY COMPANY"})
    assert not [f for f in findings if f["code"] == "unsourced_competitor_claim"]



# ---------------------------------------------------------------------------
# 3. deliver anyway, never silently
# ---------------------------------------------------------------------------
_UNREPAIRABLE = {
    "severity": "critical",
    "code": "executive_summary_budget_footing",
    "message": "Budget does not foot",
    "location": "Executive Summary!D25",
}


def test_unrepairable_critical_is_reported_with_ids_and_bytes_still_delivered():
    xlsx = _workbook({"D1": "fine"}, with_chart=False)
    with mock.patch("bundle_qa.run_bundle_qa", return_value=[_UNREPAIRABLE]):
        res = bundle_qa.gate_bundle(None, xlsx, DATA)
    assert res["xlsx_bytes"] == xlsx
    assert res["summary"]["critical_count"] == 1
    assert res["summary"]["critical_ids"] == ["executive_summary_budget_footing@Executive Summary!D25"]


def test_app_gate_logs_surviving_criticals_at_error_with_ids(caplog):
    xlsx = _workbook({"D1": "fine"}, with_chart=False)
    with mock.patch("bundle_qa.run_bundle_qa", return_value=[_UNREPAIRABLE]):
        with caplog.at_level(logging.ERROR):
            _p, _x, summary = app_module._run_bundle_qa_gate(None, xlsx, DATA, "job123")
    errors = [r.getMessage() for r in caplog.records if r.levelno >= logging.ERROR]
    assert any(
        "critical finding(s) remain after repair" in m
        and "job123" in m
        and "executive_summary_budget_footing@Executive Summary!D25" in m
        for m in errors
    ), errors
    assert summary["critical_count"] == 1
    fields = app_module._slack_qa_fields(summary)
    assert fields["qa_critical_count"] == 1
    assert fields["qa_critical_codes"] == ["executive_summary_budget_footing"]


def test_slack_message_carries_a_qa_line():
    payload = {
        "client_name": "AWP Safety",
        "qa_critical_count": 2,
        "qa_critical_codes": ["snake_case_leak", "unsourced_competitor_claim"],
        "qa_critical_ids": ["snake_case_leak@slide 8"],
        "qa_repaired_count": 3,
    }
    line = slack_plan_notifier.qa_line(payload)
    assert line.startswith("QA: 2 critical (snake_case_leak, unsourced_competitor_claim)")
    assert "3 auto-repaired" in line
    blocks = slack_plan_notifier._build_plan_notification_blocks(payload)
    assert any("QA: 2 critical" in str(b) for b in blocks)
    # no verdict passed -> no QA claim at all
    assert slack_plan_notifier.qa_line({"client_name": "x"}) == ""


# ---------------------------------------------------------------------------
# 4. fail safe
# ---------------------------------------------------------------------------
def test_gate_crash_ships_original_bytes_with_no_verdict():
    xlsx = _workbook({"D57": LEAK}, with_chart=False)
    with mock.patch("bundle_qa.run_bundle_qa", side_effect=RuntimeError("boom")):
        res = bundle_qa.gate_bundle(None, xlsx, DATA)
    assert res["xlsx_bytes"] == xlsx and res["summary"] is None
    assert "boom" in res["error"]


def test_unreadable_repair_is_discarded():
    xlsx = _workbook({"D57": LEAK}, with_chart=False)
    with mock.patch("bundle_qa._repair_xlsx", return_value=b"not a zip"):
        res = bundle_qa.gate_bundle(None, xlsx, DATA)
    assert res["xlsx_bytes"] == xlsx
    assert res["summary"]["repairs"] == {}
    assert res["summary"]["critical_count"] >= 1


# ---------------------------------------------------------------------------
# End to end: the async job path wires the verdict into job + Slack
# ---------------------------------------------------------------------------
from test_bundle_qa_gate import (  # noqa: E402
    _poll_until_done,
    _submit_async,
    live_server,  # noqa: F401 -- pytest fixture
)


def test_async_job_records_ids_and_slack_gets_the_qa_line(live_server):  # noqa: F811
    sent: list = []
    with mock.patch("bundle_qa.run_bundle_qa", return_value=[_UNREPAIRABLE]), mock.patch(
        "slack_plan_notifier.notify_plan_generated", side_effect=sent.append
    ):
        job_id = _submit_async(live_server, "content-gate")
        result = _poll_until_done(live_server, job_id)
        deadline = time.time() + 10
        while not sent and time.time() < deadline:
            time.sleep(0.05)
    assert result["status"] == "completed"
    assert result["qa_critical_ids"] == ["executive_summary_budget_footing@Executive Summary!D25"]
    assert sent, "Slack notifier was not called"
    assert sent[0]["qa_critical_count"] == 1
    assert slack_plan_notifier.qa_line(sent[0]).startswith("QA: 1 critical")


# ---------------------------------------------------------------------------
# Verifier follow-ups (2026-10-01)
# ---------------------------------------------------------------------------
def test_gate_flags_a_competitor_claim_that_follows_the_client_name():
    findings: list = []
    bundle_qa._check_unsourced_competitor_claim(
        _units("Hershey and Mars are drawing from the same maintenance talent pool."),
        findings,
        "The Hershey Company",
    )
    assert len(findings) == 1 and "'Mars'" in findings[0]["message"]


def _truncated_workbook(n: int) -> bytes:
    """n trailing-ellipsis cells, each with its full text elsewhere -- the
    shape that made the old per-unit corpus scan quadratic."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Market Intelligence"
    for i in range(n):
        full = (
            f"Row {i} describes the regional hiring landscape for maintenance "
            f"technicians in market {i}"
        )
        ws.cell(row=i + 1, column=2, value=full[: 60 + (i % 7)] + "…")
        ws.cell(row=i + 1, column=3, value=full)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_truncation_classification_has_a_per_pass_budget():
    texts = [f"text number {i} with enough characters to be a needle" for i in range(5)]
    index = bundle_qa._TruncationIndex(texts)
    index.remaining = 2
    kinds = [
        bundle_qa._truncation_kind(f"text number {i} with enough characters to be a nee…", index)
        for i in range(4)
    ]
    assert kinds[:2] == ["mid_word", "mid_word"]
    assert kinds[2:] == ["unknown", "unknown"]  # budget spent -> warn stands, no rewrite


def test_truncation_scan_scales_linearly(monkeypatch):
    """6,000 truncated cells: pre-fix gate 46.6 s on the dev Mac (verifier:
    2,500 -> 51 s, 10,000 -> 613 s), post-fix ~4 s. The 20 s bound is
    generous for a loaded CI box. (Budget raised through the env so the
    timeout path cannot mask a slow scan.)"""
    monkeypatch.setenv("BUNDLE_QA_GATE_BUDGET_S", "120")
    blob = _truncated_workbook(6000)
    t0 = time.monotonic()
    res = bundle_qa.gate_bundle(None, blob, DATA)
    elapsed = time.monotonic() - t0
    assert res["summary"]["qa_status"] != "timeout"
    assert elapsed < 20, f"gate took {elapsed:.1f}s on 6,000 truncated cells"


def test_gate_over_budget_ships_original_bytes_and_says_so(caplog):
    blob = _truncated_workbook(3000)
    with caplog.at_level(logging.ERROR):
        res = bundle_qa.gate_bundle(None, blob, DATA, budget_s=0.05)
    assert res["xlsx_bytes"] == blob
    s = res["summary"]
    assert s["qa_status"] == "timeout" and s["codes"] == ["qa_gate_timeout"]
    assert s["critical_count"] == 0 and s["timed_out"] is True
    msgs = [r.getMessage() for r in caplog.records if r.levelno >= logging.ERROR]
    assert any("qa_gate_timeout" in m and "ORIGINAL bundle" in m for m in msgs), msgs


def test_gate_budget_has_an_env_override(monkeypatch):
    assert bundle_qa.gate_budget_s() == bundle_qa.GATE_BUDGET_S_DEFAULT == 8.0
    monkeypatch.setenv("BUNDLE_QA_GATE_BUDGET_S", "0.05")
    assert bundle_qa.gate_budget_s() == 0.05
    res = bundle_qa.gate_bundle(None, _truncated_workbook(3000), DATA)
    assert res["summary"]["qa_status"] == "timeout"
    monkeypatch.setenv("BUNDLE_QA_GATE_BUDGET_S", "not-a-number")
    assert bundle_qa.gate_budget_s() == 8.0


def test_timeout_reaches_the_job_record_and_the_slack_line(caplog):
    blob = _truncated_workbook(3000)
    with mock.patch.dict("os.environ", {"BUNDLE_QA_GATE_BUDGET_S": "0.05"}):
        with caplog.at_level(logging.ERROR):
            _p, xlsx, summary = app_module._run_bundle_qa_gate(None, blob, DATA, "job-timeout")
    assert xlsx == blob
    fields = app_module._bundle_qa_response_fields(summary)
    assert fields["qa_status"] == "timeout"
    assert fields["qa_codes"] == ["qa_gate_timeout"]
    line = slack_plan_notifier.qa_line(app_module._slack_qa_fields(summary))
    assert line.startswith("QA: not checked -- qa_gate_timeout")
    msgs = [r.getMessage() for r in caplog.records if r.levelno >= logging.ERROR]
    assert any("qa_gate_timeout for job-timeout" in m for m in msgs), msgs


def test_repair_humanises_a_supply_tier_in_sentence_case():
    """design review 2026-10-01: the repair pass wrote 'a Critically Scarce
    talent-supply tier' (Title Case mid-sentence)."""
    pptx = _deck("Seattle, WA was prioritized based on a critically_scarce talent-supply tier.")
    res = bundle_qa.gate_bundle(pptx, None, DATA)
    lines = _deck_lines(res["pptx_bytes"])
    assert "Seattle, WA was prioritized based on a critically scarce talent-supply tier." in lines
