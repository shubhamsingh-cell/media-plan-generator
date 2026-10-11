"""Plan-output audit (2026-10-11) -- N2: industry-wide niche boards must be
gated on the plan's actual roles (RTI International / healthcare_medical:
Coordinator, Worker, RN)."""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402

import excel_v2  # noqa: E402

RTI_ROLES = ["Coordinator", "Worker", "RN"]


def _sheet_text(ws) -> str:
    return "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value not in (None, "")
    )


def _niche_sheet_text(roles, industry) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_niche_board_matching(
        ws,
        {
            "industry": industry,
            "target_roles": roles,
            "locations": ["Chicago, IL"],
            "client_name": "Audit Co",
        },
    )
    return _sheet_text(ws)


# -- N2 ---------------------------------------------------------------------


def test_n2_physician_boards_not_printed_without_a_physician_role():
    text = _niche_sheet_text(RTI_ROLES, "healthcare_medical")
    assert "PracticeLink" not in text.replace("practicelink.com", "")
    assert "JAMA Career Center" not in text
    # generic healthcare board + the nursing boards the RN role justifies stay
    assert "Health eCareers" in text
    assert "Vivian Health" in text


def test_n2_physician_boards_kept_when_a_physician_role_exists():
    text = _niche_sheet_text(["Hospitalist", "RN"], "healthcare_medical")
    assert "PracticeLink" in text
    assert "JAMA Career Center" in text


def test_n2_channels_strategy_niche_section_is_role_gated():
    roles = RTI_ROLES
    boards = excel_v2._industry_niche_boards_for_roles("healthcare_medical", roles)
    assert "PracticeLink" not in boards and "JAMA Career Center" not in boards
    assert "Health eCareers" in boards
    # non-gated industries are returned unchanged
    assert excel_v2._industry_niche_boards_for_roles(
        "finance_banking", ["Analyst"]
    ) == excel_v2.INDUSTRY_NICHE_CHANNELS["finance_banking"]




# -- N4 ---------------------------------------------------------------------


def _channels_rationales(industry, roles):
    data = {
        "client_name": "Audit Co",
        "industry": industry,
        "budget": "$200,000",
        "locations": ["Chicago, IL"],
        "roles": roles,
        "target_roles": roles,
        "_enriched": {},
        "_synthesized": {},
        "_budget_allocation": {
            "channel_allocations": {
                "niche_boards": {"dollar_amount": 70000, "percentage": 35, "cpc": 1.4},
                "programmatic_dsp": {
                    "dollar_amount": 60000,
                    "percentage": 30,
                    "cpc": 1.2,
                },
            }
        },
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    excel_v2._build_sheet_channels(ws, data)
    out = {}
    for row in ws.iter_rows():
        label = next((c.value for c in row if c.value == "Niche / Industry Boards"), None)
        for c in row:
            if isinstance(c.value, str) and "% of budget" in c.value:
                out["niche" if label else "other"] = c.value
    return out


def test_n4_niche_channel_does_not_claim_to_target_unmatched_roles():
    out = _channels_rationales("healthcare_medical", ["Coordinator", "Worker"])
    assert "niche" in out, out
    assert "targets" not in out["niche"], out["niche"]
    # non-niche channels keep the plan-wide roles clause
    assert "targets Coordinator, Worker" in out["other"], out["other"]


def test_n4_niche_channel_names_only_roles_that_matched_a_board():
    out = _channels_rationales("healthcare_medical", ["Coordinator", "Worker", "RN"])
    assert "targets RN" in out["niche"], out["niche"]
    assert "Coordinator" not in out["niche"] and "Worker" not in out["niche"]


def test_gated_board_table_references_real_boards_and_categories():
    all_names = {n for v in excel_v2.INDUSTRY_NICHE_CHANNELS.values() for n in v}
    for name, cats in excel_v2._ROLE_GATED_INDUSTRY_BOARDS.items():
        assert name in all_names, name
        assert cats <= set(excel_v2.ROLE_NICHE_BOARDS), (name, cats)
