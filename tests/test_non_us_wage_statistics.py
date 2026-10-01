"""US-only wage/earnings figures never print on a plan with no US market.

Shipped defects (prod sweep 2026-10-01: UK, India, Japan decks; reproduced
locally for UK/India/Japan/Germany/Brazil with the production QCEW feed
injected):

  * Competitive Landscape slide, INDUSTRY HIRING TRENDS: "Avg weekly wage:
    US$2,751". That is BLS QCEW's US national private-sector average weekly
    wage for the industry's NAICS sector (api_enrichment.
    fetch_industry_employment reads the area_fips US000 row), printed as a
    UK/India/Japan/Germany/Brazil plan's own industry trend.
  * Market Intelligence > Workforce Trends: US pay statistics from
    data/workforce_trends_intelligence.json (US graduate starting-salary
    survey, US AI/ML engineer salary average and senior median, US gig
    workers' hourly pay / freelancer salary / earnings) survived the
    sentence-level US-only filter because none of them carries a US-only
    marker, so they printed as "US$101,500", "US$206,000", "US$16.67" ...
    on every non-US workbook.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation

import budget_engine
import data_synthesizer
import excel_v2
import ppt_generator
import tools_regen_bundles as T
import us_only_markers
from bundle_qa import run_bundle_qa

_KB_PATH = Path(__file__).resolve().parent.parent / "data" / "workforce_trends_intelligence.json"

# Every US pay figure the workforce KB carries, by KB path. Each is a US
# survey/statistic (ZipRecruiter US graduates, US AI/ML engineer pay, US
# gig workforce) -- none describes a UK/India/Japan/Germany/Brazil market.
_US_PAY_STATS = (
    ("generational_trends", "gen_z", "salary_expectations", "expected_starting_salary"),
    ("generational_trends", "gen_z", "salary_expectations", "actual_average_starting_salary"),
    ("generational_trends", "gen_z", "salary_expectations", "expectation_gap"),
    ("generational_trends", "gen_z", "salary_expectations", "financial_success_threshold"),
    ("generational_trends", "gen_z", "salary_expectations", "middle_class_perception"),
    ("generational_trends", "gen_z", "salary_expectations", "financial_health_requirement"),
    ("job_type_trends", "growing_demand", "ai_ml_engineers", "salary_average"),
    ("job_type_trends", "growing_demand", "ai_ml_engineers", "senior_level_median"),
    ("job_type_trends", "gig_economy", "us_workforce", "high_earners"),
    ("job_type_trends", "gig_economy", "us_workforce", "average_hourly_pay"),
    ("job_type_trends", "gig_economy", "us_workforce", "average_freelancer_salary"),
    ("job_type_trends", "gig_economy", "us_workforce", "skilled_freelancer_earnings"),
)

# The dollar figure each statistic prints (as it reads in the KB).
_US_PAY_FIGURES = (
    "101,500", "68,400", "33,100", "587,797", "74,000", "177,633",
    "206,000", "240,000", "100K+", "16.67", "108,028", "1.5 trillion",
)


def _kb() -> dict:
    with open(_KB_PATH, encoding="utf-8") as f:
        return json.load(f)


def _kb_leaf(kb: dict, path: tuple[str, ...]):
    node = kb
    for key in path:
        node = node[key]
    return node


# --------------------------------------------------------------------------
# The shared marker list (us_only_markers -- bundle_qa uses the same list)
# --------------------------------------------------------------------------


@pytest.mark.parametrize("path", _US_PAY_STATS, ids=lambda p: p[-1])
def test_each_us_pay_statistic_carries_a_us_only_marker(path):
    """Exactly as the workbook flattens it ("Label: value")."""
    value = _kb_leaf(_kb(), path)
    sentence = excel_v2._flatten_value({path[-1]: value})
    assert us_only_markers.has_us_only_marker(sentence), sentence


@pytest.mark.parametrize(
    "text",
    [
        "Unadjusted gender pay gap",
        "Germany median gross salary (all professions)",
        "53,900 EUR (The Stepstone Group, 2025)",
        "€51K median (€36K-€77K) - Software Engineer [payscale.com]",
        "Salary averages and advanced filters. No additional hiring services.",
        "Average Hourly Rate: $39/hour",
        "Salary Premium: 12% premium at IC level vs non-AI roles",
        "AI Ml Engineers: Hiring Growth: 88% year-on-year growth in 2025",
        "Salary Range: €60K midpoint of published band (€55K-€65K)",
    ],
)
def test_markers_leave_local_and_global_text_alone(text):
    assert not us_only_markers.has_us_only_marker(text)


# --------------------------------------------------------------------------
# Workbook: Market Intelligence > Workforce Trends
# --------------------------------------------------------------------------


def _workbook_plan(country: str, currency: str) -> dict:
    roles = [{"title": "Customer Service Rep", "count": 20, "tier": "entry"}]
    locations = [{"city": "Test City", "state": "", "country": country}]
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=100_000,
        roles=roles,
        locations=locations,
        industry="general_entry_level",
        channel_percentages={"Indeed": 40, "LinkedIn": 30, "Google Search Ads": 20, "Programmatic Job Boards": 10},
        collar_type="white",
        campaign_start_month=9,
        locations_raw=[f"Test City, {country}"],
        plan_currency=currency,
    )
    symbol = {"GBP": "£", "EUR": "€", "INR": "₹"}.get(currency, "$")
    workforce = data_synthesizer.fuse_workforce_insights(
        {}, {"workforce_trends": _kb()}, "general_entry_level"
    )
    return {
        "client_name": f"Wage Test {currency}",
        "industry": "general_entry_level",
        "budget": f"{symbol}100,000",
        "campaign_duration": "3 months",
        "campaign_start_month": 9,
        "roles": [r["title"] for r in roles],
        "target_roles": roles,
        "locations": [f"Test City, {country}"],
        "competitors": [],
        "_budget_allocation": alloc,
        "plan_currency": currency,
        "_synthesized": {"workforce_insights": workforce},
    }


def _market_intelligence(data: dict) -> tuple[bytes, str]:
    raw = excel_v2.generate_excel_v2(dict(data))
    raw = raw[0] if isinstance(raw, tuple) else raw
    ws = openpyxl.load_workbook(io.BytesIO(raw))["Market Intelligence"]
    return raw, "\n".join(str(c.value) for r in ws.iter_rows() for c in r if c.value is not None)


@pytest.mark.parametrize(
    "country,currency",
    [("United Kingdom", "GBP"), ("India", "INR"), ("Germany", "EUR")],
)
def test_non_us_workbook_prints_no_us_pay_statistic(country, currency):
    data = _workbook_plan(country, currency)
    raw, text = _market_intelligence(data)
    leaked = [fig for fig in _US_PAY_FIGURES if fig in text]
    assert not leaked, leaked
    # The genuinely-global lines of the same section still print.
    assert "AI Ml Engineers: Hiring Growth: 88% year-on-year growth in 2025" in text
    us_data = [
        f
        for f in run_bundle_qa(None, raw, data)
        if f.get("code") == "us_data_on_non_us_plan" and f.get("severity") == "critical"
    ]
    assert not us_data, us_data


def test_us_workbook_keeps_its_us_pay_statistics():
    _, text = _market_intelligence(_workbook_plan("United States", "USD"))
    for fig in _US_PAY_FIGURES:
        assert fig in text, fig


# --------------------------------------------------------------------------
# Deck: Competitive Landscape > INDUSTRY HIRING TRENDS
# --------------------------------------------------------------------------

# fetch_industry_employment's return shape (BLS QCEW, US000 national row).
_QCEW = {
    "source": "BLS QCEW",
    "total_employed": 6_120_000,
    "avg_weekly_wage": 2751,
    "avg_annual_wage": 2751 * 52,
    "establishments": 483_000,
    "sector_name": "Technology",
    "year": "2024",
    "naics": "51",
}


def _deck_text(locations: list[str], budget: str) -> str:
    brief = {
        "client_name": "Hoshino Electronics",
        "budget": budget,
        "campaign_duration": "3-6 months",
        "industry": "technology",
        "locations": locations,
        "roles": ["Software Engineer", "QA Engineer"],
        "target_roles": [{"title": r, "count": 5} for r in ("Software Engineer", "QA Engineer")],
    }
    data = T.build_plan_data(brief)
    enriched = {"industry_employment": dict(_QCEW)}
    data["_enriched"] = enriched
    data["_synthesized"] = data_synthesizer.synthesize(enriched, {}, data)
    pptx = ppt_generator.generate_pptx(dict(data))
    pptx = pptx.getvalue() if hasattr(pptx, "getvalue") else pptx
    return "\n".join(
        sh.text_frame.text
        for slide in Presentation(io.BytesIO(pptx)).slides
        for sh in slide.shapes
        if sh.has_text_frame
    )


@pytest.mark.parametrize(
    "location,budget",
    [
        ("London, UK", "£400,000"),
        ("Bengaluru, Karnataka, India", "₹25,000,000"),
        ("Tokyo, Japan", "¥45,000,000"),
        ("Stuttgart, Germany", "€300,000"),
        ("São Paulo, Brazil", "R$1,500,000"),
    ],
)
def test_non_us_deck_prints_no_us_weekly_wage(location, budget):
    text = _deck_text([location], budget)
    assert "INDUSTRY HIRING TRENDS" in text
    assert "weekly wage" not in text.lower()
    assert "2,751" not in text


def test_us_deck_keeps_the_qcew_weekly_wage():
    assert "Avg weekly wage: $2,751" in _deck_text(["Houston, TX"], "$250,000")


def test_mixed_plan_keeps_it_marked_as_us_dollars():
    """A US market is on the plan, so the US national figure describes part
    of it -- kept, and marked US$ beside the plan's GBP figures."""
    text = _deck_text(["New York, NY", "London, UK"], "£400,000")
    assert "Avg weekly wage: US$2,751" in text
