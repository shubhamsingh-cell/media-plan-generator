"""Enterprise hourly hiring pack: workbook sheets, deck appendix, request layer.

Fixtures are generic on purpose (a made-up "Example Client", Texas sites):
the repo is public, so no client, company or program name appears in code,
tests or copy.  ``test_no_client_names_in_new_files`` enforces that over every
file this feature adds, by hash so the test itself names nobody.
"""

from __future__ import annotations

import datetime
import hashlib
import io
import re
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

import budget_engine  # noqa: E402
import enterprise_pack as ep  # noqa: E402
import excel_v2  # noqa: E402
from _xl_eval import Evaluator  # noqa: E402

CHANNELS = {
    "Indeed": 30,
    "Programmatic Job Boards": 25,
    "Meta (Facebook/Instagram)": 15,
    "Google Search Ads": 10,
    "ZipRecruiter": 10,
    "Craigslist": 10,
}
TODAY = datetime.date(2026, 3, 20)


@pytest.fixture(autouse=True)
def _fixed_today(monkeypatch):
    monkeypatch.setattr(ep, "_today", lambda: TODAY)


def make_plan(
    enterprise: bool = True,
    excluded=None,
    sites: int = 3,
    weeks: int = 8,
    month: int = 4,
    channels=None,
    **extra,
) -> dict:
    roles = [{"title": "Warehouse Associate", "count": 40, "tier": "mid"}]
    names = ["Dallas", "Houston", "Austin", "San Antonio", "El Paso"][:sites]
    locs = [{"city": c, "state": "TX", "country": "United States"} for c in names]
    chans = dict(channels or CHANNELS)
    for ex in excluded or []:
        chans.pop(ex, None)
    alloc = budget_engine.calculate_budget_allocation(
        total_budget=150_000,
        roles=roles,
        locations=locs,
        industry="logistics",
        channel_percentages=chans,
        collar_type="blue",
        campaign_start_month=month,
    )
    data = {
        "client_name": "Example Client",
        "industry": "logistics",
        "budget": "$150,000",
        "campaign_duration": f"{weeks} weeks",
        "campaign_weeks": weeks,
        "campaign_start_month": month,
        "hire_volume": "40 hires",
        "locations": [f"{c}, TX" for c in names],
        "roles": ["Warehouse Associate"],
        "target_roles": roles,
        "_budget_allocation": alloc,
    }
    if excluded:
        data["excluded_publishers"] = list(excluded)
    if enterprise:
        data["enterprise_pack"] = True
    data.update(extra)
    return data


def build_wb(data: dict, data_only: bool = False):
    raw = excel_v2.generate_excel_v2(data)
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only)


def sheet_values(wb) -> dict:
    return {
        ws.title: [[c.value for c in row] for row in ws.iter_rows()]
        for ws in wb.worksheets
    }


def find_row(ws, text: str, col: int = 2) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == text:
            return r
    raise AssertionError(f"{text!r} not found in {ws.title}")


# ---------------------------------------------------------------------------
# Toggle off: nothing changes
# ---------------------------------------------------------------------------
def test_toggle_off_workbook_text_identical_to_a_standard_plan():
    base = sheet_values(build_wb(make_plan(enterprise=False)))
    stray = make_plan(
        enterprise=False,
        enterprise_pack=False,
        enterprise_week_start="monday",
        enterprise_system_of_record="x",
    )
    off = sheet_values(build_wb(stray))
    assert not set(ep.SHEET_TITLES) & set(base)
    assert off == base


def test_toggle_on_adds_exactly_the_five_sheets_and_leaves_the_rest():
    base = sheet_values(build_wb(make_plan(enterprise=False)))
    on = sheet_values(build_wb(make_plan()))
    assert [t for t in on if t in ep.SHEET_TITLES] == list(ep.SHEET_TITLES)
    for title, rows in base.items():
        assert on[title] == rows, f"{title} changed when the pack was switched on"


# ---------------------------------------------------------------------------
# Weekly Pacing: 3 sites x 8 weeks
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def pacing_wb():
    mp = pytest.MonkeyPatch()  # module scope runs before the autouse fixture
    mp.setattr(ep, "_today", lambda: TODAY)
    try:
        return build_wb(make_plan())
    finally:
        mp.undo()


def test_pacing_week_rows_dates_and_columns(pacing_wb):
    ws = pacing_wb[ep.SHEET_PACING]
    hdr = find_row(ws, "Week")
    headers = [ws.cell(row=hdr, column=c).value for c in range(2, 13)]
    assert headers == [
        "Week",
        "Calendar weight",
        "Recommended spend",
        "Minimum-test spend",
        "Surge spend",
        "Cumulative (recommended)",
        "Apply starts, low",
        "Apply starts, high",
        "Weekly goal",
        "Actual apply starts (enter)",
        "Status",
    ]
    first = hdr + 1
    labels = [ws.cell(row=first + i, column=2).value for i in range(8)]
    # 1 Apr 2026 is a Wednesday: week 1 is the first Sunday on/after it.
    assert labels[0].startswith("W01  Sun 5 Apr")
    assert labels[7].startswith("W08  Sun 24 May")
    assert ws.cell(row=first + 8, column=2).value == "Total"
    assert ws.cell(row=first + 9, column=2).value is None


def test_pacing_uses_real_formulas_and_thresholds_are_input_cells(pacing_wb):
    ws = pacing_wb[ep.SHEET_PACING]
    hdr = find_row(ws, "Week")
    r = hdr + 1
    rec, mn, sg, cum = (ws.cell(row=r, column=c).value for c in (4, 5, 6, 7))
    assert rec.startswith("=$C$") and "/SUM($C$" in rec
    assert mn == f"=D{r}*$C$10" and sg == f"=D{r}*$C$11"
    assert cum == f"=SUM($D${hdr + 1}:D{r})"
    status = ws.cell(row=r, column=12).value
    assert status.startswith("=IF(K") and '"GREEN"' in status and '"RED"' in status
    # thresholds + factors are shaded input cells with a "default" label
    for label, value in (
        ("Green threshold (% of weekly goal)", 0.9),
        ("Amber threshold (% of weekly goal)", 0.85),
        ("Minimum-test factor (x recommended)", 0.5),
        ("Surge factor (x recommended)", 1.5),
    ):
        row = find_row(ws, label)
        cell = ws.cell(row=row, column=3)
        assert cell.value == value
        assert cell.fill.start_color.rgb.endswith("FEF3C7"), label
        assert "default" in ws.cell(row=row, column=4).value
    assert find_row(ws, "Planned campaign budget")


def test_pacing_budget_ladder_rows(pacing_wb):
    ws = pacing_wb[ep.SHEET_PACING]
    hdr = find_row(ws, "Step")
    steps = [ws.cell(row=hdr + i, column=2).value for i in (1, 2, 3)]
    assert steps == ["Minimum test", "Recommended weekly", "Surge"]
    heads = [ws.cell(row=hdr, column=c).value for c in range(2, 10)]
    assert (
        "Apply starts / week, low" in heads and "Apply starts, campaign high" in heads
    )
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Basis: apply starts = spend / cost per apply start" in text


def test_pacing_site_matrix_splits_each_week_by_plan_weights(pacing_wb):
    ws = pacing_wb[ep.SHEET_PACING]
    share_row = find_row(ws, "Share of each week's spend")
    shares = [ws.cell(row=share_row, column=c).value for c in (3, 4, 5)]
    assert abs(sum(shares) - 1) < 0.0005
    assert (
        len({round(s, 3) for s in shares}) > 1
    )  # the plan's own weights, not 1/3 each
    hdr = share_row - 1
    assert [ws.cell(row=hdr, column=c).value for c in (2, 3, 4, 5, 6)] == [
        "Week",
        "Dallas, TX",
        "Houston, TX",
        "Austin, TX",
        "Check",
    ]
    assert ws.cell(row=share_row + 1, column=3).value.startswith("=$D")
    tot = find_row(ws, "Total", 2)  # first Total is the weekly table
    # the matrix total row follows the 8 week rows
    grid_total = share_row + 9
    assert ws.cell(row=grid_total, column=2).value == "Total"
    assert tot < grid_total


def test_every_cached_formula_value_equals_what_excel_computes():
    data_wb = build_wb(make_plan(excluded=["Craigslist"]), data_only=True)
    wb = build_wb(make_plan(excluded=["Craigslist"]))
    ev = Evaluator(wb)
    checked = 0
    for title in ep.SHEET_TITLES:
        ws = wb[title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    want = ev.cell(title, cell.coordinate)
                    got = data_wb[title][cell.coordinate].value
                    if want == "" or want is False:
                        assert got in (None, ""), (title, cell.coordinate, got)
                    elif isinstance(want, str):
                        assert got == want, (title, cell.coordinate, got, want)
                    else:
                        assert got == pytest.approx(float(want), rel=1e-6, abs=1e-6), (
                            title,
                            cell.coordinate,
                            cell.value,
                            got,
                            want,
                        )
                    checked += 1
    assert checked > 150


def test_editing_an_input_recalculates_the_sheet():
    wb = build_wb(make_plan())
    ws = wb[ep.SHEET_PACING]
    bud_row = find_row(ws, "Planned campaign budget")
    hdr = find_row(ws, "Week")
    before = Evaluator(wb).cell(ep.SHEET_PACING, f"D{hdr + 1}")
    ws.cell(row=bud_row, column=3).value = ws.cell(row=bud_row, column=3).value * 2
    after = Evaluator(wb).cell(ep.SHEET_PACING, f"D{hdr + 1}")
    assert after == pytest.approx(before * 2)
    # status turns RED/AMBER/GREEN from the editable thresholds
    goal_row = hdr + 1
    ev = Evaluator(wb)
    goal = ev.cell(ep.SHEET_PACING, f"J{goal_row}")
    for actual, word in (
        (goal * 0.95, "GREEN"),
        (goal * 0.87, "AMBER"),
        (goal * 0.5, "RED"),
    ):
        ws.cell(row=goal_row, column=11).value = actual
        assert Evaluator(wb).cell(ep.SHEET_PACING, f"L{goal_row}") == word


def test_week_start_day_is_configurable():
    wb = build_wb(make_plan(enterprise_week_start="monday"))
    ws = wb[ep.SHEET_PACING]
    hdr = find_row(ws, "Week")
    assert ws.cell(row=hdr + 1, column=2).value.startswith("W01  Mon 6 Apr")
    assert "Monday to Sunday" in " ".join(
        str(c.value) for c in ws["B"] if c.value and "Fiscal week" in str(c.value)
    )


def test_no_cpa_range_omits_apply_starts_with_a_note():
    one = make_plan(channels={"Indeed": 100})
    ws = build_wb(one)[ep.SHEET_PACING]
    hdr = find_row(ws, "Week")
    heads = [ws.cell(row=hdr, column=c).value for c in range(2, 13)]
    assert "Apply starts, low" not in heads and "Status" not in heads
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Expected apply starts are omitted" in text


def test_single_site_plan_has_no_site_matrix():
    ws = build_wb(make_plan(sites=1))[ep.SHEET_PACING]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Single-site plan" in text
    with pytest.raises(AssertionError):
        find_row(ws, "Share of each week's spend")


# ---------------------------------------------------------------------------
# Publisher Tiers & Gates
# ---------------------------------------------------------------------------
def test_tiers_respect_client_exclusions_and_availability():
    import publisher_filter

    plan = make_plan(excluded=["Craigslist", "ZipRecruiter"])
    disclosure = publisher_filter.disclosure_line(plan)
    wb = build_wb(plan)
    text = {
        t: " ".join(str(c.value) for row in wb[t].iter_rows() for c in row if c.value)
        for t in ep.SHEET_TITLES
    }
    for title in ep.SHEET_TITLES:
        body = text[title].replace(disclosure, "").lower()
        assert "ziprecruiter" not in body, title
        assert "craigslist" not in body, title
    ws = wb[ep.SHEET_TIERS]
    names = [ws.cell(row=r, column=2).value for r in range(14, 19)]
    assert names[0] == "Indeed" and "ZipRecruiter" not in names
    assert disclosure in text[ep.SHEET_TIERS]


def test_tiers_sheet_formulas_and_default_tiers():
    wb = build_wb(make_plan(excluded=["Craigslist"]))
    ws = wb[ep.SHEET_TIERS]
    hdr = find_row(ws, "Publisher")
    row = hdr + 1
    assert ws.cell(row=row, column=7).value.startswith("=IF(F")
    assert ws.cell(row=row, column=9).fill.start_color.rgb.endswith(
        "FEF3C7"
    )  # editable minimum
    assert "'Weekly Pacing'!" in ws.cell(row=row, column=9).value
    tiers = [
        wb_v
        for wb_v in (
            Evaluator(wb).cell(ep.SHEET_TIERS, f"G{r}") for r in range(row, row + 5)
        )
    ]
    assert tiers[0] == "Tier 1" and set(tiers) <= {"Tier 1", "Tier 2", "Tier 3"}
    assert "Tier 2" in tiers
    # cut-offs are editable and drive the tier column
    ws["C9"].value = 0.1
    assert Evaluator(wb).cell(ep.SHEET_TIERS, f"G{row + 1}") == "Tier 2"


def test_switch_triggers_have_editable_thresholds_and_days():
    ws = build_wb(make_plan())[ep.SHEET_TIERS]
    r = find_row(ws, "Trigger")
    for i in (1, 2, 3):
        thr, days = ws.cell(row=r + i, column=3), ws.cell(row=r + i, column=4)
        assert thr.fill.start_color.rgb.endswith("FEF3C7")
        assert days.fill.start_color.rgb.endswith("FEF3C7") and isinstance(
            days.value, int
        )
    assert find_row(ws, "Gate")


# ---------------------------------------------------------------------------
# Tracking & Measurement
# ---------------------------------------------------------------------------
def test_tracking_convention_and_system_of_record():
    plan = make_plan(enterprise_system_of_record="client ATS export")
    ws = build_wb(plan)[ep.SHEET_TRACKING]
    r = find_row(ws, "System of record for applies")
    assert ws.cell(row=r, column=3).value == "client ATS export"
    assert (
        ws.cell(row=find_row(ws, "utm_campaign"), column=4).value
        == "example-client_dal_w01"
    )
    params = [ws.cell(row=x, column=4).value for x in range(1, ws.max_row + 1)]
    assert (
        "utm_source=indeed&utm_medium=jobboard&utm_campaign=example-client_dal_w01"
        in params
    )
    assert any(p and "utm_medium=paid-social" in p for p in params)
    assert ws.cell(row=find_row(ws, "Tolerance between systems"), column=3).value == 0.1
    # default text when the wizard gave none
    ws2 = build_wb(make_plan())[ep.SHEET_TRACKING]
    assert (
        ws2.cell(row=find_row(ws2, "System of record for applies"), column=3).value
        == ep.DEFAULT_SYSTEM_OF_RECORD
    )


def test_tracking_checklist_carries_the_lead_times():
    ws = build_wb(
        make_plan(enterprise_request_lock_day="wednesday", enterprise_setup_sla_days=3)
    )[ep.SHEET_TRACKING]
    lock = ws.cell(row=find_row(ws, "Request lock day"), column=4).value
    sla = ws.cell(row=find_row(ws, "Setup SLA"), column=4).value
    assert lock.startswith("Wednesday") and sla.startswith("3 business days")


# ---------------------------------------------------------------------------
# Feasibility + Job ad quality
# ---------------------------------------------------------------------------
def test_feasibility_statuses_are_cached_and_gate_result():
    wb = build_wb(make_plan(), data_only=True)
    ws = wb[ep.SHEET_FEASIBILITY]
    statuses = {
        ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value
        for r in range(1, ws.max_row + 1)
    }
    assert statuses["Supply in market: United States"] == "PASS"
    assert statuses["ATS redirect and tracking support"] == "CONFIRM"
    assert statuses["Scale vs target"] == "CHECK"  # projected 24-48 vs a goal of 40
    assert statuses["System of record named"] == "CONFIRM"
    assert statuses["Gate result"] == "READY AFTER CONFIRMATIONS"


def test_feasibility_inputs_flip_the_gate():
    wb = build_wb(make_plan())
    ws = wb[ep.SHEET_FEASIBILITY]
    for label, value in (("ATS redirect and tracking support", "No"),):
        r = find_row(ws, label)
        ws.cell(row=r, column=5).value = value
    assert (
        Evaluator(wb).cell(
            ep.SHEET_FEASIBILITY,
            f"C{find_row(ws, 'ATS redirect and tracking support')}",
        )
        == "FAIL"
    )
    gate = Evaluator(wb).cell(ep.SHEET_FEASIBILITY, f"C{find_row(ws, 'Gate result')}")
    assert gate == "NOT READY"


def test_feasibility_flags_a_thin_market():
    plan = make_plan(channels={"Indeed": 60, "ZipRecruiter": 40})
    ws = build_wb(plan, data_only=True)[ep.SHEET_FEASIBILITY]
    row = find_row(ws, "Supply in market: United States")
    assert ws.cell(row=row, column=3).value == "CHECK"  # 2 publishers < 3 minimum


def test_job_ad_quality_flags_detectable_issues():
    plan = make_plan(
        use_case="Hiring now. Sign-on bonus and flexible schedule at every site.",
        roles=["Sr. Associate (L4) / Packer"],
        target_roles=["Sr. Associate (L4) / Packer"],
    )
    ws = build_wb(plan)[ep.SHEET_JOBAD]
    rows = {
        ws.cell(row=r, column=2).value: ws.cell(row=r, column=4).value
        for r in range(1, ws.max_row + 1)
    }
    assert "No pay figure" in rows["Pay is shown"]
    assert "3 sites" in rows["No promise a site cannot deliver"]
    assert "sign-on bonus" in rows["No promise a site cannot deliver"]
    assert "Check these titles" in rows["Plain-language job title"]
    first = find_row(ws, "Pay is shown")
    assert ws.cell(row=first, column=5).fill.start_color.rgb.endswith("FEF3C7")


# ---------------------------------------------------------------------------
# Deck appendix
# ---------------------------------------------------------------------------
def build_deck(data: dict):
    import ppt_generator
    from pptx import Presentation

    return Presentation(io.BytesIO(ppt_generator.generate_pptx(dict(data))))


def slide_texts(prs) -> list:
    return [
        [
            sh.text_frame.text
            for sh in sl.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        for sl in prs.slides
    ]


def _collide_findings(prs, first_slide: int) -> list:
    sys.path.insert(0, str(PROJECT_ROOT / "scripts" / "deck_qa"))
    import collide

    out = []
    for i, slide in enumerate(prs.slides, 1):
        if i >= first_slide:
            out.extend(collide.analyze_slide(slide, i))
    return out


def test_deck_off_is_text_identical_and_on_adds_three_appendix_slides():
    base = build_deck(make_plan(enterprise=False))
    stray = build_deck(
        make_plan(
            enterprise=False, enterprise_pack=False, enterprise_week_start="monday"
        )
    )
    on = build_deck(make_plan())
    assert slide_texts(stray) == slide_texts(base)
    assert len(on.slides) == len(base.slides) + 3
    assert slide_texts(on)[: len(base.slides)] == slide_texts(base)
    titles = [" ".join(t[:3]) for t in slide_texts(on)[-3:]]
    assert (
        "Weekly Pacing" in titles[0]
        and "Publisher Tiers" in titles[1]
        and "Tracking" in titles[2]
    )


def test_deck_appendix_has_zero_collisions_and_reads_the_model():
    prs = build_deck(make_plan(excluded=["Craigslist"]))
    first = len(prs.slides) - 2
    assert _collide_findings(prs, first) == []
    text = " ".join(" ".join(t) for t in slide_texts(prs)[-3:])
    assert "GREEN: at or above 90%" in text and "RED: below 85%" in text
    assert "Indeed" in text and "Craigslist" not in text
    assert "example-client_dal_w01" in text
    # ladder numbers agree with the workbook's formulas
    model = ep.build_model(make_plan(excluded=["Craigslist"]), TODAY)
    weekly = ep.ladder(model)[1]["weekly"]
    wb = build_wb(make_plan(excluded=["Craigslist"]))
    ws = wb[ep.SHEET_PACING]
    row = find_row(ws, "Recommended weekly")
    assert Evaluator(wb).cell(ep.SHEET_PACING, f"D{row}") == pytest.approx(weekly)


def test_deck_appendix_stress_long_names_many_publishers_no_cpa():
    channels = {
        f"Regional Employment Network Number {i} Of The Group": 5 for i in range(1, 15)
    }
    channels["Indeed"] = 30
    prs = build_deck(
        make_plan(
            channels=channels,
            sites=5,
            enterprise_system_of_record="x" * 110,
        )
    )
    assert _collide_findings(prs, len(prs.slides) - 2) == []
    one = build_deck(make_plan(channels={"Indeed": 100}))
    assert _collide_findings(one, len(one.slides) - 2) == []
    text = " ".join(" ".join(t) for t in slide_texts(one)[-3:])
    assert "apply starts are omitted" in text


# ---------------------------------------------------------------------------
# Request layer
# ---------------------------------------------------------------------------
def test_request_validation_and_caps():
    ok = {
        "enterprise_pack": True,
        "enterprise_week_start": "Sunday",
        "enterprise_setup_sla_days": "5",
    }
    assert ep.request_problem(ok) == ("", "")
    assert ep.request_problem({"enterprise_week_start": "someday"})[0] == "week_start"
    assert ep.request_problem({"enterprise_request_lock_day": "9"})[0] == "lock_day"
    assert ep.request_problem({"enterprise_setup_sla_days": "abc"})[0] == "sla"
    assert ep.request_problem({"enterprise_setup_sla_days": 31})[0] == "sla"
    assert ep.request_problem({"enterprise_setup_sla_days": -1})[0] == "sla"
    assert (
        ep.request_problem({"enterprise_system_of_record": "x" * 121})[0] == "too_long"
    )
    assert ep.request_problem({"enterprise_system_of_record": ["x"]})[0] == "type"


def test_normalize_request_canonicalises_when_on_and_clears_when_off():
    d = {
        "enterprise_pack": "true",
        "enterprise_week_start": "MONDAY",
        "enterprise_system_of_record": "  ATS  ",
        "enterprise_request_lock_day": "",
        "enterprise_setup_sla_days": "4",
    }
    ep.normalize_request(d)
    assert d == {
        "enterprise_pack": True,
        "enterprise_week_start": "monday",
        "enterprise_system_of_record": "ATS",
        "enterprise_request_lock_day": "",
        "enterprise_setup_sla_days": 4,
    }
    off = {
        "enterprise_pack": False,
        "enterprise_week_start": "monday",
        "client_name": "X",
    }
    ep.normalize_request(off)
    assert off == {"client_name": "X"}


# ---------------------------------------------------------------------------
# Public-repo hygiene
# ---------------------------------------------------------------------------
_BANNED_SHA256 = {
    "cbc62794911ff31b2864ecd3dbbbee7ebcb7ea41c5a42e2cba377f3cfdb42811",
    "77909277a1ffc4b06a39d95f00d94b495a5c099351155c60d552c4dc8822cf9a",
    "800a7de021cad393effb24d5817c31d10361c56c1131d73672813f2312a35a88",
    "222f930b87524636da58b1987eabd52f1285cb0ef7dad72f28fcd026fbb3f73f",
    "adccfa16b4fc2c779ad3c361d448aedddb000cd45ab3f91bef0bc4de706306c9",
    "38816a4bbe78aa3dbff468dc0d08c217df8cc3b4f3c3fd9f139bab84d69dff5c",
}


def test_no_client_names_in_new_files():
    files = [
        PROJECT_ROOT / "enterprise_pack.py",
        PROJECT_ROOT / "tests" / "test_enterprise_pack.py",
        PROJECT_ROOT / "tests" / "_xl_eval.py",
    ]
    files += sorted((PROJECT_ROOT / "tests").glob("test_enterprise_*.py"))
    bad = []
    for path in dict.fromkeys(files):
        for word in set(
            re.findall(r"[A-Za-z0-9]+", path.read_text(encoding="utf-8").lower())
        ):
            if hashlib.sha256(word.encode()).hexdigest() in _BANNED_SHA256:
                bad.append((path.name, "<redacted>"))
    assert not bad, bad
