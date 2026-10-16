"""Export and delivery route handlers (GET + POST).

Extracted from app.py to reduce its size.  Handles:
- POST /api/export/sheets
- POST /api/export/status (GET-like, via POST)
- POST /api/export/pdf  -- generate PDF report from plan data
- POST /api/deliver
- POST /api/report/html
- POST /api/nova/export
- GET  /api/plan/export/pdf?plan_id=xxx  -- download PDF for stored plan
"""

from __future__ import annotations

import json
import logging
import os
import re
import sys
import time
from typing import Any

from shared_utils import format_industry_label

logger = logging.getLogger(__name__)


def _allowed_zip_dirs() -> list[str]:
    """Return realpath'd directories from which ZIP attachments may be read.

    Generated ZIPs are written to ``<repo>/data/generated_docs`` (see app.py).
    Only files resolving inside that directory are permitted as email
    attachments -- a client-controlled ``zip_file_path`` must not be able to
    exfiltrate arbitrary files (e.g. /etc/passwd).
    """
    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return [os.path.realpath(os.path.join(repo_root, "data", "generated_docs"))]


def _safe_zip_attachment_path(zip_file_path: Any) -> str | None:
    """Validate a client-supplied ZIP path against the allowed output dir.

    Resolves symlinks/``..`` with ``os.path.realpath`` and accepts the path
    only if it lives inside an allowed directory. On violation, logs a warning
    and returns ``None`` (caller proceeds WITHOUT an attachment rather than
    raising). Returns the original path string when it is legitimate.
    """
    if not zip_file_path or not isinstance(zip_file_path, str):
        return None
    try:
        resolved = os.path.realpath(zip_file_path)
    except (OSError, ValueError) as exc:
        logger.warning(
            "deliver: could not resolve zip_file_path %r: %s", zip_file_path, exc
        )
        return None
    for allowed in _allowed_zip_dirs():
        allowed_prefix = allowed + os.sep
        if resolved == allowed or resolved.startswith(allowed_prefix):
            return zip_file_path
    logger.warning(
        "deliver: rejected out-of-bounds zip_file_path %r (resolved=%r); "
        "sending email WITHOUT attachment",
        zip_file_path,
        resolved,
    )
    return None


# ---------------------------------------------------------------------------
# Route dispatch
# ---------------------------------------------------------------------------


def handle_export_get_routes(handler: Any, path: str, parsed: Any) -> bool:
    """Dispatch export-related GET routes.  Returns True if handled."""
    if path == "/api/plan/export/pdf":
        _handle_pdf_export_get(handler, path, parsed)
        return True
    return False


def handle_export_post_routes(handler: Any, path: str, parsed: Any) -> bool:
    """Dispatch export/delivery POST routes.  Returns True if handled."""
    _fn = _EXPORT_POST_ROUTE_MAP.get(path)
    if _fn is not None:
        _fn(handler, path, parsed)
        return True
    return False


# ---------------------------------------------------------------------------
# Individual route handlers
# ---------------------------------------------------------------------------


def _track_export_event(plan_data: dict, export_format: str) -> None:
    """Fire PostHog media_plan_exported event (non-blocking)."""
    try:
        from posthog_tracker import track_plan_export

        track_plan_export(
            email=plan_data.get("requester_email", "anonymous"),
            export_format=export_format,
            client=plan_data.get("client_name") or "",
        )
    except Exception:
        pass


# ---------------------------------------------------------------------------
# CSV / XLSX export of the GENERATED plan (wizard audit D-07)
# ---------------------------------------------------------------------------
# The wizard used to post its own REQUEST (budget text, role/location tags, no
# channels) as ``plan_data``: the CSV came out header-only and the XLSX carried
# "Number of Channels 0", empty Channels/Budget sheets and a hard-coded 4-week
# ramp. These two formats now export the plan the generator STORED under
# ``plan_id`` (app._store_plan_result, 24 h) -- its funded channels, dollar
# amounts and recommendations -- and refuse a payload that is not a plan
# instead of writing an empty file.

_PLAN_ID_RE = re.compile(r"^([a-f0-9]{12}|[a-f0-9]{32})\Z")

_PLAN_EXPORT_CHANNEL_HEADER = [
    "Channel",
    "Share of budget (%)",
    "Budget",
    "CPC",
    "CPA",
    "Projected clicks",
    "Rationale",
]


def _stored_plan_json(plan_id: str) -> dict | None:
    """The plan JSON the generator stored under ``plan_id``, or None when it
    is unknown or past its 24 h life. Reads app.py's own store (this worker's
    dict, then the shared layers) -- no app.py change needed."""
    _app = sys.modules.get("app") or sys.modules.get("__main__")
    lookup = getattr(_app, "_plan_result_lookup", None)
    if lookup is None:
        return None
    entry = lookup(plan_id)
    if not entry or time.time() - (entry.get("created") or 0) > 86400:
        return None
    return entry.get("data") or None


def _as_float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _channel_display_name(stored_name: str) -> str:
    """The workbook's name for a stored channel ("Programmatic Dsp" ->
    "Programmatic (DSP)"). The stored name is the engine key, spaced and
    title-cased, so lower-casing and re-joining recovers the key."""
    try:
        from display_format import channel_label

        return channel_label(stored_name.strip().lower().replace(" ", "_")) or (
            stored_name
        )
    except (ImportError, AttributeError):
        return stored_name


def _plan_export_sheets(
    plan_json: dict, request: dict
) -> dict[str, list[list[Any]]]:
    """Sheet name -> rows for the stored plan ``plan_json``.

    ``request`` supplies what the stored plan does not carry (client name,
    industry label). Text cells go through ``sheets_export._safe_str`` so a
    value starting with = + - @ cannot run as a spreadsheet formula."""
    from datetime import datetime, timezone

    from sheets_export import _safe_str

    meta = plan_json.get("metadata") or {}
    channels = [c for c in (plan_json.get("channels") or []) if isinstance(c, dict)]
    total = _as_float(
        meta.get("total_budget") or (plan_json.get("budget_summary") or {}).get("total")
    )
    industry = (
        request.get("industry_label")
        or meta.get("industry")
        or request.get("industry")
        or ""
    )

    def _join(values: Any) -> str:
        if not isinstance(values, list):
            values = [values] if values else []
        return ", ".join(str(v) for v in values if v)

    summary: list[list[Any]] = [
        ["Field", "Value"],
        ["Client", _safe_str(request.get("client_name") or "Client")],
        ["Industry", _safe_str(industry)],
        ["Total campaign budget", total],
        ["Location(s)", _safe_str(_join(meta.get("locations")))],
        ["Role(s)", _safe_str(_join(meta.get("roles")))],
        ["Funded channels", len(channels)],
        ["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")],
    ]

    chan_rows: list[list[Any]] = [list(_PLAN_EXPORT_CHANNEL_HEADER)]
    for ch in channels:
        chan_rows.append(
            [
                _safe_str(_channel_display_name(ch.get("name") or "")),
                round(_as_float(ch.get("allocation_pct")), 1),
                round(_as_float(ch.get("budget")), 2),
                _safe_str(ch.get("cpc_range") or ""),
                _safe_str(ch.get("cpa_range") or ""),
                int(_as_float(ch.get("projected_clicks"))),
                _safe_str(ch.get("recommended_reason") or ""),
            ]
        )
    if channels:
        funded = sum(_as_float(c.get("budget")) for c in channels)
        # Each row's share is rounded to 0.1, so their sum can read 100.1;
        # the total share comes from the dollars instead.
        total_share = (
            round(funded / total * 100, 1)
            if total > 0
            else round(sum(_as_float(c.get("allocation_pct")) for c in channels), 1)
        )
        chan_rows.append(
            [
                "Total",
                total_share,
                round(funded, 2),
                "",
                "",
                int(sum(_as_float(c.get("projected_clicks")) for c in channels)),
                "",
            ]
        )

    recs: list[list[Any]] = [["#", "Recommendation"]]
    for i, rec in enumerate(plan_json.get("recommendations") or [], start=1):
        recs.append([i, _safe_str(rec)])

    return {"Summary": summary, "Channels": chan_rows, "Recommendations": recs}


def _plan_export_csv(sheets: dict[str, list[list[Any]]]) -> bytes:
    """One CSV: each sheet as a titled block (UTF-8 with BOM for Excel)."""
    import csv
    import io

    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Media Plan Export"])
    for name, rows in sheets.items():
        writer.writerow([])
        writer.writerow([name.upper()])
        writer.writerows(rows)
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


def _plan_export_xlsx(sheets: dict[str, list[list[Any]]]) -> bytes | None:
    """XLSX workbook, one tab per sheet, real numeric cells. None when
    openpyxl is not installed (the caller falls back to CSV)."""
    try:
        import io

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    import joveo_brand_2026 as brand

    indigo = brand.INDIGO.lstrip("#")
    header_fill = PatternFill(start_color=indigo, end_color=indigo, fill_type="solid")
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        widths: dict[int, int] = {}
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                widths[c_idx] = max(widths.get(c_idx, 0), len(str(value or "")))
                if r_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                elif isinstance(value, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(value, int):
                    cell.number_format = "#,##0"
        for c_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = min(width + 4, 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _send_download(
    handler: Any, body: bytes, content_type: str, filename: str
) -> None:
    handler.send_response(200)
    handler.send_header("Content-Type", content_type)
    handler.send_header("Content-Disposition", f'attachment; filename="{filename}"')
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _export_plan_file(handler: Any, data: dict[str, Any], export_format: str) -> bool:
    """Handle ``format`` csv/xlsx for a stored plan.

    Returns True when a response was sent. Returns False for the legacy
    shape (a finished plan with channels posted as ``plan_data``) so the
    caller keeps its existing behaviour for API users; a request that is
    neither is answered 400 instead of exporting an empty file.
    """
    plan_id = str(data.get("plan_id") or "").strip()
    posted_raw = data.get("plan_data")
    posted: dict[str, Any] = posted_raw if isinstance(posted_raw, dict) else {}
    if not plan_id:
        posted_channels = posted.get("channels") or posted.get(
            "channel_recommendations"
        )
        if isinstance(posted_channels, list) and posted_channels:
            return False
        handler._send_json(
            {
                "error": (
                    "Nothing to export: send the plan_id of a generated plan "
                    "(the wizard request alone has no channels)."
                )
            },
            status_code=400,
        )
        return True
    if not _PLAN_ID_RE.match(plan_id):
        handler._send_json({"error": "Invalid plan_id"}, status_code=400)
        return True
    plan_json = _stored_plan_json(plan_id)
    if not plan_json:
        handler._send_json(
            {
                "error": (
                    "Plan not found or expired. Generate the plan again to "
                    "export it."
                )
            },
            status_code=404,
        )
        return True

    request = dict(posted)
    for key in ("client_name", "industry_label", "industry"):
        if data.get(key):
            request[key] = data[key]
    sheets = _plan_export_sheets(plan_json, request)
    slug = re.sub(r"[^a-zA-Z0-9_\-]", "_", str(request.get("client_name") or "Client"))
    _track_export_event(request, export_format)
    if export_format == "xlsx":
        xlsx = _plan_export_xlsx(sheets)
        if xlsx:
            _send_download(
                handler,
                xlsx,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"{slug}_Media_Plan_Export.xlsx",
            )
            return True
    _send_download(
        handler,
        _plan_export_csv(sheets),
        "text/csv; charset=utf-8",
        f"{slug}_Media_Plan_Export.csv",
    )
    return True


def _handle_export_sheets(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/export/sheets -- Sheets/CSV/XLSX export."""
    try:
        content_len = int(handler.headers.get("Content-Length") or 0)
        body = handler.rfile.read(content_len) if content_len > 0 else b"{}"
        data = json.loads(body)

        plan_data = data.get("plan_data") or data
        export_format = data.get("format") or "sheets"

        # CSV/XLSX export the generated plan (by plan_id), never the wizard
        # request (audit D-07). A posted finished plan falls through below.
        if export_format in ("csv", "xlsx") and _export_plan_file(
            handler, data, export_format
        ):
            return

        # PostHog: track export event
        _track_export_event(plan_data, export_format)

        from sheets_export import (
            export_media_plan,
            export_to_csv,
            export_to_xlsx,
        )

        if export_format == "sheets":
            sheet_url = export_media_plan(plan_data)
            if sheet_url:
                handler._send_json({"url": sheet_url, "format": "sheets"})
            else:
                # Fall back to XLSX or CSV
                xlsx_bytes = export_to_xlsx(plan_data)
                if xlsx_bytes:
                    client_name = re.sub(
                        r"[^a-zA-Z0-9_\-]",
                        "_",
                        plan_data.get("client_name") or "Client",
                    )
                    handler.send_response(200)
                    handler.send_header(
                        "Content-Type",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                    handler.send_header(
                        "Content-Disposition",
                        f'attachment; filename="{client_name}_Media_Plan_Export.xlsx"',
                    )
                    handler.send_header("Content-Length", str(len(xlsx_bytes)))
                    handler.send_header("X-Export-Fallback", "xlsx")
                    handler.end_headers()
                    handler.wfile.write(xlsx_bytes)
                else:
                    _send_csv_fallback(handler, plan_data)

        elif export_format == "xlsx":
            xlsx_bytes = export_to_xlsx(plan_data)
            if xlsx_bytes:
                client_name = re.sub(
                    r"[^a-zA-Z0-9_\-]",
                    "_",
                    plan_data.get("client_name") or "Client",
                )
                handler.send_response(200)
                handler.send_header(
                    "Content-Type",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                handler.send_header(
                    "Content-Disposition",
                    f'attachment; filename="{client_name}_Media_Plan_Export.xlsx"',
                )
                handler.send_header("Content-Length", str(len(xlsx_bytes)))
                handler.end_headers()
                handler.wfile.write(xlsx_bytes)
            else:
                _send_csv_fallback(handler, plan_data)

        elif export_format == "csv":
            _send_csv_response(handler, plan_data)

        else:
            handler.send_response(400)
            handler.send_header("Content-Type", "application/json")
            handler.end_headers()
            handler.wfile.write(
                json.dumps(
                    {
                        "error": f"Unknown format: {export_format}. Use 'sheets', 'xlsx', or 'csv'."
                    }
                ).encode()
            )

    except json.JSONDecodeError:
        handler.send_response(400)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "Invalid JSON body"}).encode())
    except Exception as e:
        logger.error("Export endpoint error: %s", e, exc_info=True)
        handler.send_response(500)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "Export failed"}).encode())


def _send_csv_fallback(handler: Any, plan_data: dict) -> None:
    """Send CSV as fallback when XLSX is unavailable."""
    from sheets_export import export_to_csv

    csv_bytes = export_to_csv(plan_data)
    client_name = re.sub(
        r"[^a-zA-Z0-9_\-]",
        "_",
        plan_data.get("client_name") or "Client",
    )
    handler.send_response(200)
    handler.send_header("Content-Type", "text/csv; charset=utf-8")
    handler.send_header(
        "Content-Disposition",
        f'attachment; filename="{client_name}_Media_Plan_Export.csv"',
    )
    handler.send_header("Content-Length", str(len(csv_bytes)))
    handler.send_header("X-Export-Fallback", "csv")
    handler.end_headers()
    handler.wfile.write(csv_bytes)


def _send_csv_response(handler: Any, plan_data: dict) -> None:
    """Send a CSV export response."""
    from sheets_export import export_to_csv

    csv_bytes = export_to_csv(plan_data)
    client_name = re.sub(
        r"[^a-zA-Z0-9_\-]",
        "_",
        plan_data.get("client_name") or "Client",
    )
    handler.send_response(200)
    handler.send_header("Content-Type", "text/csv; charset=utf-8")
    handler.send_header(
        "Content-Disposition",
        f'attachment; filename="{client_name}_Media_Plan_Export.csv"',
    )
    handler.send_header("Content-Length", str(len(csv_bytes)))
    handler.end_headers()
    handler.wfile.write(csv_bytes)


def _handle_export_status(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/export/status -- check export service status."""
    try:
        from sheets_export import get_status as sheets_status

        handler._send_json(sheets_status())
    except Exception as e:
        logger.error("Export status error: %s", e, exc_info=True)
        handler._send_json({"configured": False, "error": str(e)}, status_code=500)


def _handle_deliver(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/deliver -- email delivery of media plan."""
    try:
        content_len = int(handler.headers.get("Content-Length") or 0)
        body = handler.rfile.read(content_len) if content_len > 0 else b"{}"
        data = json.loads(body)
        from plan_delivery import send_plan_email

        client_ip = handler._get_client_ip()
        # SECURITY (MPG-F5): client-controlled zip_file_path is an
        # arbitrary-file-read/exfil vector -- contain it to the generated
        # output dir; on violation proceed without attachment (do not raise).
        safe_zip_path = _safe_zip_attachment_path(data.get("zip_file_path"))
        result = send_plan_email(
            recipient_email=data.get("email") or "",
            client_name=data.get("client_name") or "",
            plan_summary=data.get("plan_summary", {}),
            zip_file_path=safe_zip_path,
            sender_ip=client_ip,
        )
        handler._send_json(result)
    except Exception as e:
        logger.error("Plan delivery error: %s", e, exc_info=True)
        handler._send_json({"success": False, "message": "Delivery failed"})


def _handle_report_html(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/report/html -- generate HTML report."""
    try:
        content_len = int(handler.headers.get("Content-Length") or 0)
        body = handler.rfile.read(content_len) if content_len > 0 else b"{}"
        data = json.loads(body)
        from pdf_generator import generate_plan_html_report

        _plan_data = data.get("plan_data", data)
        _industry = (
            data.get("industry")
            or _plan_data.get("industry_label")
            or _plan_data.get("industry")
            or "Technology"
        )
        _industry = format_industry_label(
            _industry,
            data.get("naics_selected_code") or _plan_data.get("naics_selected_code"),
            data.get("naics_selected_title")
            or _plan_data.get("naics_selected_title"),
        )
        html_content = generate_plan_html_report(
            plan_data=_plan_data,
            client_name=data.get("client_name", "Client"),
            industry=_industry,
        )
        html_bytes = html_content.encode("utf-8")
        handler.send_response(200)
        handler.send_header("Content-Type", "text/html; charset=utf-8")
        handler.send_header(
            "Content-Disposition", "attachment; filename=media_plan_report.html"
        )
        handler.send_header("Content-Length", str(len(html_bytes)))
        handler.end_headers()
        handler.wfile.write(html_bytes)
    except Exception as e:
        logger.error("HTML report error: %s", e, exc_info=True)
        handler.send_response(500)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "Report generation failed"}).encode())


def _handle_nova_export(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/nova/export -- export Nova chat conversation as HTML."""
    try:
        content_len = int(handler.headers.get("Content-Length") or 0)
        body = handler.rfile.read(content_len) if content_len > 0 else b"{}"
        data = json.loads(body)
        from nova_export import export_conversation_html

        html_content = export_conversation_html(
            data.get("conversation_history") or [],
            data.get("metadata", {}),
        )
        html_bytes = html_content.encode("utf-8")
        handler.send_response(200)
        handler.send_header("Content-Type", "text/html; charset=utf-8")
        handler.send_header(
            "Content-Disposition",
            "attachment; filename=nova-conversation-export.html",
        )
        handler.send_header("Content-Length", str(len(html_bytes)))
        cors_origin = (
            handler._get_cors_origin() if hasattr(handler, "_get_cors_origin") else None
        )
        if cors_origin:
            handler.send_header("Access-Control-Allow-Origin", cors_origin)
        handler.end_headers()
        handler.wfile.write(html_bytes)
    except Exception as e:
        logger.error("Nova export error: %s", e, exc_info=True)
        handler.send_response(500)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "Export failed"}).encode())


# ---------------------------------------------------------------------------
# PDF Export handlers
# ---------------------------------------------------------------------------


def _handle_pdf_export_get(handler: Any, path: str, parsed: Any) -> None:
    """GET /api/plan/export/pdf?plan_id=xxx -- download PDF for a stored plan.

    Looks up plan data from _plan_results_store by plan_id, generates a
    branded PDF report, and returns it as a file download.
    """
    try:
        from urllib.parse import parse_qs

        qs = parse_qs(parsed.query) if hasattr(parsed, "query") else {}
        plan_id = (qs.get("plan_id") or [""])[0]

        # Both widths _plan_id has ever had: 32-char (128-bit) now, 12-char
        # legacy still inside its 24h TTL. Keep in step with app.py's
        # /api/plan-results/<id> validator -- they key the same store.
        if not plan_id or not re.match(r"^([a-f0-9]{12}|[a-f0-9]{32})\Z", plan_id):
            handler._send_json({"error": "Missing or invalid plan_id"}, status_code=400)
            return

        # Look up plan data: this worker's dict, else the shared layers
        # (the plan may have been generated on another worker / pre-deploy).
        _app = sys.modules.get("app") or sys.modules.get("__main__")
        _plan_results_store = getattr(_app, "_plan_results_store", {})
        _plan_results_lock = getattr(_app, "_plan_results_lock", None)

        _lookup = getattr(_app, "_plan_result_lookup", None)
        if _lookup is None:
            # app module not loaded (or renamed the helper): say so instead of
            # an AttributeError turning into a generic 500.
            handler._send_json(
                {"error": "Plan store unavailable"}, status_code=503
            )
            return
        entry = _lookup(plan_id)
        if entry and time.time() - (entry.get("created") or 0) > 86400:
            if _plan_results_lock:
                with _plan_results_lock:
                    _plan_results_store.pop(plan_id, None)
            entry = None

        if not entry:
            handler._send_json({"error": "Plan not found or expired"}, status_code=404)
            return

        plan_json = entry.get("data") or {}
        metadata = plan_json.get("metadata") or {}
        summary = plan_json.get("summary") or {}
        client_name = metadata.get("client_name") or "Client"
        industry = (
            metadata.get("industry_label") or metadata.get("industry") or "Technology"
        )

        # Build plan_data in the format pdf_report expects
        plan_data = {
            "budget": metadata.get("total_budget") or summary.get("total_budget") or 0,
            "channels": summary.get("channels") or [],
            "roles": metadata.get("roles") or [],
            "locations": metadata.get("locations") or [],
            "market_intelligence": plan_json.get("market_intelligence") or {},
            "recommendations": summary.get("recommendations")
            or plan_json.get("recommendations")
            or [],
            "timeline": plan_json.get("timeline") or [],
            "risk_analysis": plan_json.get("risk_analysis") or [],
            "competitive_landscape": plan_json.get("competitive_landscape") or [],
        }

        _generate_and_send_pdf(handler, plan_data, client_name, industry)

    except Exception as e:
        logger.error("PDF export GET error: %s", e, exc_info=True)
        handler._send_json({"error": "PDF export failed"}, status_code=500)


def _handle_pdf_export_post(handler: Any, path: str, parsed: Any) -> None:
    """POST /api/export/pdf -- generate PDF from posted plan data."""
    try:
        content_len = int(handler.headers.get("Content-Length") or 0)
        body = handler.rfile.read(content_len) if content_len > 0 else b"{}"
        data = json.loads(body)

        plan_data = data.get("plan_data") or data
        client_name = (
            data.get("client_name") or plan_data.get("client_name") or "Client"
        )
        industry = (
            data.get("industry")
            or plan_data.get("industry_label")
            or plan_data.get("industry")
            or "Technology"
        )
        # S94: append the user's precise NAICS pick (if any) -- this POST path
        # is fed straight from the wizard's client-side payload rather than
        # the server-enriched gen_data, so the suffix must be applied here too.
        industry = format_industry_label(
            industry,
            data.get("naics_selected_code") or plan_data.get("naics_selected_code"),
            data.get("naics_selected_title") or plan_data.get("naics_selected_title"),
        )

        # The wizard posts its own REQUEST here (budget text, role/location
        # tags, no channels). Rendered as-is that is a PDF reading "Total
        # Budget $0.00 / Channels 0 / No channel data available" -- a
        # document that looks like a plan and is empty (audit D-07). Refuse
        # it; the page then prints its on-screen plan instead.
        _channels = plan_data.get("channels") if isinstance(plan_data, dict) else None
        if not (isinstance(_channels, list) and _channels):
            handler._send_json(
                {
                    "error": (
                        "The PDF report needs the generated plan's channels; "
                        "this request has none."
                    )
                },
                status_code=422,
            )
            return

        # PostHog: track export event
        _track_export_event(plan_data, "pdf")

        _generate_and_send_pdf(handler, plan_data, client_name, industry)

    except json.JSONDecodeError:
        handler.send_response(400)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "Invalid JSON body"}).encode())
    except ImportError:
        handler.send_response(501)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(
            json.dumps(
                {"error": "PDF export unavailable -- reportlab not installed"}
            ).encode()
        )
    except Exception as e:
        logger.error("PDF export POST error: %s", e, exc_info=True)
        handler.send_response(500)
        handler.send_header("Content-Type", "application/json")
        handler.end_headers()
        handler.wfile.write(json.dumps({"error": "PDF generation failed"}).encode())


def _generate_and_send_pdf(
    handler: Any,
    plan_data: dict,
    client_name: str,
    industry: str,
) -> None:
    """Generate PDF and send as download response.

    Args:
        handler: HTTP request handler instance.
        plan_data: Plan data dict.
        client_name: Client/company name.
        industry: Industry vertical.
    """
    from pdf_report import generate_pdf_report

    pdf_bytes = generate_pdf_report(
        plan_data=plan_data,
        client_name=client_name,
        industry=industry,
    )

    safe_name = re.sub(r"[^a-zA-Z0-9_\-]", "_", client_name)
    filename = f"{safe_name}_Media_Plan_Report.pdf"

    handler.send_response(200)
    handler.send_header("Content-Type", "application/pdf")
    handler.send_header(
        "Content-Disposition",
        f'attachment; filename="{filename}"',
    )
    handler.send_header("Content-Length", str(len(pdf_bytes)))
    handler.end_headers()
    handler.wfile.write(pdf_bytes)


# ---------------------------------------------------------------------------
# Route map
# ---------------------------------------------------------------------------

_EXPORT_POST_ROUTE_MAP: dict[str, Any] = {
    "/api/export/sheets": _handle_export_sheets,
    "/api/export/status": _handle_export_status,
    "/api/export/pdf": _handle_pdf_export_post,
    "/api/deliver": _handle_deliver,
    "/api/report/html": _handle_report_html,
    "/api/nova/export": _handle_nova_export,
}
